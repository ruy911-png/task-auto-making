# -*- coding: utf-8 -*-
# v24 (2026-10-27): PR #2·이슈 #5 검토 반영 + 독립 검토 10회 — 합계형: 빈 칸 '누락'(구분 제목 줄은 원본 행 포함 관계로 가림, 하위 계정 일부만 담으면 알림), 저장값 없는 수식 '계산값 없음', 같은 계정 두 번 '확인 필요', 위쪽 소계형(값 있는 구분 줄)·합계 열 이중 합산 방지, 값 있는 구분 칸 아래 계정 일부 누락 알림 / 행 대조: 시트 관계를 층(분할·같은 내용 보기·발췌)으로 판정해 빠진 거래·중복 기재 보고, 작은 분할 시트 대조, 메모 시트는 '대조 안 한 시트' / 16자리 넘는 정수 번호 글자 비교, 실수 끝자리 허용, '재검산 불가 N칸', 한글 Windows 출력 보호
# v23 (2026-10-26): 행 대조에서 16자리 이상 숫자 글자(승인·카드·계좌 번호)를 글자로 비교 — 소수 변환으로 끝자리가 사라져 바뀐 번호를 일치로 보던 문제 수정. 검토 지적 28건을 자동 시험으로 고정
# v22 (2026-10-26): 구조 파악 후 행 대조 추가 — 원본 행을 옮겨 정리한 보고서·글자 줄 원본(PDF 추출 등)을 거래 단위로 맞춰 값 차이·누락·추가 찾기, 열 이름이 달라도 값으로 열 대응, 같은 행 계산식(단가=금액÷수량 등) 위반 찾기. 합계형 대조 결과는 v21과 동일
# v21 (2026-10-23): 동작 변경 없는 코드 정리 — 안 쓰는 상수·중복 import·중복 함수 제거, 보고서 시트 분석 1회로 줄임 (결과 v20과 동일)
# v20 (2026-10-23): 놓친 오류·거짓 오류 수정 — SUM 범위에서 빠진 칸·소계 이중 계산 찾기, 날짜·비율·소수 칸 재검산 오류 수정, 반복계산·쓰기암호 파일에서 Excel 멈춤 해결, .xlsm 지원, ₩·원·뒤 마이너스 금액, 원본 옆·아래 다른 표 제외, 코드 열·이름 붙은 소계·월별 열·차대 분리 원본 대조, 추가기능 함수·3D 참조·암시적 교차 처리, 결과 엑셀 서식(천 단위·머리글·줄바꿈·틀 고정)
# v19 (2026-10-22): 수식 글자 상수 보존, 배열 빈 글자 저장값 구분·출력 범위 영향 추적, 결과표 행/열 제목 분리
# v18 (2026-10-22): 기본 모드(formulas)에서도 배열수식 검산 — 배열 범위를 칸마다 다시 계산해 저장값과 비교. FILTER 등 _xlws 함수, 동적 배열 참조(C1#) 계산. 계산 못 한 배열수식만 '검사 불가'
# v17 (2026-10-22): Excel 재계산(USE_EXCEL=True)이 배열수식을 비교했는데도 '배열수식 검사 불가' 주의가 함께 나오던 것 제거
# v16 (2026-10-22): Excel 재계산(USE_EXCEL=True)에서 배열수식도 검산 — 배열 범위의 칸마다 저장값과 Excel 재계산값 비교. 그 외 v15와 같음
# v15 (2026-10-22): v14 속도 수정에서 생긴 차이 복구 — 표 중간의 병합 구분 줄 아래 칸이 열 제목을 잃던 것. 그 외 v14와 같음
# v14 (2026-10-22): 속도 — 보고서 칸이 많을 때 시간이 칸 수의 제곱으로 늘던 두 곳(보고서 칸 찾기, 결과 저장) 수정. 같은 파일을 여러 번 다시 읽던 것을 한 번만 읽음. 결과는 v13과 같음
# v13 (2026-10-22): 코드 정리 — 안 쓰는 함수 3개·중복 import 제거, 동작 변화 없음
# v12 (2026-10-22): 규칙표로 지정한 금액 열도 오류값·계산값 없는 수식 검사에 포함(숫자 열로 자동 분류되지 않아도). 정상 저장된 ="" 는 '계산값 없음'으로 보지 않음. 데이터 표 수식도 수식 개수에 포함('검사 불가' 표시)
# v11 (2026-10-22): 원본 금액의 #N/A·계산값 없는 수식을 0으로 더하지 않음(검산 불가 표시). 첫 데이터 줄의 오류값을 제목으로 오인하던 것 수정. 보고서 가로 소계(열 방향)는 구간 합으로 대조. 배열수식을 참조하는 SUBTOTAL에서 멈추던 것 수정
# v10 (2026-10-22): 보고서 라벨에 코드·괄호만 붙은 경우('1010 현금', '현금(원화)' ↔ 원본 '현금')는 '유사 확인'이 아니라 '일치'. 보고서 합계·소계 열은 합계 행처럼 '제외'
# v9 (2026-10-22): 원본 제목 줄 수정 — 띄어쓴 제목('금 액'), 빈·공백 제목 열은 '(C열)'로 이름 붙여 사용, 같은 제목 중복은 '금액(D열)'로 구분, 두 줄 제목은 숫자 바로 윗줄을 제목으로. 원본 제목이 '금 액'처럼 띄어 쓴 경우 '금액 열 없음'으로 멈추던 문제 수정. 금액 칸의 공백 글자(' '·전각공백·탭)를 오류값으로 보던 v7 버그 수정(빈칸으로 처리)
# v8 (2026-10-22): 실행 중 진행 상황을 터미널에 표시 (단계 시작·끝, 걸린 시간). 검사 내용·결과 파일은 v7과 같음
# v7 (2026-10-22): 원본 금액 칸의 오류값(#VALUE! 등)·글자를 0으로 더하지 않음 — 그 값이 들어간 집계만 '검산 불가'로 막고 원본 셀 위치 표시. 배열수식을 수식 개수에 포함하고 '배열수식 검사 불가' 위치 표시
# v6 (2026-10-22): Excel 재계산 — INDIRECT·TODAY 등 휘발성 함수·외부링크가 1개만 있어도 통합문서 전체를 '검산 불가'로 처리하던 것을 그 셀과 의존 셀로 한정. Excel 작업자 실패 이유(표준출력) 보존, 단계별 소요시간 표시
# v5 (2026-10-21): nan/inf 숫자 오인 수정(대조에서 불일치가 숨던 버그), 텍스트서식·비표준 오류값 가짜 경고 제거, x: 접두어 XML(.NET 저장 형식)에서도 깨진 서식 복구·위치 찾기 동작
# v4 (2026-10-21): 결과 파일명에 실행 시각(_YYYYMMDD_HHMMSS) — 재실행·결과 파일 열림 충돌 없음, 이전 결과 보존
# v3 (2026-10-21): Excel 실행 없이 원본·보고서 검산, 실패 위치·이유 확인
# 실행 파일은 이것 하나입니다. 원본은 수정하지 않습니다.
# 기본 검산에는 Microsoft Excel·pywin32가 필요하지 않습니다.
# 필요: openpyxl. 선택: formulas 1.3.4 (수식 계산 범위 확대).
# 숫자 값 보고서는 원본 집계와 대조하며, 계산할 수 없는 수식은 사유를 표시합니다.
# 사용: 같은 폴더에 xlsx 1개 → 수식 검사 / 2개 → 원본·보고서 대조 및 보고서 수식 검사
# 명시: py excel_automation_v24.py 원본.xlsx 보고서.xlsx
# 채움: py excel_automation_v24.py 원본.xlsx 보고서.xlsx --채우기
# 수식: py excel_automation_v24.py 보고서.xlsx --수식
# 원본은 수정하지 않으며 결과와 규칙표를 별도 파일로 저장합니다.
# 양식 자동 판별은 한계가 있으므로 출력 매핑과 규칙표를 확인하세요.
# ===== 여기만 수정 =====
FOLDER = r""       # 비우면 이 파이썬 파일이 있는 폴더
RAW_FILE = r""     # 원본 파일명 또는 전체 경로 (자동 판별이 틀리면 지정)
REPORT_FILE = r""  # 보고서 파일명 또는 전체 경로
MODE = "자동"     # 자동 / 검산 / 수식 / 채우기
TOLERANCE = 0.5    # 보고서 단위 기준 허용 오차 (규칙표 값이 우선)
RECALC_TOL = 0.5   # 수식 재계산 허용 오차 (단순 합계 재검산)
RECALC_EPS = 1e-6  # 전체 재검산 허용 오차: 저장값과 재계산값은 같은 정밀도라 아주 작게 (0.5면 비율·소수 칸 변조를 놓침)
cf_FORMULAS_UNRELIABLE = {"GROUPBY", "PIVOTBY", "SINGLE", "PERCENTRANK", "PERCENTRANK.INC", "PERCENTRANK.EXC", "LINEST", "LOGEST", "MUNIT", "BAHTTEXT", "TBILLEQ", "TBILLPRICE", "BYROW", "BYCOL"}   # 2026-10-23 실제 Excel 대비 결과가 달랐던 함수
USE_EXCEL = False  # 기본은 Excel 실행 안 함. True는 Windows·Excel·pywin32가 있을 때만 사용
EXCEL_TIMEOUT = 120  # Excel 작업 최대 대기 초. 실패하면 이유를 남기고 대체 검사
EXCLUDE_DEFAULT = "소계, 합계, 계, 총계, 누계, 총합계, subtotal, total"
# =======================
try:
    import openpyxl
except ImportError:
    print("최종: 실행 불가 — openpyxl이 필요합니다. py -m pip install openpyxl")
    raise SystemExit(1)

# 수식 검사
import os, re, sys, glob, datetime, zipfile, math, io, json, numbers, posixpath, shutil, subprocess, tempfile, contextlib, functools, time
import xml.etree.ElementTree as ET

_PROGRESS = False          # main()에서 켬. 라이브러리처럼 불러 쓸 때는 출력 없음
_PROGRESS_T0 = None


def _progress(msg):
    """진행 상황 한 줄 (경과 시간 포함). 결과에는 영향 없음"""
    if not _PROGRESS:
        return
    elapsed = time.monotonic() - (_PROGRESS_T0 or time.monotonic())
    print(f"[{elapsed:6.1f}초] {msg}", flush=True)
from collections import Counter
from collections import defaultdict
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.formula import Tokenizer
cf_FORMULAS_IMPORT_ERROR = None
try:
    import formulas as _formulas
    cf_HAS_FORMULAS = True
except Exception as import_error:
    _formulas = None
    cf_HAS_FORMULAS = False
    cf_FORMULAS_IMPORT_ERROR = f"{type(import_error).__name__}: {import_error}"

# 외부 파일 링크: [1]Sheet1!A1  또는  '[파일.xlsx]Sheet1'!A1  (표 참조 매출표[금액]과 구분)
cf_EXT = re.compile(r"(?:^|[^\w가-힣.\]])\[(?:\d+|[^\[\]]*\.xls[xmb]?)\]", re.I)
cf_ERRORS_ALL = {"#NULL!", "#DIV/0!", "#VALUE!", "#REF!", "#NAME?", "#NUM!", "#N/A", "#GETTING_DATA", "#SPILL!", "#CALC!", "#FIELD!", "#BLOCKED!", "#CONNECT!", "#BUSY!", "#UNKNOWN!", "#PYTHON!"}
# 시트참조:  'RAW 데이터'!A1  /  RAW!$B$2:$B$9  /  [1]Sheet1!A1
cf_REF = re.compile(r"(?:'((?:[^']|'')+)'|(\[[^\]]+\][^'!=(),\s]+|[^'!=(),\s+\-*/^&<>]+))!\$?([A-Z]{1,3})\$?(\d+)(?::\$?([A-Z]{1,3})\$?(\d+))?")


def cf__clean_copy(src, styles):
    """임시 복사본 생성. styles=True면 셀·행·열 서식 번호와 조건부서식까지 제거 (openpyxl IndexError 우회)"""
    fd, tmp = tempfile.mkstemp(suffix=".xlsx")
    os.close(fd)  # Windows에서는 열린 NamedTemporaryFile을 다시 열 수 없음
    cf = re.compile(rb"<(?:\w+:)?conditionalFormatting\b.*?</(?:\w+:)?conditionalFormatting>|<(?:\w+:)?conditionalFormatting\b[^>]*/>", re.S)
    ext = re.compile(rb"<(?:\w+:)?extLst>.*?</(?:\w+:)?extLst>", re.S)
    cell_s = re.compile(rb'(<(?:\w+:)?c\b[^>]*?)\s+s="\d+"')
    row_s = re.compile(rb'(<(?:\w+:)?row\b[^>]*?)\s+s="\d+"')
    col_s = re.compile(rb'(<(?:\w+:)?col\b[^>]*?)\s+style="\d+"')
    try:
        with zipfile.ZipFile(src) as zi, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zo:
            for n in zi.namelist():
                d = zi.read(n)
                if n.startswith("xl/worksheets/sheet") and n.endswith(".xml"):
                    d = ext.sub(b"", cf.sub(b"", d))
                    if styles:
                        d = col_s.sub(rb"\1", row_s.sub(rb"\1", cell_s.sub(rb"\1", d)))
                zo.writestr(n, d)
    except Exception:
        os.remove(tmp)
        raise
    return tmp


_WB_CACHE = {}


def cf_safe_load(path, warnings=None, fresh=False, **kw):
    """load_workbook. 실패하면 1차(조건부서식 제거) → 2차(서식 번호까지 제거)로 재시도하고 경고 기록.
    같은 실행 안에서 같은 파일을 같은 방식으로 다시 읽으면 처음 읽은 것을 그대로 씀 (fresh=True면 새로 읽음: 수정해서 저장할 때)"""
    try:
        st = os.stat(path)
        key = (os.path.abspath(path), st.st_mtime_ns, st.st_size, bool(kw.get("data_only")), bool(kw.get("read_only")))
    except OSError:
        key = None
    if key is not None and not fresh and key in _WB_CACHE:
        wb, warns = _WB_CACHE[key]
        if warnings is not None:
            warnings.extend(w for w in warns if w not in warnings)
        return wb
    warns = []
    wb = _cf_load_uncached(path, warns, **kw)
    if warnings is not None:
        warnings.extend(warns)
    if key is not None and not fresh:
        _WB_CACHE[key] = (wb, warns)
    return wb


def _cf_load_uncached(path, warnings, **kw):
    try:
        if kw.get("read_only"):
            return load_workbook(path, **kw)
        with open(path, "rb") as stream:
            return load_workbook(stream, **kw)
    except Exception as e:
        for level, msg in ((False, "조건부서식을 읽지 못해 무시하고 진행 (결과 파일에서 조건부서식 사라짐)"),
                           (True, "서식 정보가 깨져 있어 서식 제외하고 진행 (결과 파일에서 셀 서식·조건부서식 사라짐, 데이터·수식은 유지)")):
            tmp = None
            try:
                tmp = cf__clean_copy(path, level)
                # 읽기 전용도 파일 핸들을 남기지 않도록 복구본은 메모리에서 읽음.
                with open(tmp, "rb") as stream:
                    wb = load_workbook(io.BytesIO(stream.read()), **kw)
            except Exception:
                continue
            finally:
                if tmp is not None:
                    os.remove(tmp)
            if warnings is not None:
                warnings.append(f"{msg}: {os.path.basename(path)}")
            return wb
        raise e


def cf__sheet_part_names(z):
    """{zip 내부 시트 경로: 시트 이름} — 네임스페이스 인식 XML 파싱 (x: 접두어·속성 순서 무관)"""
    main = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    rel_ns = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
    out = {}
    try:
        wb = ET.fromstring(z.read("xl/workbook.xml"))
        rels = ET.fromstring(z.read("xl/_rels/workbook.xml.rels"))
    except (KeyError, ET.ParseError):
        return out
    tgt = {r.get("Id"): r.get("Target", "") for r in rels if r.get("TargetMode") != "External"}
    for sh in wb.iter(main + "sheet"):
        t = tgt.get(sh.get(rel_ns), "")
        if not t:
            continue
        t = t.lstrip("/") if t.startswith("/") else "xl/" + t
        out[t.replace("xl/../", "")] = sh.get("name")
    return out


def cf_broken_styles(path, limit=50):
    """서식 번호가 서식표(styles.xml) 범위를 벗어난 위치 목록. [(위치, 상세), ...], 전체 개수"""
    with zipfile.ZipFile(path) as z:
        names = z.namelist()
        st = z.read("xl/styles.xml") if "xl/styles.xml" in names else b""
        m = re.search(rb'<(?:\w+:)?cellXfs[^>]*count="(\d+)"', st); n_xf = int(m.group(1)) if m else len(re.findall(rb"<(?:\w+:)?xf\b", st))
        m = re.search(rb'<(?:\w+:)?dxfs[^>]*count="(\d+)"', st); n_dxf = int(m.group(1)) if m else len(re.findall(rb"<(?:\w+:)?dxf\b", st))
        # 시트 파일명 → 시트 이름
        sheet_name = cf__sheet_part_names(z)
        out, total = [], 0
        def push(loc, detail):
            nonlocal total
            total += 1
            if len(out) < limit:
                out.append((loc, detail))
        for n in names:
            if not (n.startswith("xl/worksheets/sheet") and n.endswith(".xml")):
                continue
            sn = sheet_name.get(n, n.split("/")[-1])
            d = z.read(n)
            for mm in re.finditer(rb'<(?:\w+:)?col\b[^>]*\bstyle="(\d+)"[^>]*>', d):
                idx = int(mm.group(1))
                if idx >= n_xf:
                    mn = re.search(rb'min="(\d+)"', mm.group(0)); col = get_column_letter(int(mn.group(1))) if mn else "?"
                    push(f"{sn}!{col}열", f"열 서식번호 {idx} / 서식표 {n_xf}개")
            for mm in re.finditer(rb'<(?:\w+:)?row\b[^>]*\br="(\d+)"[^>]*\bs="(\d+)"[^>]*>', d):
                if int(mm.group(2)) >= n_xf:
                    push(f"{sn}!{mm.group(1).decode()}행", f"행 서식번호 {mm.group(2).decode()} / 서식표 {n_xf}개")
            for mm in re.finditer(rb'<(?:\w+:)?c\b[^>]*>', d):
                tag = mm.group(0)
                si = re.search(rb'\bs="(\d+)"', tag); ri = re.search(rb'\br="([A-Z]+\d+)"', tag)
                if si and int(si.group(1)) >= n_xf:
                    push(f"{sn}!{ri.group(1).decode() if ri else '?'}", f"셀 서식번호 {si.group(1).decode()} / 서식표 {n_xf}개")
            for blk in re.finditer(rb'<(?:\w+:)?conditionalFormatting\b[^>]*sqref="([^"]+)"[^>]*>(.*?)</(?:\w+:)?conditionalFormatting>', d, re.S):
                for mm in re.finditer(rb'dxfId="(\d+)"', blk.group(2)):
                    if int(mm.group(1)) >= n_dxf:
                        push(f"{sn}!{blk.group(1).decode()} 조건부서식", f"조건부서식 서식번호 {mm.group(1).decode()} / 조건부서식표 {n_dxf}개")
        return out, total


def cf__last_data_row(ws, col_letter):
    c = column_index_from_string(col_letter)
    last = 0
    for r in range(1, ws.max_row + 1):
        if ws.cell(r, c).value not in (None, ""):
            last = r
    return last


# 구조적 참조: 매출표[금액] / 매출표[[#Data],[금액]] / 매출표[#All]   ([@금액]은 표 안 자기참조 → 제외)
cf_TREF = re.compile(r"(?<![\w가-힣.'])([A-Za-z_가-힣][\w가-힣.]*)\[")
cf_SPECIAL = {"#All", "#Data", "#Headers", "#Totals", "#This Row"}


def cf__tables(wb):
    """{표이름: {"sheet", "ref", "cols": [...], "first_row", "last_row", "col_idx": [..]}}"""
    from openpyxl.utils import range_boundaries
    out = {}
    for ws in wb.worksheets:
        for name, t in ws.tables.items():
            ref = t if isinstance(t, str) else t.ref        # openpyxl 버전에 따라 문자열/객체
            c1, r1, c2, r2 = range_boundaries(ref)
            cols = [ws.cell(r1, c).value for c in range(c1, c2 + 1)]
            out[name] = {"sheet": ws.title, "ref": ref, "cols": cols, "first_row": r1, "last_row": r2,
                         "col_idx": list(range(c1, c2 + 1))}
    return out


def cf__defined_names(wb):
    """{이름: 'RAW!$B$2:$B$10'}  (시트!범위 형태만)"""
    out = {}
    try:
        items = wb.defined_names.items()
    except AttributeError:                       # 구버전 openpyxl
        items = [(d.name, d) for d in wb.defined_names.definedName]
    for name, d in items:
        if d.attr_text and "!" in d.attr_text and "[" not in d.attr_text:
            out[name] = d.attr_text
    return out


def cf__bracket(formula, start):
    """formula[start]가 '[' 일 때 짝이 맞는 ']'까지의 내용 반환"""
    depth, i = 0, start
    while i < len(formula):
        if formula[i] == "[": depth += 1
        elif formula[i] == "]":
            depth -= 1
            if depth == 0:
                return formula[start + 1:i]
        i += 1
    return formula[start + 1:]


def cf__without_text_literals(formula):
    """Excel 큰따옴표 문자열은 참조·오류 토큰이 아니다 ("" 이스케이프 포함)."""
    return re.sub(r'"(?:[^"]|"")*"', '""', formula)


def cf__table_refs(formula):
    """[(표이름, [열이름,...]), ...]"""
    res = []
    formula = cf__without_text_literals(formula)
    for m in cf_TREF.finditer(formula):
        inner = cf__bracket(formula, m.end() - 1)
        cols = [c for c in re.findall(r"\[([^\[\]]+)\]", inner)] or ([inner] if inner else [])
        cols = [c.lstrip("@") for c in cols if c.lstrip("@") not in cf_SPECIAL and not c.startswith("#")]
        res.append((m.group(1), cols))
    return res


def cf__refs(formula):
    for m in cf_REF.finditer(cf__without_text_literals(formula)):
        sheet = (m.group(1) or m.group(2)).replace("''", "'")
        yield sheet, m.group(3), int(m.group(4)), m.group(5) or m.group(3), int(m.group(6) or m.group(4))


# ---------------------------------------------------------------- v6: 값·수식 패턴 검사
def cf__norm_saved(v, epoch=None):
    """저장값 정규화: 날짜·시간 서식 칸(openpyxl이 datetime으로 읽음) → 엑셀 일련번호, '_x0003_' 같은 제어문자 표기 → 실제 글자"""
    if isinstance(v, (datetime.datetime, datetime.date)) and (v.year, v.month) < (1900, 3):
        return None                                          # 1900-02-29 버그 구간: 일련번호 하루 어긋남 → 대조하지 않음
    if isinstance(v, (datetime.datetime, datetime.date, datetime.time, datetime.timedelta)):
        try:
            from openpyxl.utils.datetime import to_excel, CALENDAR_WINDOWS_1900
            return to_excel(v, CALENDAR_WINDOWS_1900)       # formulas는 1900 날짜 체계로 계산 → 1904 체계 파일도 1900 일련번호로 맞춤
        except Exception:
            return v
    if isinstance(v, str) and "_x" in v:
        return re.sub(r"_x([0-9A-Fa-f]{4})_", lambda m: chr(int(m.group(1), 16)), v)
    return v


def cf__differs(cached, val):
    """저장값과 재계산값이 다른가 (숫자는 허용 오차 + 아주 큰 수는 상대 오차)"""
    cv, vv = cf__to_num(cached), cf__to_num(val)
    if cv is not None and vv is not None:
        return abs(cv - vv) > max(RECALC_EPS, abs(cv) * 1e-9)
    return str(cf__norm_saved(cached)).strip() != str(cf__norm_saved(val)).strip()   # '_x0003_' 표기는 양쪽 다 실제 글자로


def cf__to_num(v):
    """'1,000' → 1000.0 / 숫자 → float / 아니면 None"""
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v) if math.isfinite(v) else None
    if isinstance(v, str):
        t = v.replace(",", "").strip()
        if t == "":
            return None
        try:
            f = float(t)
        except ValueError:
            return None
        return f if math.isfinite(f) else None              # 'nan'·'inf' 글자는 숫자 아님
    return None


cf_CELL_RE = re.compile(r"(\$?)([A-Z]{1,3})(\$?)(\d+)")


def cf__canon(formula, row, col):
    """수식을 '현재 셀 기준 상대 위치' 형태로 바꿈 (엑셀 R1C1 과 같은 원리). 같은 패턴이면 같은 문자열"""
    try:
        toks = Tokenizer(formula).items
    except Exception:
        return formula
    out = []
    for t in toks:
        if t.type == "OPERAND" and t.subtype == "RANGE":
            def rep(m):
                c_abs, c, r_abs, r = m.groups()
                ci, ri = column_index_from_string(c), int(r)
                cs = f"C{ci}" if c_abs else f"C[{ci - col}]"
                rs = f"R{ri}" if r_abs else f"R[{ri - row}]"
                return rs + cs
            out.append(cf_CELL_RE.sub(rep, t.value))
        else:
            out.append(t.value)
    return "".join(out)


def cf__pattern_checks(ws, add):
    """같은 열(세로 연속 구간)에서 대부분이 같은 수식인데 일부만 값·다른 수식 → 표시 (엑셀 '영역 수식 불일치' 규칙)"""
    max_r, max_c = ws.max_row, ws.max_column
    for c in range(1, max_c + 1):
        cf_run = []
        for r in range(1, max_r + 2):
            v = ws.cell(r, c).value if r <= max_r else None
            is_f = isinstance(v, str) and v.startswith("=")
            is_n = cf__to_num(v) is not None and not isinstance(v, str)
            if is_f or is_n:
                cf_run.append((r, v, is_f))
            else:
                cf__judge_run(ws, c, cf_run, add); cf_run = []


_AGG_RE = re.compile(r"^=\s*(SUM|SUBTOTAL\(\s*(?:9|109)\s*,)\s*\(?\s*\$?([A-Z]{1,3})\$?(\d+)\s*:\s*\$?([A-Z]{1,3})\$?(\d+)\s*\)\s*$", re.I)


def cf__omitted_cells(ws, add):
    """합계 수식 범위와 수식 칸 사이에 숫자가 있는데 범위에서 빠짐 → 표시 (엑셀 '인접 칸 빠짐' 규칙과 같은 취지).
    예: B12 =SUM(B2:B10) 인데 B11 이 숫자. 사이 칸이 소계 수식이면 제외"""
    for row in ws.iter_rows():
        for cell in row:
            v = cell.value
            if not (isinstance(v, str) and v.startswith("=")):
                continue
            m = _AGG_RE.match(v.replace(" ", ""))
            if not m:
                continue
            c1, r1, c2, r2 = column_index_from_string(m.group(2)), int(m.group(3)), column_index_from_string(m.group(4)), int(m.group(5))
            r, c = cell.row, cell.column
            if m.group(1).upper() == "SUM":                      # SUM 범위 안에 소계 수식이 있으면 이중 계산 (SUBTOTAL은 안쪽 SUBTOTAL을 자동 제외)
                inner = []
                for rr in range(min(r1, r2), max(r1, r2) + 1):
                    for cc in range(min(c1, c2), max(c1, c2) + 1):
                        g = ws.cell(rr, cc).value
                        mm = _AGG_RE.match(g.replace(" ", "")) if isinstance(g, str) and g.startswith("=") else None
                        if mm:
                            a1, b1, a2, b2 = column_index_from_string(mm.group(2)), int(mm.group(3)), column_index_from_string(mm.group(4)), int(mm.group(5))
                            if min(c1, c2) <= a1 <= a2 <= max(c1, c2) and min(r1, r2) <= b1 <= b2 <= max(r1, r2):
                                inner.append(f"{get_column_letter(cc)}{rr}")
                if inner:
                    add("주의", "합계에 소계 중복 포함", f"{ws.title}!{cell.coordinate}",
                        f"{v} 범위 안의 소계 칸 {', '.join(inner)} 이 그 소계의 상세 칸과 함께 더해짐 (이중 계산)")
            gap = []
            if c1 == c2 == c and r > r2 >= r1:                  # 세로 합계: 범위 끝 ~ 수식 칸 사이
                gap = [(rr, c) for rr in range(r2 + 1, r)]
            elif r1 == r2 == r and c > c2 >= c1:                # 가로 합계
                gap = [(r, cc) for cc in range(c2 + 1, c)]
            if not gap or len(gap) > 3:
                continue
            nums = []
            for rr, cc in gap:
                g = ws.cell(rr, cc).value
                if isinstance(g, str) and g.startswith("=") and re.search(r"\b(SUM|SUBTOTAL)\s*\(", g, re.I):
                    nums = []; break                            # 사이 칸이 소계 → 의도된 구조
                if isinstance(g, (int, float)) and not isinstance(g, bool) or (isinstance(g, str) and g.startswith("=")):
                    nums.append(f"{get_column_letter(cc)}{rr}")
            if nums:
                add("주의", "합계 범위에서 빠진 인접 칸", f"{ws.title}!{cell.coordinate}",
                    f"{v} 의 범위 바로 다음 칸 {', '.join(nums)} 에 숫자가 있는데 합계에서 빠짐 (행·열을 추가하고 범위를 안 늘린 경우)")


def cf__shape(formula, row, col):
    """수식 '모양': 상대 위치 + 따옴표 글자 제거 (HYPERLINK("…") 처럼 글자만 다른 수식을 같은 모양으로)"""
    try:
        toks = Tokenizer(formula).items
    except Exception:
        return formula
    if any(t.type == "OPERAND" and t.subtype == "TEXT" for t in toks):
        return re.sub(r'"(?:[^"]|"")*"', '""', cf__canon(formula, row, col))
    return cf__canon(formula, row, col)


def cf__judge_run(ws, c, cf_run, add):
    if len(cf_run) < 3:
        return
    forms = [(r, cf__canon(v, r, c)) for r, v, is_f in cf_run if is_f]
    if len(forms) < 2:
        return
    pat, n = Counter(p for _, p in forms).most_common(1)[0]
    if n < 2 or n < len(forms) * 0.5:                 # 정확한 패턴 과반 없음 → 모양(글자 무시)으로 '수식 자리에 값'만 판단
        shapes = Counter(cf__shape(v, r, c) for r, v, is_f in cf_run if is_f)
        shp, ns = shapes.most_common(1)[0]
        if ns < 3 or ns < len(forms) * 0.8:
            return
        f_rows = [r for r, _, is_f in cf_run if is_f]
        lo, hi = min(f_rows), max(f_rows)
        for r, v, is_f in cf_run:
            if not is_f and lo < r < hi:
                add("오류", "수식 자리에 값", f"{ws.title}!{get_column_letter(c)}{r}",
                    f"위아래는 같은 모양의 수식인데 이 칸만 숫자 {v} (따옴표 안 글자만 다른 수식 {ns}개)")
        return
    f_rows = [r for r, _, is_f in cf_run if is_f]
    lo, hi = min(f_rows), max(f_rows)
    for r, v, is_f in cf_run:
        loc = f"{ws.title}!{get_column_letter(c)}{r}"
        if not (lo < r < hi):                          # 구간 위·아래 끝(입력값·합계 줄)은 제외
            continue
        if not is_f:
            add("오류", "수식 자리에 값", loc, f"위아래는 수식인데 이 칸만 숫자 {v} (열의 수식 {n}/{len(cf_run)}개가 같은 패턴)")
        elif cf__canon(v, r, c) != pat:
            add("주의", "주변 수식과 불일치", loc, f"{v}  (열의 다른 수식 {n}개와 패턴이 다름)")




def cf__range_cells(wbv, default_ws, ref):
    """'RAW'!A2:B4 → (시트, [(r,c),...]) / 못 읽으면 None"""
    m = re.match(r"^(?:'((?:[^']|'')+)'|([^'!]+))!(.+)$", ref.strip())
    if m:
        sheet = (m.group(1) or m.group(2)).replace("''", "'"); rng = m.group(3)
    else:
        sheet, rng = default_ws.title, ref.strip()
    if sheet not in wbv.sheetnames:
        return None
    ws = wbv[sheet]
    rng = rng.replace("$", "")
    m2 = re.match(r"^([A-Z]{1,3})(\d+)(?::([A-Z]{1,3})(\d+))?$", rng)
    if not m2:
        return None
    c1, r1 = column_index_from_string(m2.group(1)), int(m2.group(2))
    c2 = column_index_from_string(m2.group(3)) if m2.group(3) else c1
    r2 = int(m2.group(4)) if m2.group(4) else r1
    r2 = min(r2, ws.max_row)
    return ws, [(r, c) for r in range(r1, r2 + 1) for c in range(c1, c2 + 1)]


def cf__recalc_simple(wbf, wbv, add, skip=None):
    """지원이 확실한 SUM·SUBTOTAL(9/109)·덧셈뺄셈만 캐시와 대조."""
    n_done, n_nocache, gaps = 0, 0, []
    for ws in wbf.worksheets:
        wv = wbv[ws.title]
        for row in ws.iter_rows():
            for cell in row:
                f = cell.value
                if not (isinstance(f, str) and f.startswith("=")):
                    continue
                loc = f"{ws.title}!{cell.coordinate}"
                total, ok, reason = 0.0, True, ""
                m = re.fullmatch(r"=\s*(SUM|SUBTOTAL)\s*\((.*)\)\s*", f, re.I)
                if m:
                    args = [a.strip() for a in m.group(2).split(",")]
                    subtotal = m.group(1).upper() == "SUBTOTAL"
                    exclude_hidden = False
                    if subtotal:
                        if not args or args[0] not in ("9", "109"):
                            continue
                        exclude_hidden = args.pop(0) == "109"
                    if not args or not all(args):
                        continue
                    for a in args:
                        rc = cf__range_cells(wbv, ws, a)
                        if rc is None:
                            ok = False; break
                        tws, cells = rc
                        source = wbf[tws.title]
                        if subtotal and (source.auto_filter.filterColumn or any(
                                t.autoFilter and t.autoFilter.filterColumn for t in source.tables.values())):
                            ok, reason = False, "필터로 제외된 행은 자동 판정할 수 없음"; break
                        blank = [(r, c) for r, c in cells if tws.cell(r, c).value in (None, "")
                                 and source.cell(r, c).data_type != "f"]
                        if blank and len(blank) < len(cells):
                            filled = [(r, c) for r, c in cells if tws.cell(r, c).value not in (None, "")]
                            if filled:
                                lo, hi = min(filled), max(filled)
                                inner = [(r, c) for r, c in blank if lo <= (r, c) <= hi]
                                if inner:
                                    gaps.append((loc, f"{tws.title}!" + ", ".join(f"{get_column_letter(c)}{r}" for r, c in inner[:5])))
                        for r, c in cells:
                            original = source.cell(r, c)
                            if original.data_type == "f" and not isinstance(original.value, str):
                                ok, reason = False, "참조 범위에 배열수식 있음 (검사 불가)"; break
                            if subtotal and original.data_type == "f" and re.search(r"\bSUBTOTAL\s*\(", original.value, re.I):
                                continue
                            if exclude_hidden and source.row_dimensions[r].hidden:
                                continue
                            value = tws.cell(r, c).value
                            if original.data_type == "f" and value is None:
                                ok, reason = False, "참조한 수식에 저장된 계산값 없음"; break
                            if tws.cell(r, c).data_type == "e":
                                ok, reason = False, "참조 범위에 오류값 있음"; break
                            # Excel SUM의 범위 인수는 텍스트 숫자와 논리값을 무시함.
                            if isinstance(value, (int, float)) and not isinstance(value, bool):
                                total += value
                        if not ok:
                            break
                elif re.fullmatch(r"\$?[A-Z]{1,3}\$?\d+(\s*[+\-]\s*(\$?[A-Z]{1,3}\$?\d+|\d+(\.\d+)?))+", f[1:].strip()):
                    for sign, ref in re.findall(r"([+\-]?)\s*(\$?[A-Z]{1,3}\$?\d+|\d+(?:\.\d+)?)", f[1:]):
                        if re.match(r"^\d", ref):
                            value = float(ref)
                        else:
                            rc = cf__range_cells(wbv, ws, ref)
                            if rc is None:
                                ok = False; break
                            tws, cells = rc
                            if not cells:
                                value = 0  # 사용 범위 밖의 실제 빈 셀
                            else:
                                original = wbf[tws.title].cell(*cells[0])
                                value = tws.cell(*cells[0]).value
                                if original.data_type == "f" and value is None:
                                    ok, reason = False, "참조한 수식에 저장된 계산값 없음"; break
                                if value in (None, ""):
                                    value = 0
                                elif isinstance(value, bool):
                                    value = int(value)
                                elif not isinstance(value, (int, float)):
                                    ok, reason = False, "덧셈뺄셈 참조값의 숫자 변환을 확인할 수 없음"; break
                        total += -value if sign == "-" else value
                else:
                    continue
                if skip is not None and loc in skip:
                    continue
                if not ok:
                    if reason:
                        add("정보", "재검산 불가", loc, reason)
                    continue
                cached = wv[cell.coordinate].value
                if cached is None:
                    n_nocache += 1; continue
                cv = cf__to_num(cached)
                if cv is None:
                    continue
                n_done += 1
                if abs(cv - total) > RECALC_TOL:
                    add("오류", "합계 재검산 불일치", loc, f"저장된 값 {cv:,.2f} / 직접 더한 값 {total:,.2f}  수식: {f}")
    for loc, detail in gaps:
        add("주의", "참조 범위 안 빈 칸", loc, f"범위 중간이 비어 있음: {detail}")
    if n_nocache:
        add("정보", "재검산 불가", "전체", f"수식 {n_nocache}개에 저장된 계산값 없음 → 엑셀에서 열어 저장한 뒤 다시 실행하면 재검산됨")
    if skip is None or n_done:
        add("정보", "합계 재검산", "전체", f"{n_done}개 수식 검산 (SUM·SUBTOTAL·덧셈뺄셈만)")


# ---------------------------------------------------------------- v7: 전체 재계산 (formulas)
def cf__implicit_range(formula):
    """함수 밖에서 범위를 바로 계산에 쓰는 옛 방식 수식인가 (예: =B1:B10*2 → Excel은 같은 행/열 한 칸만 씀)"""
    try:
        depth = 0
        for t in Tokenizer(formula).items:
            if t.type == "FUNC":
                depth += 1 if t.subtype == "OPEN" else -1
            elif t.type == "OPERAND" and t.subtype == "RANGE" and depth == 0 and ":" in t.value and not t.value.endswith("#"):
                return True
    except Exception:
        return False
    return False


def cf__multi(v):
    """formulas 결과가 2칸 이상 배열인가 (일반 수식인데 배열이 나오면 Excel은 암시적 교차로 한 칸만 씀)"""
    try:
        v = getattr(v, "value", v)
        return hasattr(v, "size") and getattr(v, "ndim", 0) > 0 and v.size > 1
    except Exception:
        return False


def cf__xl_scalar(v):
    """formulas 결과(배열) → 스칼라. (값, 오류문자열 or None)"""
    try:
        import numpy as np
        while hasattr(v, "shape") and getattr(v, "ndim", 0) > 0:
            if v.size == 0:
                return None, None
            v = v.flat[0] if hasattr(v, "flat") else v[0]
        if hasattr(v, "value"):
            v = v.value
            return cf__xl_scalar(v)
        if isinstance(v, np.generic):
            v = v.item()
    except Exception:
        pass
    s = str(v)
    if s.startswith("#"):
        return None, s
    return v, None


def cf__recalc_full(path, wbf, wbv, add):
    with tempfile.TemporaryDirectory(prefix="excel_recalc_") as tmpdir:
        return cf__recalc_full_in_dir(path, wbf, wbv, add, tmpdir)


def cf__recalc_full_in_dir(path, wbf, wbv, add, tmpdir):
    """formulas 로 파일 전체를 다시 계산해 저장값과 대조. 반환: 불일치가 발견된 셀 집합 (중복 표시 방지용) / 실패 시 None"""
    if not cf_HAS_FORMULAS:
        if cf_FORMULAS_IMPORT_ERROR == "ModuleNotFoundError: No module named 'formulas'":
            detail = "formulas 라이브러리 미설치. 설치: py -m pip install formulas"
        elif cf_FORMULAS_IMPORT_ERROR:
            detail = f"formulas 불러오기 실패: {cf_FORMULAS_IMPORT_ERROR}"
        else:
            detail = "formulas 라이브러리 사용 중지"
        add("정보", "전체 재검산 불가", "전체", detail + " → SUM·SUBTOTAL·덧셈뺄셈만 재검산")
        return None
    import logging
    logging.getLogger("formulas").setLevel(logging.ERROR)
    # formulas 가 지원하지 않는 수식(표 구조적 참조 등)은 저장값으로 바꾼 임시본을 만들어 나머지를 계산
    excluded, excluded_funcs, missing = [], Counter(), []
    try:
        supported = {k.upper() for k in _formulas.get_functions().keys()}
    except Exception:
        supported = None
    def unsupported_func(v):
        if supported is None:
            return None
        if re.search(r"(?:'[^']+:[^']+'|\b[^\s'!(),:=+\-*/&]+:[^\s'!(),:=+\-*/&]+)!", v or ""):
            return "3D 참조"                                     # formulas가 '1월:3월'!A1 형식을 읽지 못함 → 통합문서 전체 재계산 실패 방지
        if re.search(r'\bTEXT\s*\([^"]*"[^"]*[ymdhs]', v or "", re.I):
            return "TEXT(날짜 서식)"                             # formulas가 'm'(월)을 분으로 계산 → 거짓 불일치 (노트북 Excel로 확인)
        try:
            for t in Tokenizer(v).items:
                if t.type == "FUNC" and t.subtype == "OPEN":
                    raw_name = t.value.rstrip("(").upper()
                    name = raw_name.replace("_XLFN.", "")
                    if raw_name not in supported and name not in supported and ("_XLFN." + name) not in supported:
                        return name
                    if name.replace("_XLWS.", "") in cf_FORMULAS_UNRELIABLE:   # formulas 결과가 실제 Excel과 다른 함수 (노트북 Excel로 확인)
                        return name
        except Exception:
            return None
        return None
    spill = {}                                            # (시트, 앵커) → 배열 범위 (동적 배열의 C1# 참조 = ANCHORARRAY(C1) 풀이용)
    for ws_a in wbf.worksheets:
        for row_a in ws_a.iter_rows():
            for c_a in row_a:
                if type(c_a.value).__name__ == "ArrayFormula" and c_a.value.ref:
                    spill[(ws_a.title.casefold(), c_a.coordinate)] = c_a.value.ref
    anchor_re = re.compile(r'''("(?:[^"]|"")*"|'(?:[^']|'')*'!)|_xlfn\.ANCHORARRAY\(\s*((?:'(?:[^']|'')+'|[^'!(),\s]+)!)?\$?([A-Z]{1,3})\$?(\d+)\s*\)''', re.I)
    def resolve_anchor(text, sheet_title):
        def sub(mm):
            if mm.group(1):                            # 글자 상수·따옴표 시트명은 원문 보존
                return mm.group(0)
            sheet_part = mm.group(2) or ""
            target = sheet_part[:-1].strip("'").replace("''", "'") if sheet_part else sheet_title
            ref = spill.get((target.casefold(), f"{mm.group(3).upper()}{mm.group(4)}"))
            return (sheet_part + ref) if ref else mm.group(0)
        return anchor_re.sub(sub, text)
    empty_caches = excel_empty_string_caches(path)
    def saved_value(sheet, coord):
        value = wbv[sheet][coord].value
        return "" if value is None and (sheet, coord) in empty_caches else value
    def make_copy(bad_formulas):
        excluded.clear(); excluded_funcs.clear(); missing.clear()
        wb2 = cf_safe_load(path, None, fresh=True)        # 셀을 바꿔 저장하므로 새로 읽음
        for ws2 in wb2.worksheets:
            for row in ws2.iter_rows():
                for c in row:
                    v = c.value
                    if isinstance(v, str) and "ANCHORARRAY" in v.upper():
                        c.value = resolve_anchor(v, ws2.title)
                    elif type(v).__name__ == "ArrayFormula" and v.text and "ANCHORARRAY" in v.text.upper():
                        v.text = resolve_anchor(v.text, ws2.title)
        for ws2 in wb2.worksheets:
            for row in ws2.iter_rows():
                for c in row:
                    v = c.value
                    if type(v).__name__ == "ArrayFormula":
                        uf = unsupported_func(v.text or "") or ("XLOOKUP" if "XLOOKUP" in (v.text or "").upper() else None)
                        if uf or (v.text or "") in bad_formulas:
                            if uf:
                                excluded_funcs[uf] += 1
                            cached = saved_value(ws2.title, c.coordinate)
                            c.value = cached
                            if cached is None:
                                missing.append(f"{ws2.title}!{c.coordinate}")
                            excluded.append(f"{ws2.title}!{c.coordinate}")
                        continue
                    if not (isinstance(v, str) and v.startswith("=")):
                        continue
                    uf = unsupported_func(v)
                    if uf:
                        excluded_funcs[uf] += 1
                    if re.search(r"[\w가-힣.]+\[|\[@", v) or v in bad_formulas or uf:
                        cached = saved_value(ws2.title, c.coordinate)
                        c.value = cached
                        if cached is None:
                            missing.append(f"{ws2.title}!{c.coordinate}")
                        excluded.append(f"{ws2.title}!{c.coordinate}")
        tmp = os.path.join(tmpdir, "recalc.xlsx")
        wb2.save(tmp)
        wb2.close()
        return tmp
    bad, sol = set(), None
    for attempt in range(6):
        try:
            src = make_copy(bad)
            if missing:
                add("정보", "전체 재검산 불가", "전체", f"미지원 수식 {len(missing)}개에 저장값 없음 → 0으로 대체하지 않고 단순 재검산만 진행")
                return None
            with contextlib.redirect_stderr(io.StringIO()), contextlib.redirect_stdout(io.StringIO()):
                model = _formulas.ExcelModel().loads(src).finish()
                sol = model.calculate()
            break
        except Exception as e:
            m = re.search(r"'(=[^']*)'\)?$", str(e)) or re.search(r"(=[^\n']+)", str(e))
            if m and m.group(1) not in bad:
                bad.add(m.group(1)); excluded = []; continue              # 그 수식만 빼고 재시도
            add("주의", "전체 재검산 실패", "전체", f"formulas 계산 중 오류: {type(e).__name__}: {str(e)[:120]} → SUM·덧셈뺄셈만 재검산")
            return None
    if sol is None:
        add("주의", "전체 재검산 실패", "전체", "formulas 계산 반복 실패 → SUM·덧셈뺄셈만 재검산")
        return None
    if excluded:
        fn = ", ".join(f"{k}({n})" for k, n in excluded_funcs.most_common(5))
        add("정보", "전체 재검산 제외", "전체", f"formulas 미지원 수식 {len(excluded)}개는 전체 재검산에서 제외" + (f" — 함수: {fn}" if fn else " (표 참조 등)") + " → SUM·SUBTOTAL·덧셈뺄셈이면 단순 재검산으로 대조")
    computed = {}
    for k, v in sol.items():
        m = re.match(r"^'\[[^\]]*\](.+)'!([A-Z]{1,3}\d+)$", str(k))
        if m:
            computed[(m.group(1).replace("''", "'").casefold(), m.group(2))] = v
            continue
        m = re.match(r"^'\[[^\]]*\](.+)'!([A-Z]{1,3})(\d+):([A-Z]{1,3})(\d+)$", str(k))
        if m:                                             # 배열수식 결과 범위 → 칸별 값
            sheet_key = m.group(1).replace("''", "'").casefold()
            c1, r1 = column_index_from_string(m.group(2)), int(m.group(3))
            c2, r2 = column_index_from_string(m.group(4)), int(m.group(5))
            arr = getattr(v, "value", v)
            try:
                for ri in range(r2 - r1 + 1):
                    for ci in range(c2 - c1 + 1):
                        computed.setdefault((sheet_key, f"{get_column_letter(c1 + ci)}{r1 + ri}"), arr[ri][ci])
            except Exception:
                pass
    n_done = n_nocache = 0; handled = set(); engine_limited = set()
    for ws in wbf.worksheets:
        for row in ws.iter_rows():
            for c in row:
                f = c.value
                if type(f).__name__ == "ArrayFormula":
                    loc = f"{ws.title}!{c.coordinate}"
                    if loc in excluded:
                        continue
                    ref = f.ref or c.coordinate
                    cells = ws[ref] if ":" in ref else ((c,),)
                    results = []
                    for row_cells in cells:
                        for rc in row_cells:
                            key = (ws.title.casefold(), rc.coordinate)
                            if key not in computed:
                                results = None; break
                            results.append((rc.coordinate, cf__xl_scalar(computed[key]), saved_value(ws.title, rc.coordinate)))
                        if results is None:
                            break
                    if not results:
                        continue
                    handled.add(loc)
                    for coord_a, (val, err), cached in results:
                        cell_loc = f"{ws.title}!{coord_a}"
                        cached = cf__norm_saved(cached)
                        if isinstance(cached, str) and cached.strip().upper() in cf_ERRORS_ALL:
                            if coord_a == c.coordinate:
                                add("오류", "오류값", cell_loc, f"{cached}  배열수식 {loc}: {f.text}")
                            n_done += 1; continue
                        if err:
                            if str(cached).strip().upper() == err:     # 저장값도 같은 오류 → 재계산 결과와 같음 (오류값 자체는 '오류값'으로 따로 표시됨)
                                n_done += 1; continue
                            if cached not in (None, ""):
                                if loc not in engine_limited:            # 배열 하나에 한 줄만 (칸마다 반복하면 읽기 어려움)
                                    engine_limited.add(loc)
                                    add("주의", "재검산 불가(계산 엔진 한계)", loc, f"배열수식 {ref} — Excel 저장값은 있는데 formulas는 {err} (formulas가 계산하지 못한 수식, USE_EXCEL=True면 Excel로 검산)  수식: {f.text}")
                                continue
                            add("오류", "재계산 시 오류값", cell_loc, f"다시 계산하면 {err}  배열수식 {loc}: {f.text}"); continue
                        if cached is None:
                            n_nocache += 1; continue
                        n_done += 1
                        cv, vv = cf__to_num(cached), cf__to_num(val)
                        if cf__differs(cached, val):
                            if cv is not None and vv is not None:
                                add("오류", "전체 재검산 불일치", cell_loc, f"저장된 값 {cv:,.2f} / 다시 계산한 값 {vv:,.2f}  배열수식 {loc} ({ref}): {f.text}")
                            else:
                                add("오류", "전체 재검산 불일치", cell_loc, f"저장된 값 '{cached}' / 다시 계산한 값 '{val}'  배열수식 {loc} ({ref}): {f.text}")
                    continue
                if not (isinstance(f, str) and f.startswith("=")):
                    continue
                key = (ws.title.casefold(), c.coordinate)
                loc = f"{ws.title}!{c.coordinate}"
                if key not in computed or loc in excluded:
                    continue
                handled.add(loc)
                if cf__multi(computed[key]) or cf__implicit_range(f):                   # 예: =B1:B10*2 (옛 방식 수식) → Excel은 같은 행 값만 씀. formulas는 배열 → 비교하면 거짓 불일치
                    add("주의", "재검산 불가(암시적 교차)", loc, f"범위를 한 칸처럼 쓰는 옛 방식 수식이라 formulas로 재계산하지 않음 (USE_EXCEL=True면 Excel로 검산)  수식: {f}")
                    continue
                val, err = cf__xl_scalar(computed[key])
                cached = cf__norm_saved(saved_value(ws.title, c.coordinate))
                if err:
                    if cached not in (None, "") and not (isinstance(cached, str) and cached.strip().upper() in cf_ERRORS_ALL):
                        add("주의", "재검산 불가(계산 엔진 한계)", loc, f"Excel 저장값 {cached} 인데 formulas는 {err} — formulas가 계산하지 못한 수식 (USE_EXCEL=True면 Excel로 검산)  수식: {f}")
                    elif str(cached).strip().upper() != err:            # 저장값도 같은 오류면 재계산과 같음
                        add("오류", "재계산 시 오류값", loc, f"다시 계산하면 {err}" + (f" (저장된 값 {cached})" if cached not in (None, "") and not str(cached).startswith("#") else "") + f"  수식: {f}")
                    n_done += 1; continue
                if cached is None:
                    n_nocache += 1; continue
                n_done += 1
                cv, vv = cf__to_num(cached), cf__to_num(val)
                if cf__differs(cached, val):
                    if cv is not None and vv is not None:
                        add("오류", "전체 재검산 불일치", loc, f"저장된 값 {cv:,.2f} / 다시 계산한 값 {vv:,.2f}  수식: {f}")
                    else:
                        add("오류", "전체 재검산 불일치", loc, f"저장된 값 '{cached}' / 다시 계산한 값 '{val}'  수식: {f}")
    if n_nocache:
        add("정보", "재검산 불가", "전체", f"수식 {n_nocache}개에 저장된 계산값 없음 (엑셀에서 열어 저장하면 대조됨)")
    add("정보", "전체 재검산", "전체", f"{n_done}개 수식을 formulas 로 다시 계산해 대조 (미지원 수식은 제외)")
    return handled


def cf__text_numbers(wbf, add, group_over=20):
    """숫자가 글자로 저장된 칸. 텍스트 서식(@)으로 일부러 넣은 칸은 제외, 한 열에 많으면 한 줄로 묶음"""
    for ws in wbf.worksheets:
        by_col = {}
        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                if isinstance(v, str) and not v.startswith("=") and cf__to_num(v) is not None and cell.number_format != "@":
                    by_col.setdefault(cell.column_letter, []).append((cell.coordinate, v))
        for col, items in by_col.items():
            if len(items) > group_over:
                add("주의", "텍스트 숫자", f"{ws.title}!{items[0][0]}:{items[-1][0]}",
                    f"{col}열 {len(items)}칸이 글자로 저장됨 (예: '{items[0][1]}'). 코드·번호 열이면 정상, 금액 열이면 합계에서 빠짐")
            else:
                for coord, v in items:
                    add("주의", "텍스트 숫자", f"{ws.title}!{coord}", f"'{v}' 가 글자로 저장됨 (합계에서 빠질 수 있음)")


def cf_check(path, recalculate=True):
    """[{"등급","항목","위치","상세"}, ...]  등급: 오류/주의/정보"""
    cfw = []
    wbf = cf_safe_load(path, cfw)                      # 수식 원문
    wbv = cf_safe_load(path, None, data_only=True)     # 캐시된 계산값
    sheets = set(wbf.sheetnames)
    tables = cf__tables(wbf)
    names = cf__defined_names(wbf)
    res = []
    add = lambda lv, item, loc, detail="": res.append({"등급": lv, "항목": item, "위치": loc, "상세": detail})
    n_formula = 0

    nonstd = []                                          # 엑셀 오류코드가 아닌 '오류' 저장값 (파일 만든 프로그램이 계산 못함)
    array_cells = []                                     # 배열수식: 재계산에서 검산하지 못한 것만 '검사 불가'로 표시
    for ws in wbf.worksheets:
        wv = wbv[ws.title]
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                loc = f"{ws.title}!{c.coordinate}"
                if c.data_type == 'e' and str(v).strip() in cf_ERRORS_ALL:
                    add("오류", "오류값", loc, v)
                if v is not None and not isinstance(v, (str, int, float, bool, bytes)) and not hasattr(v, "year") and (hasattr(v, "text") or hasattr(v, "ref") or "Formula" in type(v).__name__):
                    n_formula += 1                                  # 배열수식·데이터 표 수식 — openpyxl이 글자가 아닌 객체로 읽음
                    kind_name = "데이터 표 수식" if "DataTable" in type(v).__name__ else "배열수식"
                    detail = (f"  수식: {v.text}" if getattr(v, "text", None) else "") + (f"  범위: {v.ref}" if getattr(v, "ref", None) else "")
                    if kind_name == "배열수식":
                        array_cells.append((loc, detail))
                    else:
                        add("주의", f"{kind_name} 검사 불가", loc, f"{kind_name}은 참조·재계산 검사를 하지 않음" + detail)
                    continue
                if not (isinstance(v, str) and v.startswith("=")):
                    continue
                n_formula += 1
                cached = wv[c.coordinate].value
                if wv[c.coordinate].data_type == 'e':
                    if str(cached).strip() in cf_ERRORS_ALL:
                        add("오류", "오류값", loc, f"{cached}  수식: {v}")
                    else:
                        nonstd.append((loc, str(cached)[:60]))
                code = cf__without_text_literals(v)
                if "#REF!" in code:
                    add("오류", "깨진 참조(#REF!)", loc, v)
                if cf_EXT.search(code):
                    add("주의", "외부 파일 링크", loc, v)
                    continue
                for tname, cols in cf__table_refs(v):
                    if tname not in tables:
                        add("오류", "없는 표 참조", loc, f"'{tname}'  수식: {v}")
                        continue
                    for col in cols:
                        if col not in tables[tname]["cols"]:
                            add("오류", "표에 없는 열", loc, f"{tname}[{col}]  수식: {v}")
                for sheet, c1, r1, c2, r2 in cf__refs(v):
                    if sheet not in sheets and ":" in sheet and all(part in sheets for part in sheet.split(":", 1)):
                        continue                                  # 3D 참조 '1월:3월'!A1 — 양 끝 시트가 있으면 정상
                    if sheet not in sheets:
                        add("오류", "없는 시트 참조", loc, f"'{sheet}'  수식: {v}")
                        continue
                    tgt = wbv[sheet]
                    cols = range(column_index_from_string(c1), column_index_from_string(c2) + 1)
                    r2c = min(r2, tgt.max_row)
                    has_val = any(tgt.cell(r, cc).value not in (None, "") for cc in cols for r in range(r1, r2c + 1))
                    if not has_val:
                        add("주의", "빈 범위 참조", loc, f"{sheet}!{c1}{r1}:{c2}{r2}  수식: {v}")
                        continue
                    if r2 > r1:                                   # 범위 참조 → 데이터 끝과 비교
                        last = max(cf__last_data_row(tgt, cc_letter) for cc_letter in {c1, c2})
                        if last > r2:
                            add("주의", "참조 범위 부족", loc, f"범위 {r2}행까지 / 데이터 {last}행까지  수식: {v}")
    # 표 자체 점검: 표 바로 아래에 데이터가 이어져 있으면 표가 안 늘어난 것
    for tname, t in tables.items():
        ws = wbv[t["sheet"]]
        below = [r for r in range(t["last_row"] + 1, ws.max_row + 1)
                 if any(ws.cell(r, c).value not in (None, "") for c in t["col_idx"])]
        if below:
            add("주의", "표 범위 부족", f"{t['sheet']}!{t['ref']}",
                f"표 '{tname}'은 {t['last_row']}행까지 / 데이터 {max(below)}행까지 → 표 범위 확장 필요")
    # 이름정의 점검
    for nm, txt in names.items():
        for sheet, c1, r1, c2, r2 in cf__refs(txt):
            if sheet not in sheets:
                add("오류", "이름정의 시트 없음", nm, txt)
            elif r2 > r1:
                last = max(cf__last_data_row(wbv[sheet], cc) for cc in {c1, c2})
                if last > r2:
                    add("주의", "이름정의 범위 부족", nm, f"'{nm}' = {txt} / 데이터 {last}행까지")
    for w in cfw:
        add("주의", "서식 깨짐(무시)" if "서식 정보" in w else "조건부서식 무시", "전체", w)
    try:
        locs, total = cf_broken_styles(path)
    except Exception:
        locs, total = [], 0
    for loc, detail in locs:
        add("주의", "서식 깨짐", loc, detail)
    if total:
        add("주의", "서식 깨짐 합계", "전체", f"{total}곳 (위 목록은 최대 {len(locs)}곳) → 엑셀에서 열어 '다른 이름으로 저장'하면 보통 정리됨")
    if nonstd:
        add("주의", "계산 안 된 저장값", "전체",
            f"{len(nonstd)}개 수식의 저장값이 엑셀 오류코드가 아님 — 파일을 만든 프로그램이 계산하지 못한 것 (예: {nonstd[0][0]} '{nonstd[0][1]}'). 엑셀에서 열어 저장 후 다시 검사하세요")
    # v6 추가 검사 (구조 무관: 셀·수식 전체를 훑음)
    _progress(f"  수식 {n_formula}개 확인. 텍스트 숫자·수식 자리의 값 검사 중")
    cf__text_numbers(wbf, add)
    for ws in wbf.worksheets:
        cf__pattern_checks(ws, add)
        cf__omitted_cells(ws, add)
    if recalculate and n_formula:                        # 수식이 없으면 재계산 단계 없음 (v13과 같은 결과)
        _t = time.monotonic()
        _progress(f"  formulas 재계산 중 (수식 {n_formula}개) — 이 단계는 진행률 표시 불가, 수식이 많으면 수 분 걸릴 수 있음")
        flagged = cf__recalc_full(path, wbf, wbv, add)
        cf__recalc_simple(wbf, wbv, add, skip=flagged)
        _progress(f"  재계산 끝 ({time.monotonic() - _t:.1f}초)")
    else:
        flagged = None
    for loc, detail in array_cells:
        if not flagged or loc not in flagged:
            add("주의", "배열수식 검사 불가", loc, "배열수식을 다시 계산하지 못함 (formulas 미설치·미지원 함수·저장값 없음 등)" + detail)
    try:
        with zipfile.ZipFile(path) as _z:
            _wbxml = _z.read("xl/workbook.xml").decode("utf-8", "replace")
        if re.search(r'<calcPr[^>]*calcMode="manual"', _wbxml):
            add("주의", "수동 계산 모드로 저장됨", "전체", "이 파일은 Excel '수동 계산'으로 저장됨 → 저장된 값이 최신이 아닐 수 있음. 재검산 불일치가 나오면 F9로 다시 계산한 뒤 저장해 보세요")
    except Exception:
        pass
    add("정보", "수식 개수", "전체", str(n_formula))
    add("정보", "표 개수", "전체", str(len(tables)))
    add("정보", "이름정의 개수", "전체", str(len(names)))
    return res


def cf_sheet_links(path):
    """{참조하는 시트: {참조되는 시트: 수식 수}}  (외부 링크 제외, 자기 참조 제외)"""
    wb = cf_safe_load(path)
    tables = cf__tables(wb)
    m = defaultdict(lambda: defaultdict(int))
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, str) and v.startswith("=") and not cf_EXT.search(v):
                    for sheet, *_ in cf__refs(v):
                        if sheet != ws.title:
                            m[ws.title][sheet] += 1
                    for tname, _ in cf__table_refs(v):
                        if tname in tables and tables[tname]["sheet"] != ws.title:
                            m[ws.title][tables[tname]["sheet"]] += 1
    return {k: dict(v) for k, v in m.items()}


# ---- v4: 결과 파일명 (실행 시각) ----
_RUN_STAMP = None


def run_stamp():
    return datetime.datetime.now().strftime("%Y%m%d_%H%M%S")


def out_path(base, tag, ext=".xlsx"):
    """<base><tag>_<시각>.xlsx. 한 번 실행 안에서는 같은 시각, 같은 이름이 있으면 _2, _3 …"""
    global _RUN_STAMP
    if _RUN_STAMP is None:
        _RUN_STAMP = run_stamp()
    p = f"{base}{tag}_{_RUN_STAMP}{ext}"
    n = 2
    while os.path.exists(p):
        p = f"{base}{tag}_{_RUN_STAMP}_{n}{ext}"; n += 1
    return p


def cf_run(path, result=None):
    res = cf_check(path) if result is None else result
    links = cf_sheet_links(path)
    base, _ = os.path.splitext(path)
    out = out_path(base, "_수식검증")

    wb = Workbook()
    ws = wb.active; ws.title = "수식검증"
    red, yel, bold = PatternFill("solid", fgColor="FFC7CE"), PatternFill("solid", fgColor="FFEB9C"), Font(bold=True)
    n_err = sum(r["등급"] == "오류" for r in res); n_warn = sum(r["등급"] == "주의" for r in res)
    lim = [r for r in res if "불가" in r["항목"] or "제외" in r["항목"] or "실패" in r["항목"]]
    lim_cells = {r["위치"] for r in lim if r["위치"] != "전체"}
    lim_whole = 0
    whole_fail = False
    for r in lim:
        if r["위치"] == "전체":
            m_ = re.search(r"(\d+)개", str(r.get("상세", "")))
            if m_:
                lim_whole += int(m_.group(1))
            else:
                whole_fail = True
    n_limited = len(lim_cells) + lim_whole
    status = "검사 범위 내 오류·주의 없음" if n_err + n_warn == 0 else f"오류 {n_err} / 주의 {n_warn}"
    if n_limited or whole_fail:
        status += f" (재검산 불가 {n_limited}칸" + (" · 전체 재검산 실패" if whole_fail else "") + " — 검사 안 된 범위)"
    ws.append(["최종", status])
    ws["A1"].font = bold; ws["B1"].font = bold
    if n_err: ws["B1"].fill = red
    elif n_warn: ws["B1"].fill = yel
    ws.append(["파일", os.path.basename(path)])
    ws.append(["실행일시", datetime.datetime.now().strftime("%Y-%m-%d %H:%M")])
    ws.append([])
    ws.append(["등급", "항목", "위치", "상세"])
    for c in ws[ws.max_row]: c.font = bold
    row_no = ws.max_row
    for r in sorted(res, key=lambda r: {"오류": 0, "주의": 1, "정보": 2}[r["등급"]]):
        ws.append([r["등급"], r["항목"], r["위치"], r["상세"]]); row_no += 1
        if r["등급"] == "오류": ws.cell(row_no, 1).fill = red
        elif r["등급"] == "주의": ws.cell(row_no, 1).fill = yel
    head_row = 5
    _style_table(ws, head_row, wrap=("상세",), max_width=90, judge=("등급",))
    _style_labels(ws, 1, 3)
    ws.column_dimensions["A"].width = max(ws.column_dimensions["A"].width or 0, 10)
    ws["B1"].font = Font(bold=True, size=12)

    ls = wb.create_sheet("시트연결")
    ls.append(["참조하는 시트", "→ 참조되는 시트", "수식 수"])
    for c in ls[1]: c.font = bold
    src_wb = cf_safe_load(path)
    for s in src_wb.sheetnames:
        if s in links:
            for tgt, n in sorted(links[s].items()):
                ls.append([s, tgt, n])
        else:
            ls.append([s, "(다른 시트 참조 없음)", 0])
    _style_table(ls, 1)
    ts = wb.create_sheet("표·이름정의")
    ts.append(["구분", "이름", "위치", "열 이름"])
    for c in ts[1]: c.font = bold
    src = cf_safe_load(path)
    for tname, t in cf__tables(src).items():
        ts.append(["표", tname, f"{t['sheet']}!{t['ref']}", ", ".join(str(c) for c in t["cols"])])
    for nm, txt in cf__defined_names(src).items():
        ts.append(["이름정의", nm, txt, ""])
    _style_table(ts, 1, wrap=("열 이름", "위치"), max_width=70)
    _diagnostic_text(wb)                 # 진단표의 수식 원문은 결과 파일에서 실행되지 않도록 텍스트로 저장
    wb.save(out)
    return out

# 원본·보고서 대조

RESULT_TAGS = ("_대사결과", "_수식검증", "_PPT대조", "_결과", "_구조비교", "_원본대사", "_채움", "규칙표_")
YELLOW = PatternFill("solid", fgColor="FFFF00")


# ---------------------------------------------------------------- 값·글자
def to_num(v):
    """숫자 / '1,000' / '(1,000)'=음수 → float.  '-'·빈칸·글자·nan → None"""
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v) if math.isfinite(v) else None
    if isinstance(v, str):
        return _text_to_num(v)
    return None


_NUM_START = set("0123456789+-.(△▲₩￦$\u200b\u200c\u200d\u2060\ufeff \t\u00a0\u3000０１２３４５６７８９") | set("KkUu")


@functools.lru_cache(maxsize=200000)
def _text_to_num(v):
    """글자 → 숫자 (같은 글자가 많아 결과를 기억). 숫자로 시작할 수 없는 글자(적요 등)는 바로 None"""
    t0 = v.lstrip()
    if not t0 or t0[0] not in _NUM_START:
        return None
    t = re.sub(r"[\u200b\u200c\u200d\u2060\ufeff]", "", v).replace(",", "").strip()   # 보이지 않는 글자 제거
    if t in ("", "-", "—", "–"):
        return None
    if t[0] in "+-₩￦$KkUu" or t.endswith(("원", "KRW", "krw")):
        t = re.sub(r"^([+\-]?)\s*(?:[₩￦$]|KRW|USD)\s*", r"\1", t, flags=re.I)          # 통화 기호 ('-$18.85' 포함)
        t = re.sub(r"\s*(?:원|KRW)$", "", t, flags=re.I).strip()                          # '원' 단위
    if t.endswith("-") and re.fullmatch(r"\(?[\d.]+\)?-", t):                            # 회계 표기 뒤 마이너스 '1,000-'
        t = "-" + t[:-1]
    neg = t.startswith("(") and t.endswith(")")
    t = t.strip("()").replace("△", "-").replace("▲", "-").strip()
    try:
        f = float(t)
    except ValueError:
        return None
    if not math.isfinite(f):                               # 'nan'·'inf' 글자는 숫자 아님 (합계 오염 방지)
        return None
    return -f if neg else f


PLACEHOLDERS = {"", "nan", "none", "null", "n/a", "na", "-", "—", "–"}


def is_placeholder(v):
    """빈 값 표시용 글자 (pandas 내보내기의 'nan' 등) — 열 종류 판단에서 제외"""
    return isinstance(v, str) and v.strip().lower() in PLACEHOLDERS


def is_dash(v):
    return isinstance(v, str) and v.strip() in ("-", "—", "–")


def is_label(v):
    return (isinstance(v, str) and v.strip() != "" and to_num(v) is None and not is_dash(v)) or hasattr(v, "year")


def nl(v):
    """정규화: 공백 제거·소문자"""
    return _nl_str(str(v))


@functools.lru_cache(maxsize=500000)
def _nl_str(v):
    return re.sub(r"\s+", "", v).lower()


@functools.lru_cache(maxsize=65536)
def nl2(v):
    """유사 매칭용: 괄호·대괄호 안 제거, 앞의 코드(숫자/영문) 제거"""
    s = re.sub(r"\([^)]*\)|\[[^\]]*\]", "", str(v))
    s = nl(s)
    s = re.sub(r"^[0-9a-z\-\._/]+(?=[가-힣])", "", s)
    return s or nl(v)


def rule_key(filename):
    """보고서 파일명 → 규칙표 이름 (날짜·연월 등 숫자 묶음 제거). 보고서_법인별_202610.xlsx → 보고서_법인별"""
    base = os.path.splitext(os.path.basename(filename))[0]
    base = re.sub(r"\d{4}[-._/]?\d{0,2}[-._/]?\d{0,2}", "", base)
    base = re.sub(r"\d{6,}", "", base)
    base = re.sub(r"^[\s_\-.]+|[\s_\-.]+$", "", base)
    base = re.sub(r"[\s_\-.]{2,}", "_", base)
    return base or "보고서"


# ---------------------------------------------------------------- 파일 열기
def safe_load(path, warnings=None, **kw):
    try:
        return load_workbook(path, **kw)
    except Exception as original_error:
        tmp = None
        try:
            tmp = cf__clean_copy(path, True)
            # 메모리에서 열면 read_only 모드에서도 임시 파일을 즉시 삭제할 수 있다.
            with open(tmp, "rb") as source:
                wb = load_workbook(io.BytesIO(source.read()), **kw)
        except Exception:
            raise original_error
        finally:
            if tmp is not None and os.path.exists(tmp):
                os.unlink(tmp)
        if warnings is not None:
            warnings.append(f"서식 깨짐 → 서식 제외하고 읽음: {os.path.basename(path)}")
        return wb


# ---------------------------------------------------------------- 결과 파일 서식 (값은 그대로, 보기만)
_HEAD_FILL = PatternFill("solid", fgColor="1F4E78")
_HEAD_FONT = Font(bold=True, color="FFFFFF")
_LABEL_FILL = PatternFill("solid", fgColor="F2F2F2")
_THIN = Side(style="thin", color="BFBFBF")
_BOX = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_JUDGE_FILL = {"불일치": "FFC7CE", "오류": "FFC7CE", "누락": "FFC7CE", "계산값 없음": "FFEB9C", "근거 없음": "FFEB9C", "주의": "FFEB9C", "유사 확인": "FCE4D6", "제외": "E7E6E6", "일치": "E2EFDA"}


def _text_width(v):
    """화면 폭 추정: 한글·전각은 2칸"""
    t = str(v) if v is not None else ""
    return sum(2 if ord(ch) > 0x2E80 else 1 for ch in t)


def _style_table(ws, header_row=1, wrap=(), max_width=60, judge=()):
    """header_row 줄을 머리글로 보고 아래 표에 서식: 머리글 색·굵게, 테두리, 숫자 천단위(음수 빨강), 긴 글자 줄바꿈, 열 너비, 틀 고정, 필터, 인쇄 설정.
    큰 표도 빠르도록 서식 객체는 미리 만들어 같이 씀"""
    last_row, last_col = ws.max_row, ws.max_column
    if last_col == 0:
        return
    head_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    top, top_wrap, center = Alignment(vertical="top"), Alignment(vertical="top", wrap_text=True), Alignment(horizontal="center", vertical="top")
    judge_fills = {k: PatternFill("solid", fgColor=v) for k, v in _JUDGE_FILL.items()}
    heads = [ws.cell(header_row, c).value for c in range(1, last_col + 1)]
    for c in range(1, last_col + 1):
        h = ws.cell(header_row, c)
        if h.value in (None, ""):
            continue
        h.fill = _HEAD_FILL; h.font = _HEAD_FONT; h.border = _BOX; h.alignment = head_align
    ws.row_dimensions[header_row].height = 20
    big = last_row - header_row > 2000                         # 큰 표는 테두리·정렬 생략, 숫자 서식·판정 색만 (속도)
    cols = list(ws.iter_cols(min_row=header_row + 1, max_row=last_row, min_col=1, max_col=last_col))
    for c, cells in enumerate(cols, start=1):
        name = heads[c - 1]
        if name in (None, ""):
            continue
        long_text = name in wrap
        decimal = False; width = _text_width(name)
        for i, cell in enumerate(cells):
            v = cell.value
            if isinstance(v, float) and not isinstance(v, bool) and not v.is_integer():
                decimal = True
            if i < 3000:
                width = max(width, len(f"{v:,.2f}") if isinstance(v, (int, float)) and not isinstance(v, bool) else _text_width(v))
        fmt = "#,##0.00;[Red]-#,##0.00" if decimal else "#,##0;[Red]-#,##0"
        floor = {"행 제목": 24, "열 제목": 24, "이유": 40, "상세": 40, "확인 방법": 40, "위치": 14}.get(name, 8)   # 짧은 내용이어도 읽기 편한 최소 너비
        ws.column_dimensions[get_column_letter(c)].width = max(floor, min(max_width if long_text else 40, width + 2))
        is_judge = name in judge
        for cell in cells:
            if not big:
                cell.border = _BOX
            if cell.data_type == "e":                        # '#DIV/0!' 같은 설명 글자가 오류 칸으로 저장되지 않게
                cell.data_type = "s"
            v = cell.value
            if isinstance(v, bool):
                if not big:
                    cell.alignment = center
            elif isinstance(v, (int, float)):
                cell.number_format = fmt
                if not big:
                    cell.alignment = top
            else:
                if long_text and not big:
                    cell.alignment = top_wrap
                elif not big:
                    cell.alignment = top
                if is_judge and v in judge_fills and cell.fill.fill_type is None and v != "일치":
                    cell.fill = judge_fills[v]
    ws.freeze_panes = ws.cell(header_row + 1, 1).coordinate
    if last_row > header_row:
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(last_col)}{last_row}"
    ws.page_setup.orientation = "landscape"; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{header_row}:{header_row}"


def _style_labels(ws, first_row, last_row):
    """'최종·파일·실행일시' 같은 항목·값 줄: 항목 열 굵게·회색"""
    for r in range(first_row, last_row + 1):
        a = ws.cell(r, 1)
        if a.value not in (None, ""):
            a.font = Font(bold=True); a.fill = _LABEL_FILL; a.border = _BOX


def _diagnostic_text(wb):
    """진단 결과와 규칙표에는 수식이 아닌 원문 문자열을 저장한다."""
    for ws in wb.worksheets:
        for row in ws:
            for cell in row:
                if cell.data_type == "f":
                    cell.data_type = "s"


def fill_merged(ws):
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    for rng in ws.merged_cells.ranges:
        try:
            v = rows[rng.min_row - 1][rng.min_col - 1]
        except IndexError:
            continue
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                if r - 1 < len(rows) and c - 1 < len(rows[r - 1]) and rows[r - 1][c - 1] in (None, ""):
                    rows[r - 1][c - 1] = v
    return rows


# ---------------------------------------------------------------- 규칙표
SETTING_KEYS = [("원본 시트", "비워두면 숫자가 가장 많은 시트"), ("보고서 시트", "비워두면 전부"), ("행 글자 열", "보고서 왼쪽 글자가 들어 있는 원본 열 (여러 개면 쉼표)"),
                ("열 글자 열", "보고서 위쪽 글자가 들어 있는 원본 열 (여러 개면 쉼표)"), ("금액 열", "원본 금액 열. 차변-대변 처럼 식도 가능"),
                ("집계", "합계 또는 건수"), ("단위 배수", "보고서가 천원이면 1000, 백만원이면 1000000. 비워두면 자동"),
                ("제외 글자", "집계 대상이 아닌 줄 (쉼표 구분)"), ("허용 오차", "이 차이 이하는 일치")]


def rule_path(report_path):
    return os.path.join(os.path.dirname(os.path.abspath(report_path)), f"규칙표_{rule_key(report_path)}.xlsx")


def load_rules(report_path):
    """{"settings": {항목: 값}, "map": {보고서글자(정규화): 원본글자}, "path": ...}  (파일 없으면 빈 규칙)"""
    p = rule_path(report_path)
    rules = {"settings": {}, "map": {}, "pending_fuzzy": {}, "path": p, "exists": os.path.exists(p)}
    if not rules["exists"]:
        return rules
    wb = load_workbook(p)
    if "설정" in wb.sheetnames:
        for r in wb["설정"].iter_rows(values_only=True):
            if r and r[0] and len(r) > 1 and r[1] not in (None, ""):
                rules["settings"][str(r[0]).strip()] = str(r[1]).strip()
    if "글자대응" in wb.sheetnames:
        for i, r in enumerate(wb["글자대응"].iter_rows(values_only=True)):
            if i == 0 or not r or r[0] in (None, ""):
                continue
            if len(r) > 1 and r[1] not in (None, ""):
                rules["map"][nl(r[0])] = str(r[1]).strip()
                if len(r) > 2 and str(r[2] or "").startswith("자동(유사)"):
                    rules["pending_fuzzy"][str(r[0]).strip()] = str(r[1]).strip()
    return rules


def save_rules(report_path, detected, new_labels, fuzzy):
    """규칙표 생성(없으면) 또는 보강(있으면: 설정은 비어 있는 값만, 글자대응은 새 글자만 추가)"""
    p = rule_path(report_path)
    bold = Font(bold=True)
    if os.path.exists(p):
        wb = load_workbook(p)
        ws = wb["설정"] if "설정" in wb.sheetnames else wb.create_sheet("설정")
        ms = wb["글자대응"] if "글자대응" in wb.sheetnames else wb.create_sheet("글자대응")
    else:
        wb = Workbook(); ws = wb.active; ws.title = "설정"; ms = wb.create_sheet("글자대응")
        ws.append(["항목", "값", "설명"]); ms.append(["보고서 글자", "원본 글자", "비고"])
        for c in list(ws[1]) + list(ms[1]):
            c.font = bold
        for k, desc in SETTING_KEYS:
            ws.append([k, "", desc])
        ws.column_dimensions["A"].width = 14; ws.column_dimensions["B"].width = 24; ws.column_dimensions["C"].width = 60
        ms.column_dimensions["A"].width = 32; ms.column_dimensions["B"].width = 24; ms.column_dimensions["C"].width = 30
        for sh_ in (ws, ms):
            for c_ in sh_[1]:
                if c_.value not in (None, ""):
                    c_.fill = _HEAD_FILL; c_.font = _HEAD_FONT; c_.alignment = Alignment(horizontal="center")
            sh_.freeze_panes = "A2"
    existing = {}
    for r in ws.iter_rows(min_row=2):
        if r[0].value:
            existing[str(r[0].value).strip()] = r
    for k, v in detected.items():
        if k in existing and existing[k][1].value in (None, ""):
            existing[k][1].value = v
    have = {nl(r[0].value) for r in ms.iter_rows(min_row=2) if r[0].value}
    for lab in new_labels:
        if nl(lab) not in have:
            ms.append([lab, None, "원본 글자를 적으세요 (원본에 없으면 비워 둠 → 근거 없음)"]); have.add(nl(lab))
    for lab, raw_val in fuzzy.items():
        if nl(lab) not in have:
            ms.append([lab, raw_val, "자동(유사) — 확인 후 비고를 확인완료로 변경"]); have.add(nl(lab))
    _diagnostic_text(wb)
    wb.save(p)
    return p


# ---------------------------------------------------------------- 원본 읽기
def read_raw(path, sheet=None, warnings=None):
    wb = safe_load(path, warnings, data_only=True)
    if sheet:
        if sheet not in wb.sheetnames:
            raise RuntimeError(f"원본에 '{sheet}' 시트 없음. 있는 시트: {wb.sheetnames}")
        ws = wb[sheet]
    else:
        ws = max(wb.worksheets, key=lambda w: sum(1 for row in w.iter_rows(values_only=True) for v in row if to_num(v) is not None and not isinstance(v, str)))
    grid = [(i, r) for i, r in enumerate(fill_merged(ws), start=1) if any(v not in (None, "") for v in r)]
    h = None
    for i, (_, r) in enumerate(grid[:50]):
        strs = {nl(v) for v in r if is_label(v)}
        vals = [v for v in r if v not in (None, "")]
        below = [rr for _, rr in grid[i + 1:i + 6]]
        if len(strs) >= 2 and len(strs) >= len(vals) * 0.6 and any(to_num(v) is not None for rr in below for v in rr):
            h = i; break
    if h is None:
        # 제목처럼 보이는 줄은 있는데 아래가 전부 글자면 → '숫자 열 없음'으로 정확히 안내
        head_like = next((r for _, r in grid[:50] if len({nl(v) for v in r if is_label(v)}) >= 2), None)
        if head_like is not None and len(grid) > 1:
            cand = [str(v) for _, r in grid[1:6] for v in r if isinstance(v, str) and re.search(r"\d", v)]
            sample = next((c for c in cand if not re.fullmatch(r"\d{4}[-./]\d{1,2}[-./]\d{1,2}.*", c.strip())), cand[0] if cand else "")
            raise RuntimeError(f"원본 '{ws.title}' 시트에 더할 숫자 열이 없음 — 값이 글자로 저장됨" + (f" (예: '{sample}')" if sample else "")
                               + ". 금액이 숫자로 된 원본이 필요합니다")
        raise RuntimeError(f"원본 '{ws.title}' 시트에서 열 제목 줄을 찾지 못함 (글자 2개 이상 + 아래에 숫자)")
    # 두 줄 제목: 첫 숫자 줄 바로 윗줄이 글자만 있고 제목 칸 수가 더 많거나 같으면 그 줄이 진짜 제목
    hl = len({nl(v) for v in grid[h][1] if is_label(v)})
    k = next((j for j in range(h + 1, len(grid)) if any(to_num(v) is not None and not isinstance(v, str) for v in grid[j][1])), None)
    if k is not None and k - 1 > h:
        cand = grid[k - 1][1]
        cand_labels = {nl(v) for v in cand if is_label(v)}
        has_error = any(isinstance(v, str) and v.strip().upper() in cf_ERRORS_ALL for v in cand)
        if not has_error and not any(to_num(v) is not None for v in cand) and len(cand_labels) > hl:
            h = k - 1
    # 빈 줄 뒤에 원본 제목을 다시 쓰는 줄이 나오면 그 아래는 다른 표 (예산표·메모표 등) → 원본 자료에서 제외
    head_names = {nl(v) for v in grid[h][1] if is_label(v)}
    for j in range(h + 1, len(grid)):
        if grid[j][0] - grid[j - 1][0] > 1:
            labs = {nl(v) for v in grid[j][1] if is_label(v)}
            if len(labs) >= 2 and labs & head_names and not any(to_num(v) is not None and not isinstance(v, str) for v in grid[j][1]):
                if warnings is not None:
                    warnings.append(f"원본 '{ws.title}' {grid[j][0]}행부터는 제목 줄이 다시 나와 다른 표로 보고 제외")
                grid = grid[:j]
                break
    head = []
    data_cols = {j for _, r in grid[h + 1:] for j, v in enumerate(r) if v not in (None, "") and not (isinstance(v, str) and not v.strip())}
    # 제목도 자료도 없는 빈 열 뒤에 다시 제목이 나오면 그 오른쪽은 다른 표 (예산표 등) → 원본 자료에서 제외
    hdr = grid[h][1]
    first_label = next((j for j, v in enumerate(hdr) if v not in (None, "")), 0)
    gap = next((j for j in range(first_label, len(hdr)) if hdr[j] in (None, "") and j not in data_cols), None)
    if gap is not None and any(v not in (None, "") for v in hdr[gap + 1:]):
        if warnings is not None:
            warnings.append(f"원본 '{ws.title}' {get_column_letter(gap + 2)}열부터는 빈 열 뒤의 다른 표로 보고 제외")
        grid = [(i, r[:gap]) for i, r in grid]
        data_cols = {j for j in data_cols if j < gap}
    for j, v in enumerate(grid[h][1]):
        name = str(v).strip() if v not in (None, "") else ""
        letter = get_column_letter(j + 1)
        if not name and j in data_cols:
            name = f"({letter}열)"                          # 제목이 비었거나 공백뿐인 열도 데이터가 있으면 사용
        if name and name in head:
            name = f"{name}({letter}열)"                    # 같은 제목 중복 → 몰래 덮어쓰지 않고 구분
        head.append(name)
    rows = [{head[j]: (r[j] if j < len(r) else None) for j in range(len(head)) if head[j]} for _, r in grid[h + 1:]]
    row_nums = [i for i, _ in grid[h + 1:]]
    col_letter = {head[j]: get_column_letter(j + 1) for j in range(len(head)) if head[j]}
    text_cols, num_cols = [], []
    for c in [x for x in head if x]:
        vals = [row[c] for row in rows if row.get(c) not in (None, "") and not is_placeholder(row.get(c))]
        if not vals:
            continue
        if sum(1 for v in vals if to_num(v) is not None) >= len(vals) * 0.8:
            num_cols.append(c)
        elif sum(1 for v in vals if is_label(v)) >= len(vals) * 0.8:
            text_cols.append(c)
    # 금액으로 쓰일 수 있는 모든 열의 비숫자 값(오류값·글자) 위치: 집계할 때 그 열이 쓰이면 '검산 불가'로 막음
    bad = {}
    for i, row in enumerate(rows):
        for c in [x for x in head if x]:
            v = row.get(c)
            if v in (None, "") or (isinstance(v, str) and not v.strip()) or is_placeholder(v) or to_num(v) is not None:
                continue                                   # 빈칸·공백 글자·'-'·'nan'은 빈 값 (0)
            bad.setdefault(i, {})[c] = (f"{ws.title}!{col_letter[c]}{row_nums[i]}", str(v).strip())
    # 계산값이 저장되지 않은 수식(엑셀로 저장한 적 없는 파일 등)은 빈칸처럼 보이지만 0이 아님 → 수식 원문을 읽어 표시.
    # 정상적으로 저장된 빈 글자 결과(="" 등)는 제외
    try:
        empty_ok = {coord for sheet_name, coord in excel_empty_string_caches(path) if sheet_name == ws.title}
        wb_f = load_workbook(path, read_only=True)
        ws_f = wb_f[ws.title]
        col_idx = {hname: j for j, hname in enumerate(head) if hname}
        row_of = {rn: i for i, rn in enumerate(row_nums)}
        for r_i, r in enumerate(ws_f.iter_rows(values_only=False), start=1):
            i = row_of.get(r_i)
            if i is None:
                continue
            for c, j in col_idx.items():
                if j < len(r) and r[j].data_type == "f" and rows[i].get(c) in (None, "") and f"{col_letter[c]}{r_i}" not in empty_ok:
                    bad.setdefault(i, {}).setdefault(c, (f"{ws.title}!{col_letter[c]}{r_i}", "계산값 없는 수식 (엑셀에서 열어 저장 필요)"))
        wb_f.close()
    except Exception:
        pass
    index = {c: defaultdict(set) for c in text_cols}
    index2 = {c: defaultdict(set) for c in text_cols}          # 유사 매칭용 (코드·괄호 제거)
    raw_label = {c: {} for c in text_cols}
    for i, row in enumerate(rows):
        for c in text_cols:
            v = row.get(c)
            if v not in (None, ""):
                k = nl(v); index[c][k].add(i); raw_label[c].setdefault(k, str(v).strip())
                index2[c][nl2(v)].add(k)
    return {"sheet": ws.title, "head": [x for x in head if x], "rows": rows, "text_cols": text_cols, "num_cols": num_cols,
            "index": index, "index2": index2, "raw_label": raw_label, "bad": bad}


# ---------------------------------------------------------------- 보고서 칸 (꼬리표)
_CODE_HEAD_RE = re.compile(r"(코드|번호|순번|^no\.?$|^id$|code$|^#$)", re.I)


def report_cells(ws, include_empty=False):
    grid = fill_merged(ws)
    title_rows = {rng.min_row - 1 for rng in ws.merged_cells.ranges
                  if rng.min_col == 1 and rng.max_col >= ws.max_column and rng.max_col > 1
                  and rng.min_row == rng.max_row}
    out = []
    ncol = max((len(r) for r in grid), default=0)
    col_block = [[] for _ in range(ncol)]       # 열마다 '바로 위에 이어진 머리글 묶음' (숫자·제목 줄을 만나면 새로 시작)
    col_last = ["empty"] * ncol
    for i, row in enumerate(grid):
        if i in title_rows:                          # 병합된 구분 줄: 위 열 제목은 그대로 두고, 아래에 새 제목이 나오면 그때 교체 (v13과 같은 동작)
            for j in range(ncol):
                col_last[j] = "title"
            continue
        labels = [is_label(v) for v in row]
        for j, v in enumerate(row):
            if labels[j]:
                if col_last[j] == "num" or col_last[j] == "title":
                    col_block[j] = []
                col_block[j] = col_block[j] + [(i + 1, str(v).strip())]
                col_last[j] = "label"
                continue
            num = 0.0 if is_dash(v) else to_num(v)
            header_pairs = list(col_block[j])
            if num is not None and header_pairs and _CODE_HEAD_RE.search(nl(header_pairs[-1][1])):
                col_last[j] = "num"                          # 코드·번호 열의 숫자는 금액이 아님 → 대조 대상에서 뺌
                continue
            if num is not None:
                col_last[j] = "num"
            if num is None and not (include_empty and v in (None, "")):
                continue
            row_pairs = [(k + 1, str(row[k]).strip()) for k in range(j) if labels[k]]
            row_labels = [lab for _, lab in row_pairs]
            col_labels = [lab for _, lab in header_pairs]
            if not row_labels and not col_labels:
                continue
            if num is None and (not row_labels or not col_labels):
                continue
            out.append({"r": i + 1, "c": j + 1, "coord": f"{get_column_letter(j + 1)}{i + 1}", "value": num,
                        "row_labels": row_labels, "col_labels": col_labels,
                        "row_positions": [pos for pos, _ in row_pairs],
                        "col_positions": [pos for pos, _ in header_pairs],
                        "block": tuple(pos for pos, _ in header_pairs)})
    return out


# ---------------------------------------------------------------- 매핑·집계
def _amount_fn(expr, raw):
    """'금액' 또는 '차변-대변' → row → 값"""
    if expr.strip() in raw["head"]:                    # 열 이름 그대로 (띄어쓰기 포함, 예: '금 액')
        parts = [("", expr.strip())]
    else:
        parts = re.findall(r"([+\-]?)\s*([^+\-]+)", expr)
    cols = []
    for sg, nm in parts:
        nm = nm.strip()
        if not nm:
            continue
        if nm not in raw["head"]:                       # 띄어쓰기만 다른 열 이름은 같은 열 ('금액' ↔ '금 액')
            same = [h for h in raw["head"] if nl(h) == nl(nm)]
            if len(same) == 1:
                nm = same[0]
            else:
                raise RuntimeError(f"원본에 금액 열 '{nm}' 없음. 있는 열: {raw['head']}")
        cols.append((sg, nm))
    def fn(row):
        t = 0.0
        for sg, nm in cols:
            v = to_num(row.get(nm)) or 0
            t += -v if sg == "-" else v
        return t
    fn.cols = [nm for _, nm in cols]
    return fn


def _resolve(label, raw, label_map, fuzzy_log, allowed=None):
    """보고서 글자 → [(원본 열, 정규화 키)...]  (규칙표 대응 → 정확 일치 → 유사 일치)"""
    k = nl(label)
    if k in label_map:
        k = nl(label_map[k])
    columns = raw["text_cols"] if allowed is None else allowed
    hits = [(c, k) for c in columns if c in raw["index"] and k in raw["index"][c]]
    if hits:
        return hits
    if nl(label) in label_map:
        return []  # 명시한 대응이 맞지 않으면 원래 글자로 유사 재매칭하지 않음
    k2 = nl2(label)
    candidates = []
    for c in columns:
        cand = raw["index2"].get(c, {}).get(k2)
        if cand and len(cand) == 1:
            rk = next(iter(cand))
            candidates.append((c, rk))
    if len(candidates) == 1:
        c, rk = candidates[0]
        # 한쪽은 그대로이고 다른 쪽에만 코드·괄호가 붙은 경우('1010 현금' ↔ '현금')는 확실한 대응 → 유사 확인으로 남기지 않음
        plain = (rk == k2) or (nl2(rk) == k)
        if rk != k and not plain:
            fuzzy_log[label] = raw["raw_label"][c][rk]
        return [(c, rk)]
    return []


_PERIOD_LABEL_RE = re.compile(r"^(\d{1,2}월|\d{4}[-./]\d{1,2}|\d{4}년)$")


def _add_date_parts(raw):
    """보고서에 '1월'·'2026-09'·'2026년' 같은 기간 제목이 있을 때만: 원본 날짜 열에서 월·연월·연도 글자 열을 만들어 대조에 씀 (원본 값은 그대로)"""
    for c in list(raw["head"]):
        vals = [row.get(c) for row in raw["rows"][:2000] if row.get(c) not in (None, "")]
        if not vals or sum(isinstance(v, (datetime.date, datetime.datetime)) for v in vals) < len(vals) * 0.8:
            continue
        for suffix, fmt in (("(월)", lambda d: f"{d.month}월"), ("(연월)", lambda d: f"{d.year}-{d.month:02d}"), ("(연도)", lambda d: f"{d.year}년")):
            name = f"{c}{suffix}"
            if name in raw["head"]:
                continue
            idx, idx2, labels = defaultdict(set), defaultdict(set), {}
            for i, row in enumerate(raw["rows"]):
                v = row.get(c)
                t = fmt(v) if isinstance(v, (datetime.date, datetime.datetime)) else None
                row[name] = t
                if t:
                    k = nl(t); idx[k].add(i); labels.setdefault(k, t); idx2[nl2(t)].add(k)
            raw["head"].append(name); raw["text_cols"].append(name)
            raw["index"][name] = idx; raw["index2"][name] = idx2; raw["raw_label"][name] = labels


@functools.lru_cache(maxsize=65536)
def _matches_marker(label, marker):
    # 한 글자 '계'의 접미 일치는 '기계' 등 일반 계정까지 제외하므로 금지. 괄호 설명이 붙은 '소계(현금성)'도 소계
    for key in (nl(label), nl2(label)):
        if key == marker or (len(marker) > 1 and key.endswith(marker)):
            return True
    return False


def _total_kind(cell):
    for label in cell["row_labels"]:
        if any(_matches_marker(label, marker) for marker in ("소계", "subtotal")):
            return "소계"
        if any(_matches_marker(label, marker) for marker in ("합계", "총계", "총합계", "total", "계")):
            return "합계"
    return None


def aggregate(cell, raw, amount_fn, count_mode, label_map, fuzzy_log, exclude):
    """필수 라벨이 해결되지 않으면 부분 합산하지 않는다."""
    cell["missing_labels"] = []
    for lab in cell["row_labels"]:
        if any(_matches_marker(lab, marker) for marker in exclude):
            return None, "제외", f"'{lab}' 은 집계 줄(제외 글자)"
    filters, hit_row, hit_col = [], False, False
    for pos, lab in zip(cell.get("row_positions", range(len(cell["row_labels"]))), cell["row_labels"]):
        hits = _resolve(lab, raw, label_map, fuzzy_log, cell.get("allowed_row"))
        if hits:
            filters += hits; hit_row = True
        elif cell.get("row_rates", {}).get(pos, 0) >= .5:
            cell["missing_labels"].append(lab)
    if cell["missing_labels"]:
        return None, "근거 없음", f"행 글자 {cell['missing_labels']} 가 원본에 없음"
    amount_override = None
    for pos, lab in zip(cell.get("col_positions", range(len(cell["col_labels"]))), cell["col_labels"]):
        hits = _resolve(lab, raw, label_map, fuzzy_log, cell.get("allowed_col"))
        amount_col = next((c for c in raw["num_cols"] if nl(c) == nl(label_map.get(nl(lab), lab))), None)
        if hits:
            filters += hits; hit_col = True
        if amount_col:
            amount_override = amount_col; hit_col = True
        if not hits and not amount_col and cell.get("col_rates", {}).get(pos, 0) >= .5:
            cell["missing_labels"].append(lab)
    if cell["missing_labels"]:
        return None, "근거 없음", f"열 글자 {cell['missing_labels']} 가 원본에 없음"
    if cell["row_labels"] and not hit_row:
        cell["missing_labels"] = list(cell["row_labels"])
        return None, "근거 없음", f"행 글자 {cell['row_labels']} 가 원본에 없음"
    if cell["col_labels"] and not hit_col:
        cell["missing_labels"] = list(cell["col_labels"])
        return None, "근거 없음", f"열 글자 {cell['col_labels']} 가 원본에 없음"
    if not filters:
        return None, "근거 없음", "걸러낼 글자 없음"
    idx = None
    for c, key in filters:
        indices = raw["index"][c].get(key, set())
        idx = indices if idx is None else idx & indices
    rows = [raw["rows"][i] for i in (idx or ())]
    if count_mode:
        return float(len(rows)), filters
    fn = (lambda row: to_num(row.get(amount_override)) or 0) if amount_override else amount_fn
    used_cols = [amount_override] if amount_override else list(getattr(amount_fn, "cols", []))
    bad = raw.get("bad") or {}
    problems = [bad[i][c] for i in sorted(idx or ()) if i in bad for c in used_cols if c in bad[i]]
    if problems:
        shown = "; ".join(f"{loc} = '{val}'" for loc, val in problems[:5]) + (f" 외 {len(problems) - 5}개" if len(problems) > 5 else "")
        return None, "검산 불가", f"원본 금액 칸에 숫자가 아닌 값(오류값·글자)이 있어 집계하지 않음: {shown}"
    return sum(fn(row) for row in rows), filters


def _auto_pick(labels, raw, manual):
    if manual:
        cols = [c.strip() for c in manual.split(",") if c.strip()]
        for c in cols:
            if c not in raw["head"]:
                raise RuntimeError(f"원본에 '{c}' 열 없음. 있는 열: {raw['head']}")
        return ", ".join(cols)
    hits = sorted([(sum(1 for l in labels if nl(l) in raw["index"][c] or nl2(l) in raw["index2"][c]), c) for c in raw["text_cols"]], reverse=True)
    hits = [(n, c) for n, c in hits if n > 0]
    return ", ".join(c for _, c in hits) if hits else None


def _auto_amount_scale(cells, raw, settings, label_map, exclude, warnings):
    """금액 열(식)과 단위 배수 결정. 규칙표 값 우선, 없으면 보고서 숫자와 가장 많이 맞는 조합"""
    expr = settings.get("금액 열") or None
    scale_s = settings.get("단위 배수") or None
    count_mode = (settings.get("집계") or "합계").strip() == "건수"
    if scale_s is not None:
        try:
            scale_value = float(scale_s)
        except (TypeError, ValueError):
            raise RuntimeError("단위 배수는 0보다 큰 숫자로 지정하세요")
        if not math.isfinite(scale_value) or scale_value <= 0:
            raise RuntimeError("단위 배수는 0보다 큰 유한한 숫자로 지정하세요")
    cands_amt = [expr] if expr else list(raw["num_cols"])
    if not expr:
        dr = [c for c in raw["num_cols"] if "차변" in nl(c)]; cr = [c for c in raw["num_cols"] if "대변" in nl(c)]
        if len(dr) == 1 and len(cr) == 1:                    # 차변·대변이 나뉜 원본: 순액 후보도 시험
            cands_amt += [f"{dr[0]}-{cr[0]}", f"{cr[0]}-{dr[0]}"]
        months = [c for c in raw["num_cols"] if re.fullmatch(r"(\d{1,2}월|\d{4}[-./]\d{1,2}|[1-4]분기|q[1-4])", nl(c), re.I)]
        if len(months) >= 2 and all("+" not in c and "-" not in c for c in months):   # 월이 열인 원본: 기간 합계 후보도 시험
            cands_amt.append("+".join(months))
    cands_scale = [float(scale_s)] if scale_s else [1.0, 1000.0, 1000000.0]
    if not cands_amt:
        raise RuntimeError(f"원본에 숫자 열이 없음. 열: {raw['head']}")
    numeric = [c for c in cells if c["value"] not in (None, 0)]
    best = None
    if numeric and (len(cands_amt) > 1 or len(cands_scale) > 1):
        for a in cands_amt:
            fn = _amount_fn(a, raw)
            for sc in cands_scale:
                n = 0
                for cell in numeric[:400]:
                    res = aggregate(cell, raw, fn, count_mode, label_map, {}, exclude)
                    if res[0] is not None and abs(res[0] - cell["value"] * sc) <= TOLERANCE * sc:
                        n += 1
                if best is None or n > best[0]:
                    best = (n, a, sc)
        if best[0] == 0:
            pref = next((c for c in cands_amt if any(k in c for k in ("금액", "잔액", "합계"))), cands_amt[0])
            warnings.append(f"금액 열·단위 자동 선택 근거 부족 → '{pref}', 배수 1 사용. 규칙표에서 지정 가능")
            return pref, 1.0, count_mode
        return best[1], best[2], count_mode
    return cands_amt[0], cands_scale[0], count_mode


def _report_sheets(wb, sheet):
    """[(시트, 빈 칸 포함 대조 칸 목록)] — 시트마다 한 번만 분석"""
    if sheet:
        if sheet not in wb.sheetnames:
            raise RuntimeError(f"보고서에 '{sheet}' 시트 없음. 있는 시트: {wb.sheetnames}")
        return [(wb[sheet], report_cells(wb[sheet], include_empty=True))]
    # 숨긴 시트(숨김·아주 숨김)는 보고서 본문이 아닌 계산·참고용이 대부분 → 대조에서 제외 (규칙표 '보고서 시트'로 지정하면 대조)
    pairs = [(ws, report_cells(ws, include_empty=True)) for ws in wb.worksheets if ws.sheet_state == "visible"]
    return [(ws, cells) for ws, cells in pairs if cells]


def _prepare(raw_path, report_path, include_empty=False, blank_check=False):
    warnings = []
    rules = load_rules(report_path)
    st = rules["settings"]
    raw = read_raw(raw_path, st.get("원본 시트") or None, warnings)
    for part in re.split(r"[+\-]", st.get("금액 열") or ""):      # 규칙표로 지정한 금액 열은 숫자 열로 인정 (오류값이 많아 자동 분류에서 빠졌어도)
        part = part.strip()
        match = next((h for h in raw["head"] if h == part or nl(h) == nl(part)), None) if part else None
        if match and match not in raw["num_cols"]:
            raw["num_cols"].append(match)
            raw["text_cols"] = [c for c in raw["text_cols"] if c != match]
    wb = safe_load(report_path, warnings, data_only=True)
    # 빈 칸은 상태를 바꾸지 않으므로 빈 칸 제외 목록 = 빈 칸 포함 목록에서 값 없는 칸을 뺀 것
    cells = [(ws.title, c) for ws, all_cells in _report_sheets(wb, st.get("보고서 시트") or None)
             for c in all_cells if include_empty or blank_check or c["value"] is not None]
    # 매핑·단위 추정은 값 있는 칸만으로 (빈 칸 검사를 켜도 기존 판단이 바뀌지 않게)
    ref = cells if not blank_check else [(sn, c) for sn, c in cells if c["value"] is not None]
    exclude = [nl(x) for x in (st.get("제외 글자") or EXCLUDE_DEFAULT).split(",") if x.strip()]
    if any(_PERIOD_LABEL_RE.match(str(l).strip()) for _, c in ref for l in c["row_labels"] + c["col_labels"]):
        _add_date_parts(raw)
    for c in raw["text_cols"]:                                # 통화·차대 구분 열: 자동 대조는 이 구분을 무시하므로 알림
        vals = {nl(row.get(c)) for row in raw["rows"] if row.get(c) not in (None, "")}
        if 2 <= len(vals) <= 4 and (vals <= {"krw", "usd", "jpy", "eur", "cny", "원화", "외화", "원", "달러", "엔"}
                                    or vals <= {"차", "대", "차변", "대변", "d", "c", "dr", "cr", "+", "-"}):
            warnings.append(f"원본 '{c}' 열에 구분값({', '.join(sorted(vals))})이 있음 — 자동 대조는 이 구분 없이 더하므로 "
                            "불일치가 이 때문일 수 있음 (통화별·차대별 보고서면 결과를 그대로 믿지 말 것)")
    rl = [l for _, c in ref for l in c["row_labels"]]; cl = [l for _, c in ref for l in c["col_labels"]]
    rc = _auto_pick(rl, raw, st.get("행 글자 열")); cc = _auto_pick(cl, raw, st.get("열 글자 열"))
    if rc is None and cc is None:
        warnings.append(f"보고서 글자가 원본의 어느 열 값에도 없음 → 근거 없음(False). 원본 열: {raw['head']} / 보고서 글자 예: {sorted(set(rl))[:5]} {sorted(set(cl))[:5]}")
    # 같은 보고서 열의 서로 다른 라벨로 해결률을 계산한다. 숫자 칸 수로 가중하지 않는다.
    row_labels, col_labels = defaultdict(set), defaultdict(set)
    for sheet_name, cell in cells:
        cell["allowed_row"] = [v.strip() for v in st["행 글자 열"].split(",")] if st.get("행 글자 열") else None
        cell["allowed_col"] = [v.strip() for v in st["열 글자 열"].split(",")] if st.get("열 글자 열") else None
    for sheet_name, cell in ref:
        cell["allowed_row"] = [v.strip() for v in st["행 글자 열"].split(",")] if st.get("행 글자 열") else None
        cell["allowed_col"] = [v.strip() for v in st["열 글자 열"].split(",")] if st.get("열 글자 열") else None
        if not _total_kind(cell):
            for pos, lab in zip(cell["row_positions"], cell["row_labels"]):
                row_labels[(sheet_name, pos)].add(lab)
            for pos, lab in zip(cell["col_positions"], cell["col_labels"]):
                col_labels[(sheet_name, pos)].add(lab)
    rate_cache = {}
    for sheet_name, cell in cells:
        if sheet_name not in rate_cache:
            row_rates = {pos: sum(bool(_resolve(lab, raw, rules["map"], {}, cell["allowed_row"])) for lab in labs) / len(labs)
                         for (sn, pos), labs in row_labels.items() if sn == sheet_name}
            col_rates = {pos: sum(bool(_resolve(lab, raw, rules["map"], {}, cell["allowed_col"])) or any(nl(c) == nl(rules["map"].get(nl(lab), lab)) for c in raw["num_cols"]) for lab in labs) / len(labs)
                         for (sn, pos), labs in col_labels.items() if sn == sheet_name}
            rate_cache[sheet_name] = row_rates, col_rates
        cell["row_rates"], cell["col_rates"] = rate_cache[sheet_name]
    try:
        tol = float(st.get("허용 오차") or TOLERANCE)
    except (TypeError, ValueError):
        raise RuntimeError("허용 오차는 0 이상의 숫자로 지정하세요")
    if not math.isfinite(tol) or tol < 0:
        raise RuntimeError("허용 오차는 0 이상의 유한한 숫자로 지정하세요")
    n_row_cols = len([x for x in str(rc or "").split(",") if x.strip()])
    report_rows = len({(sn, c["r"]) for sn, c in ref if not _total_kind(c)})
    if n_row_cols >= 4 and not st.get("행 글자 열") and report_rows >= max(5, 0.5 * len(raw["rows"])):
        # 행 글자가 원본의 여러 열(일자·종목·수량…)에 걸침 = 보고서가 원본을 행 단위로 옮긴 표. 합계형 대조를 하면 거짓 불일치가 나옴
        raise RuntimeError(f"보고서가 원본을 행 단위로 옮긴 표로 보임 (행 글자가 원본 {n_row_cols}개 열 '{rc}'에 걸침). "
                           "합계형 보고서가 아니라 자동 대조하지 않음 — 대조하려면 규칙표의 '행 글자 열'을 지정")
    amt, scale, count_mode = _auto_amount_scale([c for _, c in ref], raw, st, rules["map"], exclude, warnings)
    return {"warnings": warnings, "rules": rules, "raw": raw, "cells": cells, "exclude": exclude, "row_col": rc, "col_col": cc,
            "amount": amt, "scale": scale, "count": count_mode, "tol": tol, "fn": _amount_fn(amt, raw)}


def _report_formula_cells(path, sheets):
    """보고서에서 수식이 있고 저장된 계산값이 아예 없는 칸 {(시트, 주소)} — 빈 칸이 '계산값 없는 수식'인지 구분용.
    수식 결과가 빈 글자("")로 저장된 칸은 엑셀에서 다시 저장해도 비어 있으므로 여기 넣지 않음(→ 누락)"""
    import zipfile
    import xml.etree.ElementTree as ET
    out = set()
    try:
        wb = load_workbook(path, read_only=True)
        names = wb.sheetnames
        wb.close()
        with zipfile.ZipFile(path) as z:
            rels = ET.fromstring(z.read("xl/_rels/workbook.xml.rels"))
            target = {r.get("Id"): r.get("Target") for r in rels}
            book = ET.fromstring(z.read("xl/workbook.xml"))
            ns = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
            rid = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
            for sh in book.find("m:sheets", ns):
                name = sh.get("name")
                if name not in sheets or name not in names:
                    continue
                t = target.get(sh.get(rid), "")
                t = t.lstrip("/") if t.startswith("/") else "xl/" + t
                root = ET.fromstring(z.read(t))
                for c in root.iter("{%s}c" % ns["m"]):
                    f = c.find("m:f", ns)
                    v_ = c.find("m:v", ns)
                    if f is not None and (v_ is None or not (v_.text or "")) and c.get("t") != "str" and c.find("m:is", ns) is None:
                        out.add((name, c.get("r")))
    except Exception:
        return set()
    return out


def verify(raw_path, report_path):
    P = _prepare(raw_path, report_path, blank_check=True)
    raw, rules = P["raw"], P["rules"]
    mismatch, unmatched, matched, fuzzy = [], [], 0, dict(rules["pending_fuzzy"])
    fuzzy_cells, new_labels = set(), set()
    checks = []
    states = {}
    row_states = {}
    pending_out = []
    seen_rowsets = {}
    parent_rows = set()
    def _col_total_kind(cell):
        for label in cell["col_labels"]:
            if any(_matches_marker(label, marker) for marker in ("소계", "subtotal")):
                return "소계"
            if any(_matches_marker(label, marker) for marker in ("합계", "총계", "총합계", "total", "계")):
                return "합계"
        return None
    # 구분 제목 줄 판정 재료: 같은 열 아래 칸·같은 줄 오른쪽 칸, 칸마다 원본 행 집합
    by_col, by_row = {}, {}
    for sn, c in P["cells"]:
        by_col.setdefault((sn, c["c"], c["block"]), []).append(c)
        by_row.setdefault((sn, c["r"]), []).append(c)
    after = {}
    for key_, v in list(by_col.items()) + list(by_row.items()):
        v.sort(key=lambda c: (c["r"], c["c"]))
        for i_, c in enumerate(v):
            after[(id(c), "v" if key_ in by_col and len(key_) == 3 else "h")] = (v, i_)
    detail_valued = [(sn, c) for sn, c in P["cells"] if c["value"] is not None and not _total_kind(c) and not _col_total_kind(c)]
    line_has_value = {(sn, "v", c["r"]) for sn, c in detail_valued} | {(sn, "h", c["c"]) for sn, c in detail_valued}
    rowset_cache, field_cache = {}, {}
    def _cell_rows(c):
        """칸의 행·열 글자에 해당하는 원본 행 번호 집합 (글자 대응이 없으면 None)"""
        key = (tuple(c["row_labels"]), tuple(c["col_labels"]), id(c.get("allowed_row")), id(c.get("allowed_col")))
        if key not in rowset_cache:
            sets = []
            for labs, allowed in ((c["row_labels"], c.get("allowed_row")), (c["col_labels"], c.get("allowed_col"))):
                for lab_ in labs:
                    hits = _resolve(lab_, raw, rules["map"], {}, allowed) or []
                    if hits:
                        sets.append(set().union(*(raw["index"][col].get(k_, set()) for col, k_ in hits)))
            rowset_cache[key] = frozenset(set.intersection(*sets)) if sets else None
        return rowset_cache[key]
    label_cache = {}
    def _label_rows(c, d_):
        """행 방향(v)은 행 글자만, 열 방향(h)은 열 글자만으로 본 원본 행 집합 — 상하위 관계 판정용 (그 열에 거래가 없는 하위 계정도 하위로 봄)"""
        labs, allowed = (c["row_labels"], c.get("allowed_row")) if d_ == "v" else (c["col_labels"], c.get("allowed_col"))
        key = (tuple(labs), id(allowed), d_)
        if key not in label_cache:
            sets = []
            for lab_ in labs:
                hits = _resolve(lab_, raw, rules["map"], {}, allowed) or []
                if hits:
                    sets.append(set().union(*(raw["index"][col].get(k_, set()) for col, k_ in hits)))
            label_cache[key] = frozenset(set.intersection(*sets)) if sets else None
        return label_cache[key]
    amt_cache = {}
    def _amt_col(c):
        """칸의 금액 열(열 글자가 원본 숫자 열 이름과 같으면 그 열, 아니면 기본 금액)"""
        key = tuple(c["col_labels"])
        if key not in amt_cache:
            amt_cache[key] = next((col for lab_ in c["col_labels"] for col in raw["num_cols"]
                                   if nl(col) == nl(rules["map"].get(nl(lab_), lab_))), None)
        return amt_cache[key]
    def _fields(c, d_="v"):
        labs, allowed = (c["row_labels"], c.get("allowed_row")) if d_ == "v" else (c["col_labels"], c.get("allowed_col"))
        key = (tuple(labs), id(allowed))
        if key not in field_cache:
            field_cache[key] = frozenset(col for lab_ in labs for col, _ in (_resolve(lab_, raw, rules["map"], {}, allowed) or []))
        return field_cache[key]
    # 같은 단계(같은 원본 열로 대응되는) 다른 줄·열에 값을 적은 양식인지 — 행 방향은 같은 열, 열 방향은 같은 줄 (블록이 나뉘어도 같은 열이면 같은 양식)
    valued_by_line = {}
    for sn, c in detail_valued:
        valued_by_line.setdefault((sn, "v", c["c"]), []).append(c)
        valued_by_line.setdefault((sn, "h", c["r"]), []).append(c)
    form_cache = {}
    def _first_child(sn, c, d_, R):
        """바로 다음 칸(방향 d_)이 이 칸의 하위(글자 기준 원본 행 부분집합)이면 그 칸. 소계·합계 줄이나 원본에 없는 제목 줄을 만나면 하위 없음
        (요약 구간 '유동자산 300 / 비유동자산 500 / 합계' 다음의 명세 구간을 하위로 잡지 않게)"""
        Pc = _label_rows(c, d_)
        if not Pc:
            return None
        v, i_ = after[(id(c), d_)]
        for j_ in range(i_ + 1, len(v)):
            n = v[j_]
            if d_ == "h" and _col_total_kind(n) or d_ == "v" and _total_kind(n):
                return None
            if d_ == "v" and _col_total_kind(n) or d_ == "h" and _total_kind(n):
                continue
            Pn = _label_rows(n, d_)
            if not Pn:
                if n["value"] is None:
                    continue                                 # 원본에 없는 빈 중분류 줄('(1)당좌자산')은 건너뜀
                return None
            return n if (Pn < Pc or (Pn == Pc and _fields(n, d_) != _fields(c, d_))) else None
        return None
    def _valued_heading_form(sn, c, d_):
        """같은 줄(열)에 같은 원본 열로 대응되고 하위 칸을 거느린 '값 있는 제목'이 있음 = 위쪽 소계형 양식 (요약 블록의 값은 제외)"""
        line = c["c"] if d_ == "v" else c["r"]
        key = (sn, d_, line, _fields(c, d_))
        if key not in form_cache:
            form_cache[key] = any(_fields(x, d_) == key[3] and _cell_rows(x) and _first_child(sn, x, d_, _cell_rows(x)) is not None
                                  for x in valued_by_line.get((sn, d_, line), ()))
        return form_cache[key]
    nocache_cells = _report_formula_cells(report_path, {sn for sn, _ in P["cells"]}) if any(c["value"] is None for _, c in P["cells"]) else set()
    def _coverage(sn, c, d_, R):
        """방향 d_로 이 칸의 하위 칸(원본 행 부분집합)이 이어지는 동안 원본 행을 얼마나 덮는지: (하위 있음, 'full'|'partial')
        덮지 못한 원본 행은 하위 계정별·전체 금액 합이 0이면(잔액 0 계정 생략, 원본 안 소계 줄) 덮은 것으로 봄"""
        v, i_ = after[(id(c), d_)]
        covered, child, child_fields = set(), False, set()
        for j_ in range(i_ + 1, len(v)):
            n = v[j_]
            if _total_kind(n) or _col_total_kind(n):
                continue
            Pn, Pc = _label_rows(n, d_), _label_rows(c, d_)
            if not Pn or not Pc:
                continue
            if not (Pn < Pc or (Pn == Pc and (child or _fields(n, d_) != _fields(c, d_)))):
                break                                    # 이 칸의 원본 행 밖 → 구분 끝 (한 계정뿐인 대분류는 같은 행 집합이어도 다른 원본 열이면 하위)
            child = True
            covered |= (_cell_rows(n) or set())
            child_fields |= _fields(n, d_)
        if not child:
            return False, None
        rest = R - covered                           # 아래 칸이 담지 않은 원본 행 — 하위 계정마다 금액 합이 0이면(잔액 0 계정 생략) 다 담은 것으로 봄
        if not rest:
            return True, "full"
        if P["count"]:
            return True, "partial"
        child_cols = sorted(child_fields) or []
        amt_col = _amt_col(c)                            # 이 칸의 금액 열(당기·전기 등)
        fn_ = (lambda row_: to_num(row_.get(amt_col)) or 0) if amt_col else P["fn"]
        by_acct, total_ = {}, 0.0
        for i in rest:
            row_ = raw["rows"][i]
            if any(isinstance(row_.get(col), str) and _matches_marker(row_.get(col), mk) for col in child_cols for mk in ("소계", "합계", "총계")):
                continue                                 # 원본 안의 소계·합계 줄
            k_ = tuple(nl(str(row_.get(col))) if row_.get(col) not in (None, "") else None for col in child_cols)
            a_ = fn_(row_)
            by_acct[k_] = by_acct.get(k_, 0.0) + a_; total_ += a_
        tol_ = P["tol"] * P["scale"]
        return True, ("full" if abs(total_) <= tol_ and all(abs(v_) <= tol_ for v_ in by_acct.values()) else "partial")
    def _valued_partial(sn, c):
        """값을 적은 구분 칸(위쪽 소계형 '유동자산 350')의 하위 칸들이 원본 계정 일부만 담는지"""
        R = _cell_rows(c)
        if not R:
            return False
        for d_ in ("v", "h"):                               # 상하위는 글자 기준(열 방향 '금액 | 영업부 | 관리부'의 금액 열은 글자 구분이 아니라 하위 아님)
            if _first_child(sn, c, d_, R) is None:
                continue
            child, verdict = _coverage(sn, c, d_, R)
            if child and verdict == "partial":
                return True
        return False
    def _heading(sn, c):
        """빈 칸이 구분 제목(대분류 '유동자산' 등)인지: 'full' = 아래(또는 오른쪽) 칸들이 이 칸의 원본 행을 모두 나눠 담음,
        'partial' = 제목 모양인데 아래 칸이 원본 일부만 담음(빠진 계정 행 의심), None = 제목 아님(누락)"""
        R = _cell_rows(c)
        if not R:
            return None
        for d_ in ("v", "h"):
            if (sn, d_, c["r"] if d_ == "v" else c["c"]) in line_has_value:
                continue                                     # 제목 줄(열)은 상세 숫자가 하나도 없어야 함
            child, verdict = _coverage(sn, c, d_, R)
            if child:
                if _valued_heading_form(sn, c, d_):
                    return None                              # 같은 단계의 다른 줄(열)엔 값을 적은 양식(위쪽 소계형) → 비운 칸은 누락
                return verdict
        return None
    for sname, cell in P["cells"]:
        loc = f"{sname}!{cell['coord']}"
        lab = " | ".join(cell["row_labels"]) + " × " + " ".join(cell["col_labels"])
        titles = {"행 제목": " | ".join(cell["row_labels"]), "열 제목": " | ".join(cell["col_labels"])}
        state = states.setdefault((sname, cell["c"], cell["block"]), {"segment": [], "all": [], "segment_locations": [], "all_locations": [], "fuzzy_segment": {}, "fuzzy_all": {}})
        rstate = row_states.setdefault((sname, cell["r"]), {"vals": [], "locs": [], "fuzzy": {}, "seg_vals": [], "seg_locs": [], "seg_fuzzy": {}})
        kind = _total_kind(cell)
        col_kind = _col_total_kind(cell)
        cell_fuzzy_map = {}
        basis = ""
        dup_of = None
        if col_kind and not kind:
            # 합계·소계 열: 같은 줄의 상세 칸(원본 집계값) 합과 대조. 소계 열은 직전 소계 이후 구간만
            vals, locs = (rstate["seg_vals"], rstate["seg_locs"]) if col_kind == "소계" else (rstate["vals"], rstate["locs"])
            cell_fuzzy_map = dict(rstate["seg_fuzzy"] if col_kind == "소계" else rstate["fuzzy"])
            basis = f"{col_kind} 열 대상: {', '.join(locs)}; "
            missing_locations = [location for location, value in zip(locs, vals) if value is None]
            if not vals:
                res = (None, "근거 없음", f"{col_kind} 열 왼쪽에 집계할 상세 숫자 칸이 없음")
            elif missing_locations:
                res = (None, "근거 없음", f"{col_kind} 열 검산 불가: 원본 근거가 없는 상세 칸 {', '.join(missing_locations)}")
            else:
                res = (sum(vals), [])
            rstate["seg_vals"] = []; rstate["seg_locs"] = []; rstate["seg_fuzzy"] = {}
            if col_kind == "소계":
                rstate["seg_id"] = rstate.get("seg_id", 0) + 1
            elif col_kind == "합계":                          # 합계 열 다음은 새 표(영업 합계 | 관리팀 … | 관리 합계) — 세로 합계 줄과 같은 처리
                rstate["seg_id"] = rstate.get("seg_id", 0) + 1; rstate["tot_id"] = rstate.get("tot_id", 0) + 1
                rstate["vals"] = []; rstate["locs"] = []; rstate["fuzzy"] = {}
            if cell["value"] is None and res[0] is None:
                continue                                     # 빈 합계 열 칸 + 왼쪽 근거 없음 (구분 제목 줄 등) → 여백
            if (sname, cell["r"]) not in parent_rows:          # 구분 줄(위쪽 소계형)의 합계 열 칸은 아래 계정 줄과 같은 돈 → 열 합계 기대값에 안 넣음
                state["segment_locations"].append(loc); state["all_locations"].append(loc)
                state["segment"].append(res[0]); state["all"].append(res[0])
        elif kind:
            entries = state["segment"] if kind == "소계" else state["all"]
            locations = state["segment_locations"] if kind == "소계" else state["all_locations"]
            cell_fuzzy_map = dict(state["fuzzy_segment"] if kind == "소계" else state["fuzzy_all"])
            basis = f"{kind} 대상: {', '.join(locations)}; "
            missing_locations = [location for location, value in zip(locations, entries) if value is None]
            if not entries:
                res = (None, "근거 없음", f"{kind} 이전에 집계할 상세 숫자 칸이 없음")
            elif missing_locations:
                res = (None, "근거 없음", f"{kind} 검산 불가: 원본 근거가 없는 상세 칸 {', '.join(missing_locations)}")
            else:
                res = (sum(entries), [])
                # 겹친 소계('소계(현금성)' 아래 '소계(유동)'): 직전 소계 이후 구간 합이 안 맞고, 직전 합계 이후 전체 상세 합과 맞으면 누적 소계로 판정
                qualified = any(nl(l) not in ("소계", "subtotal") and any(_matches_marker(l, mk) for mk in ("소계", "subtotal")) for l in cell["row_labels"])
                if kind == "소계" and qualified and abs(res[0] - cell["value"] * P["scale"]) > P["tol"] * P["scale"] \
                        and state["all"] and None not in state["all"] and len(state["all"]) > len(entries) \
                        and abs(sum(state["all"]) - cell["value"] * P["scale"]) <= P["tol"] * P["scale"]:
                    # 단순 오류(소계 자리에 누적값을 잘못 넣음)와 구분할 수 없으므로 통과시키지 않고 '사람 확인'으로 남김
                    res = (sum(state["all"]), [])
                    basis = f"누적 소계로 보면 일치 (대상: {', '.join(state['all_locations'])}), 직전 소계 이후 구간 합과는 다름; "
                    cell_fuzzy_map = {" ".join(cell["row_labels"]): "누적 소계인지 확인"}
            state["segment"] = []
            state["segment_locations"] = []
            state["fuzzy_segment"] = {}
            state["seg_id"] = state.get("seg_id", 0) + 1
            if kind == "합계":
                state["tot_id"] = state.get("tot_id", 0) + 1
                state["all"] = []
                state["all_locations"] = []
                state["fuzzy_all"] = {}
        else:
            local_fuzzy = {}
            res = aggregate(cell, raw, P["fn"], P["count"], rules["map"], local_fuzzy, P["exclude"])
            blank_missing = None
            if cell["value"] is None:
                # 빈 상세 칸: 원본 근거 없으면 여백, 집계 0이면 0으로 봄, 집계≠0이면 '누락'(원본 집계값은 소계·합계 기대값에 그대로 넣음).
                # 단 구분 제목 줄(아래 칸들이 이 칸의 원본 행을 나눠 담음 — 대분류 → 계정)은 대조하지 않음
                if res[0] is None:
                    continue
                if abs(res[0]) > P["tol"] * P["scale"]:
                    pend = {"loc": loc, "lab": lab, "titles": titles, "val": res[0], "cell": cell, "sheet": sname}
                    head = _heading(sname, cell)
                    if head == "full":
                        parent_rows.add((sname, cell["r"]))
                    if (sname, cell["coord"]) in nocache_cells:   # 저장값 없는 수식 → '계산값 없음'. 제목 칸이면 구간에 안 넣고, 상세 칸이면 원본값을 넣음
                        if head:
                            pending_out.append(pend)
                            continue
                        blank_missing = pend
                        head = "nocache"
                    if head == "full":
                        continue
                    if head == "partial":                    # 제목 줄인데 아래에 원본 계정 일부가 없음 → 빠진 행 의심
                        pend["why"] = f"구분 제목 줄로 보이나 아래 칸들이 원본 집계 {res[0] / P['scale']:g} 중 일부만 담음 — 빠진 계정 행 확인"
                        pending_out.append(pend)
                        continue
                    if head != "nocache":
                        blank_missing = pend
                else:
                    if (sname, cell["coord"]) in nocache_cells:   # 원본 0인 칸의 저장값 없는 수식 — 값을 모르므로 '계산값 없음'
                        pending_out.append({"loc": loc, "lab": lab, "titles": titles, "val": res[0], "cell": cell, "sheet": sname})
                        continue
                    if _heading(sname, cell) == "partial":     # 순합계 0인 구분 제목(±상계 계정)인데 아래에 원본 계정 일부가 없음
                        pending_out.append({"loc": loc, "lab": lab, "titles": titles, "val": res[0], "cell": cell, "sheet": sname,
                                            "why": "구분 제목 줄로 보이나 아래 칸들이 원본 계정 일부만 담음 (구분 합계는 0) — 빠진 계정 행 확인"})
                        continue
                    res = (0.0, [])
            fuzzy.update(local_fuzzy)
            cell_fuzzy_map = {label: fuzzy[label] for label in cell["row_labels"] + cell["col_labels"] if label in fuzzy}
            state["fuzzy_segment"].update(cell_fuzzy_map)
            state["fuzzy_all"].update(cell_fuzzy_map)
            dup_of = dup_h = dup_v = dup_all = dup_h_all = None
            if cell["value"] is not None and res[0] is not None:   # 같은 구간·같은 열(또는 같은 줄)에 같은 원본 집계를 담은 칸이 이미 있음 = 같은 계정 두 번
                rs_ = _cell_rows(cell)
                if rs_:
                    sig_ = (_fields(cell, "v"), _fields(cell, "h"), _amt_col(cell), rs_)
                    # 같은 소계 구간 안 두 번 = 확인 필요 + 기대값에 한 번만 (보고서가 두 번 더했으면 불일치)
                    # 같은 합계 표의 다른 구간에 또 나옴 = 확인 필요만 (주요/전체 구간처럼 정상일 수도, 잘못 적었을 수도 있음 — 기대값은 보고서대로 둠)
                    dkey = (sname, "v", cell["c"], cell["block"], state.get("seg_id", 0)) + sig_
                    dup_v = seen_rowsets.get(dkey); seen_rowsets.setdefault(dkey, loc)
                    akey = (sname, "va", cell["c"], cell["block"], state.get("tot_id", 0)) + sig_
                    dup_all = seen_rowsets.get(akey); seen_rowsets.setdefault(akey, loc)
                    hkey = (sname, "h", cell["r"], rstate.get("seg_id", 0)) + sig_
                    dup_h = seen_rowsets.get(hkey); seen_rowsets.setdefault(hkey, loc)
                    hakey = (sname, "ha", cell["r"], rstate.get("tot_id", 0)) + sig_
                    dup_h_all = seen_rowsets.get(hakey); seen_rowsets.setdefault(hakey, loc)
                    dup_of = dup_v or dup_h or dup_all or dup_h_all
            # 값을 적은 구분 칸(위쪽 소계형 '유동자산 300' 아래 계정 줄)은 하위 칸과 같은 돈이라 소계·합계 기대값에 넣지 않음
            Rv = _cell_rows(cell) if cell["value"] is not None and res[0] is not None else None
            v_parent = bool(Rv) and _first_child(sname, cell, "v", Rv) is not None
            h_parent = bool(Rv) and _first_child(sname, cell, "h", Rv) is not None
            if v_parent:
                parent_rows.add((sname, cell["r"]))
            if not v_parent and not dup_v:                     # 같은 구간에 같은 계정을 두 번 적었으면 소계·합계 기대값에는 한 번만
                state["segment_locations"].append(loc); state["segment"].append(res[0])
                state["all_locations"].append(loc); state["all"].append(res[0])
            if not h_parent and not dup_h:
                rstate["vals"].append(res[0]); rstate["locs"].append(loc); rstate["fuzzy"].update(cell_fuzzy_map)
                rstate["seg_vals"].append(res[0]); rstate["seg_locs"].append(loc); rstate["seg_fuzzy"].update(cell_fuzzy_map)
            new_labels.update(cell.get("missing_labels", []))
            if blank_missing:
                pending_out.append(blank_missing)
                continue
        if cell_fuzzy_map:
            fuzzy_cells.add(loc)
        check = {"위치": loc, "검증통과": False, "판정": "", "라벨": lab, **titles,
                 "보고서": cell["value"], "원본": None, "이유": ""}
        checks.append(check)
        if res[0] is None:
            if cell["value"] is None:
                checks.pop(); continue                   # 빈 칸 + 원본 근거 없음 → 서식 여백으로 보고 대조 안 함
            check.update({"판정": res[1], "이유": res[2]})
            unmatched.append({"위치": loc, "구분": res[1], "라벨": lab, **titles, "이유": res[2]})
            continue
        if cell["value"] is None:
            if abs(res[0]) <= P["tol"] * P["scale"]:
                checks.pop()
                if (sname, cell["coord"]) in nocache_cells:   # 저장값 없는 수식은 값을 모름 → '계산값 없음'
                    pending_out.append({"loc": loc, "lab": lab, "titles": titles, "val": res[0], "cell": cell, "sheet": sname})
                continue                                 # 빈 칸 + 원본 집계 0 → 정상 (0을 빈칸으로 표시)
            checks.pop()                                 # 빈 소계·합계 칸 → 아래에서 수식 칸인지 보고 '누락'·'계산값 없음'
            pending_out.append({"loc": loc, "lab": lab, "titles": titles, "val": res[0], "cell": cell, "sheet": sname})
            continue
        total = res[0]; matched += 1
        check["원본"] = total / P["scale"]
        if abs(total - cell["value"] * P["scale"]) > P["tol"] * P["scale"]:
            sign_only = total != 0 and abs(total + cell["value"] * P["scale"]) <= P["tol"] * P["scale"]
            check.update({"판정": "불일치", "이유": basis + ("부호만 반대 (비용·부채를 음수로 표시한 보고서인지 확인); " if sign_only else "")
                          + f"보고서 {cell['value']:g}, 원본 집계 {total / P['scale']:g}, 차이(보고서-원본) {cell['value'] - total / P['scale']:g}, 허용 오차 {P['tol']:g} 초과 (보고서 단위)"})
            mismatch.append({"위치": loc, "라벨": lab, **titles, "보고서": cell["value"], "원본": total / P["scale"], "차이": cell["value"] - total / P["scale"]})
        elif cell_fuzzy_map:
            mappings = "; ".join(f"'{label}' → '{target}'" for label, target in cell_fuzzy_map.items())
            check.update({"판정": "유사 확인", "이유": basis + f"금액은 일치하나 유사 글자 대응 미확인: {mappings}"})
        elif not kind and not col_kind and dup_of:
            why = (f"같은 원본 집계를 담은 칸({dup_of})이 같은 구간에 이미 있음 — 같은 계정을 두 번 적었는지 확인 (소계·합계 기대값에는 한 번만 넣음)" if (dup_v or dup_h)
                   else f"같은 원본 집계를 담은 칸({dup_of})이 같은 표의 다른 구간에 있음 — 일부러 다시 보여 준 구간인지, 잘못 적었는지 확인")
            check.update({"판정": "확인 필요", "이유": basis + why})
            unmatched.append({"위치": loc, "구분": "확인 필요", "라벨": lab, **titles, "이유": why})
        elif not kind and not col_kind and _valued_partial(sname, cell):
            why = "금액은 원본과 같으나 아래(오른쪽) 칸들이 이 구분의 원본 계정 일부만 담음 — 빠진 계정 행 확인 (서로 상계되는 계정이 함께 빠졌을 수 있음)"
            check.update({"판정": "확인 필요", "이유": basis + why})
            unmatched.append({"위치": loc, "구분": "확인 필요", "라벨": lab, **titles, "이유": why})
        else:
            check.update({"검증통과": True, "판정": "일치", "이유": "원본 근거의 집계값과 허용 오차 이내로 일치"})
    formula_cells = nocache_cells
    for p in pending_out:
        no_cache = (p["sheet"], p["cell"]["coord"]) in formula_cells
        kind_name = "계산값 없음" if no_cache else "누락"
        why = ("수식 칸인데 저장된 계산값이 없음 — 엑셀에서 열어 저장한 뒤 다시 실행" if no_cache
               else p.get("why") or f"보고서 칸이 비어 있는데 원본 집계는 {p['val'] / P['scale']:g}")
        checks.append({"위치": p["loc"], "검증통과": False, "판정": kind_name, "라벨": p["lab"], **p["titles"], "보고서": None,
                       "원본": p["val"] / P["scale"], "이유": why})
        unmatched.append({"위치": p["loc"], "구분": kind_name, "라벨": p["lab"], **p["titles"], "이유": why})
    review_locations = fuzzy_cells | {m["위치"] for m in mismatch} | {u["위치"] for u in unmatched}
    # 칸 수 = 값 있는 칸 + 누락으로 판정한 빈 칸 (원본 근거 없는 빈 칸·집계 0인 빈 칸은 세지 않음)
    n_cells = sum(1 for _, c in P["cells"] if c["value"] is not None) + sum(1 for u in unmatched if u["구분"] in ("누락", "계산값 없음"))
    return {"raw": os.path.basename(raw_path), "report": os.path.basename(report_path), "raw_sheet": raw["sheet"],
            "raw_rows": len(raw["rows"]), "raw_head": list(raw["head"]),
            "mapping": {"row": P["row_col"], "col": P["col_col"], "amount": P["amount"] + (" (건수)" if P["count"] else ""), "scale": P["scale"]},
            "checks": checks, "passed": bool(checks) and all(x["검증통과"] for x in checks),
            "cells": n_cells, "matched": matched, "mismatch": mismatch, "unmatched": unmatched,
            "excluded": sum(1 for u in unmatched if u["구분"] == "제외"), "no_basis": sum(1 for u in unmatched if u["구분"] == "근거 없음"),
            "warnings": P["warnings"], "rules_path": rules["path"], "rules_exists": rules["exists"],
            "new_labels": sorted(l for l in new_labels if l and nl(l) not in rules["map"]), "fuzzy": fuzzy,
            "fuzzy_cells": sorted(fuzzy_cells), "review_count": len(review_locations), "review_ratio": len(review_locations) / max(n_cells, 1),
            "detected": {"원본 시트": raw["sheet"], "행 글자 열": P["row_col"] or "", "열 글자 열": P["col_col"] or "", "금액 열": P["amount"],
                         "집계": "건수" if P["count"] else "합계", "단위 배수": P["scale"], "제외 글자": ", ".join(P["exclude"]), "허용 오차": P["tol"]}}


def fill(raw_path, report_path):
    P = _prepare(raw_path, report_path, include_empty=True)
    raw, rules = P["raw"], P["rules"]
    wbw = safe_load(report_path, None)
    filled = changed = 0; unmatched = []; fuzzy = {}
    for sname, cell in P["cells"]:
        res = aggregate(cell, raw, P["fn"], P["count"], rules["map"], fuzzy, P["exclude"])
        if res[0] is None:
            unmatched.append({"위치": f"{sname}!{cell['coord']}", "구분": res[1], "이유": res[2]}); continue
        tgt = wbw[sname].cell(cell["r"], cell["c"])
        if isinstance(tgt.value, str) and tgt.value.startswith("="):
            continue
        total = res[0] / P["scale"]
        old = to_num(tgt.value)
        tgt.value = int(total) if float(total).is_integer() else total; filled += 1
        if old is None or abs(old - total) > P["tol"]:
            tgt.fill = YELLOW; changed += 1
    base, ext = os.path.splitext(report_path)
    out = out_path(base, "_채움", ext); wbw.save(out)
    save_rules(report_path, {"원본 시트": raw["sheet"], "행 글자 열": P["row_col"] or "", "열 글자 열": P["col_col"] or "", "금액 열": P["amount"],
                             "집계": "건수" if P["count"] else "합계", "단위 배수": P["scale"], "제외 글자": ", ".join(P["exclude"]), "허용 오차": P["tol"]}, [], fuzzy)
    return out, {"raw": os.path.basename(raw_path), "report": os.path.basename(report_path), "mapping": {"row": P["row_col"], "col": P["col_col"], "amount": P["amount"], "scale": P["scale"]},
                 "filled": filled, "changed": changed, "unmatched": unmatched, "warnings": P["warnings"], "rules_path": rules["path"]}


# ---------------------------------------------------------------- 결과·실행
def _blocked(r):
    """원본 금액 칸의 오류값·글자 때문에 집계하지 않은 보고서 칸 수"""
    return sum(1 for u in r.get("unmatched", []) if u.get("구분") == "검산 불가")


def _missing(r):
    """빈 칸인데 원본 집계가 있는 칸"""
    return sum(1 for u in r.get("unmatched", []) if u.get("구분") == "누락")


def _flagged(r):
    """금액은 맞지만 사람 확인이 필요한 칸 (같은 계정 줄 두 번, 구분 아래 계정 일부만 표시)"""
    return sum(1 for u in r.get("unmatched", []) if u.get("구분") == "확인 필요")


def _nocache(r):
    """저장된 계산값이 없는 수식 칸 (엑셀에서 열어 저장하면 해결)"""
    return sum(1 for u in r.get("unmatched", []) if u.get("구분") == "계산값 없음")


def summary(r):
    pct = r["matched"] / max(r["cells"], 1) * 100
    L = [f"원본: {r['raw']} ('{r['raw_sheet']}' 시트)  →  보고서: {r['report']}",
         f"매핑: 행 글자 ↔ 원본 '{r['mapping']['row']}' / 열 글자 ↔ 원본 '{r['mapping']['col']}' / 금액 = '{r['mapping']['amount']}' / 단위 배수 {int(r['mapping']['scale'])}",
         f"보고서 숫자 칸 {r['cells']}개 = 일치 {r['matched'] - len(r['mismatch']) - _flagged(r)} + 불일치 {len(r['mismatch'])}" + (f" + 확인 필요 {_flagged(r)}" if _flagged(r) else "") + f" + 제외 {r['excluded']} + 근거 없음 {r['no_basis']}" + (f" + 검산 불가 {_blocked(r)}" if _blocked(r) else "") + (f" + 누락 {_missing(r)}" if _missing(r) else "") + (f" + 계산값 없음 {_nocache(r)}" if _nocache(r) else "") + f"   (원본 매칭 {pct:.0f}%)",
         f"규칙표: {os.path.basename(r['rules_path'])} " + ("(사용)" if r["rules_exists"] else "(이번에 생성)") + (f" — 원본 글자를 채울 칸 {len(r['new_labels'])}개" if r["new_labels"] else "")]
    L.append(f"사람 확인 {r['review_count']}칸 / {r['cells']}칸 ({r['review_ratio']:.1%}, 중복 제외; 유사 확인 {len(r['fuzzy_cells'])}칸)")
    if r["review_ratio"] >= .20:
        L.append("확인 비율 목표(20% 미만) 미달: 근거 없음·유사대응을 규칙표에서 확인하세요")
    for w in r["warnings"]:
        L.append("주의: " + w)
    return L


def run(raw_path, report_path, result=None):
    r = verify(raw_path, report_path) if result is None else result
    save_rules(report_path, r["detected"], r["new_labels"], r["fuzzy"])
    base, _ = os.path.splitext(report_path)
    out = out_path(base, "_원본대사")
    bold, red, yel, gray = Font(bold=True), PatternFill("solid", fgColor="FFC7CE"), PatternFill("solid", fgColor="FFEB9C"), PatternFill("solid", fgColor="E7E6E6")
    wb = Workbook(); ws = wb.active; ws.title = "요약"
    ws.append(["최종", "전부 일치" if r["cells"] and not r["review_count"] else f"확인 필요 {r['review_count']}칸"]); ws["A1"].font = bold; ws["B1"].font = bold
    if r["mismatch"] or _missing(r): ws["B1"].fill = red
    ws.append(["실행일시", datetime.datetime.now().strftime("%Y-%m-%d %H:%M")])
    for line in summary(r):
        ws.append(["", line])
    ws.column_dimensions["A"].width = 12; ws.column_dimensions["B"].width = 120
    m = wb.create_sheet("불일치"); m.append(["위치", "행 제목", "열 제목", "보고서 값", "원본 집계", "차이"])
    for n, x in enumerate(r["mismatch"], start=2):
        m.append([x["위치"], x["행 제목"], x["열 제목"], x["보고서"], x["원본"], x["차이"]]); m.cell(n, 6).fill = red
    u = wb.create_sheet("근거없음·제외"); u.append(["위치", "구분", "행 제목", "열 제목", "이유"])
    for n, x in enumerate(sorted(r["unmatched"], key=lambda z: (z["구분"] != "근거 없음", z["위치"])), start=2):
        u.append([x["위치"], x["구분"], x["행 제목"], x["열 제목"], x["이유"]]); u.cell(n, 2).fill = red if x["구분"] == "누락" else yel if x["구분"] in ("근거 없음", "계산값 없음") else gray
    fz = wb.create_sheet("유사확인"); fz.append(["위치", "확인 방법"])
    for loc in r["fuzzy_cells"]:
        fz.append([loc, "규칙표 글자대응에서 원본 글자를 확인하고 비고를 확인완료로 변경"] )
    all_checks = wb.create_sheet("전체판정")
    check_fields = ["위치", "검증통과", "판정", "행 제목", "열 제목", "보고서", "원본", "이유"]
    all_checks.append(check_fields)
    for n, item in enumerate(r["checks"], start=2):
        all_checks.append([item[field] for field in check_fields])
        if not item["검증통과"]:
            all_checks.cell(n, 2).fill = yel
    for sh in (m, u, fz, all_checks):
        _style_table(sh, 1, wrap=("이유", "확인 방법", "행 제목", "열 제목"), max_width=70, judge=("판정", "구분"))
    _style_labels(ws, 1, 2)
    ws["B1"].font = Font(bold=True, size=12)
    ws.page_setup.orientation = "landscape"; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.column_dimensions["B"].width = min(120, max(40, max(_text_width(c.value) for c in ws["B"]) + 2))
    _diagnostic_text(wb)
    wb.save(out)
    return out


def collect(folder):
    def pick(pattern):
        return sorted(f for f in glob.glob(os.path.join(folder, pattern))
                      if not os.path.basename(f).startswith("~$") and not any(x in os.path.basename(f) for x in RESULT_TAGS))
    files = pick("*.xlsx")
    if len(files) < 2:                                   # xlsx가 모자랄 때만 매크로 통합문서(.xlsm)도 대상 (기존 xlsx 2개 폴더의 선택은 그대로)
        files = sorted(files + pick("*.xlsm"))
    return files


def _order(files):
    a, b = files[0], files[1]
    def score(raw_path, rep_path):
        try:
            raw = read_raw(raw_path)
            wb = safe_load(rep_path, data_only=True)
            cells = [c for ws in wb.worksheets for c in report_cells(ws, include_empty=True)]
            labs = {nl(l) for c in cells for l in c["row_labels"] + c["col_labels"]}
            return sum(1 for c in raw["text_cols"] for l in labs if l in raw["index"][c])
        except Exception:
            return -1
    return (a, b) if score(a, b) >= score(b, a) else (b, a)

# ---------------------------------------------------------------- 구조 파악·행 대조 (v22)
# 합계형이 아닌 보고서(원본 행을 옮겨 정리한 표, 글자 줄 원본을 표로 만든 것)를 거래(행) 단위로 맞춘다.
# 합계형 대조가 '행 단위로 옮긴 표'·'맞는 칸 0개'로 끝나거나, 보고서가 원본 행을 그대로 옮긴 표로 확실할 때만 쓴다.
_ROW_DATE_RE = re.compile(r"^\s*(\d{4})\s*[./\-년]\s*(\d{1,2})\s*[./\-월]\s*(\d{1,2})\s*일?(?:\s+\d{1,2}:\d{2}(?::\d{2}(?:\.\d+)?)?)?\s*$")
_ROW_DATE8_RE = re.compile(r"^(19|20)(\d{2})(\d{2})(\d{2})$")
_ROW_DATE_HEAD_RE = re.compile(r"^(\d{4})[./\-](\d{1,2})[./\-](\d{1,2})$")
_ROW_NUM_TOKEN_RE = re.compile(r"^[-+△▲]?\(?[₩$]?\s*[\d,]*\.?\d+\)?(?:-|원|KRW)?$", re.I)
_ROW_PAREN_NUM_RE = re.compile(r"\(\s*[-+]?\s*[₩$]?\s*([-+]?[\d,]*\.?\d+)\s*\)")
_ROW_UNIT_RE = re.compile(r"^([-+]?[\d,]*\.?\d+)\s*(주|건|개|매|좌|株)$")
_ROW_TOTAL_RE = re.compile(r"(^|\s)(합\s*계|총\s*계|소\s*계|누\s*계|total|subtotal|grand\s*total)\s*$", re.I)
ROW_MATCH_MIN = 0.5      # 행을 같은 거래로 볼 최소 일치 비율 (대응된 열 중)
ROW_FORMULA_MIN = 0.97   # 같은 행 계산식으로 인정할 최소 성립 비율
ROW_FORMULA_REL = 0.001  # 같은 행 계산식 허용 오차: 표시 자릿수 반 단위 또는 값의 0.1% 중 큰 쪽 (단수 반올림·소수 주식 단가 차이)
ROW_AUTO_MAX = 50000     # 합계형 대조가 된 경우에도 행 대조 여부를 확인할 최대 원본 행 수 (속도)
_ROW_CACHE = {}


def _row_date_str(y, mo, d):
    if 1 <= mo <= 12 and 1 <= d <= 31:
        return ("d", f"{y:04d}-{mo:02d}-{d:02d}")
    return None


def _row_norm(v):
    """행 대조용 값 정규화 → ('d', 'YYYY-MM-DD') / ('n', 숫자) / ('t', 글자) / None"""
    if v is None:
        return None
    if hasattr(v, "year"):
        return ("d", f"{v.year:04d}-{v.month:02d}-{v.day:02d}")
    if isinstance(v, bool):
        return ("t", str(v).lower())
    if isinstance(v, (int, float)):
        return ("n", float(v)) if math.isfinite(v) else None
    s = str(v).strip()
    if not s:
        return None
    if is_dash(s):
        return ("n", 0.0)                                 # '-' = 0 (보고서 관행)
    s_digits = s.replace(",", "").replace(" ", "")
    if re.fullmatch(r"-?\d+", s_digits) and len(s_digits.lstrip("-")) > 15:
        return ("t", s_digits)                            # 16자리 이상 정수(번호)는 실수로 바꾸면 끝자리가 사라짐 → 쉼표·공백 뺀 글자로 비교
    m = _ROW_DATE_RE.match(s)
    if m:
        r = _row_date_str(*(int(x) for x in m.groups()))
        if r:
            return r
    m = _ROW_DATE8_RE.match(s)
    if m:
        r = _row_date_str(int(m.group(1) + m.group(2)), int(m.group(3)), int(m.group(4)))
        if r:
            return r
    n = to_num(s)
    if n is not None:
        return ("n", n)
    if s.endswith("%") and to_num(s[:-1]) is not None:
        return ("n", to_num(s[:-1]) / 100)
    m = _ROW_UNIT_RE.match(s)
    if m and to_num(m.group(1)) is not None:
        return ("n", to_num(m.group(1)))
    return ("t", nl(s))


def _row_serial_date(x):
    """숫자 → 날짜 글자 (엑셀 일련번호 또는 YYYYMMDD) / None"""
    if not float(x).is_integer():
        return None
    x = int(x)
    if 20000 <= x <= 80000:
        d = datetime.date(1899, 12, 30) + datetime.timedelta(days=x)
        return f"{d.year:04d}-{d.month:02d}-{d.day:02d}"
    m = _ROW_DATE8_RE.match(str(x))
    if m:
        r = _row_date_str(int(m.group(1) + m.group(2)), int(m.group(3)), int(m.group(4)))
        return r[1] if r else None
    return None


def _row_eq(a, b):
    if a is None or b is None:
        return a is None and b is None
    if a[0] != b[0]:
        if {a[0], b[0]} == {"d", "n"}:                   # 날짜 ↔ 엑셀 일련번호·YYYYMMDD
            d, n = (a, b) if a[0] == "d" else (b, a)
            return _row_serial_date(n[1]) == d[1]
        return False
    if a[0] == "n":
        return abs(a[1] - b[1]) <= 1e-7 + abs(a[1]) * 5e-16  # 실수 계산 끝자리(1~2ulp)만 허용 — 1조 원대에서도 1원 차이는 잡음
    return a[1] == b[1]


def _row_key(v):
    if v is None:
        return None
    if v[0] == "n":
        return ("n", round(v[1], 6))
    return v


def _row_show(v):
    if v is None:
        return None
    if v[0] == "n":
        return int(v[1]) if float(v[1]).is_integer() and abs(v[1]) < 1e15 else v[1]
    return v[1]


def _row_text_cols(rows, width):
    """열마다 주된 값 종류 (d/n/t) — 글자 열의 숫자 모양 값은 글자로 비교 (코드 '0231' 등)"""
    kinds = []
    for j in range(width):
        cnt = {}
        for r in rows:
            v = r[j]
            if v is not None:
                cnt[v[0]] = cnt.get(v[0], 0) + 1
        kinds.append(max(cnt, key=cnt.get) if cnt else None)
    return kinds


def _row_apply_kinds(rows, raw_str, kinds):
    for r, s in zip(rows, raw_str):
        for j, k in enumerate(kinds):
            v = r[j]
            if v is None or k is None or v[0] == k or is_dash(s[j]):
                continue
            if k == "t":
                r[j] = ("t", nl(s[j]))
            elif k == "d" and v[0] == "n":
                d = _row_serial_date(v[1])
                if d:
                    r[j] = ("d", d)


def _flat_tables(path):
    """시트마다 제목 줄 있는 표 1개: [{sheet, head_row, heads, rows(정규화), str, rownums, notes}] (큰 표 먼저)"""
    wb = safe_load(path, data_only=True)
    out = []
    for ws in wb.worksheets:
        if ws.sheet_state != "visible":
            continue
        grid = fill_merged(ws)
        rows = [(i, list(r)) for i, r in enumerate(grid, start=1) if any(v not in (None, "") for v in r)]
        hmerged = {rng.min_row for rng in ws.merged_cells.ranges if rng.max_col > rng.min_col}
        cands = []
        for k, (i, r) in enumerate(rows[:15]):
            labs = [v for v in r if is_label(v) and not hasattr(v, "year")]
            vals = [v for v in r if v not in (None, "")]
            if len(labs) >= 2 and len(labs) >= 0.6 * len(vals) and k + 1 < len(rows):
                cands.append((len(vals), -(i in hmerged), -k, k))
        if not cands:
            continue
        h = max(cands)[3]                               # 칸이 가장 많은 제목 줄 (위의 회사명·작성자 줄, 병합된 큰 제목 줄은 건너뜀)
        head_row, head = rows[h]
        data_rows = rows[h + 1:]
        width = max([j + 1 for j, v in enumerate(head) if v not in (None, "")] +
                    [j + 1 for _, r in data_rows for j, v in enumerate(r) if v not in (None, "")])
        heads, notes = [], []
        for j in range(width):
            v = head[j] if j < len(head) else None
            name = str(v).strip() if v not in (None, "") else f"({get_column_letter(j + 1)}열)"
            while name in heads:
                name += "_"
            heads.append(name)
        data, strs, nums, totals = [], [], [], 0
        for i, r in data_rows:
            r = (r + [None] * width)[:width]
            if any(isinstance(v, str) and _ROW_TOTAL_RE.search(v.strip()) for v in r):
                totals += 1                               # 합계 줄은 행 대조 대상 아님
                continue
            data.append([_row_norm(v) for v in r])
            strs.append(["" if v is None else str(v).strip() for v in r])
            nums.append(i)
        if totals:
            notes.append(f"'{ws.title}' 합계·소계 줄 {totals}개는 행 대조에서 제외")
        kinds = _row_text_cols(data, width)
        _row_apply_kinds(data, strs, kinds)
        out.append({"sheet": ws.title, "head_row": head_row, "heads": heads, "rows": data, "str": strs, "rownums": nums,
                    "notes": notes, "kinds": kinds})
    out.sort(key=lambda t: -len(t["rows"]) * len(t["heads"]))
    if not out:
        raise RuntimeError(f"{os.path.basename(path)}: 제목 줄이 있는 표를 찾지 못함")
    return out


def _text_source(path):
    """글자 줄 원본이면 (시트명, [(엑셀 행, 줄)]) — 아니면 None"""
    wb = safe_load(path, data_only=True)
    best = None
    for ws in wb.worksheets:
        cols = {}
        for i, r in enumerate(ws.iter_rows(values_only=True), start=1):
            for j, v in enumerate(r):
                if isinstance(v, str) and v.strip():
                    cols.setdefault(j, []).append((i, v))
        for j, items in cols.items():
            long = [s for _, s in items if len(s.split()) >= 5 and sum(1 for t in s.split() if _ROW_NUM_TOKEN_RE.match(t)) >= 3]
            if len(long) >= 5 and len(long) >= 0.3 * len(items):
                if best is None or len(long) > best[0]:
                    lines = [(i, line) for i, s in items for line in s.splitlines() if line.strip()]
                    best = (len(long), ws.title, lines)
    return None if best is None else (best[1], best[2])


def _tok_num(t):
    if not _ROW_NUM_TOKEN_RE.match(t):
        return None
    return to_num(t.replace("$", "").replace("₩", ""))


def _text_roles(segment):
    """'판매 아마존(US0231351067)' → 첫 낱말 / 이름 / 마지막 괄호 / 첫 낱말 뒤 전체"""
    seg = segment.strip()
    parts = seg.split(None, 1)
    w0 = parts[0] if parts else ""
    rest = parts[1] if len(parts) > 1 else ""
    paren, body = "", rest
    m = re.search(r"\(([^()]*)\)\s*$", rest)
    if m:
        paren, body = m.group(1), rest[:m.start()]
    return {"text:w0": w0, "text:body": body, "text:paren": paren, "text:rest": rest, "text": seg}


def _text_records(lines):
    """날짜로 시작하는 줄 = 거래 1건. 숫자는 줄 끝에서부터 n1, n2 … (이름에 숫자가 있어도 위치가 밀리지 않게).
    바로 아래 괄호 숫자 줄('($ 464.20)')은 p1, p2 …, 숫자 없는 짧은 줄은 이름이 넘어간 것으로 보고 이어 붙임."""
    recs = []
    for row, line in lines:
        toks = line.split()
        if not toks:
            continue
        m = _ROW_DATE_HEAD_RE.match(toks[0])
        if m:
            date = _row_date_str(*(int(x) for x in m.groups()))
            if not date:
                continue
            recs.append({"_row": row, "_line": line.strip(), "_toks": toks[1:], "date": date, "_cont": []})
            continue
        if not recs:
            continue
        last = recs[-1]
        s = line.strip()
        if s.startswith("(") and "p1" not in last:
            vals = [to_num(x) for x in _ROW_PAREN_NUM_RE.findall(s)]
            vals = [v for v in vals if v is not None]
            if vals and not _ROW_PAREN_NUM_RE.sub("", s).strip():
                for n, v in enumerate(reversed(vals), start=1):
                    last[f"p{n}"] = ("n", v)
            continue
        if any(len(t) > 30 for t in toks) or any(_ROW_DATE_HEAD_RE.match(t) for t in toks):
            continue                                       # 쪽 머리글·잡음 줄
        if not any(_tok_num(t) is not None for t in last["_toks"]) and any(_tok_num(t) is not None for t in toks):
            last["_toks"] = last["_toks"] + toks           # 날짜 줄에 숫자가 없고 다음 줄에 숫자 → 같은 거래
            last["_line"] += " " + s
        elif len(toks) <= 4 and not any(_tok_num(t) is not None for t in toks) and "p1" not in last and not last["_cont"]:
            last["_cont"].append(s)                        # 이름이 다음 줄로 넘어감
    # 끝에서부터 숫자 읽기 → 거래마다 개수가 가장 흔한 값을 넘으면 앞쪽 숫자는 이름의 일부 ('KODEX 200')
    counts = []
    for r in recs:
        k = len(r["_toks"])
        while k > 0 and _tok_num(r["_toks"][k - 1]) is not None:
            k -= 1
        r["_k"] = k
        counts.append(len(r["_toks"]) - k)
    typical = max(set(counts), key=counts.count) if counts else 0
    for r in recs:
        toks, k = r["_toks"], r["_k"]
        if len(toks) - k > typical > 0:
            k = len(toks) - typical
        nums = [_tok_num(t) for t in toks[k:]]
        seg = " ".join(toks[:k])
        if r["_cont"]:
            sep = "" if seg.count("(") > seg.count(")") else " "
            seg = seg + sep + sep.join(r["_cont"])
        for key, val in _text_roles(seg).items():
            r[key] = ("t", nl(val)) if val else None
        r["_text"] = seg
        r["date"] = r["date"]
        for n, v in enumerate(reversed(nums), start=1):
            r[f"n{n}"] = ("n", v)
        for key in [k2 for k2, v in r.items() if v is None]:
            del r[key]
        for key in ("_toks", "_k", "_cont"):
            r.pop(key, None)
    return recs, typical


def _text_labels(lines, rep_heads, typical):
    """원본 글자 안의 제목 줄(보고서 열 이름 3개 이상)로 필드 이름 붙이기"""
    want = {nl2(h) for h in rep_heads} | {nl(h) for h in rep_heads}
    counts = {}
    for _, line in lines:
        toks = line.split()
        if not toks or _ROW_DATE_HEAD_RE.match(toks[0]):
            continue
        if sum(1 for t in toks if nl(t) in want or nl2(t) in want) >= 3:
            counts[tuple(toks)] = counts.get(tuple(toks), 0) + 1
    if not counts:
        return {}
    head = list(max(counts, key=counts.get))
    labels = {"date": head[0]}
    if len(head) >= typical + 1:
        for k, t in enumerate(reversed(head[len(head) - typical:]), start=1):
            labels[f"n{k}"] = t
        mid = head[1:len(head) - typical]
        if mid:
            labels["text:w0"] = mid[0]
            last = mid[-1] if len(mid) > 1 else ""
            m = re.search(r"\(([^()]*)\)\s*$", last)
            if m:
                labels["text:paren"] = m.group(1)
                labels["text:body"] = last[:m.start()]
            elif last:
                labels["text:body"] = last
    return labels


def _text_derived_labels(recs, labels):
    """이름 없는 숫자(괄호 줄 등)는 이름 있는 두 숫자 필드와의 관계(예: 잔액 ÷ 환율)가 거의 모든 거래에서 성립하면 그 관계를 이름으로"""
    named = [f for f in labels if f.startswith("n")]
    unnamed = sorted({k for r in recs for k in r if k.startswith("p") and k not in labels})
    for f in unnamed:
        best = None
        for a in named:
            for b in named:
                if a == b:
                    continue
                n = ok = nz = 0
                for r in recs:
                    x, y, z = r.get(a), r.get(b), r.get(f)
                    if x is None or y is None or z is None or not y[1]:
                        continue
                    n += 1
                    if abs(x[1] / y[1] - z[1]) <= max(0.01, abs(z[1]) * 0.002):
                        ok += 1
                        nz += abs(z[1]) > 0.01
                if n >= 10 and ok >= 0.97 * n and nz >= 5 and (best is None or (nz, ok) > best[:2]):
                    best = (nz, ok, a, b)
        if best:
            labels[f] = f"{labels[best[2]]}÷{labels[best[3]]}"


def _row_fields(recs):
    keys = []
    for r in recs:
        for k in r:
            if not k.startswith("_") and k not in keys:
                keys.append(k)
    return keys


def _row_name_sim(head, label):
    a, b = nl(head), nl(label or "")
    if not b:
        return 0
    if a == b:
        return 2
    a2, b2 = nl2(head), nl2(label)
    if a2 and (a2 == b2 or a2 in b or b2 in a):
        return 1
    return 0


def _row_kind_ok(rk, fk):
    return rk == fk or {rk, fk} == {"d", "n"}


def _row_initial_map(T, P):
    """값이 겹치는 비율 + 이름으로 처음 열 대응 (열 이름이 달라도 됨)"""
    heads, rows, recs, fields, labels = T["heads"], T["rows"], P["raw_recs"], P["raw_fields"], P["labels"]
    raw_sets = {f: {_row_key(r.get(f)) for r in recs if r.get(f) is not None} for f in fields}
    scored = []
    for j, h in enumerate(heads):
        keyed = [_row_key(r[j]) for r in rows if r[j] is not None]
        if not keyed:
            continue
        for f in fields:
            if f == "text" and P["raw_kind"] == "text":
                continue
            hit = sum(1 for v in keyed if v in raw_sets[f]) / len(keyed)
            sim = _row_name_sim(h, labels.get(f, f))
            if hit >= 0.5 or (sim == 2 and hit >= 0.3):
                scored.append((hit + 0.2 * sim, j, f))
    cmap, used = {}, set()
    for sc, j, f in sorted(scored, key=lambda x: -x[0]):
        if j in cmap or f in used:
            continue
        cmap[j] = f
        used.add(f)
    return cmap


def _row_refine(T, P, rep_to_raw, use_names=True):
    """맞춘 행 짝에서 열마다 값이 가장 자주 같은 원본 필드로 대응을 다시 정함. 이름이 같은 필드는 값이 달라도 우선
    (값만 보고 엉뚱한 필드에 붙여 오류를 '일치'로 놓치지 않게). 대응을 못 찾으면 처음 대응을 유지."""
    heads, rows, recs, labels = T["heads"], T["rows"], P["raw_recs"], P["labels"]
    pairs = list(rep_to_raw.items())
    if len(pairs) < 3:
        return T["cmap"], []
    fields = [f for f in P["raw_fields"] if not (f == "text" and P["raw_kind"] == "text")]
    fkind = {}
    for f in fields:
        cnt = {}
        for _, k in pairs[:300]:
            v = recs[k].get(f)
            if v is not None:
                cnt[v[0]] = cnt.get(v[0], 0) + 1
        fkind[f] = max(cnt, key=cnt.get) if cnt else None
    claims, ranked = [], {}
    for j, h in enumerate(heads):
        filled = [(i, k) for i, k in pairs if rows[i][j] is not None]
        if not filled:
            continue
        rk = T["kinds"][j]
        cands = []
        for f in fields:
            eq = sum(1 for i, k in filled if _row_eq(rows[i][j], recs[k].get(f))) / len(filled)
            sim = _row_name_sim(h, labels.get(f, f)) if use_names is True or (use_names == "derived" and f.startswith("p")) else 0
            if (sim and fkind[f] and _row_kind_ok(rk, fkind[f])) or eq >= 0.5:
                cands.append((sim, eq, f))
        if cands:
            cands.sort(key=lambda x: (-x[0], -x[1]))
            ranked[j] = cands
            claims.append((cands[0][0], cands[0][1], j))
    cmap, used, doubtful = {}, set(), []
    for sim, eq, j in sorted(claims, key=lambda x: (-x[0], -x[1])):
        for sim2, eq2, f in ranked[j]:
            if f in used:
                continue
            cmap[j] = f
            used.add(f)
            if sim2 and eq2 < 0.5 and any(e >= 0.9 and g not in used for s, e, g in ranked[j]):
                doubtful.append((j, f, [g for s, e, g in ranked[j] if e >= 0.9 and g not in used]))
            break
    for j, f in T["cmap"].items():                       # 처음엔 값이 겹쳤는데 행 짝에서 못 찾은 열 → 버리지 않고 유지 (차이로 드러나게)
        if j not in cmap and f not in used:
            cmap[j] = f
            used.add(f)
    return cmap, doubtful


def _row_load_raw(raw_path, rep_heads):
    text = _text_source(raw_path)
    table = None
    try:
        table = _flat_tables(raw_path)[0]
    except Exception:
        pass
    out = []
    if text is not None:
        recs, typical = _text_records(text[1])
        labels = _text_labels(text[1], rep_heads, typical)
        _text_derived_labels(recs, labels)
        out.append({"raw_kind": "text", "raw_sheet": text[0], "raw_recs": recs, "raw_fields": _row_fields(recs), "labels": labels})
    if table is not None:
        recs = []
        for i, r in zip(table["rownums"], table["rows"]):
            rec = {"_row": i}
            for h, v in zip(table["heads"], r):
                if v is not None:
                    rec[h] = v
            recs.append(rec)
        out.append({"raw_kind": "table", "raw_sheet": table["sheet"], "raw_recs": recs, "raw_fields": table["heads"], "labels": {},
                    "raw_kinds": dict(zip(table["heads"], table["kinds"]))})
    if not out:
        raise RuntimeError(f"{os.path.basename(raw_path)}: 원본에서 표나 거래 줄을 찾지 못함")
    return out


def _row_match(T, P):
    """보고서 행 ↔ 원본 거래 1:1 맞추기. 대응 열 값이 많이 같은 짝부터 확정 (같은 점수면 순서가 가까운 짝).
    후보는 값이 드문 두 열에서만 찾아 큰 표에서도 빠르게."""
    cmap, rows, recs = T["cmap"], T["rows"], P["raw_recs"]
    index = {}
    for k, r in enumerate(recs):
        for j, f in cmap.items():
            v = r.get(f)
            if v is not None:
                index.setdefault((f, _row_key(v)), []).append(k)
    pairs = []
    nrep, nraw = max(len(rows), 1), max(len(recs), 1)
    distinct = {}
    for (f, _v) in index:
        distinct[f] = distinct.get(f, 0) + 1
    for i, row in enumerate(rows):
        filled = [j for j in cmap if row[j] is not None]
        if not filled:
            continue
        posts = [(j, index.get((cmap[j], _row_key(row[j])), [])) for j in filled]
        lists = sorted((lst for _, lst in posts if lst), key=len)
        if not lists:
            continue
        cand = set(lists[0])
        if len(lists) > 1:
            cand.update(lists[1])
        for k in cand:
            rec = recs[k]
            # 칸마다 무게 = 원본 후보 값이 드문 정도 (0·'-'·같은 날짜처럼 흔한 값만 같은 행끼리 짝짓지 않게,
            # 바뀐 값이 우연히 다른 행 값과 같아도 짝을 놓치지 않게 — 무게는 보고서 값이 아니라 원본 후보 값 기준)
            score = total = 0.0
            neq = 0
            for j in filled:
                v = rec.get(cmap[j])
                n = len(index.get((cmap[j], _row_key(v)), ())) if v is not None else 0
                w = math.log(1 + (nraw / n if n else distinct.get(cmap[j], 1)))
                total += w
                if _row_eq(row[j], v):
                    score += w
                    neq += 1
            if neq >= 2 and score >= ROW_MATCH_MIN * total:
                pairs.append((-score, -neq, abs(i / nrep - k / nraw), i, k))
    pairs.sort()
    rep_to_raw, raw_used = {}, set()
    for _, _, _, i, k in pairs:
        if i in rep_to_raw or k in raw_used:
            continue
        rep_to_raw[i] = k
        raw_used.add(k)
    return rep_to_raw, raw_used


def _row_formulas(heads, rows):
    """같은 행의 열 계산식 (거의 모든 행에서 맞는 식만). [(대상 열, 식 글자, 계산 함수, a, b, 소수 자리)]"""
    sample = rows[:5000]
    numeric = []
    for j in range(len(heads)):
        vals = [r[j] for r in sample if r[j] is not None]
        if len(vals) >= 10 and sum(1 for v in vals if v[0] == "n") >= 0.9 * len(vals):
            numeric.append(j)
    if len(numeric) < 3:
        return []
    col = {j: [r[j][1] if r[j] is not None and r[j][0] == "n" else None for r in sample] for j in numeric}
    def decimals(vals):
        vals = [v for v in vals if v is not None]
        for d in range(0, 7):
            if sum(1 for v in vals if abs(v * 10 ** d - round(v * 10 ** d)) < 1e-6) >= 0.95 * len(vals):
                return d
        return 7
    ops = [("*", lambda a, b: a * b, True), ("/", lambda a, b: a / b if b else None, False),
           ("+", lambda a, b: a + b, True), ("-", lambda a, b: a - b, False)]
    found = []
    for t in numeric:
        tv = col[t]
        if len({x for x in tv if x is not None}) < 5:
            continue                                     # 거의 상수인 열(0 등)은 식 대상 아님
        d = decimals(tv)
        tol = 0.5 * 10 ** (-d) + 1e-9
        best = None
        for a in numeric:
            for b in numeric:
                if len({a, b, t}) < 3:
                    continue
                av, bv = col[a], col[b]
                if len({x for x in bv if x is not None}) <= 1 or len({x for x in av if x is not None}) <= 1:
                    continue
                for sym, fn, comm in ops:
                    if comm and a > b:
                        continue
                    n = ok = big = 0
                    for x, y, z in zip(av, bv, tv):
                        if x is None or y is None or z is None:
                            continue
                        r = fn(x, y)
                        if r is None:
                            continue
                        n += 1
                        if abs(r - z) <= max(tol, abs(z) * ROW_FORMULA_REL):
                            ok += 1
                            big += abs(z) > 10 * tol
                    if n >= 10 and ok >= ROW_FORMULA_MIN * n and ok >= 0.5 * len(sample) and big >= 0.5 * ok:
                        cand = (ok / n, -("*/+-".index(sym)), a, b, sym, fn)
                        if best is None or cand[:2] > best[:2]:
                            best = cand
        if best:
            _, _, a, b, sym, fn = best
            found.append((t, f"{heads[t]} = {heads[a]} {sym} {heads[b]}", fn, a, b, d, sym))
    keep = {}
    for item in found:                                   # 같은 세 열의 식은 하나만: 소수 자리 적은 열 → 나눗셈 식 우선
        key = frozenset((item[0], item[3], item[4]))
        rank = lambda x: (x[5], "/-*+".index(x[6]))
        if key not in keep or rank(item) < rank(keep[key]):
            keep[key] = item
    return [x[:6] for x in found if keep[frozenset((x[0], x[3], x[4]))] is x]


def _row_keys(T):
    """보고서 행을 하나로 가리키는 최소 열 조합 (표시용): 날짜부터, 구분력이 큰 열을 차례로. 서술형 글자 열은 제외"""
    rows, total = T["rows"], max(len(T["rows"]), 1)
    def avg_len(j):
        vals = [len(r[j][1]) for r in rows if r[j] is not None and r[j][0] == "t"]
        return sum(vals) / len(vals) if vals else 0
    pool = [j for j in T["cmap"] if not (T["kinds"][j] == "t" and (len({r[j] for r in rows}) > 0.9 * total or avg_len(j) > 12))]
    rank = {"d": 0, "n": 1, "t": 2}
    chosen = [j for j in pool if T["kinds"][j] == "d"][:1]
    texts = [j for j in pool if T["kinds"][j] == "t"]
    if texts:                                           # 날짜 다음은 이름·코드 같은 짧은 글자 열 (거래 대상)
        def digit_share(j):
            vals = [r[j][1] for r in rows if r[j] is not None and r[j][0] == "t"]
            return sum(1 for v in vals if re.search(r"\d", v)) / max(len(vals), 1)
        # 거래 대상(이름) 열: 숫자 없는 글자 열 중 값 종류가 가장 많은 열, 같으면 왼쪽 열
        chosen.append(max(texts, key=lambda j: (digit_share(j) < 0.5, len({r[j] for r in rows}), -j)))
    while len(chosen) < 4:
        cur = len({tuple(r[c] for c in chosen) for r in rows})
        if cur / total >= 0.99:
            break
        best = None
        for j in pool:
            if j in chosen:
                continue
            gain = len({tuple(r[c] for c in chosen + [j]) for r in rows})
            cand = (gain, -avg_len(j), -rank.get(T["kinds"][j], 3), j)
            if best is None or cand > best:
                best = cand
        if best is None or best[0] <= cur:
            break
        chosen.append(best[3])
    return [T["heads"][j] for j in chosen]


def _row_fit(T, P):
    """보고서 표 하나를 원본 하나에 맞춤 → T에 cmap·짝·의심 열 기록, 맞춘 비율 반환"""
    T["cmap"] = _row_initial_map(T, P)
    T["doubtful"] = []
    rep_to_raw = {}
    for _ in range(3):
        if not T["cmap"]:
            break
        rep_to_raw, _ = _row_match(T, P)
        cmap2, doubtful = _row_refine(T, P, rep_to_raw)
        # 글자 원본 제목 줄에서 붙인 이름(n1, n2 …)이 한 칸 밀린 것으로 보이면 (제목 낱말에 띄어쓰기 등) 그 이름은 쓰지 않음.
        # 값 관계로 붙인 이름(p…, 예: 잔액÷환율)은 믿을 수 있으므로 유지 → 열을 잘못 옮긴 오류는 계속 찾음
        if any(f.startswith("n") and any(g.startswith("n") for g in alts) for _, f, alts in doubtful):
            cmap2, doubtful = _row_refine(T, P, rep_to_raw, use_names="derived")
        T["doubtful"] = [j for j, _, _ in doubtful]
        if cmap2 == T["cmap"] or not cmap2:
            break
        T["cmap"] = cmap2
    T["pairs"], T["raw_used"] = _row_match(T, P) if T["cmap"] else ({}, set())
    return len(T["pairs"]) / max(len(T["rows"]), 1)


def _row_load(raw_path, report_path):
    """(보고서 표 목록, 원본) — 표마다 cmap·짝 포함. 같은 파일이면 다시 읽지 않음."""
    try:
        key = tuple((p, os.path.getmtime(p), os.path.getsize(p)) for p in (raw_path, report_path))
    except OSError:
        key = None
    if key is not None and key in _ROW_CACHE:
        return _ROW_CACHE[key]
    tables = _flat_tables(report_path)
    sources = _row_load_raw(raw_path, tables[0]["heads"])
    best = None
    for P in sources:                                    # 글자 줄로 읽은 것과 표로 읽은 것 중 더 잘 맞는 쪽
        T = dict(tables[0])
        rate = _row_fit(T, P)
        if best is None or rate > best[0]:
            best = (rate, P)
    P = best[1]
    for T in tables:
        _row_fit(T, P)
    res = (tables, P)
    if key is not None:
        _ROW_CACHE.clear()
        _ROW_CACHE[key] = res
    return res


def _row_is_rows(T, P):
    n = len(T["rows"])
    rate = len(T["pairs"]) / max(n, 1)
    num_eq = sum(1 for j in T["cmap"] if T["kinds"][j] == "n")
    if P["raw_kind"] == "text":
        return n >= 5 and rate >= 0.5 and len(T["cmap"]) >= 3 and num_eq >= 1
    return n >= 5 and rate >= 0.5 and len(T["cmap"]) >= 3


def _row_sheets(tables, P):
    """행 대조할 시트: 행 정리 표로 판정된 시트 + (그런 시트가 있을 때) 행이 1~4개뿐이지만 거의 모든 행이 원본 거래와 맞는 작은 분할 시트"""
    big = [T for T in tables if _row_is_rows(T, P)]
    if not big:
        return []
    return [T for T in tables if T in big or (1 <= len(T["rows"]) < 5 and len(T["cmap"]) >= 3 and len(T["pairs"]) >= 0.8 * len(T["rows"]))]


def _header_only(ws, heads=()):
    """(제목 1~2줄 +) 열 제목 줄 하나만 있는 빈 양식: 마지막 줄이 짧은 글자 3칸 이상이고 그 가운데 2칸 이상이 다른 표의 열 제목과 같으며,
    그 위는 1칸짜리 제목 줄 2개 이하 (메모·조정 내용 줄은 빈 양식이 아님)"""
    rows = [[v for v in r if v not in (None, "")] for r in ws.iter_rows(values_only=True)]
    rows = [r for r in rows if r]
    if not rows or len(rows) > 3 or any(len(r) != 1 for r in rows[:-1]):
        return False
    head = rows[-1]
    return (len(head) >= 3 and all(isinstance(v, str) and len(v.strip()) <= 20 and to_num(v) is None for v in head)
            and sum(1 for v in head if nl(v) in heads) >= 2)


def _content_sheets_without_table(path, tables):
    """제목 줄 있는 표가 없어 행 대조 대상에서 빠졌지만 내용이 있는 보이는 시트 (메모 시트 등)"""
    have = {t["sheet"] for t in tables}
    heads = {nl(str(h)) for t in tables for h in t["heads"] if h not in (None, "")}
    try:
        wb = safe_load(path, data_only=True)
    except Exception:
        return []
    out = []
    for ws in wb.worksheets:
        if ws.sheet_state != "visible" or ws.title in have:
            continue
        rows = [k for k in (sum(1 for v in r if v not in (None, "")) for r in ws.iter_rows(values_only=True)) if k]
        if not rows or _header_only(ws, heads):
            continue                                       # 빈 시트, (제목 1~2줄 +) 열 제목 줄만 있는 빈 양식은 대조할 내용 없음
        out.append(ws.title)
    return out


def analyze_structure(raw_path, report_path):
    """원본·보고서 구조 파악: 원본 종류(표/글자), 보고서 종류(합계형/행 정리), 열 대응, 키 열, 같은 행 계산식"""
    try:
        tables, P = _row_load(raw_path, report_path)
    except Exception as exc:
        return {"raw_kind": "table", "report_kind": "aggregate", "column_map": {}, "keys": [], "row_formulas": [], "reason": str(exc)}
    T = tables[0]
    lab = P["labels"]
    column_map = {T["heads"][j]: (f"{lab[f]}({f})" if f in lab else f) for j, f in T["cmap"].items()}
    is_rows = _row_is_rows(T, P)
    formulas = _row_formulas(T["heads"], T["rows"]) if is_rows else []
    return {"raw_kind": P["raw_kind"], "report_kind": "rows" if is_rows else "aggregate", "column_map": column_map,
            "keys": _row_keys(T) if is_rows else [], "row_formulas": [f[1] for f in formulas],
            "matched_rate": round(len(T["pairs"]) / max(len(T["rows"]), 1), 4),
            "report_sheet": T["sheet"], "raw_sheet": P["raw_sheet"], "report_rows": len(T["rows"]), "raw_records": len(P["raw_recs"]),
            "sheets": [t["sheet"] for t in _row_sheets(tables, P)],
            "skipped_sheets": [t["sheet"] for t in tables if t not in _row_sheets(tables, P)] + _content_sheets_without_table(report_path, tables),
            "unmapped": [h for j, h in enumerate(T["heads"]) if j not in T["cmap"]]}


def _row_reason_field(P, f):
    return P["labels"].get(f, f)


def reconcile_rows(raw_path, report_path):
    """행 대조 결과 목록 (판정: 일치 / 값 차이 / 원본에 없음 / 보고서에 없음 / 확인 필요). 보고서 시트가 여럿이면 시트마다."""
    tables, P = _row_load(raw_path, report_path)
    recs = P["raw_recs"]
    use = _row_sheets(tables, P) or [tables[0]]
    if not use[0]["cmap"]:
        raise RuntimeError("보고서 열과 값이 겹치는 원본 열이 없음 — 원본·보고서 짝인지 확인하세요")
    out = []
    # 시트 간 관계 — 시트를 '층'으로 묶음. 한 층 = 원본을 서로 거의 겹치지 않게 나눠 담은 시트들(월별·법인별 분할, 또는 '전체' 시트 하나).
    # 큰 시트부터 차례로, 이미 만든 층 가운데 겹침이 가장 적은 층에 넣되 그 시트 행의 80% 이상이 겹치면 새 층(같은 내용의 다른 보기:
    # 원화·달러 기준, '전체'와 월별 분할, 발췌). 층 안에서 두 시트에 함께 있는 거래는 시트 순서상 뒤 시트 행을 '중복 기재'.
    # 보고서 전체 합(U)의 절반 이상을 담은 층은 완결본으로 보고 U에 있는데 그 층에 없는 거래를 '보고서에 없음',
    # 절반 미만인 층은 발췌로 보고 알림 1건. 원본 중 어느 시트에도 없는 거래는 '보고서에 없음 (어느 시트에도 없음)'.
    n_raw = len(recs)
    sets = {id(T): set(T["pairs"].values()) for T in use}
    try:
        _wb = load_workbook(report_path, read_only=True); order = {n: i for i, n in enumerate(_wb.sheetnames)}; _wb.close()
    except Exception:
        order = {}
    pos = {id(T): (order.get(T["sheet"], 10 ** 6), use.index(T)) for T in use}
    def _layers(limit, pool=None):
        out_ = []
        for T in sorted(use if pool is None else pool, key=lambda T: (-len(sets[id(T)]), pos[id(T)])):
            S_ = sets[id(T)]
            best = None
            for L in out_:
                ov = len(S_ & L["union"])
                if ov < limit * max(len(S_), 1) and (best is None or ov < best[0]):
                    best = (ov, L)
            if best is None:
                out_.append({"sheets": [T], "union": set(S_)})
            else:
                best[1]["sheets"].append(T); best[1]["union"] |= S_
        return out_
    U_all = set().union(*sets.values()) if sets else set()
    # 먼저 겹침이 거의 없는(2% 미만) 층으로 나눠 봄 — 월별·법인별처럼 같은 원본을 두 가지로 나눈 보고서. 완결 층(U의 80% 이상)이 있고
    # 나머지 층이 모두 완결 층 안에 80% 이상 들어가는 시트 1개(발췌)면 그대로 씀. (완결 층 = U의 80% 이상) 아니면 겹침 80% 미만은 같은 층(중복 기재)으로 봄
    layers = _layers(0.02)
    full_ = [L for L in layers if len(L["union"]) >= 0.8 * len(U_all)]   # 80% 이상 = 완결본(빠진 거래 보고), 미만 = 발췌·필터 보기(알림)
    cover = set().union(*(L["union"] for L in full_)) if full_ else set()
    if not (full_ and all(L in full_ or (len(L["sheets"]) == 1 and len(L["union"] & cover) >= 0.8 * len(L["union"])) for L in layers)):
        layers = _layers(0.8)
        full_ = None
    else:                                                   # 남은 발췌 시트끼리 다시 묶어 봄 (달러 상·하처럼 나눠 담은 다른 보기 + 그 안의 중복)
        extra = [L["sheets"][0] for L in layers if L not in full_]
        rest_ = []
        for X in sorted(extra, key=lambda T: (-len(sets[id(T)]), pos[id(T)])):   # 완결 층에 조금만 겹치는 작은 분할 시트(중복 포함) → 그 층으로
            S_ = sets[id(X)]
            cand = [(len(S_ & L["union"]), L) for L in full_ if len(S_ & L["union"]) < 0.5 * max(len(S_), 1)]   # 새 거래가 더 많을 때만
            if cand:
                L = min(cand, key=lambda t: t[0])[1]
                L["sheets"].append(X); L["union"] |= S_
            else:
                rest_.append(X)
        regroup = [L for L in _layers(0.8, rest_) if len(L["union"]) >= 0.8 * len(U_all)] if len(rest_) > 1 else []
        grouped = {id(T) for L in regroup for T in L["sheets"]}
        layers = full_ + regroup + [{"sheets": [T], "union": set(sets[id(T)])} for T in rest_ if id(T) not in grouped]
        full_ = full_ + regroup                             # 다시 묶어 완결본이 된 보기만 (아니면 시트마다 발췌 알림)
    dup, notes, inner_missing, dup_other = set(), [], [], {}
    for L in layers:
        where = {}
        for T in L["sheets"]:
            for ii, k in T["pairs"].items():
                where.setdefault(k, []).append((T, ii))
        multi = {k for k, occ in where.items() if len({id(T) for T, _ in occ}) > 1}
        profiles = {}
        def _feat(v_):                                      # 원본 값 ('d', 'YYYY-MM-DD') → 연월
            if isinstance(v_, tuple) and v_ and v_[0] == "d":
                return ("ym", str(v_[1])[:7])
            return (v_.year, v_.month) if isinstance(v_, (datetime.date, datetime.datetime)) else v_
        def _profile(T):
            """시트마다 한 번: 여러 시트에 있는 거래를 뺀 원본 위치 범위, 열마다 값 개수 (값 종류가 적은 열인지)"""
            if id(T) not in profiles:
                ks = [j for j in T["pairs"].values() if j not in multi]
                cnt = {f: Counter(_feat(recs[j].get(f)) for j in ks) for f in P["raw_fields"]}
                profiles[id(T)] = {"lo": min(ks) if ks else None, "hi": max(ks) if ks else None, "n": len(ks), "cnt": cnt,
                                   "cat": {f: bool(ks) and len(cnt[f]) * 3 <= len(ks) for f in P["raw_fields"]}}
            return profiles[id(T)]
        for k, occ in where.items():                        # 같은 층의 두 시트에 있는 거래: 한 곳만 '중복 기재'로 표시
            if len({id(T) for T, _ in occ}) < 2:
                continue
            # 원래 자리 판단 (작을수록 제자리):
            # ① 그 시트의 다른 거래(여러 시트에 있는 거래는 빼고)의 원본 위치 범위 안에 있는지 — 월별·페이지 분할
            # ② 범위로 못 가리면: 값 종류가 적은 열(부서·법인, 일자는 연월)에서 그 시트 거래들과 같은 값 비율이 다른 시트보다 얼마나 높은지
            #    (모든 거래가 같은 값인 열 — 통화·회계연도 — 은 차이가 0이라 영향 없음)
            # ③ 범위와의 거리 ④ 바로 위·아래 줄의 원본 위치 ⑤ 큰 시트·앞 시트
            rec = recs[k]
            shares = {}
            for T, ii in occ:
                prof = _profile(T)
                shares[(id(T), ii)] = {f: (prof["cnt"][f].get(_feat(rec.get(f)), 0) / prof["n"]) if prof["n"] and prof["cat"][f] else None
                                       for f in P["raw_fields"]}
            def misfit(o):
                T, ii = o
                prof = _profile(T)
                lo, hi = prof["lo"], prof["hi"]
                dist = 0 if lo is None or lo <= k <= hi else min(abs(k - lo), abs(k - hi))
                adv = 0.0
                for f in P["raw_fields"]:
                    mine = shares[(id(T), ii)][f]
                    if mine is None:
                        continue
                    rest = [v_[f] for key_, v_ in shares.items() if key_ != (id(T), ii) and v_[f] is not None]
                    if rest:
                        adv = max(adv, mine - min(rest))
                nb = [T["pairs"][j] for j in (ii - 1, ii + 1) if j in T["pairs"]]
                strong = adv >= 0.5                          # 부서·법인처럼 뚜렷한 값 차이는 범위보다 먼저
                return (0 if strong else 1, -round(adv, 1) if strong else 0, 0 if dist == 0 else 1, -round(adv, 1), dist,
                        min((abs(j - k) for j in nb), default=10 ** 9))
            big = max(len(sets[id(T)]) for T, _ in occ)
            keep = min(occ, key=lambda o: misfit(o) + (0 if len(sets[id(o[0])]) >= 0.8 * big else 1, pos[id(o[0])], o[1]))
            for T, ii in occ:
                if (T, ii) != keep and T is not keep[0]:
                    dup.add((T["sheet"], ii)); dup_other[(T["sheet"], ii)] = keep[0]["sheet"]
        names = " + ".join(T["sheet"] for T in sorted(L["sheets"], key=lambda T: pos[id(T)]))
        if len(layers) == 1:
            continue
        if (L in full_) if full_ is not None else len(L["union"]) >= 0.5 * len(U_all):
            inner_missing.append((names, U_all - L["union"]))
        else:
            notes.append((names, len(L["union"]), len(U_all)))
    all_used = U_all
    for T in use:
        heads, rows, cmap = T["heads"], T["rows"], T["cmap"]
        sheet = T["sheet"]
        formulas = _row_formulas(heads, rows)
        matched_rows = {tuple(rows[i][j] for j in sorted(cmap)): recs[k]["_row"] for i, k in T["pairs"].items()}
        for j, h in enumerate(heads):
            if j in cmap:
                continue
            empty = all(r[j] is None for r in rows)
            out.append({"판정": "확인 필요", "시트": sheet, "보고서 행": None, "원본 행": None, "열": h, "보고서 값": None, "원본 값": None,
                        "이유": "열 전체가 비어 있음 (수식이면 엑셀에서 열어 저장 후 다시 실행)" if empty else
                                "원본에서 대응하는 열을 찾지 못해 이 열은 대조하지 않음"})
        for j in T.get("doubtful", []):
            out.append({"판정": "확인 필요", "시트": sheet, "보고서 행": None, "원본 행": None, "열": heads[j], "보고서 값": None, "원본 값": None,
                        "이유": f"이름이 같은 원본 '{_row_reason_field(P, cmap[j])}'과 값이 대부분 다름 (열을 잘못 옮겼을 수 있음)"})
        for i, row in enumerate(rows):
            rn = T["rownums"][i]
            k = T["pairs"].get(i)
            if k is None:
                if all(v is None for v in row):
                    continue
                twin = matched_rows.get(tuple(row[j] for j in sorted(cmap)))
                if twin is not None:
                    out.append({"판정": "확인 필요", "시트": sheet, "보고서 행": rn, "원본 행": twin, "열": None, "보고서 값": None, "원본 값": None,
                                "이유": "같은 거래가 보고서에 두 번 있음 (중복 기재)"})
                    continue
                out.append({"판정": "원본에 없음", "시트": sheet, "보고서 행": rn, "원본 행": None, "열": None, "보고서 값": None, "원본 값": None,
                            "이유": "같은 거래를 원본에서 찾지 못함 (대응 열 값이 절반 이상 같은 원본 행 없음)"})
                continue
            rec = recs[k]
            diffs = []
            if (sheet, i) in dup:
                diffs.append(None)
                out.append({"판정": "확인 필요", "시트": sheet, "보고서 행": rn, "원본 행": rec["_row"], "열": None, "보고서 값": None, "원본 값": None,
                            "이유": f"같은 원본 거래가 나눠 담은 다른 시트('{dup_other.get((sheet, i), '?')}')에도 있음 (중복 기재)"})
            for j, f in cmap.items():
                if not _row_eq(row[j], rec.get(f)):
                    diffs.append(j)
                    out.append({"판정": "값 차이", "시트": sheet, "보고서 행": rn, "원본 행": rec["_row"], "열": heads[j],
                                "보고서 값": _row_show(row[j]), "원본 값": _row_show(rec.get(f)),
                                "이유": f"원본 '{_row_reason_field(P, f)}' 값과 다름"})
            for t, text, fn, a, b, d in formulas:
                if t in diffs or a in diffs or b in diffs or any(row[x] is None or row[x][0] != "n" for x in (t, a, b)):
                    continue
                calc = fn(row[a][1], row[b][1])
                if calc is None or abs(calc - row[t][1]) <= max(0.5 * 10 ** (-d) + 1e-9, abs(row[t][1]) * ROW_FORMULA_REL):
                    continue
                same_as_raw = t in cmap and _row_eq(row[t], rec.get(cmap[t]))
                diffs.append(t)
                out.append({"판정": "확인 필요" if same_as_raw else "값 차이", "시트": sheet, "보고서 행": rn, "원본 행": rec["_row"], "열": heads[t],
                            "보고서 값": _row_show(row[t]), "원본 값": round(calc, d),
                            "이유": f"같은 행 계산식 '{text}' 계산값과 다름" + (" — 원본 값과는 같음 (원본 자체 확인)" if same_as_raw else "")})
            if not diffs:
                out.append({"판정": "일치", "시트": sheet, "보고서 행": rn, "원본 행": rec["_row"], "열": None, "보고서 값": None, "원본 값": None, "이유": ""})
    for sheet_name, lack in inner_missing:                  # 다른 시트(층)에는 있는데 이 시트(층)에 없는 거래
        _missing_raw(out, recs, P, set(range(n_raw)) - lack, sheet_name, "같은 내용의 다른 시트에는 있는데 이 시트에 없는 원본 거래")
    for names, n_part, n_union in notes:
        out.append({"판정": "확인 필요", "시트": names, "보고서 행": None, "원본 행": None, "열": None,
                    "보고서 값": None, "원본 값": None,
                    "이유": f"다른 시트에도 있는 거래 {n_part}건만 담은 시트 (보고서 전체 {n_union}건) — 발췌 시트인지, 같은 거래를 두 번 적은 것인지 확인 (누락은 판단하지 않음)"})
    _missing_raw(out, recs, P, all_used, use[0]["sheet"] if len(use) == 1 else "(어느 시트에도 없음)")
    return out


def _missing_raw(out, recs, P, used, sheet, why="원본 거래가 보고서에 없음"):
    for k, rec in enumerate(recs):
        if k not in used:
            show = rec.get("_line") or " / ".join(f"{f}={_row_show(rec.get(f))}" for f in P["raw_fields"] if rec.get(f) is not None)
            out.append({"판정": "보고서에 없음", "시트": sheet, "보고서 행": None, "원본 행": rec["_row"], "열": None, "보고서 값": None,
                        "원본 값": show[:200], "이유": why})


def _rows_mode_wanted(raw_path, report_path, exc):
    """합계형 대조가 '행 단위로 옮긴 표'·'원본 표 못 찾음'으로 끝났거나 맞는 칸이 0개일 때 행 대조로 전환할지"""
    if exc is not None and not any(x in str(exc) for x in ("행 단위로 옮긴", "열 제목 줄을 찾지 못함", "더할 숫자 열이 없음")):
        return False
    try:
        return analyze_structure(raw_path, report_path)["report_kind"] == "rows"
    except Exception:
        return False


def _rows_order(raw_path, report_path, explicit, exc):
    """행 대조로 전환할 (원본, 보고서) — 전환하지 않으면 None. 자동 판별일 때는 반대 순서도 확인 (글자 줄 원본 등)"""
    if _rows_mode_wanted(raw_path, report_path, exc):
        return raw_path, report_path
    if not explicit and (exc is None or any(x in str(exc) for x in ("행 단위로 옮긴", "열 제목 줄을 찾지 못함", "더할 숫자 열이 없음"))):
        if _rows_mode_wanted(report_path, raw_path, None):
            return report_path, raw_path
    return None


def _rows_prefer(raw_path, report_path, result):
    """합계형 대조가 됐어도 보고서가 원본 행을 그대로 옮긴 표가 확실하면(같은 이름 열 3개 이상, 90% 이상 행이 맞음) 행 대조"""
    try:
        if result.get("raw_rows", ROW_AUTO_MAX + 1) > ROW_AUTO_MAX:
            return False
        raw_heads = {nl2(h) for h in result.get("raw_head", [])}
        rep = _flat_tables(report_path)[0]
        if len(rep["rows"]) < 5 or sum(1 for h in rep["heads"] if nl2(h) in raw_heads) < 3:
            return False
        st = analyze_structure(raw_path, report_path)
        return st["report_kind"] == "rows" and st["matched_rate"] >= 0.9
    except Exception:
        return False


def run_rows(raw_path, report_path):
    """행 대조 결과 파일 저장 → (저장 경로, 결과 목록, 구조, 판정별 수, 값 차이 행 수)"""
    st = analyze_structure(raw_path, report_path)
    res = reconcile_rows(raw_path, report_path)
    tables, _ = _row_load(raw_path, report_path)
    base, _ = os.path.splitext(report_path)
    out = out_path(base, "_원본대사")
    cnt = {k: sum(1 for x in res if x["판정"] == k) for k in ("일치", "값 차이", "원본에 없음", "보고서에 없음", "확인 필요")}
    diff_rows = len({(x["시트"], x["보고서 행"]) for x in res if x["판정"] == "값 차이"})
    clean = cnt["일치"] and not (cnt["값 차이"] or cnt["원본에 없음"] or cnt["보고서에 없음"] or cnt["확인 필요"] or st.get("skipped_sheets"))
    wb = Workbook(); ws = wb.active; ws.title = "요약"
    ws.append(["최종", "전부 일치" if clean else
               f"확인 필요: 값 차이 {cnt['값 차이']}칸({diff_rows}행) / 원본에 없음 {cnt['원본에 없음']}행 / 보고서에 없음 {cnt['보고서에 없음']}건 / 확인 필요 {cnt['확인 필요']}건"
               + (f" / 대조 안 한 시트 {len(st['skipped_sheets'])}개" if st.get("skipped_sheets") else "")])
    ws.append(["실행일시", datetime.datetime.now().strftime("%Y-%m-%d %H:%M")])
    lines = [f"원본: {os.path.basename(raw_path)} ('{st.get('raw_sheet')}' 시트, {'글자 줄에서 거래 추출' if st['raw_kind'] == 'text' else '표'}, {st.get('raw_records')}건)  →  보고서: {os.path.basename(report_path)} ('{', '.join(st['sheets']) or st.get('report_sheet')}' 시트)",
             "방식: 행 대조 (합계형 보고서가 아니라 원본 행을 옮긴 표로 판단)",
             "열 대응: " + ", ".join(f"{k} ↔ {v}" for k, v in st["column_map"].items()),
             "거래 구분 열: " + ", ".join(st["keys"]),
             "같은 행 계산식: " + (", ".join(st["row_formulas"]) if st["row_formulas"] else "없음"),
             f"판정: 일치 {cnt['일치']}행 / 값 차이 {diff_rows}행 / 원본에 없음 {cnt['원본에 없음']}행 / 원본 중 보고서에 없음 {cnt['보고서에 없음']}건 / 확인 필요 {cnt['확인 필요']}건"]
    if st.get("unmapped"):
        lines.append("원본과 대응하지 않은 보고서 열(대조 안 함): " + ", ".join(st["unmapped"]))
    if st.get("skipped_sheets"):
        lines.append("행 대조하지 않은 보고서 시트: " + ", ".join(st["skipped_sheets"]))
    for T in tables:
        lines += ["주의: " + n for n in T["notes"]]
    for line in lines:
        ws.append(["", line])
    fields = ["판정", "시트", "보고서 행", "원본 행", "열", "보고서 값", "원본 값", "이유"]
    d = wb.create_sheet("행대조")
    d.append(fields)
    for x in res:
        if x["판정"] != "일치":
            d.append([x[f] for f in fields])
    a = wb.create_sheet("전체판정")
    a.append(fields)
    for x in res:
        a.append([x[f] for f in fields])
    for sh in (d, a):
        _style_table(sh, 1, wrap=("이유", "원본 값"), max_width=70, judge=("판정",))
    _style_labels(ws, 1, 2)
    ws["B1"].font = Font(bold=True, size=12)
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = min(120, max(40, max(_text_width(c.value) for c in ws["B"]) + 2))
    _diagnostic_text(wb)
    wb.save(out)
    return out, res, st, cnt, diff_rows


# 실제 Excel 실행


def excel_import_com():
    """지연 가져오기: Windows가 아닌 환경에서도 기본 검산은 실행 가능하다."""
    import pythoncom
    import win32com.client
    import win32api
    import win32process
    return pythoncom, win32com.client, win32api, win32process


def excel_write_json(path, value):
    temporary = str(path) + '.tmp'
    with open(temporary, 'w', encoding='utf-8') as stream:
        json.dump(value, stream, ensure_ascii=False)
    os.replace(temporary, path)


def excel_network_reason(path):
    """재계산 도중 외부 데이터를 가져오는 알려진 경로는 실행하지 않는다."""
    if not zipfile.is_zipfile(path):
        return None  # 형식 오류는 정상 열기 단계에서 구체적으로 기록한다.
    with zipfile.ZipFile(path) as archive:
        names = archive.namelist()
        if 'xl/connections.xml' in names or any(name.startswith('xl/queryTables/') for name in names):
            return '외부 데이터 연결이 있어 Excel 자동 재계산을 생략함 (연결 새로 고침을 실행하지 않음)'
        for name in names:
            if not (name.startswith('xl/worksheets/') and name.endswith('.xml')):
                continue
            try:
                root = ET.fromstring(archive.read(name))
            except ET.ParseError:
                continue
            for node in root.iter():
                if node.tag.rsplit('}', 1)[-1] != 'f' or not node.text:
                    continue
                formula = re.sub(r'"(?:[^"]|"")*"', '""', node.text)
                hit = re.search(r'(?<![\w.])(?:_xlfn\.)?(WEBSERVICE|RTD|STOCKHISTORY|IMAGE)\s*\(', formula, re.I)
                if hit or ('|' in formula and '!' in formula):
                    function = hit.group(1).upper() if hit else 'DDE'
                    return f'외부 통신 가능 수식 {function}이 있어 Excel 자동 재계산을 생략함'
    return None


def excel_worker(request):
    """별도 Python 프로세스에서만 호출. 테스트에서는 COM 객체를 대체한다."""
    result = {'status': 'failed', 'reason': '', 'warnings': []}
    try:
        pythoncom, client, api, process = excel_import_com()
    except Exception as exc:
        result.update(status='unavailable', reason=f'pywin32 가져오기 실패: {type(exc).__name__}: {exc}')
        return result
    app = workbook = None
    initialized = owned = False
    stage = 'COM 초기화'
    deadline = time.monotonic() + float(request['timeout'])
    timings = {}
    mark = [time.monotonic()]
    def lap(name):
        now = time.monotonic(); timings[name] = round(now - mark[0], 1); mark[0] = now
    try:
        pythoncom.CoInitialize()
        initialized = True
        existing = set(process.EnumProcesses())
        stage = 'Microsoft Excel 실행'
        app = client.DispatchEx('Excel.Application')
        lap('Excel 실행')
        stage = '독립 Excel 인스턴스 확인'
        _, pid = process.GetWindowThreadProcessId(app.Hwnd)
        if not pid or pid in existing:
            raise RuntimeError('독립 Excel 프로세스 소유권을 확인하지 못함; 기존 Excel은 변경하지 않음')
        handle = api.OpenProcess(0x1000, False, pid)  # PROCESS_QUERY_LIMITED_INFORMATION
        try:
            created = str(process.GetProcessTimes(handle)['CreationTime'])
        finally:
            handle.Close()
        owned = True
        excel_write_json(request['owner_path'], {'pid': pid, 'created': created, 'owned': True})
        stage = 'Excel 실행 설정'
        app.Visible = False
        app.DisplayAlerts = False
        app.EnableEvents = False
        app.AskToUpdateLinks = False
        app.AutomationSecurity = 3  # msoAutomationSecurityForceDisable
        stage = '통합문서 열기'
        workbook = app.Workbooks.Open(Filename=request['copy_path'], UpdateLinks=0, ReadOnly=False,
                                      IgnoreReadOnlyRecommended=True, AddToMru=False, Notify=False)
        # 별도 인스턴스에만 적용한다. 원본이나 사용자의 열린 Excel은 건드리지 않는다.
        lap('열기')
        app.Calculation = -4135  # xlCalculationManual
        stage = '전체 재계산'
        app.CalculateFullRebuild()
        stage = '계산 완료 대기'
        iterative = bool(app.Iteration)                 # 반복 계산(순환 참조 허용) 파일은 계산이 끝나도 상태가 '대기'로 남음 (노트북 Excel 실측)
        settle = time.monotonic() + 5
        while app.CalculationState != 0:  # xlDone
            if iterative and time.monotonic() >= settle:
                break
            if time.monotonic() >= deadline:
                raise TimeoutError('Excel 계산 완료 제한 시간 초과')
            pythoncom.PumpWaitingMessages()
            time.sleep(0.05)
        lap('계산')
        stage = '재계산 결과 저장'
        workbook.Save()
        lap('저장')
        result.update(status='ok', reason='Microsoft Excel 전체 재계산 완료')
    except Exception as exc:
        status = 'timeout' if isinstance(exc, TimeoutError) else 'failed'
        result.update(status=status, reason=f'{stage} 실패: {type(exc).__name__}: {exc}')
    finally:
        if workbook is not None:
            try:
                workbook.Close(SaveChanges=False)
            except Exception as exc:
                result['warnings'].append(f'임시 통합문서 닫기 실패: {type(exc).__name__}: {exc}')
        if app is not None and owned:
            try:
                app.Quit()
            except Exception as exc:
                result['cleanup_required'] = True
                result['warnings'].append(f'작업용 Excel 종료 실패: {type(exc).__name__}: {exc}')
        workbook = app = None
        if initialized:
            try:
                pythoncom.CoUninitialize()
            except Exception as exc:
                result['warnings'].append(f'COM 해제 실패: {type(exc).__name__}: {exc}')
        lap('닫기·종료')
        result['timings'] = timings
    return result


def excel_worker_main(request_path):
    """단일 실행 파일의 숨은 작업자 진입점. 요청/응답은 임시 폴더 안 JSON이다."""
    with open(request_path, encoding='utf-8') as stream:
        request = json.load(stream)
    result = excel_worker(request)
    excel_write_json(request['response_path'], result)
    return 0


def excel_terminate_owned(owner_path):
    """시간 초과 시 이 작업이 생성한 PID와 생성 시간이 모두 맞는 Excel만 종료."""
    if not os.path.isfile(owner_path):
        return '소유 Excel 기록 없음; 확인되지 않은 Excel 프로세스는 종료하지 않음'
    handle = None
    try:
        with open(owner_path, encoding='utf-8') as stream:
            owner = json.load(stream)
        if not owner.get('owned'):
            return '소유권 미확인; Excel 프로세스는 종료하지 않음'
        _, _, api, process = excel_import_com()
        handle = api.OpenProcess(0x1000 | 0x0001, False, int(owner['pid']))
        created = str(process.GetProcessTimes(handle)['CreationTime'])
        if created != owner['created']:
            return '프로세스 생성 시간이 다름; 다른 Excel 프로세스는 종료하지 않음'
        api.TerminateProcess(handle, 1)
        return '시간 초과/종료 실패한 작업용 Excel 프로세스만 종료함'
    except Exception as exc:
        # 이미 정상 종료된 경우도 있다. 원인을 보존하며 다른 프로세스는 시도하지 않는다.
        return f'작업용 Excel 종료 확인: {type(exc).__name__}: {exc}'
    finally:
        if handle is not None:
            handle.Close()


_SANITIZE_CELL_RE = re.compile(rb'<((?:\w+:)?)c\b([^>]*?)\s+t="e"([^>]*)>(.*?)</\1c>', re.S)
_SANITIZE_V_RE = re.compile(rb"<(?:\w+:)?v>(.*?)</(?:\w+:)?v>", re.S)
_EXCEL_STD_ERRORS = cf_ERRORS_ALL


def excel_sanitize_copy(path):
    """Excel에 넘길 '복사본'만 정리: 오류 형식(t="e") 셀에 엑셀 오류코드가 아닌 값이 저장돼 있으면
    Excel이 파일 자체를 열지 못한다(복구 열기도 실패, 노트북 Excel 실측). 그 저장값만 지워 Excel이 다시 계산하게 한다.
    반환: 정리한 셀 수. 원본 파일에는 쓰지 않는다."""
    std = {e.encode() for e in _EXCEL_STD_ERRORS}
    count = [0]
    def fix(match):
        value = _SANITIZE_V_RE.search(match.group(4))
        if value and value.group(1).strip() in std:
            return match.group(0)
        count[0] += 1
        inner = _SANITIZE_V_RE.sub(b"", match.group(4))
        return b"<" + match.group(1) + b"c" + match.group(2) + match.group(3) + b">" + inner + b"</" + match.group(1) + b"c>"
    changed = {}
    with zipfile.ZipFile(path) as archive:
        for name in archive.namelist():
            if name == 'xl/workbook.xml':                   # 쓰기 암호(읽기 전용 권장)가 있으면 Excel이 암호 창을 띄우고 멈춤 → 복사본에서만 해제
                data = archive.read(name)
                new = re.sub(rb'<fileSharing\b[^>]*/>', b'', data)
                if new != data:
                    changed[name] = new
            if name.startswith('xl/worksheets/') and name.endswith('.xml'):
                data = archive.read(name)
                if b't="e"' in data:
                    new = _SANITIZE_CELL_RE.sub(fix, data)
                    if new != data:
                        changed[name] = new
    if changed:
        tmp = path + '.tmp'
        with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                zout.writestr(item, changed.get(item.filename, zin.read(item.filename)))
        os.replace(tmp, path)
    return count[0]


@contextlib.contextmanager
def excel_recalculate(path, timeout=120):
    """재계산 복사본은 with 블록 안에서만 유효하다. 모든 COM 호출은 자식에서 수행."""
    if isinstance(timeout, bool) or not isinstance(timeout, (int, float)) or not math.isfinite(timeout) or timeout <= 0:
        raise ValueError('Excel 제한 시간은 0보다 큰 유한한 초 단위 숫자여야 함')
    result = {'status': 'failed', 'reason': '', 'warnings': []}
    if sys.platform != 'win32':
        result.update(status='unavailable', reason='Windows가 아니므로 pywin32 Excel 재계산 불가')
        yield result
        return
    folder = tempfile.mkdtemp(prefix='excel_recalc_')
    owner_path = None
    try:
        source = os.path.abspath(os.fspath(path))
        copy_path = os.path.join(folder, os.path.basename(source))
        shutil.copyfile(source, copy_path)
        result['sanitized'] = excel_sanitize_copy(copy_path)
        blocked = excel_network_reason(copy_path)
        if blocked:
            result.update(status='unavailable', reason=blocked)
        else:
            request = {'copy_path': copy_path, 'response_path': os.path.join(folder, 'response.json'),
                       'owner_path': os.path.join(folder, 'owner.json'), 'timeout': timeout}
            owner_path = request['owner_path']
            request_path = os.path.join(folder, 'request.json')
            excel_write_json(request_path, request)
            command = [sys.executable, os.path.abspath(__file__), '--excel-worker', request_path]
            try:
                child_env = dict(os.environ, PYTHONIOENCODING='utf-8', PYTHONUTF8='1')   # 한글 Windows(cp949)에서 이유 글자 깨짐 방지
                completed = subprocess.run(command, capture_output=True, text=True, encoding='utf-8',
                                           errors='replace', timeout=timeout, env=child_env,
                                           creationflags=getattr(subprocess, 'CREATE_NO_WINDOW', 0))
                if completed.returncode != 0 or not os.path.isfile(request['response_path']):
                    said = ' / '.join(t for t in ((completed.stdout or '').strip()[-500:], (completed.stderr or '').strip()[-1000:]) if t)
                    result['reason'] = f'Excel 작업자 실행 실패 (종료 코드 {completed.returncode}): {said or "출력 없음"}'
                    if os.path.isfile(owner_path):
                        result['warnings'].append(excel_terminate_owned(owner_path))
                else:
                    with open(request['response_path'], encoding='utf-8') as stream:
                        response = json.load(stream)
                    if response.get('status') not in ('ok', 'failed', 'unavailable', 'timeout'):
                        raise ValueError('Excel 작업자 응답에 유효한 상태 없음')
                    result.update(response)
                    if result.pop('cleanup_required', False):
                        result['warnings'].append(excel_terminate_owned(request['owner_path']))
                    if result['status'] == 'ok':
                        if not os.path.isfile(copy_path):
                            raise FileNotFoundError('재계산 임시 파일 없음')
                        result['cached_path'] = copy_path
            except subprocess.TimeoutExpired:
                result.update(status='timeout', reason=f'Excel 재계산 {timeout:g}초 제한 초과 (실행·열기·계산·저장·종료 포함)')
                result['warnings'].append(excel_terminate_owned(request['owner_path']))
    except Exception as exc:
        result.update(status='failed', reason=f'Excel 재계산 준비/응답 처리 실패: {type(exc).__name__}: {exc}')
        if owner_path and os.path.isfile(owner_path):
            result['warnings'].append(excel_terminate_owned(owner_path))
    try:
        yield result
    finally:
        try:
            shutil.rmtree(folder)
        except OSError as exc:
            result['warnings'].append(f'임시 파일 삭제 실패: {type(exc).__name__}: {exc}')


def excel_formula_signature(cell):
    value = cell.value
    if cell.data_type != 'f':
        return None
    if isinstance(value, str):
        return ('scalar', value)
    return (type(value).__name__, getattr(value, 'text', None), getattr(value, 'ref', None))


def excel_empty_string_caches(path):
    """openpyxl reads <c t='str'><v/></c> as None; Excel means empty text.

    Do not treat an absent <v> (or an empty numeric <v>) as a known empty string.
    Sheet names are mapped from OPC relationships, not assumed sheet ordering.
    """
    main = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
    rel_attr = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id'
    result = set()
    with zipfile.ZipFile(path) as archive:
        wb = ET.fromstring(archive.read('xl/workbook.xml'))
        relationships = ET.fromstring(archive.read('xl/_rels/workbook.xml.rels'))
        targets = {rel.attrib['Id']: rel.attrib['Target'] for rel in relationships
                   if rel.attrib.get('TargetMode') != 'External'}
        for sheet in wb.findall('{%s}sheets/{%s}sheet' % (main, main)):
            target = targets.get(sheet.attrib.get(rel_attr))
            if not target:
                continue
            target = target.lstrip('/') if target.startswith('/') else posixpath.normpath(posixpath.join('xl', target))
            root = ET.fromstring(archive.read(target))
            for cell in root.iter('{%s}c' % main):
                value = cell.find('{%s}v' % main)
                if cell.get('t') == 'str' and value is not None and value.text is None:
                    result.add((sheet.attrib['name'], cell.attrib['r']))
    return result


_SCOPE_LINK_RE = cf_EXT
_SCOPE_VOLATILE_RE = re.compile(r'(?<![\w.])(?:_xlfn\.)?(RAND|RANDBETWEEN|RANDARRAY|NOW|TODAY|RTD|WEBSERVICE|STOCKHISTORY|INDIRECT|CELL|INFO)\s*\(', re.I)
_SCOPE_CELL_RE = re.compile(r"^\$?([A-Za-z]{1,3})\$?(\d+)$")
_SCOPE_COL_RE = re.compile(r"^\$?([A-Za-z]{1,3})$")
_SCOPE_ROW_RE = re.compile(r"^\$?(\d+)$")
_SCOPE_MAX_ROW, _SCOPE_MAX_COL = 1048576, 16384


def _scope_split_sheet(text, current):
    """'시트'!A1 / 시트!A1 / A1 → (시트, 'A1'). 해석 불가면 None."""
    if '!' not in text:
        return current, text
    head, _, ref = text.rpartition('!')
    if head.startswith("'") and head.endswith("'") and len(head) >= 2:
        head = head[1:-1].replace("''", "'")
    if not head or '[' in head or ']' in head or ':' in head:
        return None  # 외부 파일·3D 참조
    return head, ref


def _scope_parse_area(ref):
    """A1 / A1:B2 / A:B / 1:3 → (r1, c1, r2, c2). 해석 불가면 None."""
    parts = ref.split(':')
    if len(parts) == 1:
        mm = _SCOPE_CELL_RE.match(parts[0])
        if not mm:
            return None
        r, c = int(mm.group(2)), column_index_from_string(mm.group(1).upper())
        return (r, c, r, c)
    if len(parts) != 2:
        return None
    a, b = parts
    ma, mb = _SCOPE_CELL_RE.match(a), _SCOPE_CELL_RE.match(b)
    if ma and mb:
        r1, r2 = sorted((int(ma.group(2)), int(mb.group(2))))
        c1, c2 = sorted((column_index_from_string(ma.group(1).upper()), column_index_from_string(mb.group(1).upper())))
        return (r1, c1, r2, c2)
    ma, mb = _SCOPE_COL_RE.match(a), _SCOPE_COL_RE.match(b)
    if ma and mb:
        c1, c2 = sorted((column_index_from_string(ma.group(1).upper()), column_index_from_string(mb.group(1).upper())))
        return (1, c1, _SCOPE_MAX_ROW, c2)
    ma, mb = _SCOPE_ROW_RE.match(a), _SCOPE_ROW_RE.match(b)
    if ma and mb:
        r1, r2 = sorted((int(ma.group(1)), int(mb.group(1))))
        return (r1, 1, r2, _SCOPE_MAX_COL)
    return None


def _scope_resolve(text, current, sheet_names):
    """참조 글자 → [(시트, r1, c1, r2, c2)] 또는 None(해석 불가)."""
    split = _scope_split_sheet(text, current)
    if split is None:
        return None
    sheet, ref = split
    if sheet not in sheet_names:
        folded = {n.casefold(): n for n in sheet_names}
        sheet = folded.get(sheet.casefold())
        if sheet is None:
            return None
    area = _scope_parse_area(ref.strip())
    return None if area is None else [(sheet,) + area]


def _scope_strip_text(expression):
    return re.sub(r'"(?:[^"]|"")*"', '""', expression or '')


def excel_formula_scope(workbook, formulas, budget=3_000_000, extra_seeds=None):
    """휘발성 함수·외부 링크의 영향을 받는 수식 셀만 찾는다.

    반환: (영향 셀 {(시트, 좌표): 근원 목록}, 통합문서 전체를 막아야 하는 사유 목록)
    해석할 수 없는 참조(표 참조, 3D, 동적 이름 등)는 영향이 있는 것으로 본다(보수적).
    계산량이 budget을 넘으면 통합문서 전체를 막는다(이전 동작).
    """
    sheet_names = list(workbook.sheetnames)
    # 이름 정의: 단순 범위만 추적, 휘발성·외부링크·해석 불가는 표시
    names = {}
    try:
        definitions = list(workbook.defined_names.values())
    except AttributeError:
        definitions = list(workbook.defined_names.definedName)
    for definition in definitions:
        text = (getattr(definition, 'attr_text', '') or '').strip()
        body = text[1:] if text.startswith('=') else text
        bare = _scope_strip_text(body)
        if _SCOPE_VOLATILE_RE.search(bare) or _SCOPE_LINK_RE.search(bare):
            names[definition.name.upper()] = ('seed', None)
            continue
        areas = []
        for piece in body.split(','):
            got = _scope_resolve(piece.strip(), None, sheet_names) if '!' in piece else None
            if got is None:
                areas = None
                break
            areas.extend(got)
        names[definition.name.upper()] = ('ok', areas) if areas else ('unknown', None)
    for ws in workbook.worksheets:  # 시트 범위 이름은 보수적으로 해석 불가
        for name in getattr(ws, 'defined_names', {}) or {}:
            names.setdefault(str(name).upper(), ('unknown', None))

    info = {}
    seeds = {}
    for (sheet, coord), signature in formulas.items():
        expression = signature[1] or ''
        bare = _scope_strip_text(expression)
        key = (sheet, coord)
        reasons = []
        if _SCOPE_VOLATILE_RE.search(bare):
            reasons.append('휘발성 함수')
        if _SCOPE_LINK_RE.search(bare):
            reasons.append('외부 링크')
        ranges, unknown = [], False
        try:
            tokens = Tokenizer(expression if expression.startswith('=') else '=' + expression).items
        except Exception:
            tokens, unknown = [], True
        for token in tokens:
            if token.type != 'OPERAND' or token.subtype != 'RANGE':
                continue
            value = token.value
            if value.lower().startswith('_xlpm.'):
                continue  # LET·LAMBDA 안의 이름(매개변수)은 셀 참조가 아님
            got = _scope_resolve(value, sheet, sheet_names)
            if got is not None:
                ranges.extend(got)
                continue
            kind, areas = names.get(value.upper(), ('unknown', None))
            if kind == 'ok':
                ranges.extend(areas)
            elif kind == 'seed':
                reasons.append('휘발성·외부링크 이름 ' + value)
            else:
                unknown = True
        if reasons:
            seeds[key] = [f'{sheet}!{coord}', '휘발성 함수·외부 링크']
        mm = _SCOPE_CELL_RE.match(coord)
        info[key] = (int(mm.group(2)), column_index_from_string(mm.group(1).upper()), ranges, unknown)

    for key, kind in (extra_seeds or {}).items():
        if key in info and key not in seeds:
            seeds[key] = [f'{key[0]}!{key[1]}', kind]
    whole = []
    if getattr(workbook, '_external_links', ()):
        if not any(_SCOPE_LINK_RE.search(_scope_strip_text(sig[1] or '')) for sig in formulas.values()):
            whole.append('통합문서 외부 링크(위치 확인 불가); 의존 범위 미확인으로 이 통합문서의 저장값 대조를 확정하지 않음')
    if whole or not seeds:
        return ({}, whole)

    tainted = dict(seeds)
    by_sheet = {}
    array_ranges = {}
    def add_cell(key):
        r, c = info[key][0], info[key][1]
        by_sheet.setdefault(key[0], {}).setdefault(c, []).append(r)
        signature = formulas[key]
        if signature[0] == 'ArrayFormula' and signature[2]:
            area = _scope_parse_area(signature[2])
            if area:
                array_ranges.setdefault(key[0], []).append((key, area))
    for key in tainted:
        add_cell(key)
    work = 0
    changed = True
    while changed:
        changed = False
        for key, (r, c, ranges, unknown) in info.items():
            if key in tainted:
                continue
            source = None
            if unknown:
                first = next(iter(seeds.values()))
                source = [first[0] + ' 등(해석 불가 참조)', first[1]]
            else:
                for (s, r1, c1, r2, c2) in ranges:
                    # 배열을 셀마다 펼치지 않고 출력 범위와 참조 범위의 겹침을 확인.
                    for array_key, (ar1, ac1, ar2, ac2) in array_ranges.get(s, ()):
                        work += 1
                        if r1 <= ar2 and ar1 <= r2 and c1 <= ac2 and ac1 <= c2:
                            source = list(tainted[array_key])
                            break
                    if source:
                        break
                    cols = by_sheet.get(s)
                    if not cols:
                        continue
                    for col, rows in cols.items():
                        work += 1 + len(rows)
                        if c1 <= col <= c2 and any(r1 <= x <= r2 for x in rows):
                            hit = next((k for k in tainted if k[0] == s and info[k][1] == col and r1 <= info[k][0] <= r2), None)
                            source = list(tainted[hit]) if hit else [f'{s}!?', '휘발성 함수·외부 링크']
                            break
                    if source:
                        break
            if work > budget:
                return ({}, ['휘발성 함수·외부 링크 의존 범위 계산량 초과; 이 통합문서의 저장값 대조를 확정하지 않음'])
            if source:
                tainted[key] = source
                add_cell(key)
                changed = True
    return (tainted, [])


_NORM_EXT_RE = re.compile(r"'?\[[^\]]+\]((?:[^'!]|'')+)'?!")


_NORM_STR_RE = re.compile(r'"(?:[^"]|"")*"')
_NORM_QUOTED_SHEET_RE = re.compile(r"'((?:[^']|'')+)'!")


def excel_formula_norm(signature):
    """Excel이 다시 저장하며 바꾸는 '표기'만 같게 본다 (글자 상수 안은 건드리지 않음).
    - 외부 파일 참조: '[파일.xlsx]시트'!A1 ↔ [1]시트!A1
    - 시트 이름 따옴표: '통제별분류'!K11 ↔ 통제별분류!K11 (노트북 Excel 실측)"""
    if not signature or signature[0] != 'scalar' or not isinstance(signature[1], str):
        return signature
    text = signature[1]
    out, last = [], 0
    def norm(part):
        part = _NORM_EXT_RE.sub(lambda m: '[X]' + m.group(1).replace("''", "'") + '!', part)
        return _NORM_QUOTED_SHEET_RE.sub(lambda m: m.group(1).replace("''", "'") + '!', part)
    for literal in _NORM_STR_RE.finditer(text):
        out.append(norm(text[last:literal.start()])); out.append(literal.group(0)); last = literal.end()
    out.append(norm(text[last:]))
    return (signature[0], ''.join(out))


def excel_value_description(value):
    if isinstance(value, bool):
        return str(value) + ' (논리값)'
    if isinstance(value, numbers.Real):
        return format(value, '.15g') + ' (숫자)'
    if isinstance(value, (datetime.datetime, datetime.date, datetime.time)):
        return value.isoformat() + ' (날짜/시간)'
    return repr(value) + (' (글자)' if isinstance(value, str) else '')


def excel_compare_cached(original_path, recalculated_path, tolerance=0.5):
    """Return diagnostics with 등급/항목/위치/상세; inputs remain untouched.

    Numerical tolerance is absolute in workbook units. Boolean, text, and dates
    retain their types, and text whitespace is significant. Changed formulas and
    their dependents are not certified. Array output cells are compared; data-table
    formulas remain unsupported.
    """
    if isinstance(tolerance, bool) or not isinstance(tolerance, numbers.Real) or not math.isfinite(tolerance) or tolerance < 0:
        raise ValueError('Excel 재검산 허용 오차는 0 이상의 유한한 숫자여야 합니다')
    rows = []
    def add(level, item, location, detail):
        rows.append({'등급': level, '항목': item, '위치': location, '상세': detail})
    books = []
    try:
        for path, data_only in ((original_path, False), (original_path, True),
                                (recalculated_path, False), (recalculated_path, True)):
            with open(path, 'rb') as stream:
                book = load_workbook(io.BytesIO(stream.read()), data_only=data_only)
            books.append(book)
        original, old_values, recalculated, new_values = books
        old_empty = excel_empty_string_caches(original_path)
        new_empty = excel_empty_string_caches(recalculated_path)
        old_formulas = {(sheet.title, cell.coordinate): excel_formula_signature(cell)
                        for sheet in original for row in sheet for cell in row if cell.data_type == 'f'}
        new_formulas = {(sheet.title, cell.coordinate): excel_formula_signature(cell)
                        for sheet in recalculated for row in sheet for cell in row if cell.data_type == 'f'}
        changed = {}
        change_text = {}
        for key in sorted(old_formulas):
            loc = f'{key[0]}!{key[1]}'
            if key not in new_formulas:
                changed[key] = '계산 전후 수식 변경'; change_text[key] = '재계산본에서 수식이 사라짐'
            elif excel_formula_norm(old_formulas[key]) != excel_formula_norm(new_formulas[key]):
                changed[key] = '계산 전후 수식 변경'
                change_text[key] = f'수식 변경 {old_formulas[key][1]!r} → {new_formulas[key][1]!r}'
        tainted, scope_reasons = excel_formula_scope(original, old_formulas, extra_seeds=changed)
        if changed:  # 배열 출력 범위가 바뀌면 새 범위에 연결된 수식도 영향 추적
            new_tainted, new_reasons = excel_formula_scope(recalculated, new_formulas, extra_seeds=changed)
            tainted.update(new_tainted)
            scope_reasons.extend(reason for reason in new_reasons if reason not in scope_reasons)
        counts = {'일치': 0, '불일치': 0, '오류값': 0, '검산 불가': 0}
        for (sheet, coord), signature in old_formulas.items():
            loc = f'{sheet}!{coord}'
            expression = signature[1] or repr(signature)
            old_cell = old_values[sheet][coord]; new_cell = new_values[sheet][coord]
            old = '' if (sheet, coord) in old_empty else old_cell.value
            new = '' if (sheet, coord) in new_empty else new_cell.value
            evidence = f'저장값 {excel_value_description(old)} / Excel 재계산 {excel_value_description(new)}; 수식: {expression}'
            if scope_reasons:
                add('정보', 'Excel 재검산 불가', loc, '; '.join(scope_reasons) + '; ' + evidence)
                counts['검산 불가'] += 1
            elif (sheet, coord) in tainted:
                origin, kind = tainted[(sheet, coord)][0], tainted[(sheet, coord)][1]
                why = f'{kind} 셀' if origin == loc else f'{kind} 셀 {origin}의 영향'
                if kind == '계산 전후 수식 변경':
                    why += ('; ' + change_text[(sheet, coord)] if (sheet, coord) in change_text else '') + '; 수식이 달라 불일치로 단정하지 않음'
                else:
                    why += '; 실행 시점·환경에 따라 값이 달라질 수 있어 불일치로 단정하지 않음'
                add('정보', 'Excel 재검산 불가', loc, why + '; ' + evidence)
                counts['검산 불가'] += 1
            elif old_cell.data_type == 'e' and isinstance(old, str) and old.strip() not in _EXCEL_STD_ERRORS:
                add('정보', 'Excel 재검산 불가', loc, '원본 저장값이 엑셀 오류코드가 아닌 비표준 값(계산 안 된 저장값)이라 대조하지 않음; ' + evidence)
                counts['검산 불가'] += 1
            elif signature[0] == 'ArrayFormula' and signature[2]:
                # 배열수식: Excel이 다시 계산한 배열 범위 전체를 칸마다 저장값과 비교
                bad, err, missing, n_cells = [], [], [], 0
                for row_cells in old_values[sheet][signature[2]] if ':' in signature[2] else ((old_values[sheet][signature[2]],),):
                    for oc in row_cells:
                        n_cells += 1
                        nc = new_values[sheet][oc.coordinate]
                        o, n = oc.value, nc.value
                        if o is None and (sheet, oc.coordinate) in old_empty:
                            o = ''
                        if n is None and (sheet, oc.coordinate) in new_empty:
                            n = ''
                        if nc.data_type == 'e':
                            err.append(f'{oc.coordinate}={n}')
                        elif o is None or n is None:
                            missing.append(oc.coordinate)
                        elif isinstance(o, numbers.Real) and isinstance(n, numbers.Real) and not isinstance(o, bool) and not isinstance(n, bool):
                            if not (math.isfinite(o) and math.isfinite(n) and abs(o - n) <= max(tolerance, abs(o) * 1e-9)):
                                bad.append((oc.coordinate, o, n))
                        elif not (type(o) is type(n) and o == n):
                            bad.append((oc.coordinate, o, n))
                if bad:
                    for coord_b, o, n in bad:
                        add('오류', 'Excel 재검산 불일치', f'{sheet}!{coord_b}',
                            f'배열수식 {loc} 범위 {signature[2]}의 칸; 저장값 {excel_value_description(o)} / Excel 재계산 {excel_value_description(n)}; 수식: {signature[1]}')
                    counts['불일치'] += 1
                elif err:
                    add('오류', 'Excel 재계산 오류값', loc, f'배열수식 범위 {signature[2]}: ' + ', '.join(err[:5]) + f'; 수식: {signature[1]}')
                    counts['오류값'] += 1
                elif missing:
                    add('정보', 'Excel 재검산 불가', loc, f'배열수식 범위 {signature[2]}에 저장값 없는 칸 ' + ', '.join(missing[:5]) + f'; 수식: {signature[1]}')
                    counts['검산 불가'] += 1
                else:
                    add('정보', 'Excel 재검산 일치', loc, f'배열수식 범위 {signature[2]} {n_cells}칸 모두 일치; 수식: {signature[1]}')
                    counts['일치'] += 1
            elif signature[0] != 'scalar':
                add('정보', 'Excel 재검산 불가', loc, '데이터 표 수식 등 출력 범위 대조 미지원; ' + evidence)
                counts['검산 불가'] += 1
            elif new_cell.data_type == 'e' and str(new).strip().upper() == '#NAME?' and old not in (None, '') \
                    and not (isinstance(old, str) and old.strip().upper() in _EXCEL_STD_ERRORS):
                # 원래 PC에는 있던 추가기능·VBA 사용자 함수가 이 PC Excel에 없어 #NAME? (매크로는 꺼서 엶) → 오류로 보지 않음
                add('주의', 'Excel 재검산 불가(추가기능·사용자 함수)', loc, '저장값은 있는데 이 PC의 Excel이 함수를 몰라 #NAME? — 추가기능·VBA 함수가 있는 PC에서 확인; ' + evidence)
                counts['검산 불가'] += 1
            elif new_cell.data_type == 'e':
                add('오류', 'Excel 재계산 오류값', loc, evidence)
                counts['오류값'] += 1
            elif old is None or new is None:
                reasons = []
                if old is None:
                    reasons.append('원본 저장값 없음')
                if new is None:
                    reasons.append('Excel 재계산 저장값 없음')
                add('정보', 'Excel 재검산 불가', loc, '; '.join(reasons) + '; ' + evidence)
                counts['검산 불가'] += 1
            else:
                old_numeric = isinstance(old, numbers.Real) and not isinstance(old, bool)
                new_numeric = isinstance(new, numbers.Real) and not isinstance(new, bool)
                if old_numeric and new_numeric:
                    if not (math.isfinite(old) and math.isfinite(new)):
                        add('정보', 'Excel 재검산 불가', loc, '유한하지 않은 숫자 저장값; ' + evidence)
                        counts['검산 불가'] += 1
                        continue
                    difference = old - new
                    equal = abs(difference) <= max(tolerance, abs(old) * 1e-9)
                    evidence += f'; 차이(저장값-Excel) {difference:.15g}; 허용 오차 {tolerance:.15g}'
                else:
                    equal = type(old) is type(new) and old == new and old_cell.data_type != 'e'
                if equal:
                    add('정보', 'Excel 재검산 일치', loc, evidence)
                    counts['일치'] += 1
                else:
                    add('오류', 'Excel 재검산 불일치', loc, evidence)
                    counts['불일치'] += 1
        suffix = ' / '.join(f'{name} {value}' for name, value in counts.items())
        add('정보', 'Excel 재검산 요약', '전체', f'수식 {len(old_formulas)}개 / {suffix}' + ('; 검사할 수식 없음' if not old_formulas else ''))
        return rows
    except Exception as error:
        add('주의', 'Excel 재검산 불가', '전체', f'원본·재계산본 읽기/대조 실패: {type(error).__name__}: {error}')
        return rows
    finally:
        for book in books:
            book.close()


def excel_formula_diagnostics(path):
    if not USE_EXCEL:
        full = cf_check(path)
        formula_count = next((int(row['상세']) for row in full if row['항목'] == '수식 개수'), 0)
        if not formula_count:
            return full + [{'등급': '정보', '항목': 'Excel 재계산 대상 없음', '위치': '전체',
                            '상세': '수식이 없어 Excel을 실행하지 않음; 값의 정확성은 원본 대조 결과 확인'}]
        return full + [{'등급': '정보', '항목': '검산 방식', '위치': '전체',
                        '상세': 'Excel 프로그램 실행 없이 formulas·단순 합계로 검사'}]
    base = cf_check(path, recalculate=False)
    formula_count = next((int(row['상세']) for row in base if row['항목'] == '수식 개수'), 0)
    if not formula_count:
        return base + [{'등급': '정보', '항목': 'Excel 재계산 대상 없음', '위치': '전체',
                        '상세': '수식이 없어 Excel을 실행하지 않음; 값의 정확성은 원본 대조 결과 확인'}]
    backend = None
    native = None
    failure = None
    _t = time.monotonic()
    _progress(f"  Excel 재계산 중 (최대 {EXCEL_TIMEOUT}초) — Excel 실행·열기·계산·저장")
    try:
        with excel_recalculate(path, timeout=EXCEL_TIMEOUT) as backend:
            if backend['status'] == 'ok':
                native = excel_compare_cached(path, backend['cached_path'], tolerance=RECALC_EPS)   # 같은 Excel이 다시 계산 → 아주 작은 오차만 허용 (0.5면 비율·소수 칸 변조를 놓침)
            else:
                failure = (backend['status'], backend['reason'])
    except Exception as exc:
        failure = ('failed', f'Excel 결과 비교 실패: {type(exc).__name__}: {exc}')
    _progress(f"  Excel 재계산 {'실패 → formulas·단순 합계로 대체 검사' if failure else '끝'} ({time.monotonic() - _t:.1f}초)")
    if failure:
        kind, reason = failure
        item = {'timeout': 'Excel 재계산 시간 초과', 'unavailable': 'Excel 재계산 불가'}.get(kind, 'Excel 재계산 실패')
        result = cf_check(path)
        result.append({'등급': '주의', '항목': item, '위치': '전체',
                       '상세': reason + ' → formulas·단순 합계로 대체 검사 (Excel 검증 성공 아님)'})
    else:
        # Excel이 배열수식까지 다시 계산해 비교했으므로 기본 검사의 '배열수식 검사 불가' 표시는 뺌 (비교 결과는 native에 있음)
        error_locs = {row['위치'] for row in base if row['항목'] == '오류값'}
        native = [row for row in native if not (row['항목'] == 'Excel 재계산 오류값' and row['위치'] in error_locs)]   # 같은 칸 오류값을 두 번 세지 않음
        result = [row for row in base if row['항목'] != '배열수식 검사 불가'] + native
        if (backend or {}).get('sanitized'):
            result.append({'등급': '정보', '항목': 'Excel 열기 전 정리', '위치': '전체',
                           '상세': f"오류 형식 셀의 비표준 저장값 {backend['sanitized']}칸을 Excel용 복사본에서만 지움 (그대로 두면 Excel이 파일을 열지 못함). 원본은 그대로"})
        timings = (backend or {}).get('timings') or {}
        if timings:
            result.append({'등급': '정보', '항목': 'Excel 재계산 소요', '위치': '전체',
                           '상세': ' / '.join(f'{k} {v}초' for k, v in timings.items()) + f' (합계 {sum(timings.values()):.1f}초)'})
    for warning in (backend or {}).get('warnings', []):
        result.append({'등급': '주의', '항목': 'Excel 종료·정리 주의', '위치': '전체', '상세': warning})
    return result

# 통합 실행부: 외부 코드 파일을 읽거나 실행하지 않음.
def _input_paths(args):
    folder = os.path.abspath(args.folder or FOLDER or os.path.dirname(__file__))
    raw_setting = args.raw or RAW_FILE
    report_setting = args.report or REPORT_FILE
    def absolute(value):
        return os.path.abspath(value if os.path.isabs(value) else os.path.join(folder, value))
    if args.files and (raw_setting or report_setting):
        raise ValueError("파일 인자와 원본·보고서 설정 중 한 방식만 사용하세요")
    if args.files:
        return [os.path.abspath(p) for p in args.files], True
    if raw_setting or report_setting:
        if raw_setting and not report_setting:
            raise ValueError("원본과 함께 보고서 파일도 지정하세요")
        return ([absolute(raw_setting)] if raw_setting else []) + [absolute(report_setting)], True
    return collect(folder), False


def _validate_files(files):
    if not 1 <= len(files) <= 2:
        raise ValueError(f"xlsx 파일이 {len(files)}개입니다. 수식 검사는 1개, 원본 대조는 2개를 지정하세요")
    for path in files:
        if not os.path.isfile(path):
            raise FileNotFoundError(f"파일 없음: {path}")
        ext = os.path.splitext(path)[1].lower()
        if ext in ('.xls', '.csv', '.ods', '.xlsb', '.txt'):
            raise ValueError(f"{ext} 파일은 직접 읽지 못합니다: {os.path.basename(path)} — 엑셀에서 열어 '다른 이름으로 저장' → "
                             "'Excel 통합 문서(*.xlsx)'로 저장한 뒤 실행하세요")
        if ext not in ('.xlsx', '.xlsm', '.xltx', '.xltm'):
            raise ValueError("이 통합본은 .xlsx(.xlsm 포함) 파일을 지원합니다")
        with open(path, 'rb') as stream:
            head = stream.read(8)
        if head == bytes.fromhex('D0CF11E0A1B11AE1'):
            raise ValueError(f"열기 암호가 걸렸거나 문서보안(DRM)이 적용된 파일입니다: {os.path.basename(path)} — "
                             "엑셀에서 암호·보안을 해제해 다른 이름으로 저장한 사본으로 실행하세요")
        if head[:4] != b'PK\x03\x04':
            raise ValueError(f"xlsx 형식이 아님: {os.path.basename(path)} — 확장자만 .xlsx인 다른 형식(csv·xls 등)일 수 있습니다. 엑셀에서 열어 .xlsx로 다시 저장하세요")
        try:
            zipfile.ZipFile(path).close()
        except zipfile.BadZipFile:
            raise ValueError(f"파일이 손상되었거나 저장이 끝나지 않은 파일입니다: {os.path.basename(path)} — 엑셀에서 열어(복구) 다시 저장한 뒤 실행하세요")
    if len(files) == 2 and os.path.samefile(files[0], files[1]):
        raise ValueError("원본과 보고서는 서로 다른 파일이어야 합니다")


def _formula_output(path):
    _t = time.monotonic()
    _progress(f"수식 검사 시작: {os.path.basename(path)} 읽는 중")
    diagnostics = excel_formula_diagnostics(path)
    _progress("  수식 검사 결과 파일 저장 중")
    out = cf_run(path, result=diagnostics)
    _progress(f"수식 검사 끝 ({time.monotonic() - _t:.1f}초)")
    wb = load_workbook(out, read_only=True, data_only=True)
    status = wb['수식검증']['B1'].value
    wb.close()
    print('수식 검사:', status)
    print('저장:', out)
    return status


def _protect_inputs(inputs, outputs):
    for target in outputs:
        for source in inputs:
            identical = os.path.normcase(os.path.realpath(target)) == os.path.normcase(os.path.realpath(source))
            if identical or (os.path.exists(target) and os.path.samefile(target, source)):
                raise ValueError(f'결과 경로가 입력 파일과 겹칩니다. 입력 파일명을 바꾸세요: {target}')


def main(argv=None):
    import argparse
    for stream in (sys.stdout, sys.stderr):                # 한글 Windows(cp949)에서 출력을 파일·파이프로 보낼 때 '—' 등으로 멈추지 않게
        try:
            stream.reconfigure(errors="replace")
        except Exception:
            pass
    parser = argparse.ArgumentParser(description='엑셀 원본 대조·수식 검사 자동화 (실행 파일 1개)')
    parser.add_argument('files', nargs='*', help='1개: 수식 검사 / 2개: 원본.xlsx 보고서.xlsx 순서')
    group = parser.add_mutually_exclusive_group()
    group.add_argument('--수식', '--formula', dest='formula', action='store_true', help='파일 1개의 수식만 검사')
    group.add_argument('--채우기', '--fill', dest='fill_mode', action='store_true', help='원본 집계로 보고서 값을 채운 새 파일 저장')
    parser.add_argument('--폴더', '--folder', dest='folder')
    parser.add_argument('--원본', '--raw', dest='raw')
    parser.add_argument('--보고서', '--report', dest='report')
    args = parser.parse_args(argv)
    global _RUN_STAMP, _PROGRESS, _PROGRESS_T0
    _RUN_STAMP = run_stamp()
    _PROGRESS, _PROGRESS_T0 = True, time.monotonic()
    _WB_CACHE.clear()
    try:
        mode = '수식' if args.formula else '채우기' if args.fill_mode else MODE
        if mode not in ('자동', '검산', '수식', '채우기'):
            raise ValueError('MODE는 자동 / 검산 / 수식 / 채우기 중 하나여야 합니다')
        files, explicit = _input_paths(args)
        _validate_files(files)
        if mode == '수식' or (mode == '자동' and len(files) == 1):
            if len(files) != 1:
                raise ValueError('수식 모드에서는 파일 1개만 지정하세요')
            status = _formula_output(files[0])
            print('최종:', status)
            return 0
        if len(files) != 2:
            raise ValueError('원본 1개와 보고서 1개가 필요합니다')
        raw_path, report_path = files if explicit else _order(files)
        _protect_inputs(files, [rule_path(report_path)])
        print('원본:', raw_path)
        print('보고서:', report_path)
        if not explicit:
            print('자동 판별입니다. 매핑이 틀리면 맨 위 RAW_FILE·REPORT_FILE을 지정하세요.')
        if mode == '채우기':
            out, result = fill(raw_path, report_path)
            for warning in result['warnings']:
                print('주의:', warning)
            print('규칙표:', result['rules_path'])
            print('저장:', out)
            print(f"최종: 채움 {result['filled']}칸 / 변경 {result['changed']}칸 / 채우지 못한 칸 {len(result['unmatched'])}개")
        else:
            _t = time.monotonic()
            _progress("원본 대조 시작: 원본 읽기·보고서 칸 찾는 중")
            try:
                result = verify(raw_path, report_path)
            except RuntimeError as exc:
                result = None
                order = _rows_order(raw_path, report_path, explicit, exc)
                if order is None:
                    raise
            if result is not None and result['matched'] == 0 and not (_missing(result) or _nocache(result)):
                order = _rows_order(raw_path, report_path, explicit, None)
                if order is not None:
                    result = None                            # 합계형으로 맞는 칸 0개 + 행 정리 표 → 행 대조
            elif result is not None and _rows_prefer(raw_path, report_path, result):
                order, result = (raw_path, report_path), None  # 원본 행을 그대로 옮긴 표가 확실 → 합계형 대신 행 대조
            if result is None:
                if order != (raw_path, report_path):
                    raw_path, report_path = order
                    print('행 대조용으로 원본·보고서를 바꿔 판별:', os.path.basename(raw_path), '→', os.path.basename(report_path))
                _progress("  합계형 보고서가 아니어서 행 대조로 전환")
                out, res, st, cnt, diff_rows = run_rows(raw_path, report_path)
                _progress(f"원본 대조 끝 ({time.monotonic() - _t:.1f}초)")
                print(f"방식: 행 대조 (원본 {'글자 줄' if st['raw_kind'] == 'text' else '표'} → 보고서 행 정리 표)")
                print("열 대응:", ", ".join(f"{k}↔{v}" for k, v in st["column_map"].items()))
                if st["row_formulas"]:
                    print("같은 행 계산식:", ", ".join(st["row_formulas"]))
                print('저장:', out)
                formula_status = _formula_output(report_path)
                if st.get("unmapped"):
                    print("대조하지 않은 보고서 열:", ", ".join(st["unmapped"]))
                if st.get("skipped_sheets"):
                    print("행 대조하지 않은 보고서 시트:", ", ".join(st["skipped_sheets"]))
                print(f"최종: 행 대조 — 값 차이 {cnt['값 차이']}칸({diff_rows}행) / 원본에 없음 {cnt['원본에 없음']}행 / 보고서에 없음 {cnt['보고서에 없음']}건 / 확인 필요 {cnt['확인 필요']}건 / 일치 {cnt['일치']}행 | 수식: {formula_status}")
                return 0
            if not result['cells']:
                raise ValueError('검산할 보고서 숫자 칸이 없습니다. 보고서·시트 설정을 확인하세요')
            _progress(f"  보고서 숫자 칸 {result['cells']}개 대조 완료. 결과 파일 저장 중")
            out = run(raw_path, report_path, result=result)
            _progress(f"원본 대조 끝 ({time.monotonic() - _t:.1f}초)")
            for line in summary(result):
                print(line)
            print('저장:', out)
            formula_status = _formula_output(report_path)
            blocked = _blocked(result)
            missing, nocache = _missing(result), _nocache(result)
            if result['matched'] == 0 and nocache and not missing and not blocked:
                print(f"최종: 계산값 없음 {nocache}칸 — 수식 칸에 저장된 계산값이 없음. 엑셀에서 열어 저장한 뒤 다시 실행하세요 | 수식: {formula_status}")
            elif result['matched'] == 0 and missing and not blocked:
                print(f"최종: 보고서 숫자 칸이 비어 있음 — 누락 {missing}칸" + (f" / 계산값 없음 {nocache}칸" if nocache else "")
                      + f" (빈 양식이면 --채우기로 원본 값을 채운 새 파일을 만드세요) | 수식: {formula_status}")
            elif result['matched'] == 0 and blocked:
                print(f"최종: 검산 불가 {blocked}칸 — 원본 금액 칸에 오류값·글자가 있어 집계하지 않음 (위치는 결과 파일 '근거없음·제외' 시트) / 근거 없음 {result['no_basis']}칸 | 수식: {formula_status}")
            elif result['matched'] == 0:
                print(f"최종: 대조 불가 — 보고서 숫자 칸 {result['cells']}개 중 원본과 맞는 칸 0개. 두 파일이 원본·보고서 짝인지 확인하세요 | 수식: {formula_status}")
            else:
                print(f"최종: 불일치 {len(result['mismatch'])}칸" + (f" / 누락 {missing}칸" if missing else "") + (f" / 계산값 없음 {nocache}칸" if nocache else "") + f" / 근거 없음 {result['no_basis']}칸" + (f" / 검산 불가 {blocked}칸" if blocked else "") + f" / 사람 확인 {result['review_ratio']:.1%} | 수식: {formula_status}")
        return 0
    except Exception as exc:
        print(f'최종: 실행 불가 — {type(exc).__name__}: {exc}')
        return 1


if __name__ == '__main__':
    if len(sys.argv) == 3 and sys.argv[1] == '--excel-worker':
        if not USE_EXCEL:
            print('최종: Excel 실행 설정이 꺼져 있습니다 (USE_EXCEL=False)')
            sys.exit(1)
        sys.exit(excel_worker_main(sys.argv[2]))
    sys.exit(main())
