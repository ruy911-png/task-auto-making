# -*- coding: utf-8 -*-
"""Independent acceptance tests for excel_automation v24 (issue #5 + reviewer comment).

TOOL      module name   (default excel_automation_v24)
TOOL_DIR  module folder (default: excel_recon, two folders up)
Inputs are built in temp dirs under the system temp folder (v24ind_work). Expectations are computed
from the generated data here, never from the engine.
"""
import datetime
import importlib
import os
import re
import subprocess
import sys
import tempfile
import unittest
import zipfile

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
TOOL = os.environ.get("TOOL", "excel_automation_v24")
TOOL_DIR = os.path.abspath(os.environ.get("TOOL_DIR", os.path.dirname(os.path.dirname(HERE))))
WORK = os.path.join(tempfile.gettempdir(), "v24ind_work")
os.makedirs(WORK, exist_ok=True)
sys.path.insert(0, TOOL_DIR)
E = importlib.import_module(TOOL)
SCRIPT = os.path.join(TOOL_DIR, TOOL + ".py")


# ----------------------------------------------------------------- helpers
def tmpdir():
    return tempfile.mkdtemp(dir=WORK)


def save(path, sheets):
    """sheets: list of (name, rows)"""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in sheets:
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append(list(r))
    wb.save(path)
    return path


def b1(path):
    wb = openpyxl.load_workbook(path, read_only=True)
    v = wb["요약"]["B1"].value
    wb.close()
    return v


# ---- aggregate fixtures
def agg_case(raw_rows, report_rows, raw_head=("일자", "계정", "금액"), report_sheet="보고서"):
    d = tmpdir()
    raw = save(os.path.join(d, "원본.xlsx"), [("원장", [list(raw_head)] + raw_rows)])
    rep = save(os.path.join(d, "보고서.xlsx"), [(report_sheet, report_rows)])
    return raw, rep


D0 = datetime.date(2026, 9, 1)


def dt(i):
    return D0 + datetime.timedelta(days=i)


# ---- row-mode fixtures
HEAD = ["일자", "거래처", "품목", "수량", "단가", "금액", "승인번호"]
NAMES = ["가나상사", "다라물산", "마바테크", "사아유통", "자차식품", "카타건설", "파하화학", "한빛전자",
         "동해수산", "서울제약", "부산철강", "대구섬유", "인천물류", "광주기계", "대전바이오", "울산조선",
         "세종미디어", "수원정밀", "창원중공", "제주관광", "포항소재", "전주한지", "청주반도", "춘천레저"]


def raw_record(i):
    qty = 10 + i
    price = 1000 + 37 * i
    return [dt(i), NAMES[i % len(NAMES)], ["볼트", "너트", "와셔"][i % 3], qty, price, qty * price,
            str(4000123412340000 + i * 7919)]


def raw_table(n):
    return [HEAD] + [raw_record(i) for i in range(n)]


def report_rows_for(idx):
    return [HEAD] + [raw_record(i) for i in idx]


def rows_case(n_raw, sheets, extra_sheets=()):
    """sheets: list of (name, list of raw indices). Returns (raw, report)."""
    d = tmpdir()
    raw = save(os.path.join(d, "원본.xlsx"), [("거래", raw_table(n_raw))])
    rep = save(os.path.join(d, "보고서.xlsx"),
               [(nm, report_rows_for(idx)) for nm, idx in sheets] + list(extra_sheets))
    return raw, rep


def raw_excel_row(i):
    return i + 2       # header in row 1


def counts(res):
    c = {}
    for x in res:
        c[x["판정"]] = c.get(x["판정"], 0) + 1
    return c


def dup_flags(res):
    return [x for x in res if x["판정"] == "확인 필요" and "중복" in (x.get("이유") or "")]


# ----------------------------------------------------------------- 1. aggregate mode
class T1Aggregate(unittest.TestCase):
    def test_1a_blank_with_raw_amount_is_missing(self):
        raw_rows = [[dt(0), "현금", 100], [dt(1), "현금", 50], [dt(2), "예금", 70], [dt(3), "예금", 30],
                    [dt(4), "매출채권", 300], [dt(5), "재고자산", 250], [dt(6), "재고자산", 150]]
        rep_rows = [["계정", "금액"], ["현금", 150], ["예금", None], ["매출채권", 300], ["재고자산", 400]]
        raw, rep = agg_case(raw_rows, rep_rows)
        r = E.verify(raw, rep)
        miss = [u for u in r["unmatched"] if u["구분"] == "누락"]
        self.assertEqual([u["위치"] for u in miss], ["보고서!B3"], r["unmatched"])
        self.assertEqual(len(r["mismatch"]), 0)
        self.assertEqual(r["cells"], 4)
        self.assertEqual(b1(E.run(raw, rep, r)), "확인 필요 1칸")

    def test_1b_blank_with_raw_zero_not_flagged(self):
        raw_rows = [[dt(0), "현금", 100], [dt(1), "현금", 50], [dt(2), "예금", 100], [dt(3), "예금", -100],
                    [dt(4), "매출채권", 300]]
        rep_rows = [["계정", "금액"], ["현금", 150], ["예금", None], ["매출채권", 300]]
        raw, rep = agg_case(raw_rows, rep_rows)
        r = E.verify(raw, rep)
        self.assertEqual(r["unmatched"], [])
        self.assertEqual(b1(E.run(raw, rep, r)), "전부 일치")

    def test_1c_blank_without_raw_basis_is_layout_gap(self):
        raw_rows = [[dt(0), "현금", 100], [dt(1), "현금", 50], [dt(2), "예금", 200], [dt(3), "매출채권", 300]]
        rep_rows = [["계정", "금액", "비고"],
                    ["현금", 150, None],
                    ["예금", 200, None],
                    [None, None, None],          # spacer row (dropped by openpyxl append? keep explicit)
                    ["매출채권", 300, None],
                    ["대여금", None, None]]      # account not in raw at all
        raw, rep = agg_case(raw_rows, rep_rows)
        r = E.verify(raw, rep)
        self.assertEqual([u for u in r["unmatched"] if u["구분"] == "누락"], [])
        self.assertEqual(b1(E.run(raw, rep, r)), "전부 일치", r["unmatched"])

    # two-level (구분/계정명) x 법인 report with blank 소계/합계 rows
    RAW2 = [["유동", "현금", "A사", 100], ["유동", "현금", "B사", 40], ["유동", "예금", "A사", 200],
            ["유동", "예금", "B사", 60], ["비유동", "토지", "A사", 500], ["비유동", "토지", "B사", 80],
            ["비유동", "건물", "A사", 300], ["비유동", "건물", "B사", 20], ["유동", "현금", "A사", 5]]

    def _two_level(self, with_totals, filled_totals=False):
        agg = {}
        for g, a, c, v in self.RAW2:
            agg[(a, c)] = agg.get((a, c), 0) + v
        def tot(accts, c):
            return sum(agg[(a, c)] for a in accts) if filled_totals else None
        rows = [["구분", "계정명", "A사", "B사"]]
        rows += [["유동", a, agg[(a, "A사")], agg[(a, "B사")]] for a in ("현금", "예금")]
        if with_totals:
            rows.append(["유동", "소계", tot(("현금", "예금"), "A사"), tot(("현금", "예금"), "B사")])
        rows += [["비유동", a, agg[(a, "A사")], agg[(a, "B사")]] for a in ("토지", "건물")]
        if with_totals:
            rows.append(["비유동", "소계", tot(("토지", "건물"), "A사"), tot(("토지", "건물"), "B사")])
            rows.append(["합계", None, tot(("현금", "예금", "토지", "건물"), "A사"), tot(("현금", "예금", "토지", "건물"), "B사")])
        return agg_case([list(x) for x in self.RAW2], rows, raw_head=("구분", "계정명", "법인", "금액"))

    def test_1d_mapping_same_with_blank_totals(self):
        r_blank = E.verify(*self._two_level(True))
        r_none = E.verify(*self._two_level(False))
        r_full = E.verify(*self._two_level(True, filled_totals=True))
        self.assertEqual(r_blank["mapping"], r_none["mapping"])
        self.assertEqual(r_blank["mapping"], r_full["mapping"])
        self.assertTrue(r_full["passed"], r_full["unmatched"] + r_full["mismatch"])   # control: correct report is clean
        self.assertEqual(r_none["unmatched"], [])

    def test_1e_blank_subtotal_total_with_detail_is_missing(self):
        raw, rep = self._two_level(True)
        r = E.verify(raw, rep)
        miss = sorted(u["위치"] for u in r["unmatched"] if u["구분"] == "누락")
        # 소계 rows 4 and 7, 합계 row 8, columns C/D
        self.assertEqual(miss, sorted(["보고서!C4", "보고서!D4", "보고서!C7", "보고서!D7", "보고서!C8", "보고서!D8"]))
        self.assertEqual(len(r["mismatch"]), 0)
        self.assertEqual(r["matched"], 8)
        self.assertEqual(b1(E.run(raw, rep, r)), "확인 필요 6칸")

    def test_1f_blank_total_column_is_missing(self):
        raw_rows = [[dt(0), "현금", "A사", 100], [dt(1), "현금", "B사", 40], [dt(2), "예금", "A사", 200], [dt(3), "예금", "B사", 60]]
        rep_rows = [["계정", "A사", "B사", "합계"], ["현금", 100, 40, None], ["예금", 200, 60, 260]]
        raw, rep = agg_case(raw_rows, rep_rows, raw_head=("일자", "계정", "법인", "금액"))
        r = E.verify(raw, rep)
        self.assertEqual([u["위치"] for u in r["unmatched"] if u["구분"] == "누락"], ["보고서!D2"], r["unmatched"])

    # ---- probes: layout gaps inside a subtotal block must not poison the subtotal
    def test_1g_probe_no_basis_blank_inside_subtotal_block(self):
        raw_rows = [[dt(0), "현금", 150], [dt(1), "예금", 100], [dt(2), "토지", 500]]
        rep_rows = [["계정", "금액"], ["현금", 150], ["예금", 100], ["대여금", None], ["소계", 250],
                    ["토지", 500], ["합계", 750]]
        raw, rep = agg_case(raw_rows, rep_rows)
        r = E.verify(raw, rep)
        self.assertEqual(r["unmatched"], [], "blank no-basis row made 소계/합계 unverifiable")
        self.assertEqual(b1(E.run(raw, rep, r)), "전부 일치")

    def test_1h_probe_section_heading_row_without_basis(self):
        raw_rows = [[dt(0), "현금", 150], [dt(1), "예금", 100], [dt(2), "토지", 500]]
        rep_rows = [["계정", "금액"], ["Ⅰ. 유동자산", None], ["현금", 150], ["예금", 100], ["소계", 250],
                    ["Ⅱ. 비유동자산", None], ["토지", 500], ["소계", 500], ["합계", 750]]
        raw, rep = agg_case(raw_rows, rep_rows)
        r = E.verify(raw, rep)
        self.assertEqual(r["unmatched"], [])
        self.assertEqual(b1(E.run(raw, rep, r)), "전부 일치")

    def test_1i_probe_section_heading_row_with_raw_group_label(self):
        # heading text equals a raw 구분 value; its blank amount is a heading, not an omission
        raw_rows = [["유동자산", "현금", 150], ["유동자산", "예금", 100], ["비유동자산", "토지", 500]]
        rep_rows = [["계정", "금액"], ["유동자산", None], ["현금", 150], ["예금", 100], ["소계", 250],
                    ["비유동자산", None], ["토지", 500], ["소계", 500], ["합계", 750]]
        raw, rep = agg_case(raw_rows, rep_rows, raw_head=("구분", "계정", "금액"))
        r = E.verify(raw, rep)
        self.assertEqual(len(r["mismatch"]), 0, r["mismatch"])
        self.assertEqual(r["unmatched"], [])

    def test_1j_probe_wrong_subtotal_not_judged_match(self):
        # 소계 is wrong (500, true 250). The cell itself must be 불일치, not 일치.
        raw_rows = [["유동자산", "현금", 150], ["유동자산", "예금", 100], ["비유동자산", "토지", 500]]
        rep_rows = [["계정", "금액"], ["유동자산", None], ["현금", 150], ["예금", 100], ["소계", 500],
                    ["비유동자산", None], ["토지", 500], ["소계", 500], ["합계", 750]]
        raw, rep = agg_case(raw_rows, rep_rows, raw_head=("구분", "계정", "금액"))
        r = E.verify(raw, rep)
        b5 = next(c for c in r["checks"] if c["위치"] == "보고서!B5")
        self.assertEqual(b5["판정"], "불일치", b5)


# ----------------------------------------------------------------- 2. row mode skipped sheet
class T2RowsSkipped(unittest.TestCase):
    def test_2a_control_clean(self):
        raw, rep = rows_case(20, [("거래내역", range(20))])
        out, res, st, cnt, _ = E.run_rows(raw, rep)
        self.assertEqual(cnt["일치"], 20)
        self.assertEqual(b1(out), "전부 일치")

    def test_2b_memo_sheet_table_blocks_all_match(self):
        memo = ("메모", [["항목", "내용"], ["작성자", "홍길동"], ["검토자", "김철수"], ["비고", "3분기 마감분"]])
        raw, rep = rows_case(20, [("거래내역", range(20))], [memo])
        out, res, st, cnt, _ = E.run_rows(raw, rep)
        self.assertEqual(cnt["일치"], 20)
        self.assertNotEqual(b1(out), "전부 일치")

    def test_2c_probe_memo_sheet_single_column(self):
        memo = ("메모", [["3분기 거래내역 정리"], ["작성: 홍길동"], ["원본 대비 누락 없음 확인 바람"]])
        raw, rep = rows_case(20, [("거래내역", range(20))], [memo])
        out, res, st, cnt, _ = E.run_rows(raw, rep)
        self.assertNotEqual(b1(out), "전부 일치", "sheet not examined at all, yet verdict is 전부 일치")


# ----------------------------------------------------------------- 3. _row_norm / _row_eq
class T3RowEq(unittest.TestCase):
    def eq(self, a, b):
        return E._row_eq(E._row_norm(a), E._row_norm(b))

    def test_3a_16_digits_plain(self):
        self.assertFalse(self.eq("1234567890123456", "1234567890123457"))
        self.assertTrue(self.eq("1234567890123456", "1234567890123456"))

    def test_3b_16_digits_commas(self):
        self.assertFalse(self.eq("1,234,567,890,123,456", "1,234,567,890,123,457"))
        self.assertTrue(self.eq("1,234,567,890,123,456", "1,234,567,890,123,456"))

    def test_3c_16_digits_spaces(self):
        self.assertFalse(self.eq("1234 5678 9012 3456", "1234 5678 9012 3457"))

    def test_3d_20_digits(self):
        self.assertFalse(self.eq("12345678901234567890", "12345678901234567891"))
        self.assertFalse(self.eq("12,345,678,901,234,567,890", "12,345,678,901,234,567,891"))

    def test_3e_normal_amounts(self):
        self.assertTrue(self.eq("1,000", 1000))
        self.assertTrue(self.eq(" 1,000 ", 1000.0))
        self.assertFalse(self.eq("1,001", 1000))
        self.assertFalse(self.eq("123456789012345", "123456789012346"))   # 15 digits, numeric path

    def test_3f_leading_zero_codes(self):
        self.assertTrue(self.eq("005930", 5930))

    def test_3g_probe_same_16_digit_number_formatting_only(self):
        # same number, only thousands separators / grouping differ -> should not be a value difference
        self.assertTrue(self.eq("1,234,567,890,123,456", "1234567890123456"))
        self.assertTrue(self.eq("1234 5678 9012 3456", "1234567890123456"))

    def test_3h_probe_negative_17_digits(self):
        self.assertFalse(self.eq("-12,345,678,901,234,567", "-12,345,678,901,234,568"))

    def test_3i_end_to_end_comma_number_last_digit(self):
        d = tmpdir()
        head = ["일자", "거래처", "금액", "승인번호"]
        def rec(i):
            n = 1234567890123000 + i * 101
            return [dt(i), NAMES[i], 5000 + 13 * i, f"{n:,}"]
        raw_rows = [rec(i) for i in range(12)]
        rep_rows = [list(x) for x in raw_rows]
        n7 = 1234567890123000 + 7 * 101 + 1
        rep_rows[7][3] = f"{n7:,}"
        raw = save(os.path.join(d, "원본.xlsx"), [("거래", [head] + raw_rows)])
        rep = save(os.path.join(d, "보고서.xlsx"), [("정리", [head] + rep_rows)])
        res = E.reconcile_rows(raw, rep)
        diffs = [x for x in res if x["판정"] == "값 차이"]
        self.assertEqual([(x["보고서 행"], x["열"]) for x in diffs], [(9, "승인번호")], diffs)
        self.assertEqual(counts(res).get("일치"), 11)


# ----------------------------------------------------------------- 4. split report sheets
class T4Split(unittest.TestCase):
    def run_case(self, n_raw, sheets):
        raw, rep = rows_case(n_raw, sheets)
        out, res, st, cnt, _ = E.run_rows(raw, rep)
        return res, b1(out)

    def assert_no_silent_double(self, res):
        by_raw = {}
        for x in res:
            if x["원본 행"] is not None and x["보고서 행"] is not None:
                by_raw.setdefault(x["원본 행"], {}).setdefault((x["시트"], x["보고서 행"]), set()).add(x["판정"])
        for k, reps in by_raw.items():
            if len(reps) > 1:
                flagged = sum(1 for v in reps.values() if v != {"일치"})
                self.assertGreaterEqual(flagged, len(reps) - 1, f"raw row {k} matched {len(reps)}x silently: {reps}")

    def test_4a_10_10(self):
        res, v = self.run_case(20, [("1", range(10)), ("2", range(10, 20))])
        c = counts(res)
        self.assertEqual(c.get("일치"), 20)
        self.assertEqual(c.get("보고서에 없음", 0), 0)
        self.assertEqual(v, "전부 일치")

    def test_4b_10_9(self):
        res, v = self.run_case(20, [("1", range(10)), ("2", [i for i in range(10, 20) if i != 15])])
        c = counts(res)
        self.assertEqual(c.get("일치"), 19)
        missing = [x for x in res if x["판정"] == "보고서에 없음"]
        self.assertEqual([x["원본 행"] for x in missing], [raw_excel_row(15)])
        self.assertNotEqual(v, "전부 일치")

    def test_4c_10_12_two_dups(self):
        res, v = self.run_case(20, [("1", range(10)), ("2", [3, 7] + list(range(10, 20)))])
        self.assertNotEqual(v, "전부 일치")
        flagged_raw = {x["원본 행"] for x in dup_flags(res)}
        self.assertEqual(flagged_raw, {raw_excel_row(3), raw_excel_row(7)})
        self.assert_no_silent_double(res)
        self.assertEqual(counts(res).get("보고서에 없음", 0), 0)

    def test_4d_two_full_independent(self):
        res, v = self.run_case(20, [("원화", range(20)), ("달러", range(20))])
        self.assertEqual(dup_flags(res), [])
        self.assertEqual(counts(res).get("일치"), 40)
        self.assertEqual(v, "전부 일치")

    def test_4e_three_sheets_clean(self):
        res, v = self.run_case(20, [("1", range(7)), ("2", range(7, 14)), ("3", range(14, 20))])
        self.assertEqual(counts(res).get("일치"), 20)
        self.assertEqual(dup_flags(res), [])
        self.assertEqual(v, "전부 일치")

    def test_4f_three_sheets_one_dup(self):
        res, v = self.run_case(20, [("1", range(7)), ("2", range(7, 14)), ("3", [2] + list(range(14, 20)))])
        self.assertNotEqual(v, "전부 일치")
        self.assertEqual({x["원본 행"] for x in dup_flags(res)}, {raw_excel_row(2)})
        self.assert_no_silent_double(res)

    def test_4g_probe_half_overlap_split(self):
        # 15 raw rows, sheet1 = 0..9, sheet2 = 5..14: 5 rows reported twice
        res, v = self.run_case(15, [("1", range(10)), ("2", range(5, 15))])
        self.assertNotEqual(v, "전부 일치", "5 raw rows appear on both sheets but verdict is 전부 일치")
        self.assert_no_silent_double(res)

    def test_4h_probe_two_thirds_overlap_split(self):
        res, v = self.run_case(20, [("1", range(15)), ("2", range(5, 20))])
        self.assertNotEqual(v, "전부 일치", "10 raw rows on both sheets but verdict is 전부 일치")
        self.assert_no_silent_double(res)

    def test_4i_probe_full_plus_subset(self):
        # '전체' + '주요거래' (subset): legitimate layout, no duplicate flags expected
        res, v = self.run_case(20, [("전체", range(20)), ("주요거래", [1, 4, 9, 12, 18])])
        self.assertEqual(dup_flags(res), [])
        self.assertEqual(counts(res).get("일치"), 25)

    def test_4j_probe_independent_reports_one_missing(self):
        # 원화 sheet complete, 달러 sheet misses raw row 6 -> omission must surface
        res, v = self.run_case(20, [("원화", range(20)), ("달러", [i for i in range(20) if i != 6])])
        self.assertNotEqual(v, "전부 일치", "달러 sheet lacks a raw row but verdict is 전부 일치")

    def test_4k_probe_dup_inside_same_sheet(self):
        res, v = self.run_case(20, [("1", list(range(20)) + [5])])
        self.assertNotEqual(v, "전부 일치")
        self.assertEqual(counts(res).get("일치"), 20)
        self.assert_no_silent_double(res)

    def test_4m_probe_monthly_split_plus_full_sheet(self):
        # '1월' + '2월' split plus '전체' sheet holding everything: no duplicate flags expected
        res, v = self.run_case(20, [("1월", range(10)), ("2월", range(10, 20)), ("전체", range(20))])
        self.assertEqual(dup_flags(res), [])
        self.assertEqual(counts(res).get("일치"), 40)

    def test_4l_probe_split_with_one_dup_and_one_missing(self):
        res, v = self.run_case(20, [("1", range(10)), ("2", [4] + [i for i in range(10, 20) if i != 13])])
        self.assertNotEqual(v, "전부 일치")
        self.assertEqual([x["원본 행"] for x in res if x["판정"] == "보고서에 없음"], [raw_excel_row(13)])
        self.assertEqual({x["원본 행"] for x in dup_flags(res)}, {raw_excel_row(4)})


# ----------------------------------------------------------------- 5. formula status
def formula_book(path, cells, cached):
    """cells: {coord: value or '=formula'}; cached: {coord: number} injected as saved values"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    for k, v in cells.items():
        ws[k] = v
    wb.save(path)
    with zipfile.ZipFile(path) as z:
        items = {n: z.read(n) for n in z.namelist()}
    xml = items["xl/worksheets/sheet1.xml"].decode("utf-8")
    for coord, val in cached.items():
        xml, n = re.subn(rf'(<c r="{coord}"[^>]*><f>[^<]*</f>)<v\s*/?>(</v>)?', rf"\g<1><v>{val}</v>", xml)
        assert n == 1, coord
    items["xl/worksheets/sheet1.xml"] = xml.encode("utf-8")
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for n, b in items.items():
            z.writestr(n, b)
    return path


class T5Formula(unittest.TestCase):
    def status(self, cells, cached):
        p = formula_book(os.path.join(tmpdir(), "수식.xlsx"), cells, cached)
        return E._formula_output(p)

    def base(self):
        return {f"A{i}": i for i in range(1, 6)}

    def test_5a_three_uncached(self):
        cells = dict(self.base(), B1="=SUM(A1:A5)", B2="=A1+A2", B3="=SUM(A1:A3)")
        s = self.status(cells, {})
        self.assertIn("재검산 불가 3칸", s)

    def test_5b_two_of_four_uncached(self):
        cells = dict(self.base(), B1="=SUM(A1:A5)", B2="=A1+A2", B3="=SUM(A1:A3)", B4="=A4+A5")
        s = self.status(cells, {"B1": 15, "B2": 3})
        self.assertIn("재검산 불가 2칸", s)
        self.assertNotIn("오류", s.split("(")[0].replace("오류·주의 없음", ""))

    def test_5c_all_cached_ok(self):
        cells = dict(self.base(), B1="=SUM(A1:A5)", B2="=A1+A2")
        s = self.status(cells, {"B1": 15, "B2": 3})
        self.assertNotIn("재검산 불가", s)
        self.assertNotIn("제한", s)

    def test_5e_two_indirect_cached(self):
        cells = dict(self.base(), B1='=INDIRECT("A1")', B2='=INDIRECT("A2")', B3="=SUM(A1:A5)")
        s = self.status(cells, {"B1": 1, "B2": 2, "B3": 15})
        self.assertIn("재검산 불가 2칸", s)

    def test_5d_control_wrong_cache_detected(self):
        cells = dict(self.base(), B1="=SUM(A1:A5)", B2="=A1+A2")
        s = self.status(cells, {"B1": 16, "B2": 3})
        self.assertRegex(s, r"오류 [1-9]")


# ----------------------------------------------------------------- 6/7. console encoding, coding line
class T6Encoding(unittest.TestCase):
    def run_script(self, args):
        env = dict(os.environ, PYTHONIOENCODING="cp949")
        p = subprocess.run([sys.executable, SCRIPT] + args, stdout=subprocess.PIPE, stderr=subprocess.PIPE,
                           env=env, cwd=tmpdir(), timeout=600)
        out = p.stdout.decode("cp949", "replace")
        return p.returncode, out, p.stderr.decode("cp949", "replace")

    def test_6a_aggregate_pipe(self):
        raw, rep = agg_case([[dt(0), "현금", 100], [dt(1), "예금", 200]], [["계정", "금액"], ["현금", 100], ["예금", None]])
        rc, out, err = self.run_script([raw, rep])
        self.assertEqual(rc, 0, out[-800:] + err[-800:])
        self.assertIn("최종", out)

    def test_6b_formula_pipe(self):
        p = formula_book(os.path.join(tmpdir(), "수식.xlsx"), {"A1": 1, "A2": 2, "B1": "=SUM(A1:A2)"}, {})
        rc, out, err = self.run_script([p])
        self.assertEqual(rc, 0, out[-800:] + err[-800:])
        self.assertIn("최종", out)

    def test_6c_rows_pipe(self):
        raw, rep = rows_case(20, [("1", range(10)), ("2", range(10, 20))])
        rc, out, err = self.run_script([raw, rep])
        self.assertEqual(rc, 0, out[-800:] + err[-800:])
        self.assertIn("최종", out)


class T7Coding(unittest.TestCase):
    def test_7_first_line_utf8(self):
        with open(SCRIPT, "rb") as f:
            first = f.readline().decode("ascii", "replace")
        m = re.match(r"^[ \t\f]*#.*?coding[:=][ \t]*([-\w.]+)", first)
        self.assertIsNotNone(m, first)
        self.assertIn(m.group(1).lower().replace("_", "-"), ("utf-8", "utf8"))


if __name__ == "__main__":
    unittest.main(verbosity=2)
