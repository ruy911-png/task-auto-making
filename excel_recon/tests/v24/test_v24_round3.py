# 3차 검토(독립 리뷰어) 지적 고정. 실행: TOOL_DIR=<폴더> python3 test_v24_round3.py
import os, sys, io, re, subprocess, unittest, collections, datetime
from openpyxl import load_workbook
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from test_v24_review import m, wx, quiet, ROWS, H, TOOL_DIR, Review
D = datetime.date


class Round3(Review):
    def agg(self, raw_rows, rep_rows, name="a"):
        raw = wx(os.path.join(self.d, f"r_{name}.xlsx"), [("원본", raw_rows)])
        rep = wx(os.path.join(self.d, f"p_{name}.xlsx"), [("보고서", rep_rows)])
        r = quiet(m.verify, raw, rep)
        b1 = load_workbook(quiet(m.run, raw, rep, r))["요약"]["B1"].value
        return r, b1, {u["위치"]: u["구분"] for u in r["unmatched"]}, {x["위치"] for x in r["mismatch"]}

    # --- 빈 칸이 아래 칸 합과 우연히 같아도 누락
    def test_a1_no_subtotal_equal_below(self):
        r, b1, u, mm = self.agg([["계정과목", "금액"], ["현금", 300], ["예금", 100], ["적금", 200]],
                                [["계정과목", "금액"], ["현금", None], ["예금", 100], ["적금", 200]], "a1")
        self.assertEqual(u.get("보고서!B2"), "누락", u); self.assertNotEqual(b1, "전부 일치")

    def test_a2_blank_equals_next_and_total_wrong(self):
        r, b1, u, mm = self.agg([["계정과목", "금액"], ["보험료", 500], ["임차료", 1000], ["관리비", 1000]],
                                [["계정과목", "금액"], ["보험료", 500], ["임차료", None], ["관리비", 1000], ["합계", 1500]], "a2")
        self.assertEqual(u.get("보고서!B3"), "누락", u); self.assertIn("보고서!B5", mm)

    def test_a3_monthly_fixed_amount(self):
        rr = [["일자", "계정", "금액"]] + [[D(2026, k, 25), "임차료", 1000000] for k in range(1, 13)]
        for blanks, tot, nm in (((11,), 11000000, "a3"), ((10, 11), 10000000, "a3b")):
            pr = [["월", "임차료"]] + [[f"{k}월", None if k in blanks else 1000000] for k in range(1, 13)] + [["합계", tot]]
            r, b1, u, mm = self.agg(rr, pr, nm)
            for k in blanks:
                self.assertEqual(u.get(f"보고서!B{k + 1}"), "누락", (nm, u))
            self.assertIn("보고서!B14", mm, nm)

    def test_a5_total_column(self):
        r, b1, u, mm = self.agg([["계정과목", "구분", "금액"], ["현금", "A", 100], ["현금", "B", 50], ["예금", "A", 100], ["예금", "B", 70]],
                                [["계정과목", "A", "B", "합계"], ["현금", None, 50, 50], ["예금", 100, 70, 170]], "a5")
        self.assertEqual(u.get("보고서!B2"), "누락", u); self.assertIn("보고서!D2", mm)

    def test_a6_near_equal(self):
        r, b1, u, mm = self.agg([["계정과목", "금액"], ["현금", 1000000000], ["예금", 999999200]],
                                [["계정과목", "금액"], ["현금", None], ["예금", 999999200]], "a6")
        self.assertEqual(u.get("보고서!B2"), "누락", u)

    def test_a7_correct_subtotal_not_mismatch(self):
        r, b1, u, mm = self.agg([["계정과목", "금액"], ["현금", 100], ["예금", 200]],
                                [["계정과목", "금액"], ["현금", None], ["예금", 200], ["소계", 300]], "a7")
        self.assertEqual(u.get("보고서!B2"), "누락", u); self.assertEqual(mm, set())

    def test_a8_heading_rows_with_total_column_ok(self):
        raw = [["대분류", "계정과목", "구분", "금액"], ["유동자산", "현금", "A", 100], ["유동자산", "현금", "B", 50], ["유동자산", "예금", "A", 200],
               ["유동자산", "예금", "B", 70], ["비유동자산", "토지", "A", 500], ["비유동자산", "토지", "B", 30]]
        for h1, h2, nm in (("유동자산", "비유동자산", "a8"), ("I. 유동", "II. 비유동", "a8b")):
            rep = [["계정과목", "A", "B", "합계"], [h1, None, None, None], ["현금", 100, 50, 150], ["예금", 200, 70, 270], ["소계", 300, 120, 420],
                   [h2, None, None, None], ["토지", 500, 30, 530], ["소계", 500, 30, 530], ["합계", 800, 150, 950]]
            r, b1, u, mm = self.agg(raw, rep, nm)
            self.assertEqual(b1, "전부 일치", (nm, u, mm))

    RAW8 = [["대분류", "계정과목", "구분", "금액"], ["유동자산", "현금", "A", 100], ["유동자산", "현금", "B", 50], ["유동자산", "예금", "A", 200],
            ["유동자산", "예금", "B", 70], ["비유동자산", "토지", "A", 500], ["비유동자산", "토지", "B", 30]]

    def test_a9_group_row_partly_filled_blank_is_missing(self):
        r, b1, u, mm = self.agg(self.RAW8, [["계정과목", "A", "B"], ["유동자산", None, 120], ["현금", 100, 50], ["예금", 200, 70]], "a9")
        self.assertEqual(u.get("보고서!B2"), "누락", u)

    def test_a10_heading_whose_rows_are_incomplete(self):
        r, b1, u, mm = self.agg(self.RAW8, [["계정과목", "A", "B"], ["유동자산", None, None], ["현금", 100, 50], ["비유동자산", None, None], ["토지", 500, 30]], "a10")
        self.assertEqual(u.get("보고서!B2"), "누락", u); self.assertNotIn("보고서!B4", u)

    # --- 행 대조: 시트 관계
    def test_b1_uneven_splits_ok(self):
        big = [(D(2026, 1 + i % 12, 1 + i % 27), f"계정{i % 5}", 1000 * (i + 1), f"적요{i}") for i in range(100)]
        self.raw = wx(os.path.join(self.d, "raw100.xlsx"), [("원본", [H] + big)])
        for parts in ((90, 10), (92, 8), (90, 5, 5)):
            sheets, a = [], 0
            for i, n in enumerate(parts):
                sheets.append((f"s{i}", [H] + big[a:a + n])); a += n
            rep, res, c = self.rows_of(sheets)
            self.assertEqual(self.final(rep), "전부 일치", (parts, c))

    def test_b5_small_split_sheet_is_reconciled(self):
        big = [(D(2026, 1 + i % 12, 1 + i % 27), f"계정{i % 5}", 1000 * (i + 1), f"적요{i}") for i in range(100)]
        self.raw = wx(os.path.join(self.d, "raw100b.xlsx"), [("원본", [H] + big)])
        rep, res, c = self.rows_of([("상", [H] + big[:90]), ("중", [H] + big[90:95]), ("하", [H] + big[95:99])])
        self.assertEqual(c["보고서에 없음"], 1, c)
        self.assertEqual([x["원본 행"] for x in res if x["판정"] == "보고서에 없음"], [101])
        rep, res, c = self.rows_of([("상", [H] + big[:96]), ("하", [H] + big[96:])])
        self.assertEqual(self.final(rep), "전부 일치", c)

    def test_b9_parallel_partial_sheets_not_duplicates(self):
        rep, res, c = self.rows_of([("원화기준", [H] + ROWS[:12]), ("달러기준", [H] + ROWS[:12])])
        self.assertFalse([x for x in res if "중복" in x["이유"]], c)
        self.assertEqual(c["보고서에 없음"], 8, c)

    # --- 숫자 비교
    def _rows_final(self, raw_rows, rep_rows, nm):
        rawp = wx(os.path.join(self.d, f"r{nm}.xlsx"), [("원본", raw_rows)]); repp = wx(os.path.join(self.d, f"p{nm}.xlsx"), [("정리", rep_rows)])
        return load_workbook(quiet(m.run_rows, rawp, repp)[0])["요약"]["B1"].value

    def test_g_float_product_vs_decimal(self):
        import random
        from decimal import Decimal
        rnd = random.Random(7); h = ["일자", "거래처", "달러", "환율", "원화"]; rr, pr = [h], [h]
        for i in range(30):
            u = Decimal(rnd.randint(100_000_00, 9_999_999_99)) / 100; r = Decimal(rnd.randint(1200_0, 1450_0)) / 10
            rr.append([D(2026, 9, 1 + i % 28), f"거래처{i}", float(u), float(r), float(u * r)])
            pr.append([D(2026, 9, 1 + i % 28), f"거래처{i}", float(u), float(r), float(u) * float(r)])
        self.assertEqual(self._rows_final(rr, pr, "g"), "전부 일치")

    def test_g2_one_won_difference_still_caught(self):
        h = ["일자", "거래처", "금액"]
        rr = [h] + [[D(2026, 9, 1 + i), f"거래처{i}", 1234567890123 + i] for i in range(10)]
        pr = [h] + [[D(2026, 9, 1 + i), f"거래처{i}", 1234567890123 + i + (1 if i == 4 else 0)] for i in range(10)]
        self.assertNotEqual(self._rows_final(rr, pr, "g2"), "전부 일치")

    def test_h_long_decimal_text_vs_number(self):
        h = ["일자", "법인", "계정명", "금액"]; rr, pr = [h], [h]
        for i in range(20):
            amt = 45_123_456_789_012 + i if i % 5 == 0 else 1000 * (i + 1)
            rr.append([D(2026, 9, 1 + i), f"법인{i % 3}", f"계정{i}", f"{amt:,.2f}"])
            pr.append([D(2026, 9, 1 + i), f"법인{i % 3}", f"계정{i}", amt])
        self.assertEqual(self._rows_final(rr, pr, "h"), "전부 일치")

    # --- 안내 문구
    def test_i4_all_formulas_without_cache(self):
        raw = wx(os.path.join(self.d, "ri.xlsx"), [("원본", [["계정과목", "구분", "금액"], ["현금", "A", 100], ["현금", "B", 200], ["예금", "A", 300], ["예금", "B", 400]])])
        f = lambda a, b: f'=SUMIFS(원본!C:C,원본!A:A,"{a}",원본!B:B,"{b}")'
        rep = wx(os.path.join(self.d, "pi.xlsx"), [("보고서", [["계정과목", "A", "B"], ["현금", f("현금", "A"), f("현금", "B")], ["예금", f("예금", "A"), f("예금", "B")]])])
        src = os.path.join(TOOL_DIR, os.environ.get("TOOL", "excel_automation_v24") + ".py")
        r = subprocess.run([sys.executable, src, raw, rep], capture_output=True, text=True, cwd=self.d, env=dict(os.environ, PYTHONIOENCODING="utf-8"))
        last = [l for l in r.stdout.splitlines() if l.startswith("최종")][-1]
        self.assertIn("계산값 없음 4칸", last); self.assertNotIn("--채우기", last)

    def test_j_title_and_header_template_sheet(self):
        rep = wx(os.path.join(self.d, "pj.xlsx"), [("거래", [H] + ROWS), ("양식", [["2월 거래 내역"], H])])
        self.assertEqual(self.final(rep), "전부 일치")


if __name__ == "__main__":
    unittest.main(defaultTest="Round3")

