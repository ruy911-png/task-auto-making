# PR #2 검토·이슈 #5 지적을 자동 시험으로 고정. 실행: 이 폴더에서 python3 test_v24_review.py (다른 판은 TOOL=모듈이름 TOOL_DIR=폴더)
import os, sys, io, re, contextlib, tempfile, datetime, importlib, subprocess, unittest, collections
from openpyxl import Workbook, load_workbook
HERE = os.path.dirname(os.path.abspath(__file__))
TOOL_DIR = os.environ.get("TOOL_DIR", os.path.dirname(os.path.dirname(HERE)))   # 기본: excel_recon 폴더
sys.path.insert(0, TOOL_DIR)
m = importlib.import_module(os.environ.get("TOOL", "excel_automation_v24"))
D = datetime.date
ROWS = [(D(2026, 1, 1 + i % 27), f"계정{i % 5}", 1000 * (i + 1), f"적요{i}") for i in range(20)]
H = ["일자", "계정명", "금액", "적요"]


def wx(path, sheets):
    wb = Workbook(); first = True
    for name, rows in sheets:
        ws = wb.active if first else wb.create_sheet(); ws.title = name; first = False
        for r in rows:
            ws.append(list(r))
    wb.save(path); return path


def quiet(fn, *a):
    with contextlib.redirect_stdout(io.StringIO()):
        return fn(*a)


class Review(unittest.TestCase):
    def setUp(self):
        self.d = tempfile.mkdtemp(prefix="v24_", dir=os.path.join(HERE, "work")) if os.path.isdir(os.path.join(HERE, "work")) else tempfile.mkdtemp()
        self.raw = wx(os.path.join(self.d, "raw.xlsx"), [("원본", [H] + ROWS)])

    def rows_of(self, sheets):
        rep = wx(os.path.join(self.d, "rep.xlsx"), sheets)
        res = quiet(m.reconcile_rows, self.raw, rep)
        return rep, res, collections.Counter(x["판정"] for x in res)

    def final(self, rep):
        out = quiet(m.run_rows, self.raw, rep)[0]
        return load_workbook(out)["요약"]["B1"].value

    # 1) 합계형 빈 칸
    def test_1_aggregate_blank_cell_is_missing(self):
        raw = wx(os.path.join(self.d, "r1.xlsx"), [("원본", [["계정명", "금액"], ["현금", 100], ["현금", 50], ["예금", 300], ["기타", 0]])])
        rep = wx(os.path.join(self.d, "p1.xlsx"), [("보고서", [["계정명", "금액"], ["현금", 150], ["예금", None], ["기타", None], ["없는계정", None]])])
        r = quiet(m.verify, raw, rep)
        kinds = {u["위치"]: u["구분"] for u in r["unmatched"]}
        self.assertEqual(kinds.get("보고서!B3"), "누락", kinds)
        self.assertNotIn("보고서!B4", kinds, "원본 집계 0인 빈 칸은 정상")
        self.assertNotIn("보고서!B5", kinds, "근거 없는 빈 칸은 여백")
        self.assertEqual(r["cells"], 2)
        out = quiet(m.run, raw, rep, r)
        self.assertNotEqual(load_workbook(out)["요약"]["B1"].value, "전부 일치")

    def test_1b_aggregate_mapping_unchanged_by_blanks(self):
        """빈 칸을 검사해도 행·열 대응은 값 있는 칸 기준 (13 짝과 같은 2단 표)"""
        raw = wx(os.path.join(self.d, "r1b.xlsx"), [("원본", [["법인", "구분", "계정과목", "금액"]] +
                 [(c, g, a, 1000 * (i + 1)) for i, (c, g, a) in enumerate((c, g, a) for c in ("A법인", "B법인") for g in ("유동", "비유동") for a in (("계정1", "계정2") if g == "유동" else ("계정3", "계정4")))])])
        rep_rows = [["구분", "계정명", "A법인", "B법인"], ["유동", "계정1", 1000, 5000], [None, "계정2", 2000, 6000], ["소계", None, None, None],
                    ["비유동", "계정3", 3000, 7000], [None, "계정4", 4000, 8000], ["소계", None, None, None], ["합계", None, None, None]]
        rep = wx(os.path.join(self.d, "p1b.xlsx"), [("보고서", rep_rows)])
        r = quiet(m.verify, raw, rep)
        self.assertEqual(r["mapping"]["col"], "법인")
        self.assertEqual([u["구분"] for u in r["unmatched"]], ["누락"] * 6)
        self.assertEqual(r["mismatch"], [])

    # 2) 건너뛴 시트
    def test_2_skipped_sheet_blocks_all_match(self):
        rep, res, c = self.rows_of([("정리", [H] + ROWS), ("메모", [["항목", "값"], ["검토자", "홍길동"], ["확인", 1]])])
        self.assertEqual(c["일치"], 20)
        self.assertNotEqual(self.final(rep), "전부 일치")

    # 3) 쉼표 번호
    def test_3_comma_long_id(self):
        a, b = "1,234,567,890,123,456,021", "1,234,567,890,123,456,022"
        self.assertFalse(m._row_eq(m._row_norm(a), m._row_norm(b)))
        self.assertTrue(m._row_eq(m._row_norm(a), m._row_norm(a)))
        self.assertTrue(m._row_eq(m._row_norm("1,000"), m._row_norm(1000)))

    # 4) 분할·중복
    def test_4a_split_10_10(self):
        rep, res, c = self.rows_of([("상", [H] + ROWS[:10]), ("하", [H] + ROWS[10:])])
        self.assertEqual((c["일치"], c["보고서에 없음"], c["확인 필요"]), (20, 0, 0), c)
        self.assertEqual(self.final(rep), "전부 일치")

    def test_4b_split_10_9(self):
        rep, res, c = self.rows_of([("상", [H] + ROWS[:10]), ("하", [H] + ROWS[10:19])])
        self.assertEqual((c["일치"], c["보고서에 없음"]), (19, 1), c)
        self.assertEqual([x["원본 행"] for x in res if x["판정"] == "보고서에 없음"], [21])

    def test_4c_duplicate_across_sheets(self):
        rep, res, c = self.rows_of([("상", [H] + ROWS[:10]), ("하", [H] + ROWS[8:])])
        self.assertEqual(c["보고서에 없음"], 0, c)
        dups = [x for x in res if x["판정"] == "확인 필요" and "중복" in x["이유"]]
        self.assertEqual(len(dups), 2, [x for x in res if x["판정"] != "일치"])
        self.assertNotEqual(self.final(rep), "전부 일치")

    def test_4d_independent_sheets_same_population(self):
        """원화기준·달러기준처럼 같은 모집단을 두 시트에 각각 담은 보고서는 중복이 아님"""
        rep, res, c = self.rows_of([("원화", [H] + ROWS), ("달러", [H] + ROWS)])
        self.assertEqual((c["일치"], c["보고서에 없음"], c["확인 필요"]), (40, 0, 0), c)

    # 6) 문구
    def test_6_recalc_limit_count(self):
        wb = Workbook(); ws = wb.active; ws["A1"] = 1; ws["A2"] = "=GROUPBY(A1,A1,SUM)"; ws["A3"] = "=A1*2"
        p = os.path.join(self.d, "f.xlsx"); wb.save(p)
        buf = io.StringIO()
        with contextlib.redirect_stdout(buf):
            status = m._formula_output(p)
        self.assertRegex(status, r"재검산 불가 \d+칸")

    # 5) cp949 출력
    def test_5_cp949_output_does_not_crash(self):
        src = os.path.join(TOOL_DIR, os.environ.get("TOOL", "excel_automation_v24") + ".py")
        wb = Workbook(); ws = wb.active; ws["A1"] = "x"; ws["B1"] = 1; ws["B2"] = "=B1*2"; p = os.path.join(self.d, "one.xlsx"); wb.save(p)
        env = dict(os.environ, PYTHONIOENCODING="cp949")
        r = subprocess.run([sys.executable, src, p, "--수식"], capture_output=True, env=env, cwd=self.d)
        self.assertEqual(r.returncode, 0, r.stderr[-500:])
        self.assertIn("최종".encode("cp949"), r.stdout)
        self.assertIn("재검산 불가".encode("cp949"), r.stdout, "'—'가 든 줄까지 출력돼야 함")
        self.assertNotIn(b"UnicodeEncodeError", r.stderr)

    def test_4e_independent_sheet_missing_rows_counted(self):
        _, res, c = self.rows_of([("원화", [H] + ROWS), ("달러", [H] + ROWS[:19])])
        self.assertEqual(c["보고서에 없음"], 1, c); self.assertEqual(c["확인 필요"], 0, c)
        rep, res, c = self.rows_of([("원화", [H] + ROWS), ("달러", [H] + ROWS[:17])])   # 85%: 같은 내용의 다른 보기에서 3건 빠짐
        self.assertEqual(c["보고서에 없음"], 3, c); self.assertNotEqual(self.final(rep), "전부 일치")

    def test_4f_full_plus_extract_sheet(self):
        _, res, c = self.rows_of([("전체", [H] + ROWS), ("주요", [H] + ROWS[:6])])
        self.assertEqual(c["보고서에 없음"], 0, c); self.assertEqual(c["확인 필요"], 1, c)
        self.assertEqual(c["일치"], 26, c)

    def test_4g_same_sheet_duplicate_row(self):
        _, res, c = self.rows_of([("거래", [H] + ROWS + [ROWS[3]])])
        self.assertEqual(c["원본에 없음"], 0, c); self.assertEqual(c["확인 필요"], 1, c)

    def test_8_header_only_sheet_is_not_skipped_content(self):
        self.assertEqual(self.final(wx(os.path.join(self.d, "p8.xlsx"), [("거래", [H] + ROWS), ("양식", [H])])), "전부 일치")

    def test_8b_memo_sheet_blocks(self):
        rep = wx(os.path.join(self.d, "p8b.xlsx"), [("거래", [H] + ROWS), ("메모", [["참고"], ["1월분 일부 누락 가능"]])])
        self.assertNotEqual(self.final(rep), "전부 일치")

    def _agg(self, rows):
        raw = wx(os.path.join(self.d, "r9.xlsx"), [("원본", [["계정과목", "구분", "금액"], ["현금", "A", 100], ["현금", "B", 200], ["예금", "A", 300], ["예금", "B", 400]])])
        rep = wx(os.path.join(self.d, "p9.xlsx"), [("보고서", rows)])
        return raw, rep

    def test_9_missing_in_messages_and_color(self):
        raw, rep = self._agg([["계정과목", "A", "B"], ["현금", 100, 200], ["예금", 300, None]])
        src = os.path.join(TOOL_DIR, os.environ.get("TOOL", "excel_automation_v24") + ".py")
        r = subprocess.run([sys.executable, src, raw, rep], capture_output=True, text=True, cwd=self.d, env=dict(os.environ, PYTHONIOENCODING="utf-8"))
        last = [l for l in r.stdout.splitlines() if l.startswith("최종")][-1]
        self.assertIn("누락 1칸", last)
        self.assertIn("+ 누락 1", r.stdout)
        res = quiet(m.verify, raw, rep); out = quiet(m.run, raw, rep, res)
        self.assertEqual(load_workbook(out)["요약"]["B1"].fill.fgColor.rgb[-6:], "FFC7CE")

    def test_10_blank_template_message(self):
        raw, rep = self._agg([["계정과목", "A", "B"], ["현금", None, None], ["예금", None, None]])
        src = os.path.join(TOOL_DIR, os.environ.get("TOOL", "excel_automation_v24") + ".py")
        r = subprocess.run([sys.executable, src, raw, rep], capture_output=True, text=True, cwd=self.d, env=dict(os.environ, PYTHONIOENCODING="utf-8"))
        last = [l for l in r.stdout.splitlines() if l.startswith("최종")][-1]
        self.assertIn("누락 4칸", last); self.assertIn("--채우기", last)

    def test_11_blank_formula_cell_is_no_cache(self):
        raw, rep = self._agg([["계정과목", "A", "B", "합계"], ["현금", 100, 200, "=B2+C2"], ["예금", 300, None, None], ["합계", "=SUM(B2:B3)", 600, None]])
        kinds = {u["위치"]: u["구분"] for u in quiet(m.verify, raw, rep)["unmatched"]}
        self.assertEqual(kinds.get("보고서!C3"), "누락", kinds)
        self.assertEqual(kinds.get("보고서!D2"), "계산값 없음", kinds)
        self.assertEqual(kinds.get("보고서!B4"), "계산값 없음", kinds)

    def test_12_section_heading_with_raw_group(self):
        raw = wx(os.path.join(self.d, "r12.xlsx"), [("원본", [["대분류", "계정과목", "금액"], ["유동자산", "현금", 100], ["유동자산", "예금", 200], ["비유동자산", "토지", 500]])])
        rep = wx(os.path.join(self.d, "p12.xlsx"), [("보고서", [["계정과목", "금액"], ["유동자산", None], ["현금", 100], ["예금", 999], ["소계", 300], ["비유동자산", None], ["토지", 500], ["합계", 800]])])
        r = quiet(m.verify, raw, rep)
        kinds = {u["위치"]: u["구분"] for u in r["unmatched"]}
        self.assertNotIn("보고서!B2", kinds, "구분 제목 줄은 누락 아님")
        self.assertNotIn("보고서!B6", kinds)
        self.assertIn("보고서!B4", {x["위치"] for x in r["mismatch"]}, "틀린 상세는 그대로 잡혀야 함")

    def test_7_source_is_utf8_with_declaration(self):
        src = os.path.join(TOOL_DIR, os.environ.get("TOOL", "excel_automation_v24") + ".py")
        b = open(src, "rb").read(); b.decode("utf-8")
        self.assertRegex(b.split(b"\n")[0].decode(), r"coding[:=]\s*utf-8")


if __name__ == "__main__":
    unittest.main()
