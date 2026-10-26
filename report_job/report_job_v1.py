# v1 (2026-10-26): 원본 데이터(SAP 다운로드 등, 한 줄 = 한 건) → 규칙표대로 집계 보고서 엑셀을 새로 생성 (범용)
# 사용법: VS Code ▶ (인자 없음) — 이 파일이 있는 폴더의 최신 원본 xlsx + 규칙표_job.xlsx(없으면 원본을 보고 기본값 추정)
#   py report_job_v1.py 원본.xlsx                      # 원본 지정
#   py report_job_v1.py 원본.xlsx --규칙 규칙표.xlsx    # 규칙표 지정
# 결과: <원본이름>_보고서_<시각>.xlsx (보고서 / 근거 / 규칙 시트). 원본 파일은 수정하지 않음
# 규칙표: '규칙' 시트 A열 항목·B열 값 (아래 RULE_HELP), '행 순서' 시트 A열 목록(선택 — 목록 순서대로, 목록에 없는 새 항목은 맨 아래 노란색)
# ===== 여기만 수정 =====
RAW_FILE   = r""     # 비워두면 폴더의 최신 xlsx (결과·규칙표 제외)
RULES_FILE = r""     # 비워두면 같은 폴더 규칙표_job.xlsx, 없으면 기본값 추정
USE_EXCEL  = False   # True: Windows Excel(pywin32)로 결과를 열어 전체 재계산 후 저장 (합계 수식 저장값 채움)
# =======================
import os, re, sys, glob, math, datetime, functools
from collections import OrderedDict
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

RULE_HELP = OrderedDict([
    ("원본 시트", ("", "비우면 숫자가 가장 많은 시트")),
    ("행 기준", ("", "보고서 왼쪽 제목 열. 쉼표로 여러 열 = 왼쪽부터 계층 (예: 법인, 계정명)")),
    ("열 기준", ("", "보고서 위쪽 제목. 날짜열:월 / 날짜열:연월 / 날짜열:연도 / 글자 열 이름 / 비움(금액 열 하나)")),
    ("금액", ("", "금액 열 이름 또는 + - 식 (예: 차변-대변)")),
    ("집계", ("합계", "합계 / 건수")),
    ("필터", ("", "조건; 조건 … 모두 만족하는 행만 (예: 구분=매출; 법인<>본사; 금액>0). 연산자 = <> > < >= <=")),
    ("제외 글자", ("", "행 기준 값에 이 글자가 들어 있는 원본 행 제외 (쉼표 구분, 예: 합계, 소계, *)")),
    ("단위", ("1", "금액 ÷ 단위 (1 / 1000 / 1000000)")),
    ("소계", ("예", "행 기준이 2개 이상일 때 첫 행 기준 값마다 소계 행")),
    ("합계 행", ("예", "맨 아래 합계 행")),
    ("합계 열", ("예", "맨 오른쪽 합계 열 (열 기준이 있을 때)")),
    ("정렬", ("원본 순", "원본 순 / 가나다 / 금액 큰 순 ('행 순서' 시트가 있으면 그 순서 우선)")),
    ("빈 칸", ("빈칸", "원본 건이 없는 칸: 0 / 빈칸")),
    ("합계 수식", ("아니오", "소계·합계를 =SUM 수식으로 (아니오면 값)")),
])
RESULT_TAGS = ("_보고서_", "규칙표", "_원본대사", "_수식검증", "_채움", "_PPT대조", "_구조비교", "_대사결과")
HEAD_FILL = PatternFill("solid", fgColor="1F4E78")
SUB_FILL = PatternFill("solid", fgColor="DDEBF7")
TOTAL_FILL = PatternFill("solid", fgColor="BDD7EE")
YELLOW = PatternFill("solid", fgColor="FFFF00")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
UNIT_NAME = {1: "원", 1000: "천원", 1000000: "백만원", 100000000: "억원"}


class JobError(Exception):
    """사용자에게 그대로 보여줄 오류"""


# ---------------------------------------------------------------- 값 해석 (엑셀 검산 엔진 v23과 같은 규칙)
_NUM_START = set("0123456789+-.(△▲₩￦$​‌‍⁠﻿ \t 　") | set("KkUu")


@functools.lru_cache(maxsize=200000)
def _text_to_num(v):
    t0 = v.lstrip()
    if not t0 or t0[0] not in _NUM_START:
        return None
    t = re.sub(r"[​‌‍⁠﻿]", "", v).replace(",", "").strip()
    if t in ("", "-", "—", "–"):
        return None
    if t[0] in "+-₩￦$KkUu" or t.endswith(("원", "KRW", "krw")):
        t = re.sub(r"^([+\-]?)\s*(?:[₩￦$]|KRW|USD)\s*", r"\1", t, flags=re.I)
        t = re.sub(r"\s*(?:원|KRW)$", "", t, flags=re.I).strip()
    if t.endswith("-") and re.fullmatch(r"\(?[\d.]+\)?-", t):
        t = "-" + t[:-1]
    neg = t.startswith("(") and t.endswith(")")
    t = t.strip("()").replace("△", "-").replace("▲", "-").strip()
    try:
        f = float(t)
    except ValueError:
        return None
    if not math.isfinite(f):
        return None
    return -f if neg else f


def to_num(v):
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v) if math.isfinite(v) else None
    if isinstance(v, str):
        return _text_to_num(v)
    return None


def is_dash(v):
    return isinstance(v, str) and v.strip() in ("-", "—", "–")


def is_blank(v):
    return v is None or (isinstance(v, str) and not v.strip())


def is_label(v):
    return (isinstance(v, str) and v.strip() != "" and to_num(v) is None and not is_dash(v)) or hasattr(v, "year")


def nl(v):
    return re.sub(r"\s+", "", str(v)).lower()


_DATE_RE = re.compile(r"^\s*(\d{4})\s*[./\-년]\s*(\d{1,2})\s*[./\-월]\s*(\d{1,2})\s*일?(?:\s+\d{1,2}:\d{2}(?::\d{2}(?:\.\d+)?)?)?\s*$")
_DATE8_RE = re.compile(r"^(19|20)\d{6}$")


def to_date(v):
    """날짜 / '2026-01-05' / '2026.01.05' / '2026년 1월 5일' / 엑셀 일련번호 / 20260105 → date, 아니면 None"""
    if hasattr(v, "year") and hasattr(v, "month"):
        return datetime.date(v.year, v.month, getattr(v, "day", 1))
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        if float(v).is_integer() and 20000 <= v <= 80000:
            return datetime.date(1899, 12, 30) + datetime.timedelta(days=int(v))
        if float(v).is_integer() and _DATE8_RE.match(str(int(v))):
            v = str(int(v))
        else:
            return None
    if isinstance(v, str):
        s = v.strip()
        m = _DATE_RE.match(s)
        if m:
            y, mo, d = (int(x) for x in m.groups())
        elif _DATE8_RE.match(s):
            y, mo, d = int(s[:4]), int(s[4:6]), int(s[6:])
        else:
            return None
        try:
            return datetime.date(y, mo, d)
        except ValueError:
            return None
    return None


def label_text(v):
    """행·열 기준 값 → 글자 (앞뒤 공백 제거, 정수 숫자는 '1100', 날짜는 YYYY-MM-DD)"""
    if v is None:
        return ""
    if hasattr(v, "year"):
        return f"{v.year:04d}-{v.month:02d}-{v.day:02d}"
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


# ---------------------------------------------------------------- 원본 읽기
def _sheet_rows(ws, notes=None):
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    for rng in ws.merged_cells.ranges:                     # 병합 칸은 왼쪽 위 값으로 채움 (글자만 — 숫자를 채우면 두 번 더해짐)
        try:
            v = rows[rng.min_row - 1][rng.min_col - 1]
        except IndexError:
            continue
        if (to_num(v) is not None and not hasattr(v, "year")) and rng.max_row > rng.min_row:
            if notes is not None:
                notes.append(f"병합된 숫자 칸 {rng.coord}: 첫 칸 값만 한 번 집계 (병합을 풀고 값을 확인하세요)")
            continue
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                if r - 1 < len(rows) and c - 1 < len(rows[r - 1]) and rows[r - 1][c - 1] in (None, ""):
                    rows[r - 1][c - 1] = v
    return rows



def _formula_cells_without_value(path, sheet_title):
    """시트 안에서 계산값 없이 저장된 수식 칸 {(행, 열 번호)} (다른 도구로 만든 파일 등)"""
    import zipfile, posixpath
    from xml.etree import ElementTree as ET
    out = set()
    try:
        with zipfile.ZipFile(path) as z:
            ns = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
                  "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships"}
            wbx = ET.fromstring(z.read("xl/workbook.xml"))
            rel = ET.fromstring(z.read("xl/_rels/workbook.xml.rels"))
            rid = next((sh.get("{%s}id" % ns["r"]) for sh in wbx.iter("{%s}sheet" % ns["m"]) if sh.get("name") == sheet_title), None)
            target = next((x.get("Target") for x in rel if x.get("Id") == rid), None)
            if not target:
                return out
            part = target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join("xl", target))
            data = z.read(part).decode("utf-8", "replace")
    except Exception:
        return out
    for m in re.finditer(r"<c\b([^>]*?)(?:/>|>(.*?)</c>)", data, re.S):
        body = m.group(2) or ""
        if "<f" in body and not re.search(r"<v>[^<]", body):
            ref = re.search(r'r="([A-Z]+)(\d+)"', m.group(1))
            if ref:
                col = 0
                for ch in ref.group(1):
                    col = col * 26 + ord(ch) - 64
                out.add((int(ref.group(2)), col))
    return out


_RAW_CACHE = {}


def read_raw(path, sheet=None):
    """→ {"sheet", "head": [열 이름], "rows": [(엑셀 행, {열: 값})], "head_row", "notes", "no_value": {(행, 열 이름)}}"""
    try:
        key = (os.path.abspath(path), os.path.getmtime(path), os.path.getsize(path))
    except OSError:
        key = None
    hit = _RAW_CACHE.get(key)
    if hit is not None and (not sheet or sheet == hit["sheet"]):
        return hit                                         # 같은 원본은 한 번만 읽음
    try:
        wb = load_workbook(path, data_only=True, read_only=False)
    except Exception as exc:
        raise JobError(f"원본 파일을 열 수 없음: {os.path.basename(path)} ({type(exc).__name__}) — xlsx로 다시 저장 후 실행")
    if sheet:
        if sheet not in wb.sheetnames:
            raise JobError(f"원본에 '{sheet}' 시트 없음. 있는 시트: {wb.sheetnames}")
        ws = wb[sheet]
    else:                                                  # 숫자(글자로 저장된 숫자 포함)가 가장 많은 시트
        ws = max(wb.worksheets, key=lambda w: sum(1 for r in w.iter_rows(max_row=5000, values_only=True) for v in r
                                                  if to_num(v) is not None and not isinstance(v, bool)))
    notes = []
    grid = _sheet_rows(ws, notes)
    nonempty = [(i + 1, r) for i, r in enumerate(grid) if any(not is_blank(v) for v in r)]
    def data_like(r):
        return any((to_num(v) is not None and not isinstance(v, bool)) or hasattr(v, "year") for v in r)
    cands = []
    for k, (i, r) in enumerate(nonempty[:15]):
        labs = [v for v in r if is_label(v) and not hasattr(v, "year")]
        vals = [v for v in r if not is_blank(v)]
        if len(labs) >= 2 and len(labs) >= 0.6 * len(vals):
            cands.append(k)
    if not cands:
        raise JobError(f"원본 '{ws.title}' 시트에서 열 제목 줄을 찾지 못함 (위 15줄 안에 글자 2개 이상인 줄)")
    # 제목 줄 = 바로 아래가 자료(숫자·날짜) 줄인 첫 후보 (위의 조회조건 줄, 아래 자료 줄을 제목으로 잘못 잡지 않게)
    k = next((k for k in cands if k + 1 < len(nonempty) and data_like(nonempty[k + 1][1])), None)
    if k is None:
        k = max(cands, key=lambda x: (sum(1 for v in nonempty[x][1] if not is_blank(v)), -x))
    head_row, head = nonempty[k]
    names, idx, dups = [], [], []
    for j, v in enumerate(head):
        if is_blank(v):
            continue
        name = str(v).strip()
        if name in names:
            dups.append(name)
        while name in names:
            name += "_"
        names.append(name)
        idx.append(j)
    if dups:
        notes.append("같은 이름 열 " + ", ".join(f"'{d}'" for d in sorted(set(dups))) + " — 두 번째부터는 이름 뒤에 '_'를 붙여 구분 (규칙표에서 그 이름으로 지정)")
    rows = []
    for i, r in nonempty[k + 1:]:
        rows.append((i, {n: (r[j] if j < len(r) else None) for n, j in zip(names, idx)}))
    if not rows:
        raise JobError(f"원본 '{ws.title}' 시트에 자료 행이 없음 (제목 줄만 있음)")
    hidden = sum(1 for i, _ in rows if ws.row_dimensions[i].hidden)
    if hidden:
        notes.append(f"숨긴 행 {hidden}개도 집계에 포함 (빼려면 필터나 제외 글자 사용)")
    col_of = {n: j + 1 for n, j in zip(names, idx)}
    nv = _formula_cells_without_value(path, ws.title)
    by_col = {c: n for n, c in col_of.items()}
    no_value = {(r, by_col[c]) for r, c in nv if c in by_col}
    raw = {"sheet": ws.title, "head": names, "rows": rows, "head_row": head_row, "notes": notes, "no_value": no_value}
    if key is not None:
        _RAW_CACHE.clear()
        _RAW_CACHE[key] = raw
    return raw



_YES = {"예", "네", "y", "yes", "true", "1", "o", "○", "◯", "사용"}
_NO = {"아니오", "아니요", "아니", "n", "no", "false", "0", "x", "×", "", "미사용"}


def _yes(v, key="값"):
    t = str("" if v is None else v).strip().lower()
    if t in _YES:
        return True
    if t in _NO:
        return False
    raise JobError(f"규칙표 '{key}' 값은 예 / 아니오 중 하나: '{v}'")



def default_rules(raw_path, raw=None):
    """원본을 보고 기본 규칙 추정. 가정은 '_가정' 목록"""
    raw = raw or read_raw(raw_path)
    head, rows = raw["head"], raw["rows"]
    sample = rows[:2000]
    assume = []
    def kinds(col):
        vals = [r[col] for _, r in sample if not is_blank(r[col])]
        n = len(vals) or 1
        num = sum(1 for v in vals if to_num(v) is not None or is_dash(v)) / n
        date = sum(1 for v in vals if to_date(v) is not None and not (isinstance(v, (int, float)) and not hasattr(v, "year"))) / n
        txt = sum(1 for v in vals if is_label(v) and not hasattr(v, "year") and to_date(v) is None) / n
        return num, date, txt, len({label_text(v) for v in vals}), len(vals)
    info = {c: kinds(c) for c in head}
    date_cols = [c for c in head if info[c][1] >= 0.8]
    num_cols = [c for c in head if info[c][0] >= 0.8 and c not in date_cols and info[c][4] > 0]
    code_re = re.compile(r"(코드|번호|순번|^no\.?$|^id$|code$|^#$)", re.I)
    amt_name = re.compile(r"(금액|amount|amt|원화|대금|합계|잔액|매출|비용|손익)", re.I)
    num_cols = [c for c in num_cols if not code_re.search(c)]
    if not num_cols:
        raise JobError("원본에서 금액(숫자) 열을 찾지 못함 — 규칙표 '금액'에 열 이름을 지정하세요. 원본 열: " + ", ".join(head))
    debit = next((c for c in num_cols if re.search(r"차변|debit", c, re.I)), None)
    credit = next((c for c in num_cols if re.search(r"대변|credit", c, re.I)), None)
    if debit and credit:
        amount = f"{debit}-{credit}"                       # 전표: 차변 - 대변
    else:
        amount = next((c for c in num_cols if amt_name.search(c)), None) or max(num_cols, key=lambda c: sum(
            abs(to_num(r[c]) or 0) for _, r in sample))
    text_cols = [c for c in head if info[c][2] >= 0.8 and c not in date_cols]
    groupish = [c for c in text_cols if 2 <= info[c][3] <= max(2, 0.5 * info[c][4])]
    if not groupish:
        groupish = text_cols
    if not groupish:
        raise JobError("원본에서 행 기준으로 쓸 글자 열을 찾지 못함 — 규칙표 '행 기준'을 지정하세요. 원본 열: " + ", ".join(head))
    acct = re.compile(r"(계정|과목|항목|account)", re.I)
    row_col = next((c for c in groupish if acct.search(c)), None) or min(groupish, key=lambda c: (info[c][3], head.index(c)))
    col = f"{date_cols[0]}:월" if date_cols else ""
    rules = OrderedDict((k, v[0]) for k, v in RULE_HELP.items())
    rules.update({"원본 시트": raw["sheet"], "행 기준": row_col, "열 기준": col, "금액": amount, "제외 글자": "합계, 소계, 총계"})
    assume.append(f"행 기준 = '{row_col}' (값 종류가 적당한 글자 열)")
    assume.append(f"열 기준 = '{col}' (날짜 열의 월)" if col else "열 기준 없음 (날짜 열을 못 찾음 → 금액 열 하나)")
    assume.append(f"금액 = '{amount}' (숫자 열)")
    assume.append("제외 글자 = 합계, 소계, 총계 (원본 안 합계 줄을 두 번 더하지 않게)")
    assume.append("집계 = 합계, 단위 = 1, 정렬 = 원본 순, 소계·합계 행·합계 열 = 예")
    rules["_가정"] = assume
    return rules


def load_rules(path):
    """규칙표 엑셀 → dict ('규칙' 시트 A열 항목 / B열 값, '행 순서' 시트 A열 목록)"""
    try:
        wb = load_workbook(path, data_only=True)
    except Exception as exc:
        raise JobError(f"규칙표를 열 수 없음: {os.path.basename(path)} ({type(exc).__name__})")
    ws = wb["규칙"] if "규칙" in wb.sheetnames else wb.worksheets[0]
    rules = OrderedDict()
    for r in ws.iter_rows(values_only=True):
        if not r or is_blank(r[0]):
            continue
        k = str(r[0]).strip()
        if k in ("항목",):
            continue
        rules[k] = "" if len(r) < 2 or r[1] is None else (label_text(r[1]) if isinstance(r[1], float) else r[1])
    if "행 순서" in wb.sheetnames:
        order = [label_text(r[0]) for r in wb["행 순서"].iter_rows(values_only=True)
                 if r and not is_blank(r[0]) and not str(r[0]).strip().startswith("(")]
        if order and order[0] in ("행 순서", "항목"):
            order = order[1:]
        rules["행 순서"] = order
    return rules


def save_rules_template(path, rules):
    wb = Workbook(); ws = wb.active; ws.title = "규칙"
    ws.append(["항목", "값", "설명"])
    for k, (dv, desc) in RULE_HELP.items():
        ws.append([k, str(rules.get(k, dv)), desc])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = HEAD_FILL
    ws.column_dimensions["A"].width = 12; ws.column_dimensions["B"].width = 28; ws.column_dimensions["C"].width = 80
    o = wb.create_sheet("행 순서")
    o.append(["행 순서", "(A열 2행부터 행 기준 값을 위에서부터 적으면 그 순서. 행 기준이 여러 개면 '값1 / 값2' 또는 첫 단계 값만. 비워 두면 '정렬' 규칙)"])
    o["A1"].font = Font(bold=True)
    wb.save(path)


def _norm_rules(rules, raw):
    r = OrderedDict((k, v[0]) for k, v in RULE_HELP.items())
    key_of = {nl(k): k for k in RULE_HELP}
    key_of[nl("행 순서")] = "행 순서"
    unknown = []
    for k, v in (rules or {}).items():
        kk = key_of.get(nl(k))                              # '합계열'·'단위 ' 같은 띄어쓰기 차이는 같은 항목
        if kk is None:
            unknown.append(str(k))
            continue
        r[kk] = v
    head = raw["head"]
    def col(name, what):
        name = str(name).strip()
        if name in head:
            return name
        hit = [h for h in head if nl(h) == nl(name)]
        if hit:
            return hit[0]
        raise JobError(f"원본에 '{name}' 열 없음 ({what}). 원본 열: {', '.join(head)}")
    rows_by = [col(x, "행 기준") for x in str(r["행 기준"] or "").split(",") if x.strip()]
    if not rows_by:
        raise JobError("규칙표 '행 기준'이 비어 있음")
    cspec = str(r["열 기준"] or "").strip()
    col_by, col_part = None, None
    if cspec:
        m = re.match(r"^(.*?)\s*:\s*(월|연월|연도)\s*$", cspec)
        if m:
            col_by, col_part = col(m.group(1), "열 기준"), m.group(2)
        else:
            col_by = col(cspec, "열 기준")
    agg = nl(r["집계"] or "합계")
    if agg in ("합계", "합", "sum"):
        count = False
    elif agg in ("건수", "개수", "갯수", "count"):
        count = True
    else:
        raise JobError(f"규칙표 '집계'는 합계 / 건수 중 하나: '{r['집계']}'")
    aspec = str(r["금액"] or "").strip()
    if not aspec and not count:
        raise JobError("규칙표 '금액'이 비어 있음")
    terms = []
    if aspec:
        if aspec in head or any(nl(h) == nl(aspec) for h in head):
            terms = [(1, col(aspec, "금액"))]
        else:
            if not re.fullmatch(r"\s*[+\-]?\s*[^+\-]+?(\s*[+\-]\s*[^+\-]+?)*\s*", aspec):
                raise JobError(f"금액 식을 읽지 못함: '{aspec}' (예: 차변-대변)")
            parts = re.findall(r"([+\-]?)\s*([^+\-]+)", aspec)
            terms = [(-1 if s == "-" else 1, col(n.strip(), "금액")) for s, n in parts]
    try:
        unit = float(str(r["단위"]).replace(",", "").strip() or "1")
    except ValueError:
        raise JobError(f"단위는 0보다 큰 숫자여야 함: '{r['단위']}'")
    if not math.isfinite(unit) or unit <= 0:
        raise JobError(f"단위는 0보다 큰 숫자여야 함: '{r['단위']}'")
    filters = []
    for part in str(r["필터"] or "").split(";"):
        part = part.strip()
        if not part:
            continue
        m = re.match(r"^(.+?)\s*(<>|>=|<=|!=|=|>|<)\s*(.*)$", part)
        if not m:
            raise JobError(f"필터 조건을 읽지 못함: '{part}' (예: 구분=매출)")
        filters.append((col(m.group(1), "필터"), m.group(2).replace("!=", "<>"), m.group(3).strip().strip('"\'')))
    excl = [x.strip() for x in str(r["제외 글자"] or "").split(",") if x.strip()]
    sort = str(r["정렬"]).strip() or "원본 순"
    if sort not in ("원본 순", "가나다", "금액 큰 순"):
        raise JobError(f"정렬은 원본 순 / 가나다 / 금액 큰 순 중 하나: '{sort}'")
    blank = nl(r["빈 칸"])
    if blank in ("0", "0.0", "영", "zero"):
        zero = True
    elif blank in ("빈칸", "공란", "없음", "", "blank", "공백"):
        zero = False
    else:
        raise JobError(f"규칙표 '빈 칸'은 0 / 빈칸 중 하나: '{r['빈 칸']}'")
    order = r.get("행 순서") or []
    if isinstance(order, str):
        order = [x.strip() for x in re.split(r"[\n,]", order) if x.strip()]
    return {"rows_by": rows_by, "col_by": col_by, "col_part": col_part, "terms": terms, "count": count,
            "unit": 1.0 if count else unit, "filters": filters, "excl": excl, "sort": sort, "order": [label_text(x) for x in order],
            "subtotal": _yes(r["소계"], "소계") and len(rows_by) >= 2, "total_row": _yes(r["합계 행"], "합계 행"),
            "total_col": _yes(r["합계 열"], "합계 열") and col_by is not None, "zero": zero,
            "formula": _yes(r["합계 수식"], "합계 수식"), "raw_rules": r, "unknown_keys": unknown}



def _filter_ok(row, f):
    name, op, want = f
    v = row.get(name)
    wd = to_date(want) if re.search(r"[./\-년]", want) else None
    if wd is not None:                                     # 날짜 조건: 원본 날짜 형식이 섞여 있어도 날짜로 비교
        d = to_date(v)
        if d is None:
            return op == "<>"
        a, b = d.toordinal(), wd.toordinal()
        return {"=": a == b, "<>": a != b, ">": a > b, "<": a < b, ">=": a >= b, "<=": a <= b}[op]
    a, b = to_num(v), to_num(want)
    if op in (">", "<", ">=", "<="):
        if a is None or b is None:
            return False
        return {">": a > b, "<": a < b, ">=": a >= b, "<=": a <= b}[op]
    if re.fullmatch(r"0\d+", want):                        # 앞자리 0 코드('001100')는 글자로 비교
        same = nl(label_text(v)) == nl(want)
    elif a is not None and b is not None:
        same = abs(a - b) < 1e-9
    else:
        same = nl(label_text(v)) == nl(want)
    return same if op == "=" else not same



def _col_key(v, R):
    if R["col_by"] is None:
        return R["col_title"], None
    if R["col_part"]:
        d = to_date(v)
        if d is None:
            return None, "날짜를 읽지 못함"
        if R["col_part"] == "월":
            return f"{d.month}월", (d.year, d.month)
        if R["col_part"] == "연월":
            return f"{d.year:04d}-{d.month:02d}", (d.year, d.month)
        return f"{d.year:04d}년", (d.year,)
    t = label_text(v)
    if not t:
        return None, "열 기준 값이 빈칸"
    return t, None


def aggregate(raw, R):
    values, counts, first, col_first, col_sort = {}, {}, {}, {}, {}
    excluded, bad, why, used, years = 0, [], [], 0, set()
    raw_total = 0.0
    for excel_row, row in raw["rows"]:
        keys = [label_text(row.get(c)) for c in R["rows_by"]]
        if R["excl"] and any(tok in k for k in keys for tok in R["excl"]):
            excluded += 1
            continue
        if R["filters"] and not all(_filter_ok(row, f) for f in R["filters"]):
            excluded += 1
            continue
        if any(not k for k in keys):
            bad.append((excel_row, "(빈 행 기준)")); why.append("행 기준 빈칸: " + ", ".join(f"{c}={label_text(row.get(c)) or '(빈)'}" for c in R["rows_by"]))
            continue
        ck, sk = _col_key(row.get(R["col_by"]) if R["col_by"] else None, R)
        if ck is None:
            bad.append((excel_row, row.get(R["col_by"]))); why.append(f"{R['col_by']}: {row.get(R['col_by'])!r} ({sk})")
            continue
        if R["count"]:
            amt = 1.0
        else:
            amt, bad_val = 0.0, None
            for sign, c in R["terms"]:
                v = row.get(c)
                if (excel_row, c) in raw["no_value"]:
                    bad_val = (c, "(수식 계산값 없음 — 엑셀에서 열어 저장 필요)")
                    break
                if is_blank(v) or is_dash(v):
                    continue
                n = to_num(v)
                if n is None:
                    bad_val = (c, v)
                    break
                amt += sign * n
            if bad_val is not None:
                bad.append((excel_row, bad_val[1])); why.append(f"{bad_val[0]}: {bad_val[1]!r} (숫자 아님)")
                continue
        rk = tuple(keys)
        val = amt / R["unit"]
        key = (rk, ck)
        values[key] = values.get(key, 0.0) + val
        counts[key] = counts.get(key, 0) + 1
        first.setdefault(rk, len(first))
        col_first.setdefault(ck, len(col_first))
        if sk is not None:
            col_sort[ck] = sk
            years.add(sk[0])
        used += 1
    # 원본 합계는 칸별 집계와 따로 다시 셈 (행·열 묶음에서 빠지거나 두 번 더해진 행이 있으면 '불일치')
    badset = {x for x, _ in bad}
    for excel_row, row in raw["rows"]:
        keys = [label_text(row.get(c)) for c in R["rows_by"]]
        if excel_row in badset or (R["excl"] and any(tok in k for k in keys for tok in R["excl"])):
            continue
        if R["filters"] and not all(_filter_ok(row, f) for f in R["filters"]):
            continue
        raw_total += 1.0 if R["count"] else sum(sign * (to_num(row.get(c)) or 0.0) for sign, c in R["terms"]) / R["unit"]
    return {"values": values, "counts": counts, "first": first, "col_first": col_first, "col_sort": col_sort,
            "excluded": excluded, "bad": bad, "why": why, "used": used, "raw_total": raw_total, "years": years}


def _row_order(A, R):
    rows = sorted(A["first"], key=A["first"].get)
    multi = len(R["rows_by"]) >= 2
    tot = {}
    for (rk, _), v in A["values"].items():
        tot[rk] = tot.get(rk, 0.0) + v
    gfirst = {}
    for rk in rows:
        gfirst.setdefault(rk[0], len(gfirst))
    if R["sort"] == "가나다":
        rows.sort()
    elif R["sort"] == "금액 큰 순":
        if multi:
            gt = {}
            for rk in rows:
                gt[rk[0]] = gt.get(rk[0], 0.0) + tot[rk]
            rows.sort(key=lambda rk: (-gt[rk[0]], gfirst[rk[0]], -tot[rk], A["first"][rk]))
        else:
            rows.sort(key=lambda rk: (-tot[rk], A["first"][rk]))
    new = []
    if R["order"]:
        pos = {}
        for i, e in enumerate(R["order"]):
            pos.setdefault(nl(e), i)
        def row_pos(rk):
            for k in (nl("/".join(rk)), nl(rk[-1])) if multi else (nl(rk[0]),):
                if k in pos:
                    return pos[k]
            return None
        gpos = {}
        if multi:
            for rk in rows:
                p = pos.get(nl(rk[0]))
                if p is None:
                    p = row_pos(rk)
                if p is not None:
                    gpos[rk[0]] = min(gpos.get(rk[0], p), p)
        idx = {rk: i for i, rk in enumerate(rows)}
        INF = float("inf")
        new = [rk for rk in rows if row_pos(rk) is None and not (multi and nl(rk[0]) in pos)]
        if multi:
            rows.sort(key=lambda rk: (gpos.get(rk[0], INF), gfirst[rk[0]] if rk[0] not in gpos else 0,
                                      row_pos(rk) if row_pos(rk) is not None else INF, idx[rk]))
        else:
            rows.sort(key=lambda rk: (row_pos(rk) if row_pos(rk) is not None else INF, idx[rk]))
    if multi:                                               # 계층 유지: 첫 행 기준 값마다 붙여 둠
        gf = {}
        for rk in rows:
            gf.setdefault(rk[0], len(gf))
        idx = {rk: i for i, rk in enumerate(rows)}
        rows.sort(key=lambda rk: (gf[rk[0]], idx[rk]))
    return rows, new



def _col_order(A, R):
    cols = sorted(A["col_first"], key=A["col_first"].get)
    if R["col_part"]:
        cols.sort(key=lambda c: A["col_sort"][c] if R["col_part"] != "월" else (A["col_sort"][c][1],))
    return cols


# ---------------------------------------------------------------- 출력
def _fmt_number(R, A):
    if R["unit"] == 1 and not R["count"] and any(abs(v - round(v)) > 1e-9 for v in A["values"].values()):
        return "#,##0.00"
    return "#,##0"


def _txt(cell, v):
    cell.value = v
    if isinstance(v, str) and v.startswith("="):
        cell.data_type = "s"                               # '=소계=' 같은 글자가 수식으로 저장되지 않게
    return cell


def _looks_total(v):
    """원본 값이 합계·소계처럼 보이면 (보고서 합계 줄·열과 헷갈리지 않게 표시만 바꿈)"""
    t = re.sub(r"\(.*?\)|\s+", "", str(v)).lower()
    return t.endswith(("합계", "소계", "총계")) or t in ("total", "subtotal", "계")


def _sum_formula(refs):
    """SUM 인수 255개 제한: 넘으면 SUM(SUM(…),SUM(…))으로 나눔"""
    refs = [x for x in refs if x]
    if not refs:
        return 0
    def build(items):
        if len(items) <= 255:
            return "SUM(" + ",".join(items) + ")"
        chunks = [items[i:i + 255] for i in range(0, len(items), 255)]
        return build([build(c) for c in chunks])
    return "=" + build(refs)


def _text_width(v):
    s = "" if v is None else str(v)
    return sum(2 if ord(ch) > 0x2E80 else 1 for ch in s)


def write_report(raw_path, raw, R, A, rows, new, cols, out_path, assume, warnings):
    wb = Workbook()
    ws = wb.active; ws.title = "보고서"
    n = len(R["rows_by"])
    unit_txt = UNIT_NAME.get(int(R["unit"]) if float(R["unit"]).is_integer() else R["unit"], f"{R['unit']:g}")
    measure = "건수" if R["count"] else " ".join(f"{'-' if s < 0 else '+'}{c}" for s, c in R["terms"]).lstrip("+")
    title = f"{os.path.splitext(os.path.basename(raw_path))[0]} 집계 — {measure} ({'건' if R['count'] else '단위: ' + unit_txt})"
    if A["years"]:
        title += f" / {min(A['years'])}" + (f"~{max(A['years'])}년" if len(A["years"]) > 1 else "년")
    if R["unit"] != 1 and not R["count"]:
        title += " — 단위 미만은 반올림해 보여 합계와 단수 차이가 날 수 있음"
    ws.cell(1, 1, title).font = Font(bold=True, size=12)
    hr = 3
    show = {}
    for ck in cols:                                         # 원본 값이 '합계'·'…소계'면 보고서 합계 줄·열과 겹치지 않게 표시만 바꿈
        show[ck] = f"{ck} 항목" if _looks_total(ck) or ck in R["rows_by"] else ck
    # 원본 값에 '합계'가 있으면 보고서 합계 줄·열 이름을 '총합계'로 (검산 도구가 원본 '합계' 값과 헷갈리지 않게)
    clash = any(ck == "합계" for ck in cols) or any(lab == "합계" for rk in rows for lab in rk)
    total_name = "총합계" if clash else "합계"
    heads = list(R["rows_by"]) + [show[c] for c in cols] + ([total_name] if R["total_col"] else [])
    for j, h in enumerate(heads, start=1):
        c = _txt(ws.cell(hr, j), h)
        c.font = Font(bold=True, color="FFFFFF"); c.fill = HEAD_FILL; c.border = BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    fmt = _fmt_number(R, A)
    ncol = len(cols)
    first_val_col = n + 1
    r = hr + 1
    detail_rows, sub_rows, evidence = [], [], []
    def put_num(rr, cc, v, bold=False, fill=None):
        c = ws.cell(rr, cc, v)
        c.number_format = fmt; c.border = BORDER
        if bold:
            c.font = Font(bold=True)
        if fill:
            c.fill = fill
        return c
    def total_cells(rr, src_rows, bold, fill):
        for j, ck in enumerate(cols):
            cc = first_val_col + j
            if R["formula"]:
                put_num(rr, cc, _sum_formula(_ranges(list(src_rows), get_column_letter(cc)).split(",")), bold, fill)
            else:
                have = [(rk, ck) for rk in src_keys[rr] if (rk, ck) in A["values"]]
                v = sum(A["values"][k] for k in have) if have else (0 if R["zero"] else None)
                put_num(rr, cc, v, bold, fill)
        if R["total_col"]:
            cc = first_val_col + ncol
            if R["formula"]:
                put_num(rr, cc, _sum_formula(_ranges(list(src_rows), None, first_val_col, first_val_col + ncol - 1).split(",")), bold, fill)
            else:
                have = [(rk, ck) for rk in src_keys[rr] for ck in cols if (rk, ck) in A["values"]]
                put_num(rr, cc, sum(A["values"][k] for k in have) if have else (0 if R["zero"] else None), bold, fill)
    src_keys = {}
    groups = OrderedDict()
    for rk in rows:
        groups.setdefault(rk[0] if R["subtotal"] else None, []).append(rk)
    newset = set(new)
    for g, members in groups.items():
        g_rows = []
        for rk in members:
            for j, lab in enumerate(rk, start=1):
                shown = f"{lab} 항목" if _looks_total(lab) else lab
                c = _txt(ws.cell(r, j), shown); c.border = BORDER
                if rk in newset:
                    c.fill = YELLOW
            for j, ck in enumerate(cols):
                v = A["values"].get((rk, ck))
                cell = put_num(r, first_val_col + j, v if v is not None else (0 if R["zero"] else None))
                if v is not None:
                    evidence.append((f"보고서!{cell.coordinate}", rk, ck, A["counts"][(rk, ck)], v))
            if R["total_col"]:
                cc = first_val_col + ncol
                if R["formula"]:
                    put_num(r, cc, f"=SUM({get_column_letter(first_val_col)}{r}:{get_column_letter(first_val_col + ncol - 1)}{r})")
                else:
                    have = [A["values"][(rk, ck)] for ck in cols if (rk, ck) in A["values"]]
                    put_num(r, cc, sum(have) if have else (0 if R["zero"] else None))
            detail_rows.append(r); g_rows.append(r); src_keys[r] = [rk]
            r += 1
        if R["subtotal"]:
            c = _txt(ws.cell(r, 1), f"{g} 소계"); c.font = Font(bold=True); c.fill = SUB_FILL; c.border = BORDER
            for j in range(2, n + 1):
                ws.cell(r, j).fill = SUB_FILL; ws.cell(r, j).border = BORDER
            src_keys[r] = members
            total_cells(r, g_rows, True, SUB_FILL)
            sub_rows.append(r)
            r += 1
    if R["total_row"]:
        c = ws.cell(r, 1, total_name); c.font = Font(bold=True); c.fill = TOTAL_FILL; c.border = BORDER
        for j in range(2, n + 1):
            ws.cell(r, j).fill = TOTAL_FILL; ws.cell(r, j).border = BORDER
        src_keys[r] = rows
        total_cells(r, sub_rows if (R["subtotal"] and R["formula"]) else detail_rows, True, TOTAL_FILL)
        r += 1
    # 너비·고정·인쇄
    for j, h in enumerate(heads, start=1):
        w = max([_text_width(h)] + [_text_width(ws.cell(x, j).value if not isinstance(ws.cell(x, j).value, (int, float))
                                                  else f"{ws.cell(x, j).value:,.0f}") for x in range(hr + 1, r)] + [6])
        ws.column_dimensions[get_column_letter(j)].width = min(45, w + 3)
    ws.freeze_panes = ws.cell(hr + 1, n + 1).coordinate
    ws.print_title_rows = f"{hr}:{hr}"
    ws.page_setup.orientation = "landscape"; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    # 근거
    ev = wb.create_sheet("근거")
    ev.append(["위치"] + list(R["rows_by"]) + ["열", "원본 건수", "금액 합" + ("" if R["count"] else f" ({unit_txt})")])
    for loc, rk, ck, cnt, v in evidence:
        ev.append([loc] + list(rk) + [ck, cnt, v])
        for c in ev[ev.max_row]:
            if isinstance(c.value, str) and c.value.startswith("="):
                c.data_type = "s"
    for c in ev[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = HEAD_FILL
    for row in ev.iter_rows(min_row=2):
        row[-1].number_format = fmt; row[-2].number_format = "#,##0"
    for j in range(1, ev.max_column + 1):
        ev.column_dimensions[get_column_letter(j)].width = 16
    ev.freeze_panes = "A2"
    # 규칙
    rs = wb.create_sheet("규칙")
    rs.append(["항목", "값"])
    for k in RULE_HELP:
        rs.append([k, str(R["raw_rules"].get(k, ""))])
    if R["order"]:
        rs.append(["행 순서", " / ".join(R["order"])])
    rs.append(["원본", f"{os.path.basename(raw_path)} ('{raw['sheet']}' 시트, 제목 {raw['head_row']}행, 자료 {len(raw['rows'])}행)"])
    rs.append(["집계에 쓴 행", f"{A['used']}행"])
    rs.append(["제외한 행", f"{A['excluded']}행 (필터·제외 글자)"])
    rs.append(["읽지 못한 행", f"{len(A['bad'])}행"])
    for (x, _), why in zip(A["bad"], A["why"]):
        rs.append(["  ", f"원본 {x}행: {why}"])
    for a in assume:
        rs.append(["가정", a])
    for w in warnings:
        rs.append(["주의", w])
    if new:
        rs.append(["새 항목", ", ".join(" / ".join(rk) for rk in new) + " (행 순서 목록에 없어 맨 아래 노란색)"])
    for row in rs.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("="):
                c.data_type = "s"
    for c in rs[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = HEAD_FILL
    rs.column_dimensions["A"].width = 14; rs.column_dimensions["B"].width = 100
    wb.save(out_path)
    return detail_rows, sub_rows, evidence


def _ranges(rownums, letter, c1=None, c2=None):
    """연속한 행 번호를 묶어 'C4:C9,C12' 같은 범위 글자로"""
    if not rownums:
        return ""
    rownums = sorted(rownums)
    parts, s = [], rownums[0]
    prev = s
    for x in rownums[1:] + [None]:
        if x is not None and x == prev + 1:
            prev = x
            continue
        if letter:
            parts.append(f"{letter}{s}" if s == prev else f"{letter}{s}:{letter}{prev}")
        else:
            a, b = get_column_letter(c1), get_column_letter(c2)
            parts.append(f"{a}{s}:{b}{prev}")
        if x is not None:
            s = prev = x
    return ",".join(parts)


def excel_recalc(path):
    """Excel로 열어 전체 재계산 후 저장 (수식 저장값 채움). 실패해도 결과 파일은 그대로"""
    try:
        import pythoncom, win32com.client
    except ImportError:
        return "pywin32가 없어 Excel 재계산을 건너뜀"
    pythoncom.CoInitialize()
    app = None
    try:
        app = win32com.client.DispatchEx("Excel.Application")
        app.Visible = False; app.DisplayAlerts = False; app.EnableEvents = False
        wb = app.Workbooks.Open(os.path.abspath(path))
        app.CalculateFullRebuild()
        wb.Save(); wb.Close(False)
        return None
    except Exception as exc:
        return f"Excel 재계산 실패 ({type(exc).__name__}) — 엑셀에서 열면 자동 계산됨"
    finally:
        try:
            if app is not None:
                app.Quit()
        except Exception:
            pass
        pythoncom.CoUninitialize()


# ---------------------------------------------------------------- 실행
def _stamp():
    return datetime.datetime.now().strftime("%Y%m%d_%H%M%S")


def build_report(raw_path, rules=None, out_path=None):
    """원본 → 보고서 엑셀. (저장 경로, info) 반환"""
    raw_path = os.path.abspath(raw_path)
    assume, warnings = [], []
    if rules is None:
        rp = os.path.join(os.path.dirname(raw_path), "규칙표_job.xlsx")
        if os.path.exists(rp):
            rules = load_rules(rp)
        else:
            rules = default_rules(raw_path)
    rules = OrderedDict(rules)
    assume += list(rules.pop("_가정", []) or [])
    raw = read_raw(raw_path, str(rules.get("원본 시트") or "").strip() or None)
    R = _norm_rules(rules, raw)
    R["col_title"] = "건수" if R["count"] else str(R["raw_rules"]["금액"]).strip()
    warnings += raw["notes"]
    if R["unknown_keys"]:
        warnings.append("규칙표의 모르는 항목은 무시함: " + ", ".join(R["unknown_keys"]))
    if not R["excl"]:
        tot_like = sorted({label_text(row.get(c)) for _, row in raw["rows"] for c in R["rows_by"]
                           if re.search(r"(합\s*계|소\s*계|총\s*계)", label_text(row.get(c)))})
        if tot_like:
            warnings.append("행 기준 값에 합계·소계 글자가 있음 (원본 안 합계 줄이면 두 번 더해짐 — 규칙표 '제외 글자'에 지정): " + ", ".join(tot_like[:5]))
    A = aggregate(raw, R)
    if not A["values"]:
        raise JobError(f"집계할 원본 행이 없음 (자료 {len(raw['rows'])}행 중 제외 {A['excluded']}행, 읽지 못함 {len(A['bad'])}행)")
    rows, new = _row_order(A, R)
    cols = _col_order(A, R)
    if R["col_part"] == "월" and len(A["years"]) > 1:
        warnings.append(f"여러 해({min(A['years'])}~{max(A['years'])})의 같은 달이 한 열로 합쳐짐 — 열 기준을 '{R['col_by']}:연월'로 바꾸면 해마다 나뉨")
    if A["bad"]:
        warnings.append(f"읽지 못한 원본 {len(A['bad'])}행은 집계에서 뺌 (규칙 시트에 행 번호): " +
                        ", ".join(f"{x}행" for x, _ in A["bad"][:10]) + (" …" if len(A["bad"]) > 10 else ""))
    if new:
        warnings.append("행 순서 목록에 없는 새 항목 (맨 아래 노란색): " + ", ".join(" / ".join(rk) for rk in new))
    if out_path is None:
        base = os.path.splitext(raw_path)[0]
        out_path = f"{base}_보고서_{_stamp()}.xlsx"
        k = 2
        while os.path.exists(out_path):
            out_path = f"{base}_보고서_{_stamp()}_{k}.xlsx"; k += 1
    if os.path.abspath(out_path) == raw_path:
        raise JobError("결과 경로가 원본 파일과 같음")
    _, _, evidence = write_report(raw_path, raw, R, A, rows, new, cols, out_path, assume, warnings)
    # 저장한 보고서를 다시 읽어 상세 칸 합계 확인 (쓰기 중 빠지거나 바뀐 칸이 있으면 '불일치')
    wbc = load_workbook(out_path, data_only=False)["보고서"]
    sheet_total = 0.0
    for loc, *_ in evidence:
        v = wbc[loc.split("!")[1]].value
        sheet_total += v if isinstance(v, (int, float)) else float("nan")
    if USE_EXCEL:
        w = excel_recalc(out_path)
        if w:
            warnings.append(w)
    grand = sum(A["values"].values())
    info = {"row_keys": rows, "col_keys": cols, "values": dict(A["values"]), "grand_total": grand, "raw_total": A["raw_total"],
            "sheet_total": sheet_total,
            "new_items": new, "excluded_rows": A["excluded"], "bad_rows": A["bad"], "warnings": warnings,
            "rules": dict(R["raw_rules"]), "assume": assume}
    return out_path, info


_RESULT_NAME = re.compile(r"_(보고서|원본대사|수식검증|채움|PPT대조|구조비교|대사결과)_\d{8}_\d{6}")


def _pick_raw(folder):
    files = [f for f in glob.glob(os.path.join(folder, "*.xlsx")) + glob.glob(os.path.join(folder, "*.xlsm"))
             if not os.path.basename(f).startswith(("~$", "규칙표")) and not _RESULT_NAME.search(os.path.basename(f))]
    if not files:
        raise JobError(f"폴더에 원본 xlsx가 없음: {folder}")
    files.sort(key=os.path.getmtime)
    if len(files) > 1:
        print(f"원본 후보 {len(files)}개 중 가장 최근 파일 사용 (다른 파일은 RAW_FILE에 지정): " + ", ".join(os.path.basename(f) for f in files[:-1][-5:]))
    return files[-1]



def main(argv=None):
    for stream in (sys.stdout, sys.stderr):                # 한글 Windows(cp949)에서 출력을 파일·파이프로 보낼 때 '—' 등으로 멈추지 않게
        try:
            stream.reconfigure(errors="replace")
        except Exception:
            pass
    argv = list(sys.argv[1:] if argv is None else argv)
    here = os.path.dirname(os.path.abspath(__file__))
    try:
        rules_path = RULES_FILE or None
        if "--규칙" in argv:
            i = argv.index("--규칙")
            if i + 1 >= len(argv):
                raise JobError("--규칙 뒤에 규칙표 파일을 지정하세요")
            rules_path = argv[i + 1]
            del argv[i:i + 2]
        raw_path = argv[0] if argv else (RAW_FILE or _pick_raw(here))
        if not os.path.isfile(raw_path):
            raise JobError(f"원본 파일 없음: {raw_path}")
        if os.path.splitext(raw_path)[1].lower() not in (".xlsx", ".xlsm"):
            raise JobError(f"{os.path.splitext(raw_path)[1]} 파일은 직접 읽지 못함 — 엑셀에서 'Excel 통합 문서(*.xlsx)'로 저장 후 실행")
        rules = None
        if rules_path:
            rules = load_rules(rules_path)
        else:
            rp = os.path.join(os.path.dirname(os.path.abspath(raw_path)), "규칙표_job.xlsx")
            if os.path.exists(rp):
                rules = load_rules(rp); rules_path = rp
            else:
                rules = default_rules(raw_path)
                save_rules_template(rp, rules)
                print(f"규칙표가 없어 원본을 보고 기본값으로 만들었음: {rp}")
                for a in rules.get("_가정", []):
                    print("  가정:", a)
        print("원본:", raw_path)
        out, info = build_report(raw_path, rules)
        for w in info["warnings"]:
            print("주의:", w)
        tol = 1e-6 * max(1.0, abs(info["raw_total"]))
        ok = abs(info["grand_total"] - info["raw_total"]) <= tol and abs(info["sheet_total"] - info["raw_total"]) <= tol
        g = info["grand_total"]
        fmt = (lambda x: f"{x:,.0f}") if all(abs(v - round(v)) < 1e-9 for v in (g, info["raw_total"])) else (lambda x: f"{x:,.2f}")
        print(f"최종: 보고서 생성 — 행 {len(info['row_keys'])} / 열 {len(info['col_keys'])} / 합계 {fmt(g)} "
              f"(원본 합계 {fmt(info['raw_total'])} {'일치' if ok else '불일치'}) | 저장: {out}")
        return 0
    except JobError as exc:
        print(f"최종: 실행 불가 — {exc}")
        return 1
    except Exception as exc:
        print(f"최종: 실행 불가 — {type(exc).__name__}: {exc}")
        return 1


if __name__ == "__main__":
    sys.exit(main())
