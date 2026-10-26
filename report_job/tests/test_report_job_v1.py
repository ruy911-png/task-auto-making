# -*- coding: utf-8 -*-
"""report_job_v1.py 계약 시험 (TDD red phase) — SPEC.md 기준.

실행:  python test_report_job_v1.py            (기본 JOB=report_job_v1, JOB_DIR=이 파일 폴더)
       JOB=다른모듈 JOB_DIR=/경로 python test_report_job_v1.py

원칙
- 기대값은 모두 이 시험 안에서 직접 만든 원본 행으로 독립 계산 (작업 모듈을 기대값 계산에 쓰지 않음).
- 생성된 '보고서' 시트를 직접 읽어 칸마다 대조 (info 값만 믿지 않음) + v23 verify로 교차 검산.
- 입력·출력은 모두 job1/work 아래 임시 폴더에서만 만든다. 원본 파일은 md5로 무변경 확인.
- 의존: 표준 라이브러리 + openpyxl.
"""
import contextlib
import datetime
import hashlib
import importlib
import io
import math
import os
import random
import re
import shutil
import subprocess
import sys
import tempfile
import time
import traceback
import unittest
from collections import Counter, OrderedDict, defaultdict

from openpyxl import Workbook, load_workbook
from openpyxl.formula import Tokenizer
from openpyxl.formula.tokenizer import Token
from openpyxl.utils.cell import range_boundaries, get_column_letter

SCRATCH = "/tmp/claude-0/-home-claude/aecaabec-d78c-585a-8942-9b5b63461a12/scratchpad"
JOB_NAME = os.environ.get("JOB", "report_job_v1")
JOB_DIR = os.environ.get("JOB_DIR", os.path.join(SCRATCH, "job1"))
V23_DIR = os.path.join(SCRATCH, "v23")
WORK = os.path.join(SCRATCH, "job1", "work")
REAL_RAWS = [
    os.path.join(SCRATCH, "review", "cases", "13_원본_전표.xlsx"),
    os.path.join(SCRATCH, "review", "cases", "14_원본.xlsx"),
    os.path.join(SCRATCH, "drive", "files", "매매일지_건별원본.xlsx"),
]
SERIAL_EPOCH = datetime.date(1899, 12, 30)


# ------------------------------------------------------------------ imports
def _import(name, path):
    if path not in sys.path:
        sys.path.insert(0, path)
    try:
        with contextlib.redirect_stdout(io.StringIO()):
            return importlib.import_module(name), None
    except BaseException:  # noqa: B902  (SystemExit 등도 시험 실패로 보고)
        return None, traceback.format_exc()


JOB, JOB_ERR = _import(JOB_NAME, JOB_DIR)
V23, V23_ERR = _import("excel_automation_v23", V23_DIR)


def setUpModule():
    os.makedirs(WORK, exist_ok=True)
    for name in os.listdir(WORK):            # 이전 실행 잔재만 정리 (WORK 안에서만)
        p = os.path.join(WORK, name)
        if name.startswith("t_") and os.path.isdir(p):
            shutil.rmtree(p, ignore_errors=True)


# ------------------------------------------------------------------ 작은 도구
def md5(path):
    with open(path, "rb") as f:
        return hashlib.md5(f.read()).hexdigest()


def close(a, b, rel=1e-9, abs_=1e-6):
    if a is None or b is None:
        return a is b
    try:
        a, b = float(a), float(b)
    except (TypeError, ValueError):
        return False
    return abs(a - b) <= max(abs_, rel * max(abs(a), abs(b)))


def is_num(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


class BadAmount(Exception):
    pass


def parse_amount(v):
    """시험용 독립 해석기 (SPEC '원본 값 해석'). 빈칸·'-' = 0, 해석 불가 → BadAmount."""
    if v is None:
        return 0.0
    if is_num(v):
        return float(v)
    s = str(v).strip()
    if s in ("", "-"):
        return 0.0
    t = s.replace(",", "").replace("₩", "").replace("$", "")
    if t.endswith("원"):
        t = t[:-1]
    t = t.strip()
    neg = False
    if t.startswith("(") and t.endswith(")"):
        neg, t = True, t[1:-1]
    if t.endswith("-") and len(t) > 1:
        neg, t = True, t[:-1]
    if t.startswith("△") or t.startswith("▲"):
        neg, t = True, t[1:]
    try:
        f = float(t)
    except ValueError:
        raise BadAmount(s)
    if not math.isfinite(f):
        raise BadAmount(s)
    return -f if neg else f


def write_xlsx(path, header, rows, pre_rows=(), sheet="Sheet", extra_sheets=None, before_sheets=None):
    wb = Workbook()
    ws = wb.active
    if before_sheets:
        ws.title = before_sheets[0][0]
        for r in before_sheets[0][1]:
            ws.append(list(r))
        ws = wb.create_sheet(sheet)
    else:
        ws.title = sheet
    for r in pre_rows:
        ws.append(list(r))
    ws.append(list(header))
    for r in rows:
        ws.append(list(r))
    for name, srows in (extra_sheets or []):
        es = wb.create_sheet(name)
        for r in srows:
            es.append(list(r))
    wb.save(path)
    return path


def write_rules_xlsx(path, rules, order=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "규칙"
    ws.append(["항목", "값"])
    for k, v in rules.items():
        ws.append([k, v])
    if order is not None:
        os_ = wb.create_sheet("행 순서")
        for item in order:
            os_.append([item])
    wb.save(path)
    return path


def base_rules(**over):
    r = OrderedDict([
        ("원본 시트", ""), ("행 기준", "계정명"), ("열 기준", "일자:월"), ("금액", "금액"), ("집계", "합계"),
        ("필터", ""), ("제외 글자", ""), ("단위", "1"), ("소계", "아니오"), ("합계 행", "예"),
        ("합계 열", "예"), ("정렬", "원본 순"), ("빈 칸", "빈칸"), ("합계 수식", "아니오"),
    ])
    for k, v in over.items():
        r[k.replace("_", " ")] = v
    return r


def month_key(d):
    return f"{d.month}월"


# ------------------------------------------------------------------ 독립 기대값 (oracle)
class Agg:
    """items: [(행 튜플, 열 값, 금액)] 원본 순서. 단위 적용 전 금액."""

    def __init__(self, items, unit=1.0, count=False):
        self.unit, self.count = float(unit), count
        self.sum, self.n, self.first = defaultdict(float), Counter(), OrderedDict()
        for i, (rk, ck, a) in enumerate(items):
            rk = tuple(rk)
            self.sum[(rk, ck)] += a
            self.n[(rk, ck)] += 1
            self.first.setdefault(rk, i)
        self.values = {k: (float(self.n[k]) if count else self.sum[k] / self.unit) for k in self.n}
        self.rows_used = len(items)

    @property
    def col_keys(self):
        return {ck for (_, ck) in self.values}

    def row_total(self, rk):
        return sum(v for (r, _), v in self.values.items() if r == rk)

    def col_total(self, ck):
        return sum(v for (_, c), v in self.values.items() if c == ck)

    def group_val(self, g, ck=None):
        return sum(v for (r, c), v in self.values.items() if r[0] == g and (ck is None or c == ck))

    def group_has(self, g, ck):
        return any(r[0] == g and c == ck for (r, c) in self.values)

    @property
    def grand(self):
        return sum(self.values.values())

    def order(self, how):
        rks = list(self.first)
        if len(rks[0]) == 1:
            if how == "원본 순":
                return rks
            if how == "가나다":
                return sorted(rks)
            if how == "금액 큰 순":
                return sorted(rks, key=lambda k: -self.row_total(k))
            raise ValueError(how)
        groups = OrderedDict()
        for k in rks:
            groups.setdefault(k[0], []).append(k)
        if how == "원본 순":
            gs = list(groups)
            inner = lambda k: self.first[k]
        elif how == "가나다":
            gs = sorted(groups)
            inner = lambda k: k
        elif how == "금액 큰 순":
            gs = sorted(groups, key=lambda g: -self.group_val(g))
            inner = lambda k: -self.row_total(k)
        else:
            raise ValueError(how)
        out = []
        for g in gs:
            out += sorted(groups[g], key=inner)
        return out


# ------------------------------------------------------------------ 보고서 읽기
def _merged_map(ws):
    m = {}
    for rng in ws.merged_cells.ranges:
        tl = ws.cell(rng.min_row, rng.min_col).value
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                m[(r, c)] = tl
    return m


def _txt(v):
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def parse_report(path, row_names):
    """'보고서' 시트 → (ws, 머리글 행, {열 제목: 열 번호}, [행 정보]). 상위 계층 빈칸은 병합/앞 행 값으로 채워 읽음."""
    wb = load_workbook(path)  # 수식 원문 그대로
    if "보고서" not in wb.sheetnames:
        raise AssertionError(f"'보고서' 시트 없음: {wb.sheetnames}")
    ws = wb["보고서"]
    mm = _merged_map(ws)
    val = lambda r, c: mm.get((r, c), ws.cell(r, c).value)
    n = len(row_names)
    hdr = None
    for r in range(1, 21):
        if [_txt(val(r, c)) for c in range(1, n + 1)] == list(row_names):
            hdr = r
            break
    if hdr is None:
        raise AssertionError(f"머리글 줄(왼쪽 {row_names})을 1~20행에서 못 찾음")
    cols = OrderedDict()
    c = n + 1
    while c <= ws.max_column:
        h = _txt(val(hdr, c))
        if h is None:
            break
        if h in cols:
            raise AssertionError(f"머리글 '{h}' 중복")
        cols[h] = c
        c += 1
    rows, carry = [], [None] * n
    for r in range(hdr + 1, ws.max_row + 1):
        labels = [_txt(val(r, c)) for c in range(1, n + 1)]
        cells = OrderedDict((h, ws.cell(r, cc)) for h, cc in cols.items())
        if all(l is None for l in labels) and all(cl.value in (None, "") for cl in cells.values()):
            continue
        sub = next((l for l in labels if l and l.endswith("소계")), None)
        if sub is not None:
            rows.append({"kind": "sub", "group": sub[:-2].strip(), "labels": tuple(labels), "r": r, "cells": cells})
            carry = [None] * n
        elif any(l == "합계" for l in labels):
            rows.append({"kind": "total", "labels": tuple(labels), "r": r, "cells": cells})
            carry = [None] * n
        else:
            for i in range(n - 1):
                if labels[i] is None:
                    labels[i] = carry[i]
            carry = list(labels)
            rows.append({"kind": "detail", "labels": tuple(labels), "r": r, "cells": cells})
    return ws, hdr, cols, rows


def parse_evidence(path):
    """'근거' 시트 → [{coord, count, amount}] (위치가 칸 주소인 줄만)."""
    wb = load_workbook(path, data_only=True)
    if "근거" not in wb.sheetnames:
        raise AssertionError(f"'근거' 시트 없음: {wb.sheetnames}")
    ws = wb["근거"]
    grid = [list(r) for r in ws.iter_rows(values_only=True)]
    for h, row in enumerate(grid[:15]):
        texts = [str(v) if v is not None else "" for v in row]
        pos = next((j for j, t in enumerate(texts) if "위치" in t), None)
        cnt = next((j for j, t in enumerate(texts) if "건수" in t), None)
        amt = next((j for j, t in enumerate(texts) if "금액" in t), None)
        if pos is not None and cnt is not None and amt is not None:
            break
    else:
        raise AssertionError("'근거' 시트에서 위치/건수/금액 머리글을 못 찾음")
    out = []
    for row in grid[h + 1:]:
        if pos >= len(row) or row[pos] is None:
            continue
        m = re.search(r"([A-Z]{1,3}\d+)\s*$", str(row[pos]).replace("$", ""))
        if not m:
            continue
        out.append({"coord": m.group(1), "count": row[cnt], "amount": row[amt]})
    return out


def sheet_text(path, sheet):
    wb = load_workbook(path, data_only=True)
    if sheet not in wb.sheetnames:
        raise AssertionError(f"'{sheet}' 시트 없음: {wb.sheetnames}")
    return "\n".join(" | ".join("" if v is None else str(v) for v in r) for r in wb[sheet].iter_rows(values_only=True))


def is_yellow(cell):
    f = cell.fill
    if f is None or f.fill_type in (None, "none"):
        return False
    rgb = getattr(f.fgColor, "rgb", None)
    return isinstance(rgb, str) and rgb.upper().endswith("FFFF00")


def _expand(ref, ws):
    ref = ref.replace("$", "")
    if re.fullmatch(r"[A-Z]{1,3}", ref.split(":")[0]):          # 'C:C' 전체 열
        c1, _, c2, _ = range_boundaries(ref)
        r1, r2 = 1, ws.max_row
    else:
        c1, r1, c2, r2 = range_boundaries(ref)
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            yield (r, c)


def coverage(ws, formula, detail_set, depth=0):
    """수식이 실제로 더하는 상세 칸 Counter와 문제 목록. SUM/SUBTOTAL(9|109)/+ 만 허용, 중첩 수식은 펼침."""
    cnt, problems = Counter(), []
    if depth > 20:
        return cnt, ["수식 순환/과도 중첩"]
    stack = []
    try:
        items = Tokenizer(formula).items
    except Exception as e:
        return cnt, [f"수식 해석 불가 {formula!r}: {e}"]
    for t in items:
        if t.type == Token.FUNC:
            if t.subtype == Token.OPEN:
                name = t.value[:-1].upper().replace("_XLFN.", "")
                if name not in ("SUM", "SUBTOTAL"):
                    problems.append(f"허용 안 된 함수 {name} in {formula}")
                stack.append(name)
            elif stack:
                stack.pop()
        elif t.type == Token.OPERAND and t.subtype == Token.RANGE:
            ref = t.value
            if "!" in ref:
                sh, ref = ref.rsplit("!", 1)
                if sh.strip("'") != ws.title:
                    problems.append(f"다른 시트 참조 {t.value} in {formula}")
                    continue
            for (r, c) in _expand(ref, ws):
                v = ws.cell(r, c).value
                if isinstance(v, str) and v.startswith("="):
                    if "SUBTOTAL" in stack and "SUBTOTAL" in v.upper():
                        continue                               # SUBTOTAL은 안쪽 SUBTOTAL을 무시
                    sub, pr = coverage(ws, v, detail_set, depth + 1)
                    cnt += sub
                    problems += pr
                elif is_num(v):
                    if (r, c) in detail_set:
                        cnt[(r, c)] += 1
                    else:
                        problems.append(f"상세 칸이 아닌 숫자 {get_column_letter(c)}{r} 포함 in {formula}")
        elif t.type == Token.OPERAND and t.subtype == Token.NUMBER:
            if not (stack and stack[-1] == "SUBTOTAL" and t.value in ("9", "109")):
                problems.append(f"상수 {t.value} in {formula}")
        elif t.type == Token.OP_IN and t.value != "+":
            problems.append(f"연산자 {t.value} in {formula}")
        elif t.type == Token.OP_PRE:
            problems.append(f"단항 연산자 {t.value} in {formula}")
    return cnt, problems


# ------------------------------------------------------------------ 기본 시험 클래스
class JobCase(unittest.TestCase):
    maxDiff = 4000

    def setUp(self):
        if JOB is None:
            self.fail(f"작업 모듈 import 실패 ({JOB_NAME} @ {JOB_DIR}):\n{JOB_ERR}")
        os.makedirs(WORK, exist_ok=True)
        self.dir = tempfile.mkdtemp(prefix="t_" + self._testMethodName[5:30] + "_", dir=WORK)

    # --- 실행
    def build(self, raw, rules, out=None, name=None):
        if out is None and name is not None:
            out = os.path.join(os.path.dirname(raw), name)
        before = md5(raw)
        buf = io.StringIO()
        try:
            with contextlib.redirect_stdout(buf):
                res = JOB.build_report(raw, rules, out_path=out)
        except BaseException:  # noqa
            self.fail(f"build_report 예외:\n{traceback.format_exc()}\n출력:\n{buf.getvalue()[-2000:]}")
        self.assertEqual(md5(raw), before, "원본 파일이 수정됨")
        self.assertIsInstance(res, tuple, "build_report는 (out_path, info) 반환")
        out_path, info = res
        self.assertTrue(out_path and os.path.exists(out_path), f"보고서 파일 없음: {out_path}")
        if out is not None:
            self.assertEqual(os.path.abspath(out_path), os.path.abspath(out))
        for k in ("row_keys", "col_keys", "values", "grand_total", "raw_total", "new_items",
                  "excluded_rows", "bad_rows", "warnings", "rules"):
            self.assertIn(k, info, f"info에 '{k}' 없음")
        return out_path, info

    def expect_error(self, raw, rules, *needles):
        buf = io.StringIO()
        with contextlib.redirect_stdout(buf):
            try:
                JOB.build_report(raw, rules, out_path=os.path.join(self.dir, "should_not_exist.xlsx"))
            except Exception as e:
                msg = str(e) + buf.getvalue()
                for nd in needles:
                    self.assertIn(nd, msg, f"오류 문구에 '{nd}' 없음: {msg!r}")
                return msg
        self.fail(f"오류가 나야 하는데 정상 종료됨 (needles={needles}); 출력: {buf.getvalue()[-500:]}")

    def run_cli(self, script, args=(), cwd=None, timeout=300):
        env = dict(os.environ, PYTHONIOENCODING="utf-8", PYTHONDONTWRITEBYTECODE="1")
        p = subprocess.run([sys.executable, script] + list(args), cwd=cwd or os.path.dirname(script),
                           capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=timeout, env=env)
        lines = [l for l in p.stdout.splitlines() if l.strip()]
        return p.returncode, (lines[-1].strip() if lines else ""), p.stdout + p.stderr

    def job_copy(self, folder):
        src = os.path.join(JOB_DIR, JOB_NAME + ".py")
        if not os.path.exists(src):
            self.fail(f"작업 파일 없음: {src}")
        dst = os.path.join(folder, os.path.basename(src))
        shutil.copy2(src, dst)
        return dst

    # --- 대조
    def assert_info(self, info, agg, ncheck_order=None):
        got = {(tuple(rk), str(ck)): v for (rk, ck), v in info["values"].items()}
        exp = {(rk, str(ck)): v for (rk, ck), v in agg.values.items()}
        self.assertEqual(set(got), set(exp), "info['values'] 칸 집합이 기대와 다름 (원본 건 없는 칸은 키 없음)")
        bad = [(k, got[k], exp[k]) for k in exp if not close(got[k], exp[k])]
        self.assertEqual(bad, [], "info['values'] 값 불일치 (칸, 보고서, 기대)")
        self.assertTrue(close(info["grand_total"], agg.grand), f"grand_total {info['grand_total']} != 기대 {agg.grand}")
        self.assertTrue(close(info["raw_total"], info["grand_total"]), f"raw_total {info['raw_total']} != grand_total {info['grand_total']}")
        if ncheck_order is not None:
            self.assertEqual([tuple(k) for k in info["row_keys"]], [tuple(k) for k in ncheck_order], "info['row_keys'] 순서")

    def assert_sheet(self, out, rules, agg, info, order=None, fmt=None, formula=False):
        names = [s.strip() for s in str(rules["행 기준"]).split(",")]
        n = len(names)
        ws, hdr, cols, rows = parse_report(out, names)
        blank_zero = str(rules.get("빈 칸", "빈칸")).strip() == "0"
        details = [r for r in rows if r["kind"] == "detail"]
        got_order = [r["labels"] for r in details]
        self.assertEqual(len(got_order), len(set(got_order)), f"상세 행 중복: {got_order}")
        self.assertEqual(got_order, [tuple(k) for k in info["row_keys"]], "보고서 상세 행 순서 != info['row_keys']")
        if order is not None:
            self.assertEqual(got_order, [tuple(k) for k in order], "상세 행 순서")
        else:
            self.assertEqual(set(got_order), set(agg.first), "상세 행 집합")
        colkeys = [str(c) for c in info["col_keys"]]
        has_col_basis = bool(str(rules.get("열 기준") or "").strip())
        want_total_col = has_col_basis and str(rules.get("합계 열")).strip() == "예"
        self.assertEqual(list(cols), colkeys + (["합계"] if want_total_col else []), "머리글 열 제목")
        detail_set, numeric_details = {}, {}
        # 상세 칸
        for r in details:
            rk = r["labels"]
            for ck in colkeys:
                cell = r["cells"][ck]
                detail_set[(cell.row, cell.column)] = (rk, ck)
                exp = agg.values.get((rk, ck))
                v = cell.value
                if is_num(v):
                    numeric_details[(cell.row, cell.column)] = v
                if exp is None:
                    if blank_zero:
                        self.assertTrue(is_num(v) and v == 0, f"{cell.coordinate} 원본 건 없음 → 0 이어야 함 (빈 칸=0): {v!r}")
                    else:
                        self.assertIsNone(v, f"{cell.coordinate} 원본 건 없음 → 빈칸이어야 함: {v!r}")
                else:
                    self.assertTrue(is_num(v), f"{cell.coordinate} {rk}×{ck} 숫자 값이어야 함 (수식 아님): {v!r}")
                    self.assertTrue(close(v, exp), f"{cell.coordinate} {rk}×{ck}: 보고서 {v} != 기대 {exp}")
                    if fmt:
                        self.assertIn(cell.number_format, fmt if isinstance(fmt, (tuple, list, set)) else (fmt,),
                                      f"{cell.coordinate} 표시 형식")
        total_cells = []   # (cell, 기대값, 기대 coverage)
        dcoords = lambda rr, cks: Counter({(x["cells"][ck].row, x["cells"][ck].column): 1 for x in rr for ck in cks
                                           if (x["cells"][ck].row, x["cells"][ck].column) in numeric_details})
        if want_total_col:
            for r in details:
                total_cells.append((r["cells"]["합계"], agg.row_total(r["labels"]), dcoords([r], colkeys), True))
        # 소계
        subs = [r for r in rows if r["kind"] == "sub"]
        if n >= 2 and str(rules.get("소계")).strip() == "예":
            groups = OrderedDict()
            for i, r in enumerate(rows):
                if r["kind"] == "detail":
                    groups.setdefault(r["labels"][0], []).append(i)
            self.assertEqual([s["group"] for s in subs], list(groups), "소계 행: 그룹마다 하나, 그룹 순서대로")
            for g, idxs in groups.items():
                self.assertEqual(idxs, list(range(idxs[0], idxs[-1] + 1)), f"'{g}' 상세 행이 붙어 있지 않음")
                nxt = rows[idxs[-1] + 1] if idxs[-1] + 1 < len(rows) else None
                self.assertTrue(nxt is not None and nxt["kind"] == "sub" and nxt["group"] == g,
                                f"'{g} 소계' 행이 그룹 마지막 상세 행 바로 아래에 없음")
                grp_rows = [rows[i] for i in idxs]
                for ck in colkeys:
                    total_cells.append((nxt["cells"][ck], agg.group_val(g, ck) if agg.group_has(g, ck) else None,
                                        dcoords(grp_rows, [ck]), agg.group_has(g, ck)))
                if want_total_col:
                    total_cells.append((nxt["cells"]["합계"], agg.group_val(g), dcoords(grp_rows, colkeys), True))
                lab = next(ws.cell(nxt["r"], c) for c in range(1, n + 1) if _txt(ws.cell(nxt["r"], c).value))
                self.assertTrue(lab.font and lab.font.bold, f"소계 행 {nxt['r']} 굵게")
        else:
            self.assertEqual(subs, [], "소계=아니오(또는 행 기준 1개)인데 소계 행이 있음")
        # 합계 행
        tots = [r for r in rows if r["kind"] == "total"]
        if str(rules.get("합계 행")).strip() == "예":
            self.assertEqual(len(tots), 1, "합계 행은 정확히 1개")
            self.assertIs(rows[-1], tots[0], "합계 행은 맨 아래")
            t = tots[0]
            for ck in colkeys:
                total_cells.append((t["cells"][ck], agg.col_total(ck), dcoords(details, [ck]), True))
            if want_total_col:
                total_cells.append((t["cells"]["합계"], agg.grand, dcoords(details, colkeys), True))
            lab = next(ws.cell(t["r"], c) for c in range(1, n + 1) if _txt(ws.cell(t["r"], c).value))
            self.assertTrue(lab.font and lab.font.bold, "합계 행 굵게")
        else:
            self.assertEqual(tots, [], "합계 행=아니오인데 합계 행이 있음")
        # 소계·합계 칸 값 또는 수식
        for cell, exp, cov, has_data in total_cells:
            v = cell.value
            if formula:
                self.assertTrue(isinstance(v, str) and v.startswith("="), f"{cell.coordinate} 합계 수식=예 → 수식이어야 함: {v!r}")
                got, problems = coverage(ws, v, detail_set)
                self.assertEqual(problems, [], f"{cell.coordinate} 수식 문제 {v}")
                miss = sorted(get_column_letter(c) + str(r) for (r, c) in cov if got[(r, c)] == 0)
                dup = sorted(get_column_letter(c) + str(r) for (r, c), k in got.items() if k > 1)
                extra = sorted(get_column_letter(c) + str(r) for (r, c) in got if (r, c) not in cov)
                self.assertEqual((miss, dup, extra), ([], [], []), f"{cell.coordinate} {v}: (빠진 칸, 중복 칸, 엉뚱한 칸)")
            else:
                if exp is None or (not has_data and not blank_zero):
                    self.assertTrue(v is None or (is_num(v) and v == 0), f"{cell.coordinate} 근거 없는 소계 칸: {v!r}")
                else:
                    self.assertTrue(is_num(v), f"{cell.coordinate} 소계·합계 값이어야 함 (합계 수식=아니오): {v!r}")
                    self.assertTrue(close(v, exp), f"{cell.coordinate}: 보고서 {v} != 기대 {exp}")
        return ws, hdr, cols, rows, numeric_details

    def assert_v23(self, raw, out, numeric_details=None, require_detail=True):
        if V23 is None:
            self.fail(f"v23 엔진 import 실패:\n{V23_ERR}")
        buf = io.StringIO()
        try:
            with contextlib.redirect_stdout(buf), contextlib.redirect_stderr(io.StringIO()):
                res = V23.verify(raw, out)
        except Exception:
            self.fail(f"v23 verify 예외:\n{traceback.format_exc()}")
        self.assertEqual(res["mismatch"], [], f"v23 불일치: {res['mismatch'][:5]}")
        if require_detail and numeric_details:
            by = {c["위치"]: c["판정"] for c in res["checks"]}
            bad = [(f"보고서!{get_column_letter(c)}{r}", by.get(f"보고서!{get_column_letter(c)}{r}"))
                   for (r, c) in numeric_details if by.get(f"보고서!{get_column_letter(c)}{r}") != "일치"]
            self.assertEqual(bad, [], f"v23에서 일치가 아닌 상세 칸 (mapping={res.get('mapping')})")
        return res

    def assert_evidence(self, out, info, agg, data_rows, unit=1.0, positions=None):
        ev = parse_evidence(out)
        cnt = sum(int(e["count"] or 0) for e in ev)
        self.assertEqual(cnt + int(info["excluded_rows"]) + len(info["bad_rows"]), data_rows,
                         "근거 건수 합 + excluded_rows + bad_rows != 원본 자료 행 수 (행이 조용히 빠지거나 두 번 셈)")
        self.assertEqual(cnt, agg.rows_used, "근거 건수 합 != 집계에 쓴 원본 행 수")
        if positions:
            byc = {e["coord"]: e for e in ev}
            for (r, c), key in positions.items():
                coord = f"{get_column_letter(c)}{r}"
                self.assertIn(coord, byc, f"근거 시트에 {coord} 없음")
                e = byc[coord]
                self.assertEqual(int(e["count"]), agg.n[key], f"근거 {coord} 건수")
                self.assertTrue(close(e["amount"], agg.values[key]) or close(e["amount"], agg.sum[key]),
                                f"근거 {coord} 금액 {e['amount']} != {agg.values[key]}")


# ================================================================== 데이터 생성기
ACCTS = ["현금", "매출", "외상매출금", "설계비", "회계감사비", "이자수익", "잡손실"]
MONTHS = [1, 2, 3, 10, 11, 12]
EMPTY = {("이자수익", 2), ("이자수익", 3), ("설계비", 10), ("회계감사비", 1), ("잡손실", 1), ("잡손실", 2),
         ("잡손실", 3), ("잡손실", 10), ("잡손실", 12)}


def basic_rows(seed=1, n=240):
    rng = random.Random(seed)
    rows = []
    while len(rows) < n:
        a, m = rng.choice(ACCTS[:-1]), rng.choice(MONTHS)
        if (a, m) in EMPTY:
            continue
        d = datetime.datetime(2026, m, rng.randint(1, 28))
        rows.append([d, a, rng.choice(["갑상사", "을상사", "병상사"]), rng.randint(-50000, 900000)])
    rows.insert(5, [datetime.datetime(2026, 11, 3), "잡손실", "갑상사", 700])     # 합이 0인 칸 (원본 건 있음)
    rows.append([datetime.datetime(2026, 11, 20), "잡손실", "을상사", -700])
    return rows


ML_SPEC = {"본사": ["매출", "매출원가", "급여"], "지점": ["매출", "급여", "임차료"], "해외": ["외환차익", "매출"]}


def multilevel_rows(seed=7):
    rng = random.Random(seed)
    combos = [(b, a) for b, accs in ML_SPEC.items() for a in accs]
    rows = []
    # 첫 등장 순서를 '가나다'·'금액 큰 순'과 다르게 섞음
    for b, a in [("지점", "임차료"), ("본사", "급여"), ("해외", "매출"), ("지점", "매출"), ("본사", "매출"),
                 ("해외", "외환차익"), ("본사", "매출원가"), ("지점", "급여")]:
        rows.append([b, a, datetime.datetime(2026, 1, 2), rng.randint(1000, 9000)])
    for _ in range(150):
        b, a = rng.choice(combos)
        m = rng.choice([1, 2, 3, 4])
        if (b, a, m) in {("해외", "외환차익", 2), ("지점", "임차료", 4)}:
            continue
        rows.append([b, a, datetime.datetime(2026, m, rng.randint(1, 28)), rng.randint(1000, 500000)])
    return rows


def ml_agg(rows, unit=1.0):
    return Agg([((b, a), month_key(d), amt) for b, a, d, amt in rows], unit)


# ================================================================== 1. 기본
class T01_Basic(JobCase):
    def test_basic_month_columns_exact_and_v23(self):
        rows = basic_rows()
        raw = write_xlsx(os.path.join(self.dir, "기본원본.xlsx"), ["일자", "계정명", "거래처", "금액"], rows)
        rules = base_rules()
        out, info = self.build(raw, rules, name="기본원본_보고서.xlsx")
        agg = Agg([((a,), month_key(d), amt) for d, a, _, amt in rows])
        self.assert_info(info, agg, ncheck_order=agg.order("원본 순"))
        self.assertEqual(info["excluded_rows"], 0)
        self.assertEqual(list(info["bad_rows"]), [])
        self.assertEqual(list(info["new_items"]), [])
        # 열: 월은 1월..12월 시간 순 (문자 정렬 '10월'<'2월' 금지). 데이터 없는 달 포함 여부는 허용
        ck = [str(c) for c in info["col_keys"]]
        present = [f"{m}월" for m in sorted(MONTHS)]
        self.assertIn(ck, (present, [f"{m}월" for m in range(1, 13)]), f"월 열 순서/구성: {ck}")
        self.assertIn((("잡손실",), "11월"), {(tuple(k[0]), str(k[1])) for k in info["values"]}, "합이 0인 칸도 원본 건이 있으면 값 키가 있어야 함")
        ws, hdr, cols, prow, nd = self.assert_sheet(out, rules, agg, info, order=agg.order("원본 순"), fmt="#,##0")
        # 합이 0인 칸은 빈칸 모드에서도 0 표시
        r_ = next(r for r in prow if r["kind"] == "detail" and r["labels"] == ("잡손실",))
        self.assertEqual(r_["cells"]["11월"].value, 0, "원본 건이 있고 합이 0인 칸은 빈칸이 아니라 0")
        # 제목·서식
        t1 = " ".join(str(c.value) for c in ws[1] if c.value is not None)
        self.assertIn("기본원본", t1, "1행 제목에 원본 이름")
        self.assertGreaterEqual(hdr, 2, "머리글은 제목 아래")
        h = ws.cell(hdr, 1)
        self.assertTrue(h.font and h.font.color is not None and str(h.font.color.rgb).upper().endswith("FFFFFF"), "머리글 흰 글씨")
        self.assertTrue(h.fill and h.fill.fill_type == "solid", "머리글 채우기(진한 파랑)")
        self.assertIsNotNone(ws.freeze_panes, "틀 고정")
        # 규칙 시트, 근거 시트
        rt = sheet_text(out, "규칙")
        for k in ("행 기준", "열 기준", "금액"):
            self.assertIn(k, rt, f"'규칙' 시트에 '{k}'")
        self.assertIn("일자:월", rt)
        self.assertEqual(str(info["rules"].get("행 기준")).strip(), "계정명", "info['rules']는 실제 적용 규칙")
        positions = {(r["cells"][c].row, r["cells"][c].column): (r["labels"], c) for r in prow if r["kind"] == "detail"
                     for c in ck if (r["labels"], c) in agg.values}
        self.assert_evidence(out, info, agg, len(rows), positions=positions)
        self.assert_v23(raw, out, nd)

    def test_blank_cell_zero_mode(self):
        rows = basic_rows(seed=2)
        raw = write_xlsx(os.path.join(self.dir, "빈칸0.xlsx"), ["일자", "계정명", "거래처", "금액"], rows)
        rules = base_rules(**{"빈 칸": "0"})
        out, info = self.build(raw, rules, name="빈칸0_보고서.xlsx")
        agg = Agg([((a,), month_key(d), amt) for d, a, _, amt in rows])
        self.assert_info(info, agg)                      # 0으로 보인 칸도 info['values']에는 키 없음
        _, _, _, _, nd = self.assert_sheet(out, rules, agg, info, order=agg.order("원본 순"))
        self.assert_v23(raw, out, nd)

    def test_default_output_name_and_rules_file_in_folder(self):
        """rules=None → 같은 폴더 규칙표_job.xlsx 사용, out_path=None → <원본이름>_보고서_YYYYMMDD_HHMMSS.xlsx"""
        rows = basic_rows(seed=3, n=60)
        raw = write_xlsx(os.path.join(self.dir, "폴더규칙.xlsx"), ["일자", "계정명", "거래처", "금액"], rows)
        rules = base_rules(**{"열 기준": "", "정렬": "가나다", "단위": 1000})
        write_rules_xlsx(os.path.join(self.dir, "규칙표_job.xlsx"), rules)
        out, info = self.build(raw, None)
        self.assertEqual(os.path.dirname(os.path.abspath(out)), os.path.abspath(self.dir), "보고서는 원본 옆")
        self.assertRegex(os.path.basename(out), r"^폴더규칙_보고서_\d{8}_\d{6}\.xlsx$")
        agg = Agg([((a,), "금액", amt) for d, a, _, amt in rows], unit=1000)
        self.assertEqual(len(info["col_keys"]), 1, "열 기준 비움 → 금액 열 하나")
        info_vals = {(tuple(k[0]), "금액"): v for k, v in info["values"].items()}
        self.assertEqual(set(info_vals), set(agg.values))
        for k in agg.values:
            self.assertTrue(close(info_vals[k], agg.values[k]), f"{k}: {info_vals[k]} != {agg.values[k]}")
        self.assertEqual([tuple(k) for k in info["row_keys"]], agg.order("가나다"), "규칙표_job.xlsx의 정렬=가나다 적용")
        self.assertTrue(close(info["grand_total"], agg.grand))

    def test_raw_sheet_auto_pick_most_numbers(self):
        rows = basic_rows(seed=4, n=80)
        memo = [("메모", "값"), ("기준", 1), ("참고", 2), ("비고", 3)]
        raw = write_xlsx(os.path.join(self.dir, "여러시트.xlsx"), ["일자", "계정명", "거래처", "금액"], rows,
                         sheet="전표", before_sheets=[("메모", memo)])
        rules = base_rules()
        out, info = self.build(raw, rules, name="여러시트_보고서.xlsx")
        agg = Agg([((a,), month_key(d), amt) for d, a, _, amt in rows])
        self.assert_info(info, agg)


# ================================================================== 2. 다단 행·소계·정렬·행 순서
class T02_MultiLevel(JobCase):
    def _ml(self, how):
        rows = multilevel_rows()
        raw = write_xlsx(os.path.join(self.dir, "다단원본.xlsx"), ["법인", "계정명", "일자", "금액"], rows)
        agg = ml_agg(rows)
        # 정렬 기대값이 동점 없이 정해지는지 시험 데이터 자체 점검
        gt = [agg.group_val(g) for g in ML_SPEC]
        self.assertEqual(len(set(gt)), len(gt), "시험 데이터: 그룹 합계 동점")
        for g in ML_SPEC:
            rt = [agg.row_total(k) for k in agg.first if k[0] == g]
            self.assertEqual(len(set(rt)), len(rt), f"시험 데이터: '{g}' 안 행 합계 동점")
        rules = base_rules(**{"행 기준": "법인, 계정명", "소계": "예", "합계 행": "예", "합계 열": "예", "정렬": how})
        out, info = self.build(raw, rules, name=f"다단_{how.replace(' ', '')}.xlsx")
        order = agg.order(how)
        self.assert_info(info, agg, ncheck_order=order)
        _, _, _, _, nd = self.assert_sheet(out, rules, agg, info, order=order, fmt="#,##0")
        self.assert_v23(raw, out, nd)

    def test_multilevel_subtotals_original_order(self):
        self._ml("원본 순")

    def test_multilevel_subtotals_ganada(self):
        self._ml("가나다")

    def test_multilevel_subtotals_amount_desc(self):
        self._ml("금액 큰 순")

    def test_single_level_amount_desc(self):
        rows = [r for r in basic_rows(seed=5) if r[3] > 0]
        raw = write_xlsx(os.path.join(self.dir, "금액순.xlsx"), ["일자", "계정명", "거래처", "금액"], rows)
        rules = base_rules(**{"정렬": "금액 큰 순"})
        out, info = self.build(raw, rules, name="금액순_보고서.xlsx")
        agg = Agg([((a,), month_key(d), amt) for d, a, _, amt in rows])
        tot = [agg.row_total(k) for k in agg.first]
        self.assertEqual(len(set(tot)), len(tot), "시험 데이터 동점")
        self.assert_info(info, agg, ncheck_order=agg.order("금액 큰 순"))
        self.assert_sheet(out, rules, agg, info, order=agg.order("금액 큰 순"))

    def test_subtotal_no_and_total_row_col_no(self):
        rows = multilevel_rows(seed=8)
        raw = write_xlsx(os.path.join(self.dir, "소계없음.xlsx"), ["법인", "계정명", "일자", "금액"], rows)
        rules = base_rules(**{"행 기준": "법인, 계정명", "소계": "아니오", "합계 행": "아니오", "합계 열": "아니오"})
        out, info = self.build(raw, rules, name="소계없음_보고서.xlsx")
        agg = ml_agg(rows)
        self.assert_info(info, agg, ncheck_order=agg.order("원본 순"))
        self.assert_sheet(out, rules, agg, info, order=agg.order("원본 순"))

    def test_row_order_list_new_items_yellow(self):
        accts = ["현금", "매출", "급여", "임차료", "잡이익", "이자수익"]
        rng = random.Random(11)
        rows = [[datetime.datetime(2026, rng.choice([1, 2, 3]), 5), rng.choice(accts), rng.randint(100, 99999)] for _ in range(120)]
        for a in accts:                                   # 모든 계정 최소 1건
            rows.append([datetime.datetime(2026, 1, 9), a, 1234])
        raw = write_xlsx(os.path.join(self.dir, "행순서원본.xlsx"), ["일자", "계정명", "금액"], rows)
        listed = ["급여", "현금", "매출", "임차료", "복리후생비"]          # 복리후생비: 목록에만 있음
        rp = write_rules_xlsx(os.path.join(self.dir, "규칙표_행순서.xlsx"), base_rules(**{"정렬": "가나다"}), order=listed)
        rules = JOB.load_rules(rp)
        self.assertIsInstance(rules, dict)
        out, info = self.build(raw, rules, name="행순서_보고서.xlsx")
        agg = Agg([((a,), month_key(d), amt) for d, a, amt in rows])
        self.assert_info(info, agg)
        rk = [tuple(k) for k in info["row_keys"]]
        listed_present = [(x,) for x in listed if (x,) in agg.first]
        tail = rk[len(rk) - 2:]
        self.assertEqual(set(tail), {("잡이익",), ("이자수익",)}, f"목록 밖 새 항목은 맨 아래: {rk}")
        head = [k for k in rk[:-2] if k != ("복리후생비",)]
        self.assertEqual(head, listed_present, "행 순서 목록 순서 우선 (정렬=가나다보다 우선)")
        self.assertEqual({tuple(k) for k in info["new_items"]}, {("잡이익",), ("이자수익",)}, "new_items")
        self.assertNotIn((("복리후생비",), "1월"), {(tuple(k[0]), str(k[1])) for k in info["values"]})
        ws, hdr, cols, prow, _ = self.assert_sheet(out, rules, Agg([((a,), month_key(d), amt) for d, a, amt in rows]),
                                                   info, order=[k for k in rk])
        for r in prow:
            if r["kind"] != "detail":
                continue
            yellow = is_yellow(ws.cell(r["r"], 1))
            if r["labels"] in {("잡이익",), ("이자수익",)}:
                self.assertTrue(yellow, f"새 항목 {r['labels']} 노란색")
            else:
                self.assertFalse(yellow, f"기존 항목 {r['labels']} 은 노란색 아님")
        self.assertIn("잡이익", sheet_text(out, "규칙") + " ".join(info["warnings"]), "새 항목은 규칙 시트 또는 경고에 알림")


# ================================================================== 3. 열 기준 변형·날짜 형식
class T03_Columns(JobCase):
    def test_col_yearmonth(self):
        rng = random.Random(21)
        ms = [(2025, 11), (2025, 12), (2026, 1), (2026, 2)]
        rows = [[datetime.datetime(*rng.choice(ms), rng.randint(1, 28)), rng.choice(ACCTS[:4]), rng.randint(-1000, 50000)] for _ in range(150)]
        raw = write_xlsx(os.path.join(self.dir, "연월.xlsx"), ["일자", "계정명", "금액"], rows)
        rules = base_rules(**{"열 기준": "일자:연월"})
        out, info = self.build(raw, rules, name="연월_보고서.xlsx")
        agg = Agg([((a,), f"{d.year}-{d.month:02d}", amt) for d, a, amt in rows])
        self.assertEqual([str(c) for c in info["col_keys"]], ["2025-11", "2025-12", "2026-01", "2026-02"], "연월 시간 순")
        self.assert_info(info, agg)
        _, _, _, _, nd = self.assert_sheet(out, rules, agg, info)
        self.assert_v23(raw, out, nd)

    def test_col_year(self):
        rng = random.Random(22)
        rows = [[datetime.datetime(rng.choice([2026, 2025]), rng.randint(1, 12), 3), rng.choice(ACCTS[:4]), rng.randint(1, 90000)] for _ in range(120)]
        raw = write_xlsx(os.path.join(self.dir, "연도.xlsx"), ["일자", "계정명", "금액"], rows)
        rules = base_rules(**{"열 기준": "일자:연도"})
        out, info = self.build(raw, rules, name="연도_보고서.xlsx")
        ck = [str(c) for c in info["col_keys"]]
        self.assertIn(ck, (["2025", "2026"], ["2025년", "2026년"]), "연도 열 (SPEC 'YYYY'; v23 호환 'YYYY년'도 허용)")
        suffix = "년" if ck[0].endswith("년") else ""
        agg = Agg([((a,), f"{d.year}{suffix}", amt) for d, a, amt in rows])
        self.assert_info(info, agg)
        _, _, _, _, nd = self.assert_sheet(out, rules, agg, info)
        if suffix:
            self.assert_v23(raw, out, nd)
        # 'YYYY' 머리글은 v23이 숫자로 읽어 대조 불가 (SPEC 충돌) → 이 경우 v23 대조 생략

    def test_col_text_column(self):
        rng = random.Random(23)
        corp = ["본사", "지점", " 지점", "해외 "]                    # 열 값도 앞뒤 공백 제거되어 합쳐져야 함
        rows = [[datetime.datetime(2026, 1, 3), rng.choice(ACCTS[:5]), rng.choice(corp), rng.randint(-3000, 80000)] for _ in range(150)]
        raw = write_xlsx(os.path.join(self.dir, "법인열.xlsx"), ["일자", "계정명", "법인", "금액"], rows)
        rules = base_rules(**{"열 기준": "법인"})
        out, info = self.build(raw, rules, name="법인열_보고서.xlsx")
        agg = Agg([((a,), b.strip(), amt) for d, a, b, amt in rows])
        self.assertEqual(set(str(c) for c in info["col_keys"]), {"본사", "지점", "해외"})
        self.assert_info(info, agg)
        _, _, _, _, nd = self.assert_sheet(out, rules, agg, info)
        self.assert_v23(raw, out, nd)

    def test_col_blank_single_amount_column(self):
        rows = basic_rows(seed=24, n=100)
        raw = write_xlsx(os.path.join(self.dir, "열없음.xlsx"), ["일자", "계정명", "거래처", "금액"], rows)
        rules = base_rules(**{"열 기준": ""})
        out, info = self.build(raw, rules, name="열없음_보고서.xlsx")
        self.assertEqual([str(c) for c in info["col_keys"]], ["금액"], "열 기준 비움 → 열 제목 = 금액 열 이름")
        agg = Agg([((a,), "금액", amt) for d, a, _, amt in rows])
        self.assert_info(info, agg)
        _, _, _, _, nd = self.assert_sheet(out, rules, agg, info)    # 합계 열 없어야 함 (열 기준 없음)
        self.assert_v23(raw, out, nd)

    def test_date_formats_mixed(self):
        def serial(y, m, d):
            return (datetime.date(y, m, d) - SERIAL_EPOCH).days
        forms = [
            lambda y, m, d: datetime.datetime(y, m, d),
            lambda y, m, d: f"{y}-{m:02d}-{d:02d}",
            lambda y, m, d: f"{y}.{m:02d}.{d:02d}",
            lambda y, m, d: f"{y}년 {m}월 {d}일",
            serial,
        ]
        rng = random.Random(25)
        rows, items = [], []
        for i in range(200):
            m, d = rng.choice([1, 2, 3, 10, 12]), rng.randint(1, 28)
            a, amt = rng.choice(ACCTS[:4]), rng.randint(1, 99999)
            rows.append([forms[i % len(forms)](2026, m, d), a, amt])
            items.append(((a,), f"{m}월", amt))
        raw = write_xlsx(os.path.join(self.dir, "날짜형식.xlsx"), ["일자", "계정명", "금액"], rows)
        rules = base_rules()
        out, info = self.build(raw, rules, name="날짜형식_보고서.xlsx")
        agg = Agg(items)
        self.assertEqual(list(info["bad_rows"]), [], "정상 날짜 형식은 오류 행이 아님")
        self.assert_info(info, agg)
        self.assert_sheet(out, rules, agg, info)

    def test_unparseable_date_not_silently_dropped(self):
        """판단: 날짜를 읽을 수 없는 행은 조용히 빠지면 안 됨 → bad_rows 또는 경고에 엑셀 행 번호."""
        rows = [[datetime.datetime(2026, 1, 5), "현금", 1000], [datetime.datetime(2026, 2, 5), "현금", 2000],
                ["미정", "현금", 5000], [datetime.datetime(2026, 2, 7), "매출", 3000]]
        raw = write_xlsx(os.path.join(self.dir, "날짜오류.xlsx"), ["일자", "계정명", "금액"], rows)
        out, info = self.build(raw, base_rules(), name="날짜오류_보고서.xlsx")
        excel_row = 4
        in_bad = any(int(r) == excel_row for r, _ in info["bad_rows"])
        warned = any(re.search(rf"(?<!\d){excel_row}(?!\d)", str(w)) for w in info["warnings"])
        self.assertTrue(in_bad or warned, f"날짜 '미정' 행(엑셀 {excel_row}행)이 bad_rows/경고 어디에도 없음: {info['warnings']}")
        self.assertTrue(close(info["raw_total"], info["grand_total"]))
        self.assertIn(info["grand_total"], (6000, 11000))
        if in_bad:
            self.assertEqual(info["grand_total"], 6000, "bad_rows로 뺐으면 합계에 넣지 않음")


# ================================================================== 4. 금액 식·건수·단위
class T04_AmountUnit(JobCase):
    def test_debit_minus_credit(self):
        rng = random.Random(31)
        rows, items = [], []
        for _ in range(160):
            d = datetime.datetime(2026, rng.choice([1, 2, 3]), 10)
            a = rng.choice(ACCTS[:5])
            dr, cr = (rng.randint(1, 90000), None) if rng.random() < .5 else (None, rng.randint(1, 90000))
            if rng.random() < .1:
                dr, cr = rng.randint(1, 500), rng.randint(1, 500)
            rows.append([d, a, dr, cr])
            items.append(((a,), month_key(d), (dr or 0) - (cr or 0)))
        raw = write_xlsx(os.path.join(self.dir, "차대.xlsx"), ["일자", "계정명", "차변", "대변"], rows)
        rules = base_rules(**{"금액": "차변-대변"})
        out, info = self.build(raw, rules, name="차대_보고서.xlsx")
        agg = Agg(items)
        self.assert_info(info, agg)
        _, _, _, _, nd = self.assert_sheet(out, rules, agg, info)
        self.assert_v23(raw, out, nd)
        rules2 = base_rules(**{"금액": "차변+대변"})
        out2, info2 = self.build(raw, rules2, name="차대합_보고서.xlsx")
        agg2 = Agg([((r[1],), month_key(r[0]), (r[2] or 0) + (r[3] or 0)) for r in rows])
        self.assert_info(info2, agg2)

    def test_count_aggregation(self):
        rows = basic_rows(seed=32, n=150)
        raw = write_xlsx(os.path.join(self.dir, "건수.xlsx"), ["일자", "계정명", "거래처", "금액"], rows)
        rules = base_rules(**{"집계": "건수"})
        out, info = self.build(raw, rules, name="건수_보고서.xlsx")
        agg = Agg([((a,), month_key(d), amt) for d, a, _, amt in rows], count=True)
        self.assert_info(info, agg)          # 판단: 건수 집계면 raw_total도 건수 (= grand_total)
        self.assertEqual(info["grand_total"], len(rows))
        self.assert_sheet(out, rules, agg, info, fmt="#,##0")

    def test_unit_1000_unrounded_and_format(self):
        rng = random.Random(33)
        rows = [[datetime.datetime(2026, rng.choice([1, 2, 3]), 1), rng.choice(ACCTS[:4]), rng.choice([1234567, 89, -4501, 999, 250001, 1])]
                for _ in range(120)]
        raw = write_xlsx(os.path.join(self.dir, "천원.xlsx"), ["일자", "계정명", "금액"], rows)
        rules = base_rules(**{"단위": "1000"})
        out, info = self.build(raw, rules, name="천원_보고서.xlsx")
        agg = Agg([((a,), month_key(d), amt) for d, a, amt in rows], unit=1000)
        self.assertTrue(any(abs(v - round(v)) > 1e-9 for v in agg.values.values()), "시험 데이터: 소수 있는 칸 필요")
        self.assert_info(info, agg)          # 반올림하지 않은 값
        ws, _, _, _, nd = self.assert_sheet(out, rules, agg, info, fmt="#,##0")
        t1 = " ".join(str(c.value) for c in ws[1] if c.value is not None)
        self.assertTrue(any(s in t1 for s in ("천원", "1,000", "1000")), f"1행 제목에 단위: {t1!r}")
        self.assert_v23(raw, out, nd)

    def test_unit_1_with_decimals_format(self):
        rows = [[datetime.datetime(2026, 1, 2), "외화예금", 18.85], [datetime.datetime(2026, 1, 3), "외화예금", 0.1],
                [datetime.datetime(2026, 2, 2), "외화예금", 0.2], [datetime.datetime(2026, 2, 3), "현금", 100]]
        raw = write_xlsx(os.path.join(self.dir, "소수.xlsx"), ["일자", "계정명", "금액"], rows)
        rules = base_rules()
        out, info = self.build(raw, rules, name="소수_보고서.xlsx")
        agg = Agg([((a,), month_key(d), amt) for d, a, amt in rows])
        self.assert_info(info, agg)
        self.assert_sheet(out, rules, agg, info, fmt=("#,##0.00",))

    def test_invalid_unit_clear_error(self):
        rows = basic_rows(seed=34, n=20)
        raw = write_xlsx(os.path.join(self.dir, "단위오류.xlsx"), ["일자", "계정명", "거래처", "금액"], rows)
        for bad in ("0", "abc", "-1000"):
            with self.subTest(단위=bad):
                self.expect_error(raw, base_rules(**{"단위": bad}), "단위")


# ================================================================== 5. 필터·제외 글자
def filter_rows(seed=41):
    rng = random.Random(seed)
    rows = []
    for _ in range(300):
        rows.append([rng.choice(["매출", "비용", " 매출 "]), rng.choice(["본사", "지점", "본사 ", " 해외"]),
                     rng.choice(["상품매출", "제품매출", "용역매출", "운반비", "설계비"]),
                     datetime.datetime(2026, rng.choice([1, 2]), 3), rng.choice([1100, 1200, 2100]),
                     rng.choice([rng.randint(-5000, 20000), "1,000", "(2,500)", 0])])
    return rows


class T05_FilterExclude(JobCase):
    HEAD = ["구분", "법인", "계정명", "일자", "코드", "금액"]

    def _run_filter(self, flt, pred, tag):
        rows = filter_rows()
        raw = write_xlsx(os.path.join(self.dir, f"필터{tag}.xlsx"), self.HEAD, rows)
        rules = base_rules(필터=flt)
        out, info = self.build(raw, rules, name=f"필터{tag}_보고서.xlsx")
        kept = [r for r in rows if pred(r[0].strip(), r[1].strip(), r[4], parse_amount(r[5]))]
        self.assertTrue(0 < len(kept) < len(rows), "시험 데이터: 필터가 일부만 남겨야 함")
        agg = Agg([((r[2],), month_key(r[3]), parse_amount(r[5])) for r in kept])
        self.assert_info(info, agg)
        self.assertEqual(info["excluded_rows"], len(rows) - len(kept), "excluded_rows = 필터로 뺀 행 수")
        self.assertEqual(list(info["bad_rows"]), [])
        self.assert_sheet(out, rules, agg, info)
        self.assert_evidence(out, info, agg, len(rows))

    def test_filter_eq_ne_gt_combined(self):
        self._run_filter("구분=매출; 법인<>본사; 금액>0", lambda g, b, c, a: g == "매출" and b != "본사" and a > 0, "A")

    def test_filter_ge_le(self):
        self._run_filter("금액>=1000; 금액<=5000; 법인=해외", lambda g, b, c, a: 1000 <= a <= 5000 and b == "해외", "B")

    def test_filter_lt(self):
        self._run_filter("금액<0", lambda g, b, c, a: a < 0, "C")

    def test_filter_numeric_code_equality(self):
        """판단: 숫자로 저장된 코드(1100)도 '코드=1100' 조건에 걸려야 함."""
        self._run_filter("코드=1100", lambda g, b, c, a: c == 1100, "D")

    def test_filter_unknown_column(self):
        rows = filter_rows()
        raw = write_xlsx(os.path.join(self.dir, "필터열없음.xlsx"), self.HEAD, rows)
        self.expect_error(raw, base_rules(필터="부서=영업"), "부서")

    def test_exclude_in_raw_subtotal_lines(self):
        rows = [
            [datetime.datetime(2026, 1, 2), "매출", 1000], [datetime.datetime(2026, 1, 3), "매출", 2000],
            [None, "매출 합계", 3000],
            [datetime.datetime(2026, 1, 4), "설계비", 500], [datetime.datetime(2026, 2, 4), "설계비", 700],
            [None, "* 소계", 1200], [None, "*", 1200],
            [datetime.datetime(2026, 2, 9), "운반비", 300],
            [None, "** 합계", 4500],
        ]
        raw = write_xlsx(os.path.join(self.dir, "원본소계.xlsx"), ["일자", "계정명", "금액"], rows)
        rules = base_rules(**{"제외 글자": "합계, 소계, *"})
        out, info = self.build(raw, rules, name="원본소계_보고서.xlsx")
        kept = [r for r in rows if r[0] is not None]
        agg = Agg([((r[1],), month_key(r[0]), r[2]) for r in kept])
        self.assert_info(info, agg)
        self.assertEqual(info["grand_total"], 4500, "원본 안 소계 줄을 두 번 더하면 안 됨")
        self.assertEqual(info["excluded_rows"], 4)
        self.assertIn(("설계비",), [tuple(k) for k in info["row_keys"]], "'설계비'는 '계'를 포함해도 제외 글자(합계·소계)에 안 걸림")
        _, _, _, _, nd = self.assert_sheet(out, rules, agg, info)
        self.assert_evidence(out, info, agg, len(rows))
        self.assert_v23(raw, out, nd)

    def test_blank_row_key_not_silent(self):
        """판단: 행 기준 값이 빈 행(SAP 총계 줄 등)은 조용히 합치거나 버리면 안 됨 → 경고(또는 bad/제외)로 알림, 합계는 일관."""
        rows = [[datetime.datetime(2026, 1, 2), "매출", 1000], [datetime.datetime(2026, 1, 3), "현금", 2000],
                [datetime.datetime(2026, 1, 31), None, 3000]]
        raw = write_xlsx(os.path.join(self.dir, "빈행기준.xlsx"), ["일자", "계정명", "금액"], rows)
        out, info = self.build(raw, base_rules(), name="빈행기준_보고서.xlsx")
        self.assertTrue(close(info["raw_total"], info["grand_total"]))
        txt = " ".join(str(w) for w in info["warnings"]) + " " + " ".join(str(b) for b in info["bad_rows"])
        self.assertTrue(re.search(r"(?<!\d)4(?!\d)", txt) or "빈" in txt,
                        f"계정명 빈 행(엑셀 4행)에 대한 알림 없음: warnings={info['warnings']} bad={info['bad_rows']}")


# ================================================================== 6. 지저분한 SAP 원본
class T06_MessyRaw(JobCase):
    def test_messy_sap_like(self):
        pre = [("[전표 조회] 계정별 원장",), ("회사코드", "1000"), ("조회기간: 2026.01.01 ~ 2026.03.31",), ()]
        data = [
            ("2026.01.05", "현금 ", "입금", "1,000"),
            ("2026.01.06", " 현금", "출금", "(1,500)"),
            ("2026.01.07", "현금", "출금", "2,000-"),
            ("2026.01.08", "외상매출금", "조정", "△300"),
            ("2026.01.09", "외상매출금 ", "판매", "₩4,000"),
            ("2026.01.10", "외상매출금", "판매", "5,000원"),
            ("2026.01.11", "미지급금", "대체", "-"),
            ("2026.01.12", "미지급금", "대체", " 2,500 "),
            ("2026.01.13", "미지급금", "대체", 12345),
            ("2026.01.14", "현금", "확인", "N/A"),
            ("2026.01.15", "외상매출금", "확인", "확인필요"),
            ("2026.01.16", "현금", "입금", 678.5),
            (None, "* 합계", None, "99,999"),
        ]
        first_data_row = len(pre) + 2                                    # 머리글 = 5행, 자료 = 6행부터
        raw = write_xlsx(os.path.join(self.dir, "SAP원본.xlsx"), ["전기일", "계정명", "적요", "금액"], data, pre_rows=pre)
        rules = base_rules(**{"열 기준": "", "제외 글자": "합계"})
        out, info = self.build(raw, rules, name="SAP원본_보고서.xlsx")
        items, bad, excl = [], [], 0
        for i, (d, a, _, amt) in enumerate(data):
            if "합계" in a:
                excl += 1
                continue
            try:
                items.append(((a.strip(),), "금액", parse_amount(amt)))
            except BadAmount:
                bad.append((first_data_row + i, amt))
        agg = Agg(items)
        self.assertEqual(agg.values[(("현금",), "금액")], 1000 - 1500 - 2000 + 678.5)
        self.assert_info(info, agg)
        self.assertEqual(info["excluded_rows"], excl)
        got_bad = sorted((int(r), str(v).strip()) for r, v in info["bad_rows"])
        self.assertEqual(got_bad, sorted((r, str(v).strip()) for r, v in bad), "bad_rows (엑셀 행 번호, 값)")
        self.assertEqual({tuple(k) for k in info["row_keys"]}, {("현금",), ("외상매출금",), ("미지급금",)}, "앞뒤 공백 다른 계정명은 하나로")
        self.assert_sheet(out, rules, agg, info)
        rt = sheet_text(out, "규칙")
        for r, v in bad:
            self.assertIn(str(r), rt, f"'규칙' 시트에 오류 행 {r}")
            self.assertIn(str(v).strip(), rt, f"'규칙' 시트에 오류 값 {v}")
        self.assert_evidence(out, info, agg, len(data))
        self.assert_v23(raw, out, None, require_detail=False)


# ================================================================== 7. 합계 수식
class T07_Formulas(JobCase):
    def test_formula_multilevel_ranges_exact(self):
        rows = multilevel_rows(seed=9)
        raw = write_xlsx(os.path.join(self.dir, "수식원본.xlsx"), ["법인", "계정명", "일자", "금액"], rows)
        rules = base_rules(**{"행 기준": "법인, 계정명", "소계": "예", "합계 행": "예", "합계 열": "예", "합계 수식": "예"})
        out, info = self.build(raw, rules, name="수식_보고서.xlsx")
        agg = ml_agg(rows)
        self.assert_info(info, agg, ncheck_order=agg.order("원본 순"))
        self.assert_sheet(out, rules, agg, info, order=agg.order("원본 순"), formula=True)

    def test_formula_single_level_zero_mode(self):
        rows = basic_rows(seed=10, n=120)
        raw = write_xlsx(os.path.join(self.dir, "수식단일.xlsx"), ["일자", "계정명", "거래처", "금액"], rows)
        rules = base_rules(**{"합계 수식": "예", "빈 칸": "0"})
        out, info = self.build(raw, rules, name="수식단일_보고서.xlsx")
        agg = Agg([((a,), month_key(d), amt) for d, a, _, amt in rows])
        self.assert_info(info, agg)
        # 빈 칸=0 으로 채운 0 칸도 상세 칸 → 범위 안 포함 허용 (coverage는 숫자 상세 칸 기준)
        self.assert_sheet(out, rules, agg, info, formula=True)


# ================================================================== 8. CLI (인자 없음)
class T08_CLI(JobCase):
    def test_noarg_cli_in_folder(self):
        folder = os.path.join(self.dir, "cli")
        os.makedirs(folder)
        script = self.job_copy(folder)
        rows = [[datetime.datetime(2026, (i % 3) + 1, 5), ["현금", "매출", "급여"][i % 3], 1000 * (i + 1)] for i in range(30)]
        raw = write_xlsx(os.path.join(folder, "전표원본.xlsx"), ["일자", "계정명", "금액"], rows)
        h0 = md5(raw)
        code, last, allout = self.run_cli(script, cwd=folder)
        self.assertEqual(code, 0, allout[-3000:])
        self.assertTrue(last.startswith("최종: 보고서 생성"), f"마지막 줄: {last!r}")
        self.assertIn("일치", last)
        self.assertNotIn("불일치", last)
        self.assertEqual(md5(raw), h0, "원본 변경됨")
        xs = sorted(f for f in os.listdir(folder) if f.lower().endswith(".xlsx"))
        reports = [f for f in xs if re.fullmatch(r"전표원본_보고서_\d{8}_\d{6}\.xlsx", f)]
        others = [f for f in xs if f not in reports and f not in ("전표원본.xlsx", "규칙표_job.xlsx")]
        self.assertEqual(len(reports), 1, f"보고서 파일 정확히 1개: {xs}")
        self.assertEqual(others, [], f"예상 밖 xlsx: {others}")
        rep = os.path.join(folder, reports[0])
        m = re.search(r"저장:\s*(.+)$", last)
        self.assertTrue(m and os.path.basename(m.group(1).strip()) == reports[0], f"마지막 줄 저장 경로: {last!r}")
        unit = float(str(JOB.default_rules(raw).get("단위") or 1).replace(",", ""))
        nums = re.search(r"합계\s*([-\d,\.]+)\s*\(\s*원본 합계\s*([-\d,\.]+)", last)
        self.assertTrue(nums, f"합계 숫자 형식: {last!r}")
        x, y = (float(g.replace(",", "")) for g in nums.groups())
        expected = sum(r[2] for r in rows) / unit
        self.assertTrue(close(x, expected, abs_=0.5) and close(y, expected, abs_=0.5), f"CLI 합계 {x}/{y} != {expected}")
        # 규칙 시트에 가정 목록
        dr = JOB.default_rules(raw)
        assumptions = dr.get("_가정") or []
        self.assertTrue(len(assumptions) > 0, "default_rules는 '_가정'에 추정 근거를 적어야 함")
        rt = sheet_text(rep, "규칙")
        for a in assumptions:
            self.assertIn(str(a).strip()[:40], rt, f"'규칙' 시트에 가정 '{a}' 없음")
        # 두 번째 실행: 결과 파일을 원본으로 착각하지 않음
        time.sleep(1.2)
        code2, last2, all2 = self.run_cli(script, cwd=folder)
        self.assertEqual(code2, 0, all2[-2000:])
        self.assertTrue(last2.startswith("최종: 보고서 생성") and "불일치" not in last2, last2)
        self.assertIn("전표원본_보고서_", last2)
        self.assertNotIn("보고서_보고서", last2, "이전 보고서를 원본으로 사용함")
        self.assertEqual(md5(raw), h0)

    def test_cli_explicit_raw_and_rules(self):
        folder = os.path.join(self.dir, "cli2")
        os.makedirs(folder)
        script = self.job_copy(folder)
        rows = basic_rows(seed=81, n=90)
        raw = write_xlsx(os.path.join(folder, "명시원본.xlsx"), ["일자", "계정명", "거래처", "금액"], rows)
        rp = write_rules_xlsx(os.path.join(folder, "내규칙.xlsx"), base_rules(**{"열 기준": "", "단위": "1"}))
        code, last, allout = self.run_cli(script, [raw, "--규칙", rp], cwd=self.dir)
        self.assertEqual(code, 0, allout[-3000:])
        self.assertTrue(last.startswith("최종: 보고서 생성") and "불일치" not in last, last)
        n_acc = len({r[1] for r in rows})
        self.assertRegex(last, rf"행\s*{n_acc}\b", "행 수 N = 상세 행 수")
        x = float(re.search(r"합계\s*([-\d,\.]+)", last).group(1).replace(",", ""))
        self.assertTrue(close(x, sum(r[3] for r in rows), abs_=0.5))

    def test_cli_empty_raw_fails_cleanly(self):
        """수용 기준: 헤더만 있는 원본 → 종료코드 1 + 마지막 줄 '최종: 실행 불가 — ...' (0 짜리 보고서 생성은 불가)."""
        folder = os.path.join(self.dir, "cli3")
        os.makedirs(folder)
        script = self.job_copy(folder)
        raw = write_xlsx(os.path.join(folder, "빈원본.xlsx"), ["일자", "계정명", "금액"], [])
        code, last, allout = self.run_cli(script, [raw], cwd=folder)
        self.assertEqual(code, 1, allout[-2000:])
        self.assertTrue(last.startswith("최종: 실행 불가"), f"마지막 줄: {last!r}")
        self.assertNotIn("Traceback", allout, "실패는 Traceback이 아니라 이유 문장으로")
        self.assertFalse([f for f in os.listdir(folder) if "_보고서_" in f], "실패 시 보고서 파일을 만들지 않음")

    def test_cli_unknown_column_fails_cleanly(self):
        folder = os.path.join(self.dir, "cli4")
        os.makedirs(folder)
        script = self.job_copy(folder)
        raw = write_xlsx(os.path.join(folder, "원본x.xlsx"), ["일자", "계정명", "금액"], [[datetime.datetime(2026, 1, 1), "현금", 1]])
        rp = write_rules_xlsx(os.path.join(folder, "규칙x.xlsx"), base_rules(**{"행 기준": "계정과목"}))
        code, last, allout = self.run_cli(script, [raw, "--규칙", rp], cwd=folder)
        self.assertEqual(code, 1, allout[-2000:])
        self.assertTrue(last.startswith("최종: 실행 불가") and "계정과목" in last, f"마지막 줄: {last!r}")
        self.assertNotIn("Traceback", allout)


# ================================================================== 9. 견고성
class T09_Robust(JobCase):
    def test_unknown_columns_library(self):
        rows = basic_rows(seed=91, n=30)
        raw = write_xlsx(os.path.join(self.dir, "열오류.xlsx"), ["일자", "계정명", "거래처", "금액"], rows)
        cases = [({"행 기준": "계정과목"}, "계정과목"), ({"행 기준": "법인, 계정명"}, "법인"),
                 ({"열 기준": "거래일:월"}, "거래일"), ({"열 기준": "부서"}, "부서"),
                 ({"금액": "금액X"}, "금액X"), ({"금액": "차변-대변"}, "차변"), ({"원본 시트": "없는시트"}, "없는시트")]
        for over, needle in cases:
            with self.subTest(**{k.replace(" ", "_"): v for k, v in over.items()}):
                self.expect_error(raw, base_rules(**{k.replace(" ", "_"): v for k, v in over.items()}), needle)

    def test_50k_rows_under_60s(self):
        path = os.path.join(self.dir, "대용량.xlsx")
        rng = random.Random(99)
        accts = [f"계정{i:02d}" for i in range(40)]
        corps = ["본사", "지점", "해외"]
        wb = Workbook(write_only=True)
        ws = wb.create_sheet("전표")
        ws.append(["법인", "계정명", "일자", "금액"])
        items = []
        for _ in range(50000):
            b, a = rng.choice(corps), rng.choice(accts)
            d = datetime.datetime(2026, rng.randint(1, 12), rng.randint(1, 28))
            amt = rng.randint(-100000, 5000000)
            ws.append([b, a, d, amt])
            items.append(((b, a), month_key(d), amt))
        wb.save(path)
        rules = base_rules(**{"행 기준": "법인, 계정명", "소계": "예"})
        t0 = time.monotonic()
        out, info = self.build(path, rules, name="대용량_보고서.xlsx")
        el = time.monotonic() - t0
        self.assertLess(el, 60, f"50,000행 {el:.1f}초")
        agg = Agg(items)
        self.assert_info(info, agg, ncheck_order=agg.order("원본 순"))


# ================================================================== 10. 실제 원본 (사본)
class T10_RealRaws(JobCase):
    def _real(self, src, allow_unusable=False):
        if not os.path.exists(src):
            self.fail(f"실제 원본 없음: {src}")
        dst = os.path.join(self.dir, os.path.basename(src))
        shutil.copy2(src, dst)
        h0 = md5(src)
        try:
            with contextlib.redirect_stdout(io.StringIO()):
                rules = JOB.default_rules(dst)
        except Exception as e:
            if allow_unusable:
                return self._expect_cli_unusable(dst, str(e))
            self.fail(f"default_rules 예외:\n{traceback.format_exc()}")
        self.assertIsInstance(rules, dict)
        self.assertIn("_가정", rules)
        rules = dict(rules)
        rules["합계 수식"] = "아니오"                         # v23 대조 조건
        try:
            with contextlib.redirect_stdout(io.StringIO()):
                out, info = JOB.build_report(dst, rules, out_path=os.path.join(self.dir, "real_보고서.xlsx"))
        except Exception as e:
            if allow_unusable:
                return self._expect_cli_unusable(dst, str(e))
            self.fail(f"build_report 예외:\n{traceback.format_exc()}")
        self.assertEqual(md5(src), h0)
        self.assertTrue(close(info["raw_total"], info["grand_total"]), f"raw_total {info['raw_total']} != grand_total {info['grand_total']}")
        self.assertTrue(close(sum(info["values"].values()), info["grand_total"]), "상세 칸 합 != grand_total")
        # 독립 대조: 금액 식이 단일 열이면 원본 열을 직접 더함
        amt_col = str(info["rules"].get("금액") or "").strip()
        unit = float(str(info["rules"].get("단위") or 1).replace(",", ""))
        flt = str(info["rules"].get("필터") or "").strip()
        excl = [x.strip() for x in str(info["rules"].get("제외 글자") or "").split(",") if x.strip()]
        if str(info["rules"].get("집계") or "합계").strip() == "합계" and not flt and not excl and re.fullmatch(r"[^+\-]+", amt_col):
            wb = load_workbook(dst, data_only=True)
            for ws in wb.worksheets:
                grid = list(ws.iter_rows(values_only=True))
                h = next((i for i, r in enumerate(grid[:15]) if any(_txt(v) == amt_col for v in r)), None)
                if h is None:
                    continue
                j = [_txt(v) for v in grid[h]].index(amt_col)
                tot = 0.0
                for r in grid[h + 1:]:
                    if j < len(r) and not all(v in (None, "") for v in r):
                        try:
                            tot += parse_amount(r[j])
                        except BadAmount:
                            pass
                self.assertTrue(close(info["raw_total"] * unit, tot, abs_=0.01), f"원본 '{amt_col}' 직접 합 {tot} != raw_total×단위 {info['raw_total'] * unit}")
                break
        self.assert_v23(dst, out, None, require_detail=False)
        return info

    def _expect_cli_unusable(self, dst, why):
        folder = os.path.dirname(dst)
        script = self.job_copy(folder)
        code, last, allout = self.run_cli(script, [dst], cwd=folder)
        self.assertEqual(code, 1, f"라이브러리는 실패({why})했는데 CLI는 종료코드 {code}: {allout[-1500:]}")
        self.assertTrue(last.startswith("최종: 실행 불가"), f"마지막 줄: {last!r}")
        self.assertNotIn("Traceback", allout)

    def test_real_13(self):
        info = self._real(REAL_RAWS[0])
        unit = float(str(info["rules"].get("단위") or 1))
        self.assertTrue(close(info["raw_total"] * unit, 1505791000, abs_=0.5), f"13 원본 합 1,505,791,000 != {info['raw_total'] * unit}")

    def test_real_14(self):
        info = self._real(REAL_RAWS[1])
        unit = float(str(info["rules"].get("단위") or 1))
        self.assertTrue(close(info["raw_total"] * unit, 97585000, abs_=0.5), f"14 원본 합 97,585,000 != {info['raw_total'] * unit}")

    def test_real_trading_journal(self):
        self._real(REAL_RAWS[2], allow_unusable=True)


if __name__ == "__main__":
    unittest.main(verbosity=2)
