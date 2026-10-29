# QA 검토(독립 리뷰어) 지적을 자동 시험으로 고정 — 실행: JOB_DIR=<폴더> python3 test_qa_job_v1.py
import os, sys, re, io, shutil, tempfile, datetime, contextlib, importlib, subprocess, unittest, zipfile
from openpyxl import Workbook, load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
JOB_DIR = os.environ.get("JOB_DIR", HERE)
SCR = os.path.dirname(HERE)
sys.path.insert(0, JOB_DIR)
J = importlib.import_module(os.environ.get("JOB", "report_job_v1"))
sys.path.insert(0, os.path.join(SCR, "v23"))
import excel_automation_v23 as V23
D = datetime.date
WORK = os.path.join(HERE, "work_qa")


def wx(path, header, rows, pre=(), sheet="Sheet", extra=None, merges=()):
    wb = Workbook(); ws = wb.active; ws.title = sheet
    for r in pre:
        ws.append(list(r))
    if header is not None:
        ws.append(list(header))
    for r in rows:
        ws.append(list(r))
    for m in merges:
        ws.merge_cells(m)
    for nm, (h, rr) in (extra or {}).items():
        w = wb.create_sheet(nm); w.append(list(h)); [w.append(list(x)) for x in rr]
    wb.save(path)
    return path


def rules(**kw):
    r = {"행 기준": "계정명", "열 기준": "", "금액": "금액", "집계": "합계", "단위": "1", "소계": "예", "합계 행": "예",
         "합계 열": "예", "정렬": "원본 순", "빈 칸": "빈칸", "합계 수식": "아니오", "필터": "", "제외 글자": ""}
    r.update({k.replace("_", " "): v for k, v in kw.items()})
    return r


def sheet_rows(out, start=1):
    ws = load_workbook(out)["보고서"]
    return [r for r in ws.iter_rows(min_row=start, values_only=True)]


class QA(unittest.TestCase):
    def setUp(self):
        os.makedirs(WORK, exist_ok=True)
        self.d = tempfile.mkdtemp(prefix=self._testMethodName[5:25] + "_", dir=WORK)

    def p(self, n):
        return os.path.join(self.d, n)

    def build(self, raw, r, name="o.xlsx"):
        with contextlib.redirect_stdout(io.StringIO()):
            return J.build_report(raw, r, self.p(name))

    def err(self, raw, r, *needles):
        with contextlib.redirect_stdout(io.StringIO()):
            with self.assertRaises(Exception) as cm:
                J.build_report(raw, r, self.p("x.xlsx"))
        for n in needles:
            self.assertIn(n, str(cm.exception))

    # ---------- 심각도 상
    def test_01_count_not_divided_by_unit(self):
        raw = wx(self.p("r.xlsx"), ["일자", "계정명", "금액"], [(D(2026, 1, i + 1), "매출", 1000 * (i + 1)) for i in range(3)])
        _, info = self.build(raw, rules(집계="건수", 단위="1000"))
        self.assertEqual(list(info["values"].values()), [3])

    def test_02_default_rules_total_line_not_doubled(self):
        rows = [(D(2026, 1, 5), "매출", 100), (D(2026, 1, 6), "원가", -40), (D(2026, 2, 1), "매출", 50), (D(2026, 2, 28), "합계", 110)]
        raw = wx(self.p("r.xlsx"), ["일자", "계정명", "금액"], rows)
        _, info = self.build(raw, None)
        self.assertEqual(info["grand_total"], 110)

    def test_03_merged_amount_counted_once_and_warned(self):
        raw = wx(self.p("r.xlsx"), ["일자", "계정명", "금액"], [(D(2026, 1, 5), "매출", 1000), (D(2026, 1, 6), "매출", None), (D(2026, 1, 7), "원가", -300)], merges=["C2:C3"])
        _, info = self.build(raw, rules())
        self.assertEqual(info["grand_total"], 700)
        self.assertTrue(any("병합" in w for w in info["warnings"]), info["warnings"])

    def test_04_blank_first_header_cell(self):
        rows = [("@5B@", "2026-01-05", "1100", "매출", "본사", 1000), ("@5B@", "2026-01-06", "1200", "원가", "본사", -300), ("@5C@", "2026-02-01", "1100", "매출", "지사", 500)]
        raw = wx(self.p("r.xlsx"), [None, "전기일", "계정코드", "계정명", "법인", "금액"], rows)
        _, info = self.build(raw, None)
        self.assertEqual(info["grand_total"], 1200)
        _, info = self.build(raw, rules(), "o2.xlsx")
        self.assertEqual(info["grand_total"], 1200)

    def test_05_text_amount_sheet_beats_summary_sheet(self):
        rows = [("2026-01-0%d" % (i % 9 + 1), ["매출", "원가", "판관비"][i % 3], f"{(i + 1) * 1000:,}") for i in range(30)]
        raw = wx(self.p("r.xlsx"), ["일자", "계정명", "금액"], rows, sheet="원본",
                 extra={"요약": (["계정명", "1월", "2월", "3월"], [("매출", 1, 2, 3), ("원가", 4, 5, 6), ("판관비", 7, 8, 9)])})
        _, info = self.build(raw, None)
        self.assertEqual(info["grand_total"], sum((i + 1) * 1000 for i in range(30)))

    def test_06_condition_line_above_header(self):
        rows = [(D(2026, 1, 5), "매출", 100), (D(2026, 1, 6), "원가", -40), (D(2026, 2, 1), "매출", 50)]
        raw = wx(self.p("r.xlsx"), ["일자", "계정명", "금액"], rows, pre=[("회사코드", 1000, "기간", "2026.01", "~", "2026.12", "통화", "KRW")])
        _, info = self.build(raw, rules())
        self.assertEqual(info["grand_total"], 110)
        _, info = self.build(raw, None, "o2.xlsx")
        self.assertEqual(info["grand_total"], 110)

    def test_07_equal_filter_on_mixed_date_formats(self):
        raw = wx(self.p("r.xlsx"), ["일자", "계정명", "금액"], [(D(2026, 1, 5), "매출", 100), ("2026.01.05", "매출", 50), (46027, "매출", 25), ("2026-01-05", "매출", 7)])
        _, info = self.build(raw, rules(필터="일자=2026-01-05"))
        self.assertEqual(info["grand_total"], 182)

    def test_08_formula_amount_without_cached_value_is_bad(self):
        wb = Workbook(); ws = wb.active
        ws.append(["계정명", "수량", "단가", "금액"])
        ws.append(["매출", 2, 100, "=B2*C2"]); ws.append(["원가", 1, 50, "=B3*C3"]); ws.append(["매출", 1, 10, 10])
        raw = self.p("r.xlsx"); wb.save(raw)
        _, info = self.build(raw, rules())
        self.assertEqual(sorted(r for r, _ in info["bad_rows"]), [2, 3])

    # ---------- 심각도 중
    def test_09_rules_template_second_run_no_false_new_items(self):
        raw = wx(self.p("r.xlsx"), ["일자", "계정명", "금액"], [(D(2026, 1, 5), "매출", 100), (D(2026, 1, 6), "원가", -40)])
        shutil.copy(os.path.join(JOB_DIR, "report_job_v1.py"), self.d)
        for _ in range(2):
            subprocess.run([sys.executable, "report_job_v1.py"], cwd=self.d, capture_output=True, text=True, timeout=120)
        rules_loaded = J.load_rules(self.p("규칙표_job.xlsx"))
        _, info = self.build(raw, rules_loaded)
        self.assertEqual(info["new_items"], [])

    def test_10_pick_raw_newest_including_xlsm(self):
        import time
        wx(self.p("작년자료.xlsx"), ["계정명", "금액"], [("x", 999)])
        time.sleep(1.1)
        wx(self.p("월간_보고서_원본.xlsx"), ["계정명", "금액"], [("y", 5)])
        time.sleep(1.1)
        wb = Workbook(); ws = wb.active; ws.append(["계정명", "금액"]); ws.append(["z", 7]); wb.save(self.p("당월.xlsm"))
        self.assertEqual(os.path.basename(J._pick_raw(self.d)), "당월.xlsm")

    def test_11_default_debit_credit(self):
        raw = wx(self.p("r.xlsx"), ["일자", "계정명", "차변", "대변"], [(D(2026, 1, 5), "매출", 0, 100), (D(2026, 1, 6), "원가", 40, 0), (D(2026, 2, 1), "매출", 0, 50)])
        _, info = self.build(raw, None)
        self.assertEqual(str(info["rules"]["금액"]).replace(" ", ""), "차변-대변")
        self.assertEqual(info["grand_total"], -110)

    def test_12_total_label_clash(self):
        raw = wx(self.p("r.xlsx"), ["구분", "계정명", "금액"], [("합계", "매출", 10), ("일반", "매출", 5), ("일반", "원가", 3)])
        out, info = self.build(raw, rules(**{"열 기준": "구분"}))
        ws = load_workbook(out)["보고서"]
        hdr = next(r for r in ws.iter_rows(values_only=True) if r and r[0] == "계정명")
        heads = [h for h in hdr if h is not None]
        self.assertEqual(len(heads), len(set(heads)), heads)
        with contextlib.redirect_stdout(io.StringIO()):
            v = V23.verify(raw, out)
        self.assertEqual(v["mismatch"], [])

    def test_13_label_starting_with_equals_is_text(self):
        raw = wx(self.p("r.xlsx"), ["계정명", "금액"], [("매출", 1), ("원가", 2)])
        wb = load_workbook(raw); wb.active["A3"].value = "=소계="; wb.active["A3"].data_type = "s"; wb.save(raw)
        out, _ = self.build(raw, rules())
        ws = load_workbook(out)["보고서"]
        cells = [c for row in ws.iter_rows() for c in row if c.value == "=소계="]
        self.assertTrue(cells)
        self.assertTrue(all(c.data_type == "s" for c in cells))

    def test_14_formula_total_under_255_args(self):
        rows = [(f"G{i:03d}", a, 1) for i in range(300) for a in ("a", "b")]
        raw = wx(self.p("r.xlsx"), ["법인", "계정명", "금액"], rows)
        out, _ = self.build(raw, rules(**{"행 기준": "법인, 계정명", "합계 수식": "예"}))
        ws = load_workbook(out)["보고서"]
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    for inner in re.findall(r"SUM\(([^()]*)\)", c.value):
                        self.assertLessEqual(inner.count(",") + 1, 255, c.coordinate)

    def test_15_amount_desc_ties_keep_groups_together(self):
        rows = [("A", "x", 100), ("B", "y", 60), ("A", "z", 0), ("B", "w", 40), ("A", "q", 0)]
        raw = wx(self.p("r.xlsx"), ["법인", "계정명", "금액"], rows)
        for sub in ("예", "아니오"):
            out, info = self.build(raw, rules(**{"행 기준": "법인, 계정명", "정렬": "금액 큰 순", "소계": sub}), f"o{sub}.xlsx")
            groups = [k[0] for k in info["row_keys"]]
            self.assertEqual(groups, sorted(groups, key=lambda g: groups.index(g)), groups)
            sheet = [(r[0], r[1]) for r in sheet_rows(out, 4) if r[1] is not None and r[0] not in (None, "합계")]
            self.assertEqual(sheet, [tuple(k) for k in info["row_keys"]])

    def test_16_row_order_multilevel_variants(self):
        rows = [("본사", "매출", 100), ("지사", "매출", 30), ("본사", "원가", -60), ("지사", "원가", -10), ("본사", "신규", 5), ("해외", "매출", 1)]
        raw = wx(self.p("r.xlsx"), ["법인", "계정명", "금액"], rows)
        r = rules(**{"행 기준": "법인, 계정명"}); r["행 순서"] = ["지사", "본사"]
        _, info = self.build(raw, r)
        self.assertEqual([k[0] for k in info["row_keys"]][:2], ["지사", "지사"])
        self.assertEqual({k[0] for k in info["new_items"]}, {"해외"})
        r["행 순서"] = ["본사/원가", "본사/매출", "지사 / 매출"]
        _, info = self.build(raw, r, "o2.xlsx")
        self.assertEqual([tuple(k) for k in info["row_keys"]][:2], [("본사", "원가"), ("본사", "매출")])

    def test_17_yes_no_and_aggregate_variants(self):
        raw = wx(self.p("r.xlsx"), ["법인", "계정명", "금액"], [("A", "x", 1), ("A", "y", 2), ("B", "x", 3)])
        for v in ("네", "○", "Yes ", "TRUE"):
            out, _ = self.build(raw, rules(**{"행 기준": "법인, 계정명", "소계": v}))
            self.assertTrue(any(r[0] and str(r[0]).endswith("소계") for r in sheet_rows(out)), v)
        self.err(raw, rules(**{"행 기준": "법인, 계정명", "소계": "있음"}), "소계")
        for v in ("개수", "count", "갯수"):
            _, info = self.build(raw, rules(집계=v))
            self.assertEqual(info["grand_total"], 3, v)
        self.err(raw, rules(집계="평균"), "집계")
        self.err(raw, rules(**{"빈 칸": "있음", "열 기준": "법인"}), "빈 칸")

    def test_18_rule_key_spacing_typos(self):
        raw = wx(self.p("r.xlsx"), ["법인", "계정명", "금액"], [("A", "x", 1000), ("B", "x", 3000)])
        r = rules(**{"열 기준": "법인"}); r.pop("합계 열"); r.pop("단위"); r["합계열"] = "아니오"; r["단위 "] = "1000"
        out, info = self.build(raw, r)
        self.assertEqual(info["grand_total"], 4)
        self.assertNotIn("합계", [c for c in sheet_rows(out)[2] if c])

    def test_19_bad_amount_expression(self):
        raw = wx(self.p("r.xlsx"), ["계정명", "차변", "대변"], [("x", 100, 30)])
        self.err(raw, rules(금액="차변--대변"), "금액")

    def test_20_leading_zero_code_filter(self):
        raw = wx(self.p("r.xlsx"), ["계정코드", "계정명", "금액"], [("001100", "매출", 100), ("1100", "매출", 10), ("1100.0", "원가", -5)])
        _, info = self.build(raw, rules(필터="계정코드=001100"))
        self.assertEqual(info["grand_total"], 100)

    def test_21_hidden_rows_and_duplicate_columns_warned(self):
        raw = wx(self.p("r.xlsx"), ["계정명", "금액", "금액"], [("x", 1, 9), ("y", 2, 9), ("z", 3, 9)])
        wb = load_workbook(raw); wb.active.row_dimensions[3].hidden = True; wb.save(raw)
        _, info = self.build(raw, rules())
        txt = " ".join(info["warnings"])
        self.assertIn("숨", txt)
        self.assertIn("같은 이름", txt)

    def test_22_report_total_checked_against_written_sheet(self):
        raw = wx(self.p("r.xlsx"), ["계정명", "금액"], [("x", 1), ("y", 2)])
        _, info = self.build(raw, rules())
        self.assertIn("sheet_total", info)
        self.assertEqual(info["sheet_total"], 3)


if __name__ == "__main__":
    unittest.main(verbosity=1)
