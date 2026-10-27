"""Structural discovery contracts; fixtures have no fixed business domain."""
import datetime as dt
import json
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.utils.datetime import CALENDAR_MAC_1904

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from report_analyzer import inspect_workbook, propose_plan


class AnalyzerTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.directory = Path(self.temp.name)

    def workbook(self, name, rows, sheet="表 α", epoch=None):
        path = self.directory / name
        book = Workbook()
        if epoch:
            book.epoch = epoch
        book.active.title = sheet
        for row in rows:
            book.active.append(row)
        book.save(path)
        book.close()
        return path

    def test_shifted_headers_and_sparse_title(self):
        path = self.workbook("shifted.xlsx", [["A title"], [], [],
                             [None, "Category", "Measure"], [None, "A", 10], [None, "B", 20]])
        result = inspect_workbook(path)
        table = result["sheets"][0]["tables"][0]
        self.assertEqual((table["header_row"], table["first_data_row"], table["last_data_row"]), (4, 5, 6))
        self.assertEqual(table["column_indices"], [2, 3])
        self.assertFalse(table["ambiguous"])
        json.dumps(result)

    def test_all_rows_formula_cache_scanning_and_1904(self):
        path = self.workbook("formula.xlsx", [["Category", "Measure"], ["A", 10]], epoch=CALENDAR_MAC_1904)
        book = load_workbook(path)
        book.active["B5101"] = "=1+2"
        book.save(path)
        book.close()
        result = inspect_workbook(path)
        self.assertTrue(result["epoch"].startswith("1904"))
        self.assertEqual(result["sheets"][0]["formula_without_cache"], ["B5101"])
        self.assertEqual(len(result["sha256"]), 64)

    def test_numeric_text_does_not_consume_leading_zero_identifiers(self):
        path = self.workbook("text.xlsx", [["Identifier", "Measure"], ["001100", "1,000-"], ["001200", "₩2,000"]])
        table = inspect_workbook(path)["sheets"][0]["tables"][0]
        self.assertEqual(table["profiles"][0]["kinds"], {"text": 2})
        self.assertEqual(table["profiles"][0]["examples"], ["001100", "001200"])
        self.assertEqual(table["profiles"][1]["kinds"], {"number_text": 2})

    def test_ambiguous_text_and_duplicate_headers_require_selection(self):
        reference = self.workbook("reference.xlsx", [["Category", "Measure"], ["A", 10]])
        for name, rows in [("text.xlsx", [["One", "Two"], ["A", "B"], ["C", "D"]]),
                           ("duplicates.xlsx", [["Group", "Measure", "Measure"], ["A", 10, 20]])]:
            with self.subTest(name=name):
                source = self.workbook(name, rows)
                plan = propose_plan(reference, source)
                self.assertFalse(plan["approved"])
                self.assertTrue(plan["unresolved"])
                self.assertEqual(plan["mappings"], [])

    def test_clear_pivot_is_only_an_unapproved_proposal(self):
        source = self.workbook("raw.xlsx", [["部署", "Datum", "Measure"],
                               ["A", dt.date(2026, 1, 2), 10], ["B", dt.date(2026, 2, 2), 20]])
        reference = self.workbook("pivot.xlsx", [["部署", "1월", "2월", "합계"],
                                  ["A", 10, None, 10], ["B", None, 20, 20]])
        plan = propose_plan(reference, source)
        self.assertFalse(plan["approved"])
        self.assertFalse(plan["unresolved"])
        self.assertTrue(plan["review_required"])
        self.assertEqual(plan["mappings"][0]["pivot"], {"column": "Datum", "kind": "month"})
        self.assertEqual(plan["mappings"][0]["terms"], [{"column": "Measure", "factor": 1}])

    def test_saved_source_bound_remains_open_for_next_period_growth(self):
        source = self.workbook("raw.xlsx", [["Category", "Measure"], ["A", 10], ["B", 20]])
        reference = self.workbook("reference.xlsx", [["Category", "Measure"], ["A", 10], ["B", 20]])
        plan = propose_plan(reference, source)
        self.assertIsNone(plan["source"]["last_data_row"])
        self.assertEqual(plan["source_info"]["detected_last_data_row"], 3)
        book = load_workbook(source)
        book.active.append(["C", 30])
        book.active.append(["D", 40])
        book.save(source)
        book.close()
        changed = propose_plan(reference, source)
        self.assertIsNone(changed["source"]["last_data_row"])
        self.assertEqual(changed["source_info"]["detected_last_data_row"], 5)
        self.assertFalse(changed["unresolved"])

    def test_footer_requires_explicit_bounds_instead_of_silent_clipping(self):
        source = self.workbook("footer.xlsx", [["Category", "Measure"], ["A", 10], ["B", 20], [], ["Memo only"]])
        reference = self.workbook("reference.xlsx", [["Category", "Measure"], ["A", 10], ["B", 20]])
        plan = propose_plan(reference, source)
        self.assertFalse(plan["approved"])
        self.assertTrue(plan["unresolved"])
        self.assertIsNone(plan["source"]["last_data_row"])
        self.assertEqual(plan["source_info"]["detected_last_data_row"], 3)


if __name__ == "__main__":
    unittest.main()
