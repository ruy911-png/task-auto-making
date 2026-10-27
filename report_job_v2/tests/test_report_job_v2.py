"""Independent acceptance tests for the generic v2 engine.

All workbooks are invented fixtures in TemporaryDirectory. No private files,
developer-machine paths, Microsoft Excel or network access are required.
Expected numbers are fixed independently of the production aggregators.
"""
import copy
import datetime as dt
import hashlib
import os
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import report_job_v2 as J


class EngineAcceptance(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="generic-report-test-")
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.reference = self.root / "reference.xlsx"
        self.raw = self.root / "source.xlsx"
        self.output = self.root / "result.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = "Results"
        ws["A1"] = "Example report"
        ws["A2"] = "Summary"
        ws["B2"] = 999
        ws["B2"].number_format = '#,##0.00'
        ws["B2"].fill = PatternFill("solid", fgColor="FFFF00")
        ws["C2"] = "=B2*2"
        ws["D2"] = "Untouched"
        ws["D2"].font = Font(bold=True)
        ws["D2"].comment = Comment("Keep this annotation", "Fixture")
        ws["B3"] = 888
        ws["A10"] = "Group"
        ws["B10"] = "Amount"
        wb.save(self.reference)
        self.write_raw(["Group", "Amount"], [["A", 100], ["B", 200], ["A", 25]])

    def write_raw(self, headers, rows, *, sheet="Data", pre=()):
        wb = Workbook()
        ws = wb.active
        ws.title = sheet
        for row in pre:
            ws.append(list(row))
        ws.append(list(headers))
        for row in rows:
            ws.append(list(row))
        wb.save(self.raw)

    def plan(self, *, mappings=None, **mapping_options):
        mapping = {
            "target": {"sheet": "Results", "cell": "B2"},
            "aggregate": "sum",
            "terms": [{"column": "Amount", "factor": 1}],
            "filters": [],
            "unit": 1,
        }
        mapping.update(mapping_options)
        return {
            "version": 2,
            "approved": False,
            "source": {"sheet": "Data", "header_row": 1},
            "mappings": mappings if mappings is not None else [mapping],
            "unresolved": [],
        }

    def approve(self, plan):
        return J.approve_plan(plan, self.reference, self.raw)

    def run_plan(self, plan=None):
        plan = self.plan() if plan is None else plan
        return J.generate_report(self.reference, self.raw, self.approve(plan), self.output)

    def value(self, coordinate="B2"):
        wb = load_workbook(self.output, data_only=False)
        try:
            return wb["Results"][coordinate].value
        finally:
            wb.close()

    @staticmethod
    def digest(path):
        return hashlib.sha256(Path(path).read_bytes()).hexdigest()

    def assert_blocked(self, plan=None):
        with self.assertRaises(J.ReportError):
            self.run_plan(plan)
        self.assertFalse(self.output.exists(), "A failed generation must not leave a result workbook")

    def test_sum_and_both_inputs_are_unchanged(self):
        before = (self.digest(self.reference), self.digest(self.raw))
        result = self.run_plan()
        self.assertEqual(self.value(), 325)
        self.assertEqual(result["cells_written"], 1)
        self.assertEqual(before, (self.digest(self.reference), self.digest(self.raw)))

    def test_reference_styles_comments_and_unmapped_formulas_survive(self):
        result = self.run_plan()
        self.assertTrue(result["formula_recalculation_required"])
        self.assertEqual(result["status"], "수식 재계산 필요")
        wb = load_workbook(self.output)
        try:
            ws = wb["Results"]
            self.assertEqual(ws["C2"].value, "=B2*2")
            self.assertEqual(ws["D2"].value, "Untouched")
            self.assertTrue(ws["D2"].font.bold)
            self.assertEqual(ws["D2"].comment.text, "Keep this annotation")
            self.assertEqual(ws["B2"].number_format, '#,##0.00')
            self.assertEqual(ws["B2"].fill.fgColor.rgb[-6:], "FFFF00")
        finally:
            wb.close()

    def test_arbitrary_sheet_header_location_and_renamed_columns(self):
        self.write_raw(["Department", "Value EUR"], [["One", 7], ["Two", 11]], sheet="Export 42", pre=[["Some title"], []])
        plan = self.plan(terms=[{"column": "Value EUR", "factor": 1}])
        plan["source"] = {"sheet": "Export 42", "header_row": 3}
        self.run_plan(plan)
        self.assertEqual(self.value(), 18)

    def test_debit_minus_credit_is_explicit(self):
        self.write_raw(["Group", "Debit", "Credit"], [["A", 20, 100], ["A", 40, 5]])
        self.run_plan(self.plan(terms=[{"column": "Debit", "factor": 1}, {"column": "Credit", "factor": -1}]))
        self.assertEqual(self.value(), -45)

    def test_count_counts_rows_without_amount_inference(self):
        self.write_raw(["Group", "Memo"], [["A", "x"], ["B", "y"], ["A", "z"]])
        self.run_plan(self.plan(aggregate="count", terms=[]))
        self.assertEqual(self.value(), 3)

    def test_aggregate_mean_min_max(self):
        for operation, expected in [("avg", 325 / 3), ("min", 25), ("max", 200)]:
            with self.subTest(operation=operation):
                if self.output.exists():
                    self.output.unlink()
                self.run_plan(self.plan(aggregate=operation, round_digits=8))
                self.assertAlmostEqual(self.value(), expected)

    def test_unrounded_repeating_average_requires_explicit_precision(self):
        self.assert_blocked(self.plan(aggregate="avg"))

    def test_unit_division_keeps_fractional_values(self):
        self.run_plan(self.plan(unit=1000))
        self.assertAlmostEqual(self.value(), 0.325)

    def test_literal_text_filter_preserves_internal_spaces(self):
        self.write_raw(["Group", "Amount"], [["A B", 100], ["AB", 900]])
        self.run_plan(self.plan(filters=[{"column": "Group", "op": "=", "value": "A B", "type": "text"}]))
        self.assertEqual(self.value(), 100)

    def test_literal_code_filter_preserves_leading_zero(self):
        self.write_raw(["Group", "Amount"], [["001100", 100], ["1100", 900]])
        self.run_plan(self.plan(filters=[{"column": "Group", "op": "=", "value": "001100", "type": "text"}]))
        self.assertEqual(self.value(), 100)

    def test_text_identifiers_do_not_round_16_digit_codes(self):
        self.write_raw(["Group", "Amount"], [["1234567890123021", 10], ["1234567890123022", 90]])
        self.run_plan(self.plan(filters=[{"column": "Group", "op": "=", "value": "1234567890123021", "type": "text"}]))
        self.assertEqual(self.value(), 10)

    def test_date_filter_recognizes_same_date_in_mixed_forms(self):
        self.write_raw(["Day", "Amount"], [[dt.date(2026, 1, 5), 100], ["2026-01-05", 200], ["20260105", 300], ["2026-01-06", 900]])
        self.run_plan(self.plan(filters=[{"column": "Day", "op": "=", "value": "20260105", "type": "date"}]))
        self.assertEqual(self.value(), 600)

    def test_blank_amount_blocks_instead_of_silently_becoming_zero(self):
        self.write_raw(["Group", "Amount"], [["A", 100], ["B", None]])
        self.assert_blocked()

    def test_explicit_blank_zero_policy_is_supported(self):
        self.write_raw(["Group", "Amount"], [["A", 100], ["B", None]])
        self.run_plan(self.plan(blank_amount="zero"))
        self.assertEqual(self.value(), 100)

    def test_uncached_amount_formula_blocks(self):
        self.write_raw(["Group", "Amount"], [["A", 100], ["B", "=900+1"]])
        self.assert_blocked()

    def test_uncached_filter_formula_blocks_instead_of_excluded(self):
        self.write_raw(["Group", "Amount"], [["A", 100], ['="A"', 900]])
        self.assert_blocked(self.plan(filters=[{"column": "Group", "op": "=", "value": "A", "type": "text"}]))

    def test_excel_error_in_required_source_field_blocks(self):
        self.write_raw(["Group", "Amount"], [["A", 100], ["B", "#DIV/0!"]])
        self.assert_blocked()

    def test_unknown_source_column_blocks(self):
        self.assert_blocked(self.plan(terms=[{"column": "Unknown", "factor": 1}]))

    def test_duplicate_source_headers_block(self):
        self.write_raw(["Group", "Amount", "Amount"], [["A", 1, 2], ["B", 3, 4]])
        self.assert_blocked()

    def test_invalid_unit_blocks(self):
        for unit in [0, -1, "nan", "inf"]:
            with self.subTest(unit=unit):
                self.assert_blocked(self.plan(unit=unit))

    def test_unknown_aggregate_blocks(self):
        self.assert_blocked(self.plan(aggregate="guess"))

    def test_unapproved_plan_is_not_executed(self):
        with self.assertRaises(J.ReportError):
            J.generate_report(self.reference, self.raw, self.plan(), self.output)
        self.assertFalse(self.output.exists())

    def test_unresolved_plan_cannot_be_approved(self):
        plan = self.plan()
        plan["unresolved"] = ["Which source column feeds B2?"]
        with self.assertRaises(J.ReportError):
            J.approve_plan(plan, self.reference, self.raw)

    def test_approval_does_not_mutate_callers_plan(self):
        original = self.plan()
        snapshot = copy.deepcopy(original)
        approved = self.approve(original)
        self.assertEqual(original, snapshot)
        self.assertTrue(approved["approved"])

    def test_changing_plan_after_approval_requires_new_approval(self):
        approved = self.approve(self.plan())
        approved["mappings"][0]["unit"] = 1000
        with self.assertRaises(J.ReportError):
            J.generate_report(self.reference, self.raw, approved, self.output)
        self.assertFalse(self.output.exists())

    def test_changing_reference_after_approval_requires_new_approval(self):
        approved = self.approve(self.plan())
        wb = load_workbook(self.reference)
        wb["Results"]["A1"] = "Changed report structure"
        wb.save(self.reference)
        with self.assertRaises(J.ReportError):
            J.generate_report(self.reference, self.raw, approved, self.output)
        self.assertFalse(self.output.exists())

    def test_approved_plan_is_reusable_with_new_period_source(self):
        approved = self.approve(self.plan())
        self.write_raw(["Group", "Amount"], [["C", 17], ["D", 31]])
        J.generate_report(self.reference, self.raw, approved, self.output)
        self.assertEqual(self.value(), 48)

    def test_duplicate_target_cells_block(self):
        plan = self.plan()
        plan["mappings"].append(copy.deepcopy(plan["mappings"][0]))
        self.assert_blocked(plan)

    def test_existing_formula_target_is_not_silently_overwritten(self):
        self.assert_blocked(self.plan(target={"sheet": "Results", "cell": "C2"}))

    def test_existing_text_target_requires_explicit_permission(self):
        self.assert_blocked(self.plan(target={"sheet": "Results", "cell": "D2"}))

    def test_output_cannot_overwrite_reference(self):
        before = self.digest(self.reference)
        with self.assertRaises(J.ReportError):
            J.generate_report(self.reference, self.raw, self.approve(self.plan()), self.reference)
        self.assertEqual(self.digest(self.reference), before)

    def test_output_cannot_overwrite_raw(self):
        before = self.digest(self.raw)
        with self.assertRaises(J.ReportError):
            J.generate_report(self.reference, self.raw, self.approve(self.plan()), self.raw)
        self.assertEqual(self.digest(self.raw), before)

    def test_output_alias_cannot_overwrite_reference(self):
        alias = self.root / "reference-alias.xlsx"
        try:
            alias.symlink_to(self.reference)
        except (OSError, NotImplementedError):
            self.skipTest("This OS does not permit test symlinks")
        before = self.digest(self.reference)
        with self.assertRaises(J.ReportError):
            J.generate_report(self.reference, self.raw, self.approve(self.plan()), alias)
        self.assertEqual(self.digest(self.reference), before)

    def test_existing_output_is_not_overwritten(self):
        self.output.write_bytes(b"Existing user result must survive")
        before = self.output.read_bytes()
        with self.assertRaises(J.ReportError):
            J.generate_report(self.reference, self.raw, self.approve(self.plan()), self.output)
        self.assertEqual(self.output.read_bytes(), before)

    def test_multiple_targets_are_independently_computed(self):
        first = self.plan(filters=[{"column": "Group", "op": "=", "value": "A", "type": "text"}])["mappings"][0]
        second = self.plan(target={"sheet": "Results", "cell": "B3"}, filters=[{"column": "Group", "op": "=", "value": "B", "type": "text"}])["mappings"][0]
        self.run_plan(self.plan(mappings=[first, second]))
        self.assertEqual(self.value("B2"), 125)
        self.assertEqual(self.value("B3"), 200)

    def test_saved_cell_swap_is_detected_even_when_total_matches(self):
        first = self.plan(filters=[{"column": "Group", "op": "=", "value": "A", "type": "text"}])["mappings"][0]
        second = self.plan(target={"sheet": "Results", "cell": "B3"}, filters=[{"column": "Group", "op": "=", "value": "B", "type": "text"}])["mappings"][0]
        self.run_plan(self.plan(mappings=[first, second]))
        wb = load_workbook(self.output)
        ws = wb["Results"]
        ws["B2"], ws["B3"] = ws["B3"].value, ws["B2"].value
        wb.save(self.output)
        with self.assertRaises(J.ReportError):
            J.verify_output(self.output, {("Results", "B2"): 125, ("Results", "B3"): 200})

    def test_large_amount_does_not_expand_allowed_error(self):
        self.write_raw(["Group", "Amount"], [["A", 3_000_000_000]])
        self.run_plan()
        wb = load_workbook(self.output)
        wb["Results"]["B2"] = 3_000_001_000
        wb.save(self.output)
        with self.assertRaises(J.ReportError):
            J.verify_output(self.output, {("Results", "B2"): 3_000_000_000})

    def test_verification_failure_does_not_publish_output(self):
        approved = self.approve(self.plan())
        with patch.object(J, "verify_output", side_effect=J.ReportError("Injected saved-cell verification failure")):
            with self.assertRaises(J.ReportError):
                J.generate_report(self.reference, self.raw, approved, self.output)
        self.assertFalse(self.output.exists())
        self.assertEqual([], list(self.root.glob("*.tmp.xlsx")), "Temporary output must be cleaned up")

    def test_table_grouping_uses_explicit_columns_and_writes_totals(self):
        mapping = self.plan()["mappings"][0]
        mapping.pop("target")
        mapping.update({
            "mode": "table",
            "target": {"sheet": "Results", "start_cell": "A10", "clear_range": "A10:D30", "allow_expand": False},
            "group_by": ["Group"],
            "include_totals": True,
            "sort": "label",
        })
        self.run_plan(self.plan(mappings=[mapping]))
        self.assertEqual(self.value("A11"), "A")
        self.assertEqual(self.value("B11"), 125)
        self.assertEqual(self.value("A12"), "B")
        self.assertEqual(self.value("B12"), 200)
        self.assertEqual(self.value("B13"), 325)

    def test_pivot_table_groups_by_month_and_category(self):
        self.write_raw(["Group", "Day", "Amount"], [
            ["A", "2026-01-05", 100],
            ["A", "2026-02-05", 25],
            ["B", "2026-01-08", 200],
        ])
        mapping = self.plan()["mappings"][0]
        mapping.update({
            "mode": "table",
            "target": {"sheet": "Results", "start_cell": "A10", "clear_range": "A10:D30", "allow_expand": False},
            "group_by": ["Group"],
            "pivot": {"column": "Day", "kind": "month"},
            "include_totals": True,
            "sort": "label",
        })
        self.run_plan(self.plan(mappings=[mapping]))
        self.assertEqual(self.value("B11"), 100)
        self.assertEqual(self.value("C11"), 25)
        self.assertEqual(self.value("D11"), 125)
        self.assertEqual(self.value("B12"), 200)
        self.assertEqual(self.value("D12"), 200)
        self.assertEqual(self.value("B13"), 300)
        self.assertEqual(self.value("C13"), 25)
        self.assertEqual(self.value("D13"), 325)

    def test_mappings_can_target_multiple_reference_sheets(self):
        wb = load_workbook(self.reference)
        ws = wb.create_sheet("Second layout")
        ws["A1"] = "Another report"
        wb.save(self.reference)
        first = self.plan()["mappings"][0]
        second = self.plan(target={"sheet": "Second layout", "cell": "C7"}, aggregate="count", terms=[])["mappings"][0]
        self.run_plan(self.plan(mappings=[first, second]))
        wb = load_workbook(self.output)
        try:
            self.assertEqual(wb["Results"]["B2"].value, 325)
            self.assertEqual(wb["Second layout"]["C7"].value, 3)
        finally:
            wb.close()

    def test_numeric_serial_date_uses_source_workbook_1904_epoch(self):
        wb = Workbook()
        wb.epoch = dt.datetime(1904, 1, 1)
        ws = wb.active
        ws.title = "Data"
        ws.append(["Day", "Amount"])
        serial = (dt.date(2026, 1, 5) - dt.date(1904, 1, 1)).days
        ws.append([serial, 100])
        ws.append(["2026-01-05", 50])
        ws.append(["2026-01-06", 900])
        wb.save(self.raw)
        self.run_plan(self.plan(filters=[{"column": "Day", "op": "=", "value": "2026-01-05", "type": "date"}]))
        self.assertEqual(self.value(), 150)

    def test_repeated_header_is_not_counted_as_data(self):
        self.write_raw(["Group", "Amount"], [["A", 100], ["Group", "Amount"], ["B", 200]])
        self.assert_blocked(self.plan(aggregate="count", terms=[]))

    def test_target_range_instead_of_single_cell_is_actionable_error(self):
        self.assert_blocked(self.plan(target={"sheet": "Results", "cell": "B2:C3"}))

    def test_out_of_range_numeric_date_is_actionable_error(self):
        self.write_raw(["Day", "Amount"], [[2_999_999, 100]])
        self.assert_blocked(self.plan(filters=[{"column": "Day", "op": "=", "value": "2026-01-05", "type": "date"}]))

    def test_table_category_beginning_equals_is_written_as_text(self):
        self.write_raw(["Group", "Amount"], [["A", 10], ["B", 20]])
        wb = load_workbook(self.raw)
        wb["Data"]["A2"] = "=HYPERLINK(\"https://example.invalid\",\"A\")"
        wb["Data"]["A2"].data_type = "s"
        wb.save(self.raw)
        mapping = self.plan()["mappings"][0]
        mapping.update({"mode": "table", "target": {"sheet": "Results", "start_cell": "A10", "clear_range": "A10:D30", "allow_expand": False}, "group_by": ["Group"], "sort": "source"})
        self.run_plan(self.plan(mappings=[mapping]))
        wb = load_workbook(self.output)
        try:
            self.assertEqual(wb["Results"]["A11"].data_type, "s")
            self.assertTrue(wb["Results"]["A11"].value.startswith("="))
        finally:
            wb.close()

    def test_table_will_not_expand_without_explicit_permission(self):
        mapping = self.plan()["mappings"][0]
        mapping.update({"mode": "table", "target": {"sheet": "Results", "start_cell": "A10", "clear_range": "A10:B11", "allow_expand": False}, "group_by": ["Group"]})
        self.assert_blocked(self.plan(mappings=[mapping]))

    def test_table_expansion_does_not_overwrite_neighboring_content(self):
        wb = load_workbook(self.reference)
        wb["Results"]["A12"] = "Separate section"
        wb.save(self.reference)
        mapping = self.plan()["mappings"][0]
        mapping.update({"mode": "table", "target": {"sheet": "Results", "start_cell": "A10", "clear_range": "A10:B11", "allow_expand": True}, "group_by": ["Group"]})
        self.assert_blocked(self.plan(mappings=[mapping]))

    def test_amount_desc_sort_uses_selected_aggregate_not_raw_sum(self):
        self.write_raw(["Group", "Amount"], [["A", 10], ["A", 10], ["B", 15]])
        mapping = self.plan(aggregate="max")["mappings"][0]
        mapping.update({"mode": "table", "target": {"sheet": "Results", "start_cell": "A10", "clear_range": "A10:D30", "allow_expand": False}, "group_by": ["Group"], "sort": "amount_desc"})
        self.run_plan(self.plan(mappings=[mapping]))
        self.assertEqual(self.value("A11"), "B", "15 should sort before 10 for max aggregation")
        self.assertEqual(self.value("B11"), 15)
        self.assertEqual(self.value("A12"), "A")

    def test_untouched_old_numeric_values_are_reported_for_review(self):
        result = self.run_plan()
        self.assertEqual(self.value("B3"), 888)
        self.assertIn(("Results", "B3"), {(x["sheet"], x["cell"]) for x in result["unmapped_numeric_cells"]})

    def test_intentional_numeric_constant_can_be_explicitly_preserved(self):
        plan = self.plan()
        plan["preserve_numeric"] = [{"sheet": "Results", "cell": "B3"}]
        result = self.run_plan(plan)
        self.assertEqual(self.value("B3"), 888)
        self.assertEqual(result["unmapped_numeric_cells"], [])

    def test_misspelled_filter_option_blocks_instead_of_being_ignored(self):
        plan = self.plan()
        plan["mappings"][0]["fliters"] = [{"column": "Group", "op": "=", "value": "A", "type": "text"}]
        self.assert_blocked(plan)

    def test_unknown_plan_option_blocks_instead_of_being_ignored(self):
        plan = self.plan()
        plan["source_sheeet"] = "Another source"
        self.assert_blocked(plan)

    def test_untouched_old_numeric_text_is_reported_for_review(self):
        wb = load_workbook(self.reference)
        wb["Results"]["B3"] = "9,999"
        wb.save(self.reference)
        result = self.run_plan()
        self.assertEqual(self.value("B3"), "9,999")
        self.assertIn(("Results", "B3"), {(x["sheet"], x["cell"]) for x in result["unmapped_numeric_cells"]})


if __name__ == "__main__":
    unittest.main(verbosity=2)
