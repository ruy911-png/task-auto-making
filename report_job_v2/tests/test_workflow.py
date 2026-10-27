"""End-to-end command line and preservation tests, independent of real files."""
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
import zipfile

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
import report_job_v2 as J


class WorkflowTests(unittest.TestCase):
    def setUp(self):
        temp = tempfile.TemporaryDirectory()
        self.addCleanup(temp.cleanup)
        self.dir = Path(temp.name)
        self.ref, self.raw, self.out = (self.dir / n for n in ("reference.xlsx", "raw.xlsx", "out.xlsx"))
        w = Workbook(); w.active.title = "Target"
        w.active["A1"] = "Generic demonstration"; w.active["B2"] = 0
        w.save(self.ref)
        w = Workbook(); w.active.title = "Records"
        w.active.append(["Key", "Value"])
        w.active.append(["001", 21]); w.active.append(["002", 9])
        w.save(self.raw)
        self.plan = {"version":2,"approved":False,"source":{"sheet":"Records","header_row":1},"unresolved":[],
                     "mappings":[{"target":{"sheet":"Target","cell":"B2"},"aggregate":"sum","terms":[{"column":"Value","factor":1}],"unit":1,"filters":[]}]}

    def cli(self, *args):
        return subprocess.run([sys.executable, str(ROOT/"report_job_v2.py"), *map(str,args)], capture_output=True, text=True, encoding="utf-8", timeout=30)

    def test_complete_explicit_rules_cli_workflow(self):
        rules, approved = self.dir/"rules.json", self.dir/"approved.json"
        rules.write_text(json.dumps(self.plan), encoding="utf-8")
        p = self.cli("approve", "--reference", self.ref, "--raw", self.raw, "--rules", rules, "--out", approved)
        self.assertEqual(p.returncode, 0, p.stderr)
        p = self.cli("run", "--reference", self.ref, "--raw", self.raw, "--rules", approved, "--out", self.out)
        self.assertEqual(p.returncode, 0, p.stderr)
        self.assertEqual(load_workbook(self.out).active["B2"].value, 30)

    def test_cli_refuses_to_generate_from_unapproved_rules(self):
        rules = self.dir/"rules.json"; rules.write_text(json.dumps(self.plan), encoding="utf-8")
        p = self.cli("run", "--reference", self.ref, "--raw", self.raw, "--rules", rules, "--out", self.out)
        self.assertEqual(p.returncode, 1)
        self.assertFalse(self.out.exists())

    def test_misspelled_aggregation_option_does_not_fall_back_to_sum(self):
        self.plan["mappings"][0]["aggreate"] = "count"
        with self.assertRaises(J.ReportError): J.approve_plan(self.plan, self.ref, self.raw)

    def test_unsupported_native_control_is_rejected_before_save(self):
        with zipfile.ZipFile(self.ref, "a") as z: z.writestr("xl/activeX/control.xml", "<control/>")
        with self.assertRaises(J.ReportError): J.approve_plan(self.plan, self.ref, self.raw)

    def test_desktop_preview_is_readable_without_executing_a_gui(self):
        from report_ui import describe_plan
        s = describe_plan(self.plan)
        self.assertIn("B2", s)
        self.assertNotIn('"mappings":', s)


if __name__ == "__main__": unittest.main()
