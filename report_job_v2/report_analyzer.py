"""Discover workbook structure without deciding the user's business rules.

All suggestions are unapproved. Discovery is deliberately separate from report
generation: matching labels is evidence for a candidate, not permission to write.
"""
from __future__ import annotations

import datetime as dt
import hashlib
import math
import posixpath
import re
import unicodedata
import zipfile
from collections import Counter
from pathlib import Path
from xml.etree import ElementTree as ET

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

MAX_DISCOVERY_ROWS = 5000
MAX_DISCOVERY_COLUMNS = 256
MAX_HEADER_ROWS = 100
PROFILE_ROWS = 300


def label_key(value):
    """Only header matching ignores spacing/case. Data comparisons must not."""
    return re.sub(r"\s+", "", unicodedata.normalize("NFKC", str(value))).casefold()


def file_signature(path):
    digest = hashlib.sha256()
    with open(path, "rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _display(value):
    if value is None:
        return None
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return value.isoformat()
    if isinstance(value, float) and not math.isfinite(value):
        return str(value)
    return value if isinstance(value, (str, int, float, bool)) else str(value)


def _kind(value):
    if value is None or (isinstance(value, str) and not value.strip()):
        return "blank"
    if isinstance(value, bool):
        return "boolean"
    if isinstance(value, (dt.date, dt.datetime)):
        return "date"
    if isinstance(value, (int, float)):
        return "number" if math.isfinite(value) else "invalid_number"
    text = str(value).strip()
    if re.fullmatch(r"\d{4}[-./]\d{1,2}[-./]\d{1,2}", text):
        try:
            dt.date(*map(int, re.split(r"[-./]", text)))
            return "date_text"
        except ValueError:
            pass
    normalized = re.sub(r"[₩￦$€£,\s]", "", text)
    normalized = re.sub(r"(?:KRW|USD|원)$", "", normalized, flags=re.I)
    normalized = normalized.replace("△", "-").replace("▲", "-")
    if normalized.endswith("-"):
        normalized = "-" + normalized[:-1]
    normalized = normalized.strip("()")
    # Leading-zero identifiers remain text. Numeric-looking values are only
    # profiling evidence; a candidate still needs explicit user approval.
    if re.fullmatch(r"[+\-]?(?:0|[1-9]\d*)(?:\.\d+)?", normalized):
        return "number_text"
    return "text"


def _formula_metadata(path):
    """Read actual cached values from OOXML, including all rows, not a sample."""
    main = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    relation = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
    result = {}
    with zipfile.ZipFile(path) as archive:
        book = ET.fromstring(archive.read("xl/workbook.xml"))
        rels = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        targets = {rel.get("Id"): rel.get("Target") for rel in rels}
        for sheet in book.iter(main + "sheet"):
            target = targets.get(sheet.get(relation + "id"))
            if not target:
                continue
            part = (target.lstrip("/") if target.startswith("/") else
                    posixpath.normpath(posixpath.join("xl", target)))
            item = {"formula_without_cache": [], "formula_error_cache": [], "formula_count": 0}
            with archive.open(part) as xml:
                for _, element in ET.iterparse(xml, events=("end",)):
                    if element.tag != main + "c":
                        continue
                    if element.find(main + "f") is not None:
                        item["formula_count"] += 1
                        cached = element.find(main + "v")
                        if cached is None or cached.text is None:
                            item["formula_without_cache"].append(element.get("r"))
                        elif element.get("t") == "e":
                            item["formula_error_cache"].append(element.get("r"))
                    element.clear()
            result[sheet.get("name")] = item
    return result


def _profiles(grid, columns, first, last):
    profiles = []
    for index, name in columns:
        values = [grid[row - 1][index - 1] for row in range(first, min(last, first + PROFILE_ROWS - 1) + 1)]
        counts = Counter(_kind(value) for value in values)
        present = [value for value in values if _kind(value) != "blank"]
        profiles.append({"name": name, "column_index": index,
                         "column_letter": get_column_letter(index), "sample_count": len(values),
                         "kinds": dict(counts), "distinct_sample_count": len({str(_display(v)) for v in present}),
                         "examples": list(dict.fromkeys(str(_display(v)) for v in present))[:5]})
    return profiles


def _tables(grid, merged_ranges):
    candidates = []
    merged_headers = {row for firstrow, left, endrow, right in merged_ranges if left != right
                      for row in range(firstrow, endrow + 1)}
    for row_index, row in enumerate(grid[:MAX_HEADER_ROWS], start=1):
        columns = [(col, str(value).strip()) for col, value in enumerate(row, start=1)
                   if _kind(value) != "blank"]
        if len(columns) < 2:
            continue
        types = [_kind(row[index - 1]) for index, _ in columns]
        text_ratio = sum(kind in ("text", "date_text", "date") for kind in types) / len(types)
        if text_ratio < 0.6:
            continue
        first = next((r for r in range(row_index + 1, min(len(grid), row_index + 4) + 1)
                      if sum(_kind(grid[r - 1][c - 1]) != "blank" for c, _ in columns) >= 2), None)
        if first is None:
            continue
        last = first
        for r in range(first + 1, len(grid) + 1):
            if all(_kind(grid[r - 1][c - 1]) == "blank" for c, _ in columns):
                break
            last = r
        sample = [grid[r - 1][c - 1] for r in range(first, min(last, first + 19) + 1) for c, _ in columns]
        ratio = sum(_kind(v) in ("number", "number_text", "date", "date_text") for v in sample) / max(1, len(sample))
        unique = len({label_key(name) for _, name in columns}) == len(columns)
        score = round(0.55 * text_ratio + 0.25 * min(1, ratio * 3) +
                      0.1 * min(1, len(columns) / 4) + 0.1 * int(unique) -
                      0.18 * int(row_index in merged_headers), 4)
        observations = []
        if not unique:
            observations.append("열 제목이 중복되어 명시적인 열 선택이 필요합니다.")
        if row_index in merged_headers:
            observations.append("제목 후보에 가로 병합 셀이 있어 열 구조 확인이 필요합니다.")
        if ratio == 0:
            observations.append("자료가 모두 문자여서 제목 줄 자동 판별이 불명확합니다.")
        candidates.append({"header_row": row_index, "column_names": [name for _, name in columns],
                           "column_indices": [col for col, _ in columns], "first_data_row": first,
                           "last_data_row": last, "score": score,
                           "ambiguous": not unique or row_index in merged_headers or ratio == 0,
                           "observations": observations,
                           "profiles": _profiles(grid, columns, first, last)})
    chosen = []
    for candidate in sorted(candidates, key=lambda item: (-item["score"], item["header_row"])):
        overlaps = [table for table in chosen if not
                    (candidate["last_data_row"] < table["header_row"] or
                     candidate["header_row"] > table["last_data_row"])]
        if overlaps:
            for table in overlaps:
                if abs(table["score"] - candidate["score"]) < 0.08:
                    table["ambiguous"] = True
                    table.setdefault("alternative_header_rows", []).append(candidate["header_row"])
            continue
        chosen.append(candidate)
    for index, table in enumerate(sorted(chosen, key=lambda item: item["header_row"]), start=1):
        table["id"] = f"table_{index}"
    return sorted(chosen, key=lambda item: item["header_row"])


def inspect_workbook(path):
    """Return a JSON-serializable, bounded structural inspection of an .xlsx/.xlsm."""
    path = Path(path).expanduser().resolve(strict=True)
    if path.suffix.lower() not in (".xlsx", ".xlsm"):
        raise ValueError("분석 가능한 파일은 .xlsx 또는 .xlsm입니다.")
    formulas = _formula_metadata(path)
    workbook = load_workbook(path, data_only=True, read_only=False)
    result = {"path": str(path), "sha256": file_signature(path),
              "epoch": workbook.epoch.isoformat(), "sheets": [], "observations": []}
    try:
        for sheet in workbook.worksheets:
            height = min(sheet.max_row, MAX_DISCOVERY_ROWS)
            width = min(sheet.max_column, MAX_DISCOVERY_COLUMNS)
            grid = [list(row) for row in sheet.iter_rows(max_row=height, max_col=width, values_only=True)]
            merges = [(rng.min_row, rng.min_col, rng.max_row, rng.max_col) for rng in sheet.merged_cells.ranges]
            notes = []
            if sheet.max_row > height or sheet.max_column > width:
                notes.append("구조 분석은 처음 5,000행·256열로 제한합니다. 데이터 집계 범위는 별도로 확정해야 합니다.")
            tables = _tables(grid, merges)
            for table in tables:
                if table["last_data_row"] == height and sheet.max_row > height:
                    table["last_data_row"] = sheet.max_row
                    table["observations"].append("자료 끝 행은 시트 범위이며 실제 집계 시 빈 행·다른 표를 확인해야 합니다.")
            item = {"name": sheet.title, "state": sheet.sheet_state,
                    "dimensions": {"rows": sheet.max_row, "columns": sheet.max_column},
                    "merged_ranges": [str(rng) for rng in sheet.merged_cells.ranges],
                    "tables": tables, "observations": notes}
            item.update(formulas.get(sheet.title, {}))
            result["sheets"].append(item)
    finally:
        workbook.close()
    return result


def _pivot_header(value):
    text = str(value).strip()
    if re.fullmatch(r"(?:19|20)\d{2}[-./]\d{1,2}(?:[-./]\d{1,2})?(?:T.*)?", text):
        return "yearmonth"
    if re.fullmatch(r"(?:19|20)\d{2}년?", text):
        return "year"
    if re.fullmatch(r"(?:[1-9]|1[0-2])월", text):
        return "month"
    return None


def propose_plan(reference_path, raw_path):
    """Return executable-schema candidates, always requiring explicit review.

    Grouping, aggregation and units are only proposed choices. Business intent
    cannot be proven from workbook structure or equal numbers alone.
    """
    reference, source = inspect_workbook(reference_path), inspect_workbook(raw_path)
    sources = [{"sheet": sheet["name"], **table} for sheet in source["sheets"] for table in sheet["tables"]]
    reports = [{"sheet": sheet["name"], **table} for sheet in reference["sheets"] for table in sheet["tables"]]
    plan = {"version": 2, "approved": False, "reference": {"path": reference["path"], "sha256": reference["sha256"]},
            "source": {}, "source_candidates": sources, "report_candidates": reports,
            "mappings": [], "unresolved": [],
            "review_required": ["집계 방법·필터·단위·기간·출력 범위를 직접 확인하고 승인하세요."]}
    if len(sources) != 1 or sources[0]["ambiguous"]:
        plan["unresolved"].append("원본에서 사용할 시트·표·제목 줄을 명시적으로 선택해야 합니다.")
        return plan
    raw_table = sources[0]
    raw_sheet = next(sheet for sheet in source["sheets"] if sheet["name"] == raw_table["sheet"])
    # A saved plan must continue through a later period's larger source file.
    # The detected last row is evidence, never an implicit fixed execution bound.
    plan["source"] = {"sheet": raw_table["sheet"], "header_row": raw_table["header_row"],
                      "first_data_row": raw_table["first_data_row"], "last_data_row": None}
    plan["source_info"] = {"detected_first_data_row": raw_table["first_data_row"],
                           "detected_last_data_row": raw_table["last_data_row"]}
    if raw_table["last_data_row"] < raw_sheet["dimensions"]["rows"]:
        plan["unresolved"].append("원본 표 뒤에 다른 내용이 있습니다. 후행 주석·합계·다른 표를 확인하고 집계 범위를 명시해야 합니다.")
    if not reports:
        plan["unresolved"].append("기존 보고서의 표·출력 위치를 자동으로 찾지 못했습니다. 명시적인 셀 또는 표 매핑이 필요합니다.")
        return plan
    raw_profiles = raw_table["profiles"]
    for target in reports:
        if target["ambiguous"]:
            plan["unresolved"].append(f"'{target['sheet']}'의 제목 {target['header_row']}행이 불명확합니다.")
            continue
        target_profiles = target["profiles"]
        # A rectangular pivot starts with a category column and remaining
        # headers represent dates or category values from a different column.
        group_names = [p["name"] for p in target_profiles if p["kinds"].get("text", 0)]
        same_headers = {name: [p for p in raw_profiles if label_key(name) == label_key(p["name"])]
                        for name in target["column_names"]}
        raw_group = [same_headers[name][0]["name"] for name in group_names if len(same_headers[name]) == 1]
        if not raw_group and target_profiles:
            examples = set(target_profiles[0]["examples"])
            overlap = [p["name"] for p in raw_profiles if p["kinds"].get("text", 0)
                       and examples and examples.intersection(p["examples"])]
            if len(overlap) == 1:
                raw_group = overlap
        pivot, terms = None, []
        total_labels = {"합계", "총합계", "총계", "total", "grandtotal"}
        headers = [name for name in target["column_names"][1:] if label_key(name) not in total_labels]
        family = {_pivot_header(name) for name in headers}
        if raw_group and headers and len(family) == 1 and None not in family:
            dates = [p["name"] for p in raw_profiles if p["kinds"].get("date", 0) + p["kinds"].get("date_text", 0)]
            if len(dates) == 1:
                pivot = {"column": dates[0], "kind": next(iter(family))}
                raw_group = raw_group[:1]
        elif raw_group and headers:
            header_keys = {label_key(name) for name in headers}
            matching = [p["name"] for p in raw_profiles if p["name"] not in raw_group and
                        header_keys.issubset({label_key(v) for v in p["examples"]})]
            if len(matching) == 1:
                pivot = {"column": matching[0], "kind": "value"}
                raw_group = raw_group[:1]
        raw_numbers = [p for p in raw_profiles if (p["kinds"].get("number", 0) + p["kinds"].get("number_text", 0)) and
                       p["name"] not in raw_group and p["name"] != (pivot or {}).get("column")]
        target_number_names = [p["name"] for p in target_profiles if p["kinds"].get("number", 0) + p["kinds"].get("number_text", 0)]
        exact_measures = [p for p in raw_numbers if any(label_key(p["name"]) == label_key(name)
                                                      for name in target_number_names)]
        choices = exact_measures if exact_measures else raw_numbers
        if len(choices) == 1:
            terms = [{"column": choices[0]["name"], "factor": 1}]
        if not raw_group or not terms:
            plan["unresolved"].append(f"'{target['sheet']}' {target['header_row']}행 표의 그룹 열 또는 금액 열을 확정할 근거가 부족합니다.")
            continue
        first_col, last_col = min(target["column_indices"]), max(target["column_indices"])
        start = f"{get_column_letter(first_col)}{target['header_row']}"
        clear = f"{start}:{get_column_letter(last_col)}{target['last_data_row']}"
        mapping = {"id": f"mapping_{len(plan['mappings']) + 1}", "mode": "table",
                   "target": {"sheet": target["sheet"], "start_cell": start, "clear_range": clear, "allow_expand": False},
                   "group_by": raw_group, "aggregate": "sum", "terms": terms,
                   "filters": [], "unit": 1, "basis": "열 제목·값 종류·표 구조가 일치하는 후보이며 집계 의도는 사용자 확인 필요"}
        if pivot:
            mapping["pivot"] = pivot
        plan["mappings"].append(mapping)
    if not plan["mappings"] and not plan["unresolved"]:
        plan["unresolved"].append("자동으로 제안할 수 있는 매핑이 없습니다. 셀 또는 표 매핑을 직접 설정하세요.")
    plan["review_required"].append("표 출력은 선택한 범위를 다시 구성합니다. 기존 수식 덮어쓰기·새 항목 확장은 별도 허용이 필요합니다.")
    return plan
