"""Generic Excel report engine v2. Python 3.9+ / openpyxl.

No filename, worksheet, business account or reporting period is hardcoded.
Use the desktop wizard with no arguments, or inspect/propose/approve/run CLI.
"""
from __future__ import annotations

import argparse
import copy
import datetime as dt
from decimal import Decimal, InvalidOperation
import hashlib
import json
import math
import os
from pathlib import Path
import re
import sys
import tempfile
import zipfile

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils.cell import coordinate_to_tuple, get_column_letter, range_boundaries
from openpyxl.utils.datetime import from_excel
from openpyxl.workbook.properties import CalcProperties

from report_analyzer import inspect_workbook, propose_plan

VERSION = "2.0.0"
EXCEL_ERRORS = {"#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!", "#SPILL!", "#CALC!"}


class ReportError(Exception):
    """An actionable error; output is not published when validation fails."""


def _fidelity_guard(path):
    """Reject native objects that this openpyxl backend cannot preserve reliably."""
    from xml.etree import ElementTree as ET
    prefixes = ("xl/activeX/", "xl/ctrlProps/", "xl/slicers/", "xl/slicerCaches/",
                "xl/threadedComments/", "xl/pivotTables/", "xl/queryTables/", "xl/externalLinks/")
    with zipfile.ZipFile(path) as z:
        blocked = [n for n in z.namelist() if n.startswith(prefixes) or n == "xl/connections.xml"]
        for name in z.namelist():
            if name.startswith(("xl/worksheets/", "xl/drawings/")) and name.endswith(".xml"):
                root = ET.fromstring(z.read(name))
                if any(x.tag.rsplit("}", 1)[-1] in ("extLst", "sp", "cxnSp", "grpSp") for x in root.iter()):
                    blocked.append(name)
        if blocked:
            raise ReportError("기준 보고서에 이 저장 방식으로 보존을 보장할 수 없는 개체·외부연결이 있습니다. 해당 파일은 중단합니다: " + ", ".join(blocked[:4]))


def _sha(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for part in iter(lambda: f.read(1024 * 1024), b""):
            h.update(part)
    return h.hexdigest()


def _digest(plan):
    body = {k: v for k, v in plan.items() if k not in ("approval", "approved")}
    return hashlib.sha256(json.dumps(body, ensure_ascii=False, sort_keys=True, allow_nan=False).encode()).hexdigest()


def _text(v):
    if v is None:
        return ""
    if isinstance(v, dt.datetime):
        return v.isoformat(sep=" ")
    if isinstance(v, dt.date):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _numeric_or_date_like(v, epoch):
    if isinstance(v, bool) or v is None: return False
    if isinstance(v, (int, float, dt.date)): return True
    if isinstance(v, str) and v.strip():
        try:
            _number(v)
            return True
        except ReportError: pass
        try:
            _date(v, epoch)
            return True
        except ReportError: pass
    return False


def _number(v, location="값"):
    if isinstance(v, bool) or v is None or isinstance(v, (dt.date, dt.time)):
        raise ReportError(f"{location}: 숫자가 아님 ({v!r})")
    s = str(v).strip()
    if isinstance(v, str):
        s = re.sub(r"[\u200b-\u200d\ufeff]", "", s)
        if s in ("-", "—", "–"):
            return Decimal(0)
        negative = s.startswith("(") and s.endswith(")")
        if negative:
            s = s[1:-1].strip()
        if s.endswith("-"):
            negative, s = not negative, s[:-1].strip()
        s = s.replace("△", "-").replace("▲", "-")
        s = re.sub(r"^([+-]?)\s*(?:KRW|USD|₩|￦|\$)\s*", r"\1", s, flags=re.I)
        s = re.sub(r"\s*(?:KRW|원)$", "", s, flags=re.I).strip()
        if "," in s and not re.fullmatch(r"[+-]?\d{1,3}(?:,\d{3})+(?:\.\d+)?", s):
            raise ReportError(f"{location}: 숫자 쉼표 형식을 확인하세요 ({v!r})")
        s = s.replace(",", "")
        if negative:
            s = "-" + s
    try:
        d = Decimal(s)
    except InvalidOperation:
        raise ReportError(f"{location}: 숫자를 읽을 수 없음 ({v!r})") from None
    if not d.is_finite():
        raise ReportError(f"{location}: 유한한 숫자가 아님")
    return d


def _date(v, epoch):
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        if 10000101 <= v <= 99991231 and float(v).is_integer():
            v = str(int(v))
        elif 0 <= v < 3000000:
            try:
                x = from_excel(v, epoch=epoch)
            except (OverflowError, ValueError):
                raise ReportError(f"날짜 일련번호 범위를 확인하세요: {v!r}") from None
            if isinstance(x, dt.datetime):
                return x.date()
            raise ReportError(f"날짜 일련번호를 확인하세요: {v!r}")
    s = _text(v)
    m = re.fullmatch(r"(\d{4})[./년-]\s*(\d{1,2})[./월-]\s*(\d{1,2})일?(?:[ T].*)?", s)
    try:
        if m:
            return dt.date(*map(int, m.groups()))
        if re.fullmatch(r"\d{8}", s):
            return dt.date(int(s[:4]), int(s[4:6]), int(s[6:]))
    except ValueError:
        pass
    raise ReportError(f"날짜를 읽을 수 없음: {v!r}")


def _book(path, data=False):
    p = Path(path)
    if p.suffix.lower() not in (".xlsx", ".xlsm"):
        raise ReportError("지원 파일은 xlsx / xlsm입니다. xls는 Excel에서 xlsx로 변환하세요.")
    try:
        return load_workbook(p, data_only=data, keep_vba=(p.suffix.lower() == ".xlsm"))
    except Exception as exc:
        raise ReportError(f"엑셀 파일을 열 수 없음: {p.name} ({exc})") from exc


def _source(raw_path, spec):
    formulas, cached = _book(raw_path), _book(raw_path, data=True)
    try:
        name = spec.get("sheet")
        if name not in formulas.sheetnames:
            raise ReportError(f"원본 시트를 명시하세요: {name!r}")
        ws, ds = formulas[name], cached[name]
        hr = spec.get("header_row")
        if not isinstance(hr, int) or hr < 1:
            raise ReportError("원본 header_row는 1 이상 정수여야 합니다.")
        heads = {}
        for cell in ws[hr]:
            if cell.value is None:
                continue
            label = _text(cell.value)
            if cell.data_type == "f" or label in heads:
                raise ReportError(f"원본 제목 중복 또는 수식: {label}")
            heads[label] = cell.column
        if not heads:
            raise ReportError("원본 제목 행에 열 이름이 없습니다.")
        first = spec.get("first_data_row") or hr + 1
        last = spec.get("last_data_row") or ws.max_row
        if not isinstance(first, int) or not isinstance(last, int) or first <= hr or last < first:
            raise ReportError("원본 데이터 시작·끝 행을 확인하세요.")
        records = []
        for rn in range(first, last + 1):
            if all(ws.cell(rn, c).value is None for c in heads.values()):
                continue
            if all(_text(ws.cell(rn, c).value) == label for label, c in heads.items()):
                raise ReportError(f"원본 {rn}행: 제목 행이 반복됩니다. 원본 범위 또는 표 구분을 확인하세요.")
            vals, missing = {}, set()
            for label, col in heads.items():
                fc, dc = ws.cell(rn, col), ds.cell(rn, col)
                if fc.data_type == "f" and dc.value is None:
                    missing.add(label)
                if fc.data_type == "e" or dc.data_type == "e":
                    missing.add(label)
                vals[label] = dc.value
            records.append((rn, vals, missing))
        if not records:
            raise ReportError("원본 데이터가 없습니다.")
        return heads, records, formulas.epoch
    finally:
        formulas.close()
        cached.close()


def _requirements(m):
    cols = set(m.get("group_by", []))
    if m.get("pivot"):
        cols.add(m["pivot"]["column"])
    if m.get("aggregate", "sum") != "count":
        cols.update(t["column"] for t in m.get("terms", []))
    cols.update(f["column"] for f in m.get("filters", []))
    return cols


def _filter(row, f, epoch):
    a, b = row[f["column"]], f.get("value")
    typ, op = f.get("type", "text"), f.get("op", "=")
    if typ == "text":
        a, b = _text(a), _text(b)
    elif typ == "number":
        a, b = _number(a, f["column"]), _number(b, "필터 기준")
    elif typ == "date":
        a, b = _date(a, epoch), _date(b, epoch)
    else:
        raise ReportError(f"필터 type을 확인하세요: {typ}")
    if op == "=": return a == b
    if op in ("<>", "!="): return a != b
    if op == ">": return a > b
    if op == "<": return a < b
    if op == ">=": return a >= b
    if op == "<=": return a <= b
    if op == "contains" and typ == "text": return b in a
    raise ReportError(f"지원하지 않는 필터 연산자: {op}")


def _selected(records, m, epoch):
    required = _requirements(m)
    chosen = []
    for rn, row, missing in records:
        unavailable = required & missing
        if unavailable:
            raise ReportError(f"원본 {rn}행 {', '.join(sorted(unavailable))}: 수식 계산값 없음 또는 Excel 오류. Excel에서 재계산·저장하세요.")
        if not all(_filter(row, f, epoch) for f in m.get("filters", [])):
            continue
        chosen.append((rn, row))
    return chosen


def _amount(rn, row, m):
    if m.get("aggregate", "sum") == "count":
        return Decimal(1)
    terms = m.get("terms", [])
    if not terms:
        raise ReportError("계산할 금액 열(terms)을 지정하세요.")
    total = Decimal(0)
    for t in terms:
        v = row[t["column"]]
        if v is None or (isinstance(v, str) and not v.strip()):
            if m.get("blank_amount", "error") != "zero":
                raise ReportError(f"원본 {rn}행 {t['column']}: 금액 빈칸 (0으로 볼지는 명시해야 합니다)")
            v = 0
        total += _number(v, f"원본 {rn}행 {t['column']}") * _number(t.get("factor", 1), "계수")
    return total


def _reduce(values, m):
    op = m.get("aggregate", "sum")
    if op not in ("sum", "count", "avg", "min", "max"):
        raise ReportError(f"지원 집계: sum/count/avg/min/max ({op!r})")
    unit = _number(m.get("unit", 1), "단위")
    if unit <= 0 or (op == "count" and unit != 1):
        raise ReportError("단위는 양수여야 하며 건수 단위는 1이어야 합니다.")
    if op in ("sum", "count"):
        result = sum(values, Decimal(0))
    elif not values:
        raise ReportError(f"{op}: 조건에 맞는 원본 행이 없습니다.")
    elif op == "avg":
        result = sum(values, Decimal(0)) / len(values)
    elif op == "min":
        result = min(values)
    else:
        result = max(values)
    result /= unit
    if "round_digits" in m:
        from decimal import ROUND_HALF_UP
        digits = int(m["round_digits"])
        if not -12 <= digits <= 12:
            raise ReportError("반올림 자릿수 범위는 -12~12입니다.")
        result = result.quantize(Decimal(1).scaleb(-digits), rounding=ROUND_HALF_UP)
    return result


def _excel_number(value):
    # Excel numeric cells cannot preserve arbitrary precision identifiers/amounts.
    if len(value.normalize().as_tuple().digits) > 15:
        raise ReportError("결과가 Excel 숫자의 15자리 정밀도를 넘습니다. 반올림 자릿수(round_digits)를 명시하세요.")
    if not value.is_finite() or abs(value) > Decimal("1e308"):
        raise ReportError("Excel에 저장할 수 없는 금액입니다.")
    return int(value) if value == value.to_integral() else float(value)


def _pivot(v, spec, epoch):
    kind = spec["kind"]
    if kind == "value":
        if not _text(v): raise ReportError("열 기준 값이 빈칸입니다.")
        return _text(v)
    d = _date(v, epoch)
    if kind == "month": return f"{d.year:04d}-{d.month:02d}"  # year retained: no cross-year coalescing
    if kind == "yearmonth": return f"{d.year:04d}-{d.month:02d}"
    if kind == "year": return str(d.year)
    raise ReportError(f"열 기준 종류를 확인하세요: {kind}")


def _table(m, selected, epoch):
    groups, pivots, buckets = [], [], {}
    by, pivot = m.get("group_by", []), m.get("pivot")
    for rn, row in selected:
        key = tuple(_text(row[c]) for c in by)
        if any(not k for k in key) and m.get("blank_keys", "error") != "include":
            raise ReportError(f"원본 {rn}행: 묶음 기준이 빈칸입니다.")
        pk = _pivot(row[pivot["column"]], pivot, epoch) if pivot else "값"
        if key not in groups: groups.append(key)
        if pk not in pivots: pivots.append(pk)
        buckets.setdefault((key, pk), []).append(_amount(rn, row, m))
    if not groups:
        raise ReportError("표에 집계할 원본 행이 없습니다.")
    if pivot and pivot["kind"] != "value": pivots.sort()
    sort = m.get("sort", "source")
    if sort == "label": groups.sort()
    elif sort == "amount_desc":
        groups.sort(key=lambda k: -_reduce([v for p in pivots for v in buckets.get((k, p), [])], m))
    elif sort != "source": raise ReportError("정렬은 source/label/amount_desc 중 하나입니다.")
    totals = m.get("include_totals", True)
    total_label = {"sum": "합계", "count": "총건수", "avg": "전체 평균", "min": "전체 최소", "max": "전체 최대"}[m.get("aggregate", "sum")]
    headers = list(by) + (pivots if pivot else [m.get("value_title", m.get("aggregate", "sum"))])
    if not by: headers = ["전체"] + headers
    if totals and pivot: headers += [total_label]
    grid = [headers]
    for key in groups:
        row = list(key) if by else ["전체"]
        for pk in pivots:
            vals = buckets.get((key, pk))
            row.append(_excel_number(_reduce(vals, m)) if vals else None)
        if totals and pivot:
            row.append(_excel_number(_reduce([v for p in pivots for v in buckets.get((key, p), [])], m)))
        grid.append(row)
    if totals:
        row = [total_label] + [None] * (max(1, len(by)) - 1)
        for p in pivots:
            row.append(_excel_number(_reduce([v for k in groups for v in buckets.get((k, p), [])], m)))
        if pivot: row.append(_excel_number(_reduce([v for vs in buckets.values() for v in vs], m)))
        grid.append(row)
    return grid


def _target(wb, sheet, coordinate, mapping, planned, value):
    if sheet not in wb.sheetnames:
        raise ReportError(f"보고서 시트 없음: {sheet}")
    if not isinstance(coordinate, str) or not re.fullmatch(r"[A-Za-z]{1,3}[1-9]\d{0,6}", coordinate):
        raise ReportError(f"단일 결과 셀 주소를 지정하세요: {coordinate!r}")
    tr, tc = coordinate_to_tuple(coordinate)
    if tr > 1048576 or tc > 16384:
        raise ReportError(f"Excel 셀 범위를 넘었습니다: {coordinate}")
    try:
        cell = wb[sheet][coordinate]
    except (ValueError, TypeError):
        raise ReportError(f"결과 셀 주소를 확인하세요: {coordinate}") from None
    if isinstance(cell, MergedCell):
        raise ReportError(f"병합 셀의 왼쪽 위 칸에 연결하세요: {sheet}!{coordinate}")
    if cell.data_type == "f" and not mapping.get("allow_overwrite_formulas", False):
        raise ReportError(f"수식 덮어쓰기는 별도 확인이 필요합니다: {sheet}!{coordinate}")
    if cell.data_type != "f" and isinstance(cell.value, str) and cell.value.strip() and not mapping.get("allow_overwrite_text", False):
        raise ReportError(f"글자 덮어쓰기는 별도 확인이 필요합니다: {sheet}!{coordinate}")
    key = (sheet, cell.coordinate)
    if key in planned:
        raise ReportError(f"결과 셀 연결 중복: {sheet}!{coordinate}")
    planned[key] = value


def _validate(plan):
    def keys(obj, allowed, label):
        if not isinstance(obj, dict): raise ReportError(f"{label} 형식이 잘못됐습니다.")
        extra = obj.keys() - set(allowed.split())
        if extra: raise ReportError(f"{label}의 알 수 없는 설정: {', '.join(sorted(extra))}")
    keys(plan, "version approved approval source mappings unresolved reference source_candidates report_candidates source_info review_required preserve_numeric", "규칙")
    if plan.get("version") != 2:
        raise ReportError("규칙 version은 2여야 합니다.")
    if plan.get("unresolved"):
        raise ReportError("아직 확정하지 않은 연결 기준: " + "; ".join(map(str, plan["unresolved"])))
    if not isinstance(plan.get("source"), dict):
        raise ReportError("원본 시트·제목 행을 지정하세요.")
    keys(plan["source"], "sheet header_row first_data_row last_data_row", "원본 설정")
    if not isinstance(plan.get("mappings"), list) or not plan["mappings"]:
        raise ReportError("보고서에 연결할 항목이 없습니다.")
    for m in plan["mappings"]:
        keys(m, "id mode target aggregate terms filters unit group_by pivot blank_amount blank_keys include_totals sort value_title round_digits allow_overwrite_text allow_overwrite_formulas basis", "연결 설정")
        keys(m.get("target", {}), "sheet cell start_cell clear_range allow_expand", "결과 위치")
        for term in m.get("terms", []): keys(term, "column factor", "계산 열")
        for f in m.get("filters", []): keys(f, "column op value type", "필터")
        if m.get("pivot"): keys(m["pivot"], "column kind", "열 기준")


def _calculate(reference_path, raw_path, plan):
    _validate(plan)
    _fidelity_guard(reference_path)
    wb = _book(reference_path)
    try:
        heads, records, epoch = _source(raw_path, plan["source"])
        planned, stats = {}, []
        for index, m in enumerate(plan["mappings"], 1):
            missing = _requirements(m) - heads.keys()
            if missing: raise ReportError("원본 열 없음: " + ", ".join(sorted(missing)))
            selected = _selected(records, m, epoch)
            target = m.get("target", {})
            mode = m.get("mode", "cells")
            if mode in ("cells", "cell"):
                val = _excel_number(_reduce([_amount(rn, row, m) for rn, row in selected], m))
                _target(wb, target.get("sheet"), target.get("cell"), m, planned, val)
            elif mode == "table":
                grid = _table(m, selected, epoch)
                sheet = target.get("sheet")
                if sheet not in wb.sheetnames: raise ReportError(f"보고서 시트 없음: {sheet}")
                sr, sc = coordinate_to_tuple(target.get("start_cell", ""))
                if not target.get("clear_range"):
                    raise ReportError("표 교체 범위(clear_range)를 명시하세요.")
                c1, r1, c2, r2 = range_boundaries(target["clear_range"])
                if None in (c1, r1, c2, r2) or c2 > 16384 or r2 > 1048576 or (c2-c1+1)*(r2-r1+1) > 1000000:
                    raise ReportError("표 교체 범위를 유효한 100만 칸 이하의 영역으로 지정하세요.")
                if sr != r1 or sc != c1:
                    raise ReportError("표 시작 셀과 교체 범위 왼쪽 위가 같아야 합니다.")
                end_r, end_c = sr + len(grid) - 1, sc + len(grid[0]) - 1
                if (end_r > r2 or end_c > c2) and not target.get("allow_expand", False):
                    raise ReportError("새 항목이 교체 범위를 넘습니다. 확장 허용과 주변 내용을 확인하세요.")
                # Expansion may use empty cells only, never shift other report sections.
                for rn in range(sr, max(r2, end_r) + 1):
                    for cn in range(sc, max(c2, end_c) + 1):
                        if (rn > r2 or cn > c2) and wb[sheet].cell(rn, cn).value is not None:
                            raise ReportError(f"표 확장 위치에 기존 내용이 있습니다: {sheet}!{get_column_letter(cn)}{rn}")
                        v = grid[rn-sr][cn-sc] if rn <= end_r and cn <= end_c else None
                        table_map = {**m, "allow_overwrite_text": True}
                        _target(wb, sheet, f"{get_column_letter(cn)}{rn}", table_map, planned, v)
            else: raise ReportError(f"연결 mode를 확인하세요: {mode}")
            stats.append({"mapping": m.get("id", str(index)), "source_rows": len(records), "selected_rows": len(selected), "filtered_rows": len(records)-len(selected)})
        return planned, stats
    finally:
        wb.close()


def approve_plan(plan, reference_path, raw_path):
    """Called ONLY after user reviews the plan. Does not generate a workbook."""
    result = copy.deepcopy(plan)
    _calculate(reference_path, raw_path, result)
    result["approved"] = True
    result["approval"] = {"reference_sha256": _sha(reference_path), "digest": _digest(result), "version": VERSION}
    return result


def verify_output(path, expected):
    wb = _book(path)
    try:
        for (sheet, coord), wanted in expected.items():
            if sheet not in wb.sheetnames:
                raise ReportError(f"저장 검증: 시트 누락 {sheet}")
            actual = wb[sheet][coord].value
            if isinstance(wanted, (int, float)) and not isinstance(wanted, bool):
                if isinstance(actual, bool) or not isinstance(actual, (int, float)) or _number(actual) != _number(wanted):
                    raise ReportError(f"저장 검증 불일치: {sheet}!{coord} (기대 {wanted!r}, 실제 {actual!r})")
            elif actual != wanted:
                raise ReportError(f"저장 검증 불일치: {sheet}!{coord} (기대 {wanted!r}, 실제 {actual!r})")
    finally:
        wb.close()


def _same(a, b):
    if Path(a).resolve() == Path(b).resolve(): return True
    try: return os.path.samefile(a, b)
    except OSError: return False


def _recalculate(path):
    try:
        import pythoncom
        import win32com.client
    except ImportError as exc:
        raise ReportError("Excel 재계산에는 Windows Excel과 pywin32가 필요합니다.") from exc
    pythoncom.CoInitialize()
    app, book = None, None
    try:
        app = win32com.client.DispatchEx("Excel.Application")
        app.Visible = False
        app.DisplayAlerts = False
        app.EnableEvents = False
        app.AskToUpdateLinks = False
        app.AutomationSecurity = 3  # never execute input macros
        book = app.Workbooks.Open(str(Path(path).resolve()), UpdateLinks=0)
        app.CalculateFullRebuild()
        book.Save()
        book.Close(False)
        book = None
    except Exception as exc:
        raise ReportError(f"Excel 재계산 실패: {exc}") from exc
    finally:
        if book is not None:
            try: book.Close(False)
            except Exception: pass
        if app is not None:
            try: app.Quit()
            except Exception: pass
        pythoncom.CoUninitialize()


def generate_report(reference_path, raw_path, plan, output_path, *, use_excel=False):
    _validate(plan)
    if not plan.get("approved") or not plan.get("approval"):
        raise ReportError("연결 기준을 먼저 확인·승인하세요. 미승인 규칙으로 생성하지 않습니다.")
    approval = plan["approval"]
    if approval.get("digest") != _digest(plan):
        raise ReportError("승인 후 규칙이 변경됐습니다. 내용을 다시 확인·승인하세요.")
    if approval.get("reference_sha256") != _sha(reference_path):
        raise ReportError("기준 보고서가 변경됐습니다. 연결 기준을 다시 확인하세요.")
    output = Path(output_path)
    for inp in (reference_path, raw_path):
        if _same(inp, output): raise ReportError("결과 경로가 입력 파일과 같습니다. 원본 덮어쓰기는 금지됩니다.")
    if output.exists() or output.is_symlink(): raise ReportError("결과 파일이 이미 있습니다. 새 이름을 지정하세요.")
    if not output.parent.is_dir(): raise ReportError("결과 저장 폴더가 없습니다.")
    if output.suffix.lower() != Path(reference_path).suffix.lower():
        raise ReportError("결과 확장자는 기준 보고서와 같아야 합니다 (xlsx/xlsm).")
    before = {_text(Path(p).resolve()): _sha(p) for p in (reference_path, raw_path)}
    expected, stats = _calculate(reference_path, raw_path, plan)
    wb = _book(reference_path)
    preserved, unmapped = {}, []
    kept = {(x.get("sheet"), x.get("cell")) for x in plan.get("preserve_numeric", [])}
    try:
        for ws in wb:
            for row in ws:
                for cell in row:
                    if cell.data_type == "f" and (ws.title, cell.coordinate) not in expected:
                        preserved[(ws.title, cell.coordinate)] = cell.value
                    elif (ws.title, cell.coordinate) not in expected and (ws.title, cell.coordinate) not in kept and (_numeric_or_date_like(cell.value, wb.epoch) or cell.data_type == "e"):
                        unmapped.append({"sheet": ws.title, "cell": cell.coordinate, "value": _text(cell.value)})
        for (sheet, coord), value in expected.items():
            cell = wb[sheet][coord]
            cell.value = value
            if isinstance(value, str): cell.data_type = "s"  # no formula injection
            if isinstance(value, (int, float)) and cell.number_format == "General": cell.number_format = "#,##0.########"
        wb.calculation = CalcProperties(calcId=0, fullCalcOnLoad=True)
        fd, temp = tempfile.mkstemp(prefix=".report_v2_", suffix=output.suffix, dir=output.parent)
        os.close(fd)
        try:
            wb.save(temp)
            wb.close()
            verify_output(temp, {**preserved, **expected})
            if use_excel:
                _recalculate(temp)
                verify_output(temp, expected)
                calc = _book(temp, data=True)
                try:
                    for sh, co in preserved:
                        c = calc[sh][co]
                        if c.value is None or c.data_type == "e":
                            raise ReportError(f"Excel 수식 계산 결과 확인 필요: {sh}!{co}")
                finally: calc.close()
            for inp in (reference_path, raw_path):
                if _sha(inp) != before[_text(Path(inp).resolve())]: raise ReportError("실행 중 입력 파일이 변경됐습니다. 결과를 게시하지 않습니다.")
            # Temp and destination share a directory/filesystem. A hard link
            # publishes validated bytes atomically and fails if the name exists.
            try:
                os.link(temp, output)
            except FileExistsError:
                raise ReportError("저장 직전에 같은 이름의 결과 파일이 생겼습니다.") from None
            except OSError as exc:
                raise ReportError(f"검증한 파일을 안전하게 게시할 수 없습니다. 로컬 저장 폴더를 선택하세요: {exc}") from exc
        finally:
            if os.path.exists(temp): os.unlink(temp)
    finally:
        wb.close()
    return {"output": str(output.resolve()), "cells_written": len(expected), "mappings": stats,
            "formula_recalculation_required": bool(preserved) and not use_excel,
            "formulas_preserved": len(preserved), "verified": True,
            "unmapped_numeric_cells": unmapped,
            "verification_scope": "mapped_values_and_preserved_formula_text",
            "status": "수식 재계산 필요" if preserved and not use_excel else ("미연결 숫자 확인 필요" if unmapped else "연결 셀 작성·저장값 대조 완료")}


def main(argv=None):
    for stream in (sys.stdout, sys.stderr):
        try: stream.reconfigure(errors="replace")
        except (AttributeError, ValueError): pass
    if argv is None and len(sys.argv) == 1:
        from report_ui import launch_gui
        launch_gui(sys.modules[__name__])
        return 0
    parser = argparse.ArgumentParser(description="범용 엑셀 보고서 v2 — 먼저 연결 기준을 확인한 뒤 생성")
    sub = parser.add_subparsers(dest="command", required=True)
    p = sub.add_parser("inspect"); p.add_argument("file")
    for name in ("propose", "approve", "run"):
        p = sub.add_parser(name)
        p.add_argument("--reference", required=True); p.add_argument("--raw", required=True)
        p.add_argument("--out", required=True)
        if name != "propose": p.add_argument("--rules", required=True)
        if name == "run": p.add_argument("--excel", action="store_true")
    args = parser.parse_args(argv)
    try:
        if args.command == "inspect":
            print(json.dumps(inspect_workbook(args.file), ensure_ascii=False, indent=2)); return 0
        if args.command == "propose": result = propose_plan(args.reference, args.raw)
        else:
            with open(args.rules, encoding="utf-8") as f: plan = json.load(f)
            if args.command == "approve": result = approve_plan(plan, args.reference, args.raw)
            else: result = generate_report(args.reference, args.raw, plan, args.out, use_excel=args.excel)
        if args.command in ("propose", "approve"):
            if any(_same(args.out, p) for p in (args.reference, args.raw)):
                raise ReportError("규칙 저장 경로가 입력 파일과 같습니다.")
            with open(args.out, "x", encoding="utf-8") as f: json.dump(result, f, ensure_ascii=False, indent=2)
        print(json.dumps(result, ensure_ascii=False, indent=2))
        return 2 if result.get("formula_recalculation_required") or result.get("unmapped_numeric_cells") else 0
    except (ReportError, OSError, ValueError, KeyError, TypeError) as exc:
        print(f"최종: 실행 중단 — {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    sys.exit(main())
