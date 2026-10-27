"""Create fictional inputs and execute explicit demo rules; no private files required."""
from pathlib import Path
import sys
import json
from datetime import date
from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
import report_job_v2 as J


def main():
    dest = Path.cwd() / "v2_demo"
    dest.mkdir(exist_ok=True)
    reference, raw, output = (dest / n for n in ("reference.xlsx", "source.xlsx", "result.xlsx"))
    if any(p.exists() for p in (reference, raw, output)):
        raise SystemExit("v2_demo에 예시 파일이 이미 있습니다. 새 폴더에서 실행하세요.")
    wb = Workbook(); ws = wb.active; ws.title = "요약"
    ws["A1"] = "가상 자료 — 지역별 보고서"
    # Keep the report's header at row 3 so shifted layouts are exercised.
    for rn, values in enumerate((("지역", "2026-01", "2026-02", "합계"), ("동쪽", 0, 0, 0), ("서쪽", 0, 0, 0), ("합계", 0, 0, 0)), 3):
        for cn, value in enumerate(values, 1): ws.cell(rn, cn, value)
    wb.save(reference)
    wb = Workbook(); ws = wb.active; ws.title = "기록"
    ws.append(["가상 원본"]); ws.append([])
    ws.append(["지역", "품목", "거래일", "수량"])
    for row in (("동쪽", "X01", date(2026,1,5), 100), ("서쪽", "X01", date(2026,1,6), 50), ("동쪽", "X02", date(2026,2,2), 200), ("서쪽", "X02", date(2026,2,4), 80)):
        ws.append(row)
    wb.save(raw)
    plan = {"version":2,"approved":False,"source":{"sheet":"기록","header_row":3},"unresolved":[],
            "mappings":[{"id":"가상예시","mode":"table","target":{"sheet":"요약","start_cell":"A3","clear_range":"A3:D6","allow_expand":True},"group_by":["지역"],"pivot":{"column":"거래일","kind":"yearmonth"},"aggregate":"sum","terms":[{"column":"수량","factor":1}],"filters":[],"unit":1,"include_totals":True}]}
    approved = J.approve_plan(plan, reference, raw)
    result = J.generate_report(reference, raw, approved, output)
    expected = {("요약","B4"):100,("요약","C4"):200,("요약","D4"):300,("요약","B5"):50,("요약","C5"):80,("요약","D5"):130,("요약","D6"):430}
    J.verify_output(output, expected)
    (dest/"rules.json").write_text(json.dumps(approved, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__": main()
