"""범용 보고서 v2의 데스크톱 설정 화면. tkinter 외 별도 GUI 패키지 불필요."""
import json
from pathlib import Path


def describe_plan(plan):
    """JSON 설정을 사용자가 검토할 수 있는 한글 설명으로 바꾼다."""
    aggregate_names = {"sum": "합계", "count": "건수", "avg": "평균", "min": "최소", "max": "최대"}
    pivot_names = {"value": "값 그대로", "month": "월", "yearmonth": "연월", "year": "연도"}
    lines = []
    for i, mapping in enumerate(plan.get("mappings", []), 1):
        source = {**plan.get("source", {}), **mapping.get("source", {})}
        target = mapping.get("target", {})
        lines.extend([f"[{i}] {mapping.get('id', '보고서 항목')}",
                      f"원본: {source.get('sheet', '(미지정)')} 시트 / 제목 {source.get('header_row', '?')}행",
                      f"자료 범위: {source.get('first_data_row') or '제목 다음'}행 ~ {source.get('last_data_row') or '마지막 자료'}행"])
        if mapping.get("mode") == "table":
            lines.append(f"출력: {target.get('sheet', '?')} 시트 / {target.get('start_cell', '?')}부터 새 표 작성")
            lines.append("행 기준: " + ", ".join(mapping.get("group_by", [])))
            pivot = mapping.get("pivot")
            lines.append("열 기준: " + (f"{pivot['column']} ({pivot_names.get(pivot['kind'], pivot['kind'])})" if pivot else "없음"))
            lines.append(f"교체 범위: {target.get('clear_range') or '(명시 안 됨)'} / 빈 영역 확장: {'허용' if target.get('allow_expand') else '차단'}")
            lines.append(f"합계 행·열: {'표시' if mapping.get('include_totals', True) else '표시 안 함'}")
        else:
            lines.append(f"출력: {target.get('sheet', '?')}!{target.get('cell', '?')}")
        terms = mapping.get("terms", [])
        expression = " + ".join(f"{term.get('column', '?')} × {term.get('factor', 1)}" for term in terms)
        lines.append(f"계산: {aggregate_names.get(mapping.get('aggregate'), mapping.get('aggregate', '?'))} / {expression or '행 수 집계'} / 단위 {mapping.get('unit', 1)}로 나눔")
        lines.append(f"반올림: 소수 {mapping['round_digits']}자리" if "round_digits" in mapping else "반올림: 별도 지정 없음")
        lines.append(f"기존 수식 교체: {'허용' if mapping.get('allow_overwrite_formulas') else '차단'}")
        filters = mapping.get("filters", [])
        lines.append("필터: " + (" 그리고 ".join(f"{f['column']} {f['op']} {f['value']} ({f.get('type', 'text')})" for f in filters) if filters else "없음"))
        lines.append("")
    if plan.get("unresolved"):
        lines.append("확인이 필요한 항목 — 해결 전 생성 불가:")
        lines.extend(f"• {item}" for item in plan["unresolved"])
    if plan.get("review_required"):
        lines.append("검토 안내:")
        items = plan["review_required"]
        lines.extend(f"• {item}" for item in (items if isinstance(items, list) else [items]))
    return "\n".join(lines) or "연결 규칙이 없습니다."


def launch_gui(engine):
    """파일/집계 조건을 명시적으로 선택한 뒤 검토·확인하여 생성한다."""
    import tkinter as tk
    from tkinter import ttk, filedialog, messagebox
    from openpyxl import load_workbook

    root = tk.Tk()
    root.title("범용 엑셀 보고서 v2")
    root.geometry("1100x950")
    root.minsize(850, 680)
    outer = ttk.Frame(root, padding=14)
    outer.pack(fill="both", expand=True)
    outer.columnconfigure(1, weight=1)
    variables = {name: tk.StringVar() for name in
                 ("reference", "raw", "output", "source_sheet", "target_sheet",
                  "amount", "minus", "pivot", "clear_range", "first_data_row", "last_data_row", "round_digits")}
    variables.update(header=tk.StringVar(value="1"), start=tk.StringVar(value="A1"),
                     unit=tk.StringVar(value="1"), aggregate=tk.StringVar(value="합계"),
                     pivot_kind=tk.StringVar(value="값 그대로"), mode=tk.StringVar(value="표 전체 다시 구성"))
    totals, expand = tk.BooleanVar(value=True), tk.BooleanVar(value=False)
    use_excel, overwrite_formulas = tk.BooleanVar(value=False), tk.BooleanVar(value=False)
    aggregates = {"합계": "sum", "건수": "count", "평균": "avg", "최소": "min", "최대": "max"}
    pivot_kinds = {"값 그대로": "value", "월": "month", "연월": "yearmonth", "연도": "year"}
    status = tk.StringVar(value="기준 보고서와 원본 데이터를 선택한 뒤 [파일 구조 확인]을 누르세요.")
    inspected = {}
    filter_rules = []

    def failure(exc):
        status.set("진행하지 못했습니다. 표시된 항목을 확인하세요.")
        messagebox.showerror("확인 필요", str(exc), parent=root)

    def choose(name):
        if name == "output":
            ext = ".xlsm" if Path(variables["reference"].get()).suffix.lower() == ".xlsm" else ".xlsx"
            path = filedialog.asksaveasfilename(parent=root, title="새 보고서 저장 위치",
                                               defaultextension=ext, filetypes=[("Excel 통합 문서", "*" + ext)])
        else:
            path = filedialog.askopenfilename(parent=root, title="기준 보고서 선택" if name == "reference" else "원본 데이터 선택",
                                             filetypes=[("Excel 통합 문서", "*.xlsx *.xlsm")])
        if path:
            variables[name].set(path)
            if name != "output":
                inspected.clear()
                status.set("선택 파일이 바뀌었습니다. [파일 구조 확인]을 다시 누르세요.")

    ttk.Label(outer, text="1. 입력 파일과 새 결과 파일", font=("", 12, "bold")).grid(row=0, column=0, columnspan=3, sticky="w")
    for row, (name, label) in enumerate((("reference", "기준 보고서"), ("raw", "원본 데이터"), ("output", "새 보고서")), 1):
        ttk.Label(outer, text=label).grid(row=row, column=0, sticky="w", pady=4)
        ttk.Entry(outer, textvariable=variables[name]).grid(row=row, column=1, sticky="ew", padx=8)
        ttk.Button(outer, text="선택…", command=lambda n=name: choose(n)).grid(row=row, column=2)

    notebook = ttk.Notebook(outer)
    notebook.grid(row=5, column=0, columnspan=3, sticky="nsew", pady=10)
    outer.rowconfigure(5, weight=1)
    form, advanced = ttk.Frame(notebook, padding=12), ttk.Frame(notebook, padding=8)
    notebook.add(form, text="표 설정")
    notebook.add(advanced, text="상세 규칙(JSON)")
    form.columnconfigure(1, weight=1)
    form.columnconfigure(3, weight=1)
    combos = {}

    def entry(row, column, label, name, choices=None):
        ttk.Label(form, text=label).grid(row=row, column=column, sticky="w", padx=(0, 8), pady=5)
        if choices is not None:
            widget = ttk.Combobox(form, textvariable=variables[name], values=choices, state="readonly")
            combos[name] = widget
        else:
            widget = ttk.Entry(form, textvariable=variables[name])
        widget.grid(row=row, column=column + 1, sticky="ew", padx=(0, 18))
        return widget

    entry(0, 0, "원본 시트", "source_sheet", [])
    entry(0, 2, "원본 제목 행 번호", "header")
    entry(1, 0, "결과 시트", "target_sheet", [])
    entry(1, 2, "시작 셀 / 갱신할 셀", "start")
    entry(2, 0, "기존 표를 비울 범위", "clear_range")
    entry(2, 2, "집계 방식", "aggregate", list(aggregates))
    ttk.Label(form, text="예: A3:F20. 이 범위의 기존 내용은 새 표로 대체됩니다.", foreground="#555555").grid(row=3, column=0, columnspan=4, sticky="w")
    ttk.Label(form, text="행으로 묶을 열(여러 개 선택 가능)").grid(row=4, column=0, columnspan=4, sticky="w", pady=(10, 3))
    groups = tk.Listbox(form, selectmode="multiple", height=5, exportselection=False)
    groups.grid(row=5, column=0, columnspan=4, sticky="nsew")
    form.rowconfigure(5, weight=1)
    entry(6, 0, "금액/값 열", "amount", [])
    entry(6, 2, "뺄 값 열(선택)", "minus", [])
    entry(7, 0, "열로 나눌 기준(선택)", "pivot", [])
    entry(7, 2, "열 기준의 표시 방법", "pivot_kind", list(pivot_kinds))
    entry(8, 0, "단위(값 ÷ 단위)", "unit")
    ttk.Checkbutton(form, text="합계 행·열 표시", variable=totals).grid(row=8, column=2, columnspan=2, sticky="w")
    entry(9, 0, "원본 자료 시작 행(선택)", "first_data_row")
    entry(9, 2, "원본 자료 마지막 행(선택)", "last_data_row")
    entry(10, 0, "반올림 소수 자릿수(선택)", "round_digits")
    ttk.Checkbutton(form, text="기존 표의 수식 교체 허용", variable=overwrite_formulas).grid(row=10, column=2, columnspan=2, sticky="w")
    ttk.Checkbutton(form, text="표 크기가 커지면 비어 있는 영역으로 확장 허용", variable=expand).grid(row=11, column=0, columnspan=4, sticky="w", pady=4)
    ttk.Checkbutton(form, text="Windows Excel로 보존한 수식 재계산(pywin32 필요)", variable=use_excel).grid(row=12, column=0, columnspan=4, sticky="w", pady=4)
    filter_summary = tk.StringVar(value="필터 없음: 원본의 모든 자료 행 대상")
    ttk.Label(form, textvariable=filter_summary).grid(row=13, column=0, columnspan=3, sticky="w")
    entry(15, 0, "작성 방식", "mode", ["표 전체 다시 구성", "특정 셀 값 갱신"])
    ttk.Label(form, text="특정 셀 갱신은 행·열 기준 없이 필터로 집계 대상을 정합니다.", foreground="#555555").grid(row=16, column=0, columnspan=4, sticky="w")

    rules_text = tk.Text(advanced, wrap="none", undo=True, font=("Consolas", 10))
    scrollbar = ttk.Scrollbar(advanced, orient="vertical", command=rules_text.yview)
    rules_text.configure(yscrollcommand=scrollbar.set)
    rules_text.pack(side="left", fill="both", expand=True)
    scrollbar.pack(side="right", fill="y")

    def load_columns(event=None):
        try:
            sheet = variables["source_sheet"].get()
            if not sheet:
                raise ValueError("원본 시트를 선택하세요.")
            row = int(variables["header"].get())
            if row < 1:
                raise ValueError("제목 행 번호는 1 이상이어야 합니다.")
            workbook = load_workbook(variables["raw"].get(), data_only=False, read_only=True)
            try:
                ws = workbook[sheet]
                headers = [str(c.value).strip() for c in ws[row] if c.value is not None and str(c.value).strip()]
            finally:
                workbook.close()
            if len(set(headers)) != len(headers):
                raise ValueError("같은 제목의 열이 있습니다. 원본에서 구분하거나 상세 규칙을 검토하세요.")
            groups.delete(0, "end")
            for h in headers:
                groups.insert("end", h)
            for name in ("amount", "minus", "pivot"):
                combos[name]["values"] = [""] + headers
                variables[name].set("")
            status.set("행 기준·값 열·출력 위치를 직접 선택하세요. 자동으로 생성하지 않습니다.")
        except Exception as exc:
            failure(exc)

    def source_changed(event=None):
        data = inspected.get("raw", {})
        sheet = next((s for s in data.get("sheets", []) if s["name"] == variables["source_sheet"].get()), None)
        if sheet and sheet.get("tables"):
            tables = sheet["tables"]
            best = sorted(tables, key=lambda t: t.get("score", 0), reverse=True)[0]
            variables["header"].set(str(best["header_row"]))
            variables["first_data_row"].set(str(best.get("first_data_row") or ""))
            # Keep future periods' longer sources visible unless a user sets a bound.
            variables["last_data_row"].set("")
        load_columns()

    def inspect_files():
        try:
            for name in ("reference", "raw"):
                path = variables[name].get()
                if not Path(path).is_file():
                    raise ValueError("기준 보고서와 원본 데이터를 모두 선택하세요.")
                inspected[name] = engine.inspect_workbook(path)
            source_names = [s["name"] for s in inspected["raw"]["sheets"]]
            target_names = [s["name"] for s in inspected["reference"]["sheets"]]
            combos["source_sheet"]["values"] = source_names
            combos["target_sheet"]["values"] = target_names
            variables["source_sheet"].set(source_names[0] if len(source_names) == 1 else "")
            variables["target_sheet"].set(target_names[0] if len(target_names) == 1 else "")
            if variables["source_sheet"].get():
                source_changed()
            else:
                status.set("시트가 여러 개입니다. 사용할 원본 시트와 결과 시트를 선택하세요.")
        except Exception as exc:
            failure(exc)

    combos["source_sheet"].bind("<<ComboboxSelected>>", source_changed)
    ttk.Button(outer, text="파일 구조 확인", command=inspect_files).grid(row=4, column=0, columnspan=3, sticky="w", pady=(6, 0))
    ttk.Button(form, text="열 다시 읽기", command=load_columns).grid(row=14, column=0, sticky="w", pady=6)

    def edit_filters():
        dialog = tk.Toplevel(root)
        dialog.title("필터 조건")
        dialog.transient(root)
        column, operator, kind, value = (tk.StringVar() for _ in range(4))
        operator.set("=")
        kind.set("text")
        for i, (label, var, options) in enumerate((("원본 열", column, list(combos["amount"]["values"])), ("조건", operator, ["=", "<>", ">", "<", ">=", "<="]), ("값 종류", kind, ["text", "number", "date"]))):
            ttk.Label(dialog, text=label).grid(row=0, column=i, padx=6)
            ttk.Combobox(dialog, textvariable=var, values=options, state="readonly", width=20).grid(row=1, column=i, padx=6)
        ttk.Label(dialog, text="비교할 값(날짜: YYYY-MM-DD)").grid(row=2, column=0, columnspan=3, sticky="w", padx=6)
        ttk.Entry(dialog, textvariable=value).grid(row=3, column=0, columnspan=3, sticky="ew", padx=6)
        listed = tk.Listbox(dialog, height=6, width=75)
        listed.grid(row=4, column=0, columnspan=3, padx=6, pady=6)

        def refresh():
            listed.delete(0, "end")
            for f in filter_rules:
                listed.insert("end", f"{f['column']} {f['op']} {f['value']} ({f['type']})")
            filter_summary.set(f"필터 {len(filter_rules)}개: 모든 조건을 만족하는 행만 집계" if filter_rules else "필터 없음: 원본의 모든 자료 행 대상")

        def add():
            if not column.get():
                return
            filter_rules.append({"column": column.get(), "op": operator.get(), "type": kind.get(), "value": value.get()})
            refresh()

        def remove():
            for i in reversed(listed.curselection()):
                del filter_rules[i]
            refresh()

        ttk.Button(dialog, text="조건 추가", command=add).grid(row=5, column=0, pady=6)
        ttk.Button(dialog, text="선택 삭제", command=remove).grid(row=5, column=1)
        ttk.Button(dialog, text="닫기", command=dialog.destroy).grid(row=5, column=2)
        refresh()

    ttk.Button(form, text="필터 조건…", command=edit_filters).grid(row=13, column=3, sticky="e")

    def current_plan():
        if notebook.index(notebook.select()) == 1:
            return json.loads(rules_text.get("1.0", "end"))
        group_columns = [groups.get(i) for i in groups.curselection()]
        aggregate = aggregates[variables["aggregate"].get()]
        cell_mode = variables["mode"].get() == "특정 셀 값 갱신"
        if not group_columns and not cell_mode:
            raise ValueError("행으로 묶을 열을 하나 이상 선택하세요.")
        if not variables["target_sheet"].get() or not variables["source_sheet"].get():
            raise ValueError("원본 시트와 결과 시트를 선택하세요.")
        terms = []
        if aggregate != "count":
            if not variables["amount"].get():
                raise ValueError("집계할 값 열을 선택하세요.")
            terms.append({"column": variables["amount"].get(), "factor": 1})
            if variables["minus"].get():
                terms.append({"column": variables["minus"].get(), "factor": -1})
        mapping = {"id": "표1", "mode": "table", "target": {"sheet": variables["target_sheet"].get(),
                   "start_cell": variables["start"].get().strip(), "clear_range": variables["clear_range"].get().strip() or None,
                   "allow_expand": bool(expand.get())}, "group_by": group_columns, "aggregate": aggregate,
                   "terms": terms, "filters": list(filter_rules), "unit": variables["unit"].get().strip(),
                   "include_totals": bool(totals.get()), "sort": "source", "allow_overwrite_formulas": bool(overwrite_formulas.get())}
        if variables["round_digits"].get().strip():
            mapping["round_digits"] = int(variables["round_digits"].get())
        if cell_mode:
            mapping["mode"] = "cells"
            mapping["target"] = {"sheet": variables["target_sheet"].get(), "cell": variables["start"].get().strip()}
            for key in ("group_by", "include_totals", "sort"):
                mapping.pop(key, None)
        elif variables["pivot"].get():
            mapping["pivot"] = {"column": variables["pivot"].get(), "kind": pivot_kinds[variables["pivot_kind"].get()]}
        source = {"sheet": variables["source_sheet"].get(), "header_row": int(variables["header"].get())}
        for name in ("first_data_row", "last_data_row"):
            if variables[name].get().strip():
                source[name] = int(variables[name].get())
        return {"version": 2, "approved": False, "source": source, "mappings": [mapping], "unresolved": []}

    def set_json(plan):
        rules_text.delete("1.0", "end")
        rules_text.insert("1.0", json.dumps(plan, ensure_ascii=False, indent=2))

    def preview_and_generate():
        try:
            reference, raw, output = (variables[n].get() for n in ("reference", "raw", "output"))
            if not all((reference, raw, output)):
                raise ValueError("기준 보고서·원본·새 보고서 저장 위치를 모두 선택하세요.")
            if Path(output).exists() or Path(output).is_symlink():
                raise ValueError("새 보고서 위치에 이미 파일이 있습니다. 다른 이름을 선택하세요.")
            plan = current_plan()
            summary = describe_plan(plan)
            summary += "\nWindows Excel 재계산: " + ("사용" if use_excel.get() else "사용 안 함 — 보존한 수식은 별도 재계산 필요")
            # Human approval is bound to the exact reviewed configuration by the engine.
            review = tk.Toplevel(root)
            review.title("생성 설정 검토")
            review.geometry("850x650")
            review.transient(root)
            review.grab_set()
            ttk.Label(review, text=f"기준: {Path(reference).name}\n원본: {Path(raw).name}\n결과: {output}", padding=10, wraplength=800).pack(fill="x")
            review_text = tk.Text(review, wrap="word", font=("", 11))
            review_text.insert("1.0", summary)
            review_text.configure(state="disabled")
            review_text.pack(fill="both", expand=True, padx=10)
            approved_by_user = []

            def confirm():
                approved_by_user.append(True)
                review.destroy()

            bar = ttk.Frame(review, padding=10)
            bar.pack(fill="x")
            ttk.Label(bar, text="집계·출력 기준을 검토한 경우 [확인 후 생성]을 누르세요.").pack(side="left")
            ttk.Button(bar, text="취소", command=review.destroy).pack(side="right")
            ttk.Button(bar, text="확인 후 생성", command=confirm).pack(side="right", padx=8)
            root.wait_window(review)
            if not approved_by_user:
                return
            approved = engine.approve_plan(plan, reference, raw)
            status.set("연결한 셀 작성 및 저장값 대조 중…")
            root.update_idletasks()
            result = engine.generate_report(reference, raw, approved, output, use_excel=bool(use_excel.get()))
            set_json(approved)
            notices = []
            if result.get("formula_recalculation_required"):
                notices.append("보존한 수식의 계산 결과는 아직 대조하지 않았습니다. Excel에서 열어 전체 재계산 후 저장해야 합니다.")
            unmapped = result.get("unmapped_numeric_cells", [])
            if unmapped:
                sample = "\n".join(f"{item.get('sheet', '?')}!{item.get('cell', '?')} : {item.get('value', '')}" if isinstance(item, dict) else str(item) for item in unmapped[:20])
                notices.append(f"기존 보고서의 연결되지 않은 숫자 {len(unmapped)}칸이 그대로 남아 있습니다. 새 원본과 대조한 결과가 아니므로 확인이 필요합니다.\n{sample}" + ("\n…" if len(unmapped) > 20 else ""))
            complete = f"연결한 {result.get('cells_written', '')}칸 작성·저장값 대조 완료"
            if notices:
                status.set(f"{complete} / 추가 확인 필요: {output}")
                messagebox.showwarning("연결한 셀 작성 완료 — 추가 확인 필요", f"{complete}\n{output}\n\n" + "\n\n".join(notices), parent=root)
            else:
                status.set(f"{complete}: {output}")
                messagebox.showinfo("연결한 셀 작성 완료", f"{complete}\n{output}", parent=root)
        except Exception as exc:
            failure(exc)

    def save_rules():
        try:
            plan = current_plan()
            path = filedialog.asksaveasfilename(parent=root, title="설정 저장", defaultextension=".json", filetypes=[("JSON 설정", "*.json")])
            if path:
                # Saving a configuration never grants approval by itself.
                plan["approved"] = False
                plan.pop("approval", None)
                Path(path).write_text(json.dumps(plan, ensure_ascii=False, indent=2), encoding="utf-8")
                status.set(f"설정 저장: {path}")
        except Exception as exc:
            failure(exc)

    def open_rules():
        try:
            path = filedialog.askopenfilename(parent=root, title="설정 열기", filetypes=[("JSON 설정", "*.json")])
            if path:
                set_json(json.loads(Path(path).read_text(encoding="utf-8-sig")))
                notebook.select(advanced)
                status.set("설정을 불러왔습니다. 입력 파일을 선택하고 설정을 검토한 뒤 생성하세요.")
        except Exception as exc:
            failure(exc)

    def propose_rules():
        try:
            reference, raw = variables["reference"].get(), variables["raw"].get()
            if not Path(reference).is_file() or not Path(raw).is_file():
                raise ValueError("기준 보고서와 원본 데이터를 먼저 선택하세요.")
            plan = engine.propose_plan(reference, raw)
            plan["approved"] = False
            plan.pop("approval", None)
            set_json(plan)
            notebook.select(advanced)
            status.set("연결 후보를 가져왔습니다. 미확정 질문을 해결하고 모든 연결·계산 기준을 검토해야 생성할 수 있습니다.")
            messagebox.showinfo("분석 결과 — 미승인 초안", describe_plan(plan), parent=root)
        except Exception as exc:
            failure(exc)

    buttons = ttk.Frame(outer)
    buttons.grid(row=6, column=0, columnspan=3, sticky="ew")
    ttk.Button(buttons, text="설정 열기…", command=open_rules).pack(side="left", padx=(0, 8))
    ttk.Button(buttons, text="설정 저장…", command=save_rules).pack(side="left")
    ttk.Button(buttons, text="분석 결과에서 연결 후보 가져오기", command=propose_rules).pack(side="left", padx=8)
    ttk.Button(buttons, text="설정 검토 후 보고서 생성", command=preview_and_generate).pack(side="right")
    ttk.Label(outer, textvariable=status, wraplength=980).grid(row=7, column=0, columnspan=3, sticky="w", pady=(10, 0))
    root.mainloop()

